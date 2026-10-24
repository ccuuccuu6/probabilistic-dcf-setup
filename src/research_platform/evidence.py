"""Archive public disclosures and resolve figures from their original text/cells.

AI chooses references; this module reads the referenced numbers. Source text is
untrusted data. No company mappings, generated code, or model API calls live here.
"""

import copy
import csv
import hashlib
import io
import ipaddress
import json
import logging
import math
import os
import re
import socket
import struct
import zipfile
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from http.client import HTTPSConnection
from pathlib import Path
from urllib.parse import quote, urlsplit
from urllib.request import HTTPRedirectHandler, HTTPSHandler, ProxyHandler, Request, build_opener

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
import xlrd


MAX_BYTES = 25 * 1024 * 1024
MAX_UNITS = 3000
FORMATS = {"dot_decimal", "comma_decimal"}
LOCAL_CATEGORIES = {"visible_alpha", "financial_reports", "transcripts"}
LOCAL_FORMATS = {".pdf": "pdf", ".xlsx": "xlsx", ".xls": "xls", ".html": "html", ".htm": "html",
                 ".txt": "txt", ".csv": "csv"}


def _binary_format(payload: bytes):
    if payload.startswith(b"%PDF"):
        return "pdf"
    if payload.startswith(b"PK\x03\x04"):
        return "xlsx"
    if (payload.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
            or payload[:2] in {b"\x09\x00", b"\x09\x02", b"\x09\x04", b"\x09\x08"}):
        return "xls"
    return None


def _xls_formula_cells(workbook, sheet_index):
    """Locate formula caches without executing formulas or interpreting tokens.

    xlrd's public cells expose only cached results. Its pinned 2.0.2 BIFF stream
    and worksheet offsets let this guard distinguish a blank formula cache from
    an ordinary blank cell. Calculation syntax remains the parser's concern.
    """
    if workbook.biff_version < 50:
        raise ValueError("Legacy spreadsheet versions older than Excel 5 are unsupported; retain the original and supply a self-contained export.")
    memory = workbook.mem
    position = workbook._sh_abs_posn[sheet_index]
    cells = set()
    while position + 4 <= len(memory):
        opcode, length = struct.unpack_from("<HH", memory, position)
        position += 4
        if position + length > len(memory):
            raise ValueError("Legacy spreadsheet has an incomplete worksheet record.")
        if opcode in xlrd.biffh.XL_FORMULA_OPCODES:
            if length < 16:
                raise ValueError("Legacy spreadsheet has an incomplete formula cache record.")
            cells.add(struct.unpack_from("<HH", memory, position))
        position += length
        if opcode == xlrd.biffh.XL_EOF:
            return cells
    raise ValueError("Legacy spreadsheet worksheet has no complete end record.")


def _local_file(root: Path, relative: str) -> Path:
    """Resolve an explicit archive member, refusing links and paths outside it."""
    root = _archive_path(root)
    relative_path = Path(relative)
    if relative_path.is_absolute() or not relative_path.parts or ".." in relative_path.parts:
        raise ValueError("Local source paths must stay relative to the supplied archive.")
    candidate = root
    for part in relative_path.parts:
        candidate = candidate / part
        if candidate.is_symlink() or (hasattr(candidate, "is_junction") and candidate.is_junction()):
            raise ValueError("Local source archives must not contain symbolic links or junctions.")
    resolved = candidate.resolve(strict=True)
    resolved.relative_to(root)
    if not resolved.is_file():
        raise ValueError("Local source is not a regular file.")
    return resolved


def _local_payload(path: Path) -> bytes:
    if not 0 < path.stat().st_size <= MAX_BYTES:
        raise ValueError(f"Local source is empty or exceeds 25 MB: {path.name}")
    with path.open("rb") as source:
        payload = source.read(MAX_BYTES + 1)
    if not payload or len(payload) > MAX_BYTES:
        raise ValueError(f"Local source is empty or exceeds 25 MB: {path.name}")
    return payload


def _archive_path(path: Path) -> Path:
    """Windows extended paths keep deeply nested originals readable, unchanged."""
    path = Path(path).absolute()
    if os.name == "nt" and not str(path).startswith("\\\\?\\"):
        raw = str(path)
        path = Path("\\\\?\\UNC\\" + raw[2:] if raw.startswith("\\\\") else "\\\\?\\" + raw)
    return path


def inventory_local_sources(root: Path) -> dict:
    """Inventory user-supplied originals without network or model requests.

    sources.csv is an optional provenance index, not a requirement to prepare a
    historical model. Its statements remain supplied metadata, not verified facts.
    Content hashes and this manifest are frozen before any paid work starts.
    """
    supplied_root = _archive_path(Path(root).expanduser())
    if supplied_root.is_symlink() or (hasattr(supplied_root, "is_junction") and supplied_root.is_junction()):
        raise ValueError("The source archive must be a real directory, not a link.")
    root = supplied_root.resolve(strict=True)
    if not root.is_dir():
        raise ValueError("The source archive must be a directory.")
    metadata, index_rows, index_hash = {}, [], None
    index_path = root / "sources.csv"
    if index_path.exists():
        index_payload = _local_payload(_local_file(root, "sources.csv"))
        index_hash = hashlib.sha256(index_payload).hexdigest()
        reader = csv.DictReader(io.StringIO(index_payload.decode("utf-8-sig")))
        if not reader.fieldnames:
            raise ValueError("sources.csv must have a header row.")
        for raw_row in reader:
            if None in raw_row:
                raise ValueError("sources.csv has more values than its header row.")
            row = {re.sub(r"[^a-z0-9]+", "_", key.lower()).strip("_"): (value or "").strip()
                   for key, value in raw_row.items()}
            index_rows.append(row)
            relative = next((row[key] for key in ("relative_filename", "relative_path", "filename", "file", "path")
                             if row.get(key)), "")
            if relative:
                relative = relative.replace("\\", "/")
                if Path(relative).is_absolute() or ".." in Path(relative).parts:
                    raise ValueError("sources.csv must use relative filenames inside the archive.")
                if relative in metadata:
                    raise ValueError(f"sources.csv repeats a filename: {relative}")
                metadata[relative] = row
    sources, members = [], []
    for child in sorted(root.iterdir(), key=lambda path: path.name.casefold()):
        if child.name == "sources.csv" or child.name.startswith(".") or child.name == "Thumbs.db":
            continue
        if child.is_symlink() or (hasattr(child, "is_junction") and child.is_junction()):
            raise ValueError("Local source archives must not contain symbolic links or junctions.")
        if child.name not in LOCAL_CATEGORIES or not child.is_dir():
            raise ValueError(f"Place source files in visible_alpha/, financial_reports/ or transcripts/: {child.name}")
        for path in sorted(child.rglob("*"), key=lambda item: item.as_posix().casefold()):
            if path.is_symlink() or (hasattr(path, "is_junction") and path.is_junction()):
                raise ValueError("Local source archives must not contain symbolic links or junctions.")
            if (path.is_dir() or path.name == "Thumbs.db" or path.name.startswith("~$")
                    or any(part.startswith(".") for part in path.relative_to(root).parts)):
                continue
            relative = path.relative_to(root).as_posix()
            path = _local_file(root, relative)
            if path.suffix.lower() not in LOCAL_FORMATS:
                raise ValueError(f"Unsupported local source {relative}; use PDF, XLSX, XLS, HTML, UTF-8 TXT or CSV originals.")
            payload = _local_payload(path)
            digest = hashlib.sha256(payload).hexdigest()
            row = metadata.get(relative, {})
            supplied_hash = row.get("sha_256_hash") or row.get("sha256") or row.get("sha_256")
            if supplied_hash and supplied_hash.lower() not in {"unknown", "not_calculated", "not calculated", "n/a"}:
                if supplied_hash.lower() != digest:
                    raise ValueError(f"sources.csv hash disagrees with the actual file: {relative}")
            publication = next((row[key] for key in ("verified_publication_date", "publication_date", "published_at")
                                if row.get(key) and row[key].lower() != "unknown"), None)
            filename_date = re.search(r"_pub-(\d{4}-\d{2}-\d{2})(?:[_.-]|$)", path.name)
            publication_hint = filename_date.group(1) if filename_date else None
            url = next((row[key] for key in ("original_download_url", "download_url", "url", "source_url")
                        if row.get(key)), "local:///" + quote(relative))
            if not url.startswith("local:///"):
                parsed = urlsplit(url)
                if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
                    raise ValueError(f"Source metadata URLs must be credential-free HTTPS: {relative}")
            # The ID contains both path and bytes: different editions with the
            # same title cannot silently replace one another in a resumed run.
            identity = hashlib.sha256((relative + "\0" + digest).encode("utf-8")).hexdigest()[:24]
            source = {"id": "local_" + identity, "local_path": relative, "sha256": digest,
                      "format": LOCAL_FORMATS[path.suffix.lower()], "category": child.name,
                      "url": url, "title": row.get("original_title") or row.get("title") or path.stem,
                      "published_at": publication, "filename_publication_hint": publication_hint,
                      "reason": "User-supplied original; semantic and reporting-basis validation still required.",
                      "kind": "vendor_export" if child.name == "visible_alpha" else
                              "transcript" if child.name == "transcripts" else "company_document",
                      "supplied_metadata": row, "metadata_verification": "user_supplied_not_verified"}
            if child.name=='visible_alpha':
                role=row.get('dataset_role') or row.get('source_role')
                if role:role=role.casefold().replace('-','_').replace(' ','_').replace('standardised','standardized')
                # User-labelled dataset filenames are role designations only.
                # Actual issuer, periods, definitions and numbers still need evidence.
                if not role:
                    label=path.stem.casefold().replace('-',' ')
                    role='company_reported' if 'company reported' in label else 'standardized' if 'standardised' in label or 'standardized' in label else None
                if role:
                    if role not in ('company_reported','standardized','combined'):raise ValueError('Unsupported vendor dataset role')
                    source['dataset_role']=role
                    source['dataset_role_origin']='sources.csv' if row.get('dataset_role') or row.get('source_role') else 'user_dataset_filename_label'
            sources.append(source)
            members.append({"local_path": relative, "sha256": digest})
    if not sources:
        raise ValueError("No source documents found; place originals in the archive's visible_alpha/, financial_reports/ or transcripts/ folders.")
    manifest = {"members": members, "index_sha256": index_hash}
    manifest_hash = hashlib.sha256(json.dumps(manifest, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    return {"root": str(root), "sources": sources, "manifest_sha256": manifest_hash,
            "index_sha256": index_hash, "index_rows": index_rows}


def parse_number(value, number_format="dot_decimal"):
    """Parse a complete numeric cell; missing cells are never silently zero."""
    if number_format not in FORMATS:
        raise ValueError("Choose dot_decimal or comma_decimal explicitly.")
    if isinstance(value, bool) or value is None:
        raise ValueError("Missing or invalid financial number.")
    if isinstance(value, (int, float)):
        result = float(value)
    elif isinstance(value, str):
        text = value.strip().replace("\u00a0", " ").replace("\u202f", " ").replace("\u2212", "-")
        text = text.replace("\u2019", "'").replace("\u2018", "'")
        if text in {"-", "\u2013", "\u2014"}:
            return 0.0
        percent = text.endswith("%")
        if percent:
            text = text[:-1].strip()
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1].strip()
            if text.startswith(("+", "-")):
                raise ValueError("A parenthesized number cannot carry another sign.")
        decimal, thousands = (".", ",") if number_format == "dot_decimal" else (",", ".")
        integer = rf"(?:\d+|\d{{1,3}}(?:{re.escape(thousands)}\d{{3}})+|\d{{1,3}}(?: \d{{3}})+|\d{{1,3}}(?:'\d{{3}})+)"
        if not re.fullmatch(rf"[+-]?{integer}(?:{re.escape(decimal)}\d+)?", text):
            raise ValueError(f"Unrecognized financial number: {value!r}")
        text = text.replace(thousands, "").replace(" ", "").replace("'", "").replace(decimal, ".")
        result = float(text) * (-1 if negative else 1) / (100 if percent else 1)
    else:
        raise ValueError("Invalid financial number type.")
    if not math.isfinite(result):
        raise ValueError("Financial numbers must be finite.")
    return result


# Thousands separated by one space occur inside a PDF table cell; multiple spaces
# delimit table columns. Both decimal conventions remain explicit at resolution.
_NUMBER = re.compile(
    r"(?<![\w.,'\u2018\u2019])(?:\([+\-\u2212]?\d(?:[\d.,]*\d)?(?:[ '\u2018\u2019\u00a0\u202f]\d{3})*(?:[.,]\d+)?\)"
    r"|[+\-\u2212]?\d(?:[\d.,]*\d)?(?:[ '\u2018\u2019\u00a0\u202f]\d{3})*(?:[.,]\d+)?)(?:\s*%)?(?![\w.,'\u2018\u2019])"
)


def _tokens(text):
    tokens = []
    for match in _NUMBER.finditer(text):
        raw = match.group()
        try:
            parse_number(raw)
        except ValueError:
            try:
                parse_number(raw, "comma_decimal")
            except ValueError:
                continue
        tokens.append({"index": len(tokens), "raw": raw, "start": match.start(), "end": match.end()})
    # A dash is nil only in a delimited table cell, never a prose dash or list mark.
    for match in re.finditer(r"(?<=\s)[\-\u2013\u2014](?=\s|$)", text):
        before, after = text[:match.start()], text[match.end():]
        if ((before.endswith("\t") or before.endswith("  "))
                and (after.startswith("\t") or after.startswith("  ") or not after or after.startswith("\n"))):
            tokens.append({"raw": match.group(), "start": match.start(), "end": match.end()})
    tokens.sort(key=lambda token: token["start"])
    for index, token in enumerate(tokens):
        token["index"] = index
    return tokens


def _public_url(url):
    """Reject local destinations before each connection, including redirects."""
    parsed = urlsplit(url)
    if (parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password
            or parsed.port not in (None, 443)):
        raise ValueError("Sources must use public HTTPS URLs without credentials or custom ports.")
    _public_addresses(parsed.hostname)
    return url


def _public_addresses(hostname):
    try:
        addresses = socket.getaddrinfo(hostname, 443, type=socket.SOCK_STREAM)
    except OSError as exc:
        raise ValueError("Source host could not be resolved.") from exc
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise ValueError("A source URL resolved to a nonpublic address.")
    return [item[4][0] for item in addresses]


class _PublicHTTPSConnection(HTTPSConnection):
    def connect(self):
        # Connect to the validated IP itself; a second DNS lookup must not turn a
        # public hostname into an internal destination between validation and use.
        address = _public_addresses(self.host)[0]
        self.sock = socket.create_connection((address, self.port), self.timeout, self.source_address)
        self.sock = self._context.wrap_socket(self.sock, server_hostname=self.host)


class _PublicHTTPSHandler(HTTPSHandler):
    def https_open(self, request):
        return self.do_open(_PublicHTTPSConnection, request, context=self._context)


class _SafeRedirect(HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, message, headers, newurl):
        _public_url(newurl)
        return super().redirect_request(request, fp, code, message, headers, newurl)


def download(url):
    """Fetch a bounded disclosure; never follow a redirect to a private service."""
    _public_url(url)
    request = Request(url, headers={"User-Agent": "Mozilla/5.0 ResearchPlatform/0.1"})
    # Proxies can redirect connections to local services, so source retrieval uses
    # direct validated HTTPS connections, retaining normal TLS certificate checks.
    with build_opener(ProxyHandler({}), _SafeRedirect(), _PublicHTTPSHandler()).open(request, timeout=30) as response:
        final_url = _public_url(response.url)
        payload = response.read(MAX_BYTES + 1)
        if len(payload) > MAX_BYTES:
            raise ValueError("Source exceeds the 25 MB download limit.")
        return payload, final_url, response.headers.get_content_type()


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.hidden += 1
        if not self.hidden and tag in {"p", "div", "br", "tr", "h1", "h2", "h3", "li", "table"}:
            self.parts.append("\n")
        if not self.hidden and tag in {"td", "th"}:
            self.parts.append("\t")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.hidden:
            self.hidden -= 1
        if not self.hidden and tag in {"p", "div", "tr", "h1", "h2", "h3", "li", "table"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _units(payload, file_format):
    if file_format == "pdf":
        pdf = PdfReader(io.BytesIO(payload))
        if len(pdf.pages) > 500:
            raise ValueError("PDF exceeds the 500-page source limit.")
        for index, page in enumerate(pdf.pages, 1):
            messages = []
            capture = logging.Handler()
            capture.emit = lambda record: messages.append(record.getMessage())
            logger = logging.getLogger("pypdf")
            logger.addHandler(capture)
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            finally:
                logger.removeHandler(capture)
            yield {"unit_id": f"page:{index}", "text": text, "numbers": _tokens(text), "warnings": messages}
            if messages:
                # Rotated headings can be omitted by layout extraction. Give the
                # reviewer the complementary plain extraction instead of hiding it.
                plain = page.extract_text(extraction_mode="plain") or ""
                yield {"unit_id": f"page:{index}:plain", "text": plain, "numbers": _tokens(plain),
                       "warnings": ["Alternative text order; verify column meanings against the original page."]}
    elif file_format == "xls":
        parser_log = io.StringIO()
        workbook = xlrd.open_workbook(file_contents=payload, formatting_info=True,
                                     on_demand=True, logfile=parser_log)
        try:
            for sheet_index in range(workbook.nsheets):
                formula_cells = _xls_formula_cells(workbook, sheet_index)
                sheet = workbook.sheet_by_index(sheet_index)
                if sheet.nrows > 10000 or sheet.ncols > 300:
                    raise ValueError("Spreadsheet dimensions exceed the bounded extraction limit.")
                for row_index in range(sheet.nrows):
                    populated, numbers, formula_addresses, errors = [], [], [], []
                    for column_index in range(sheet.ncols):
                        cell = sheet.cell(row_index, column_index)
                        coordinate = f"{get_column_letter(column_index + 1)}{row_index + 1}"
                        formula = (row_index, column_index) in formula_cells
                        if formula:
                            formula_addresses.append(coordinate)
                            if (cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR)
                                    or cell.value == ""):
                                raise ValueError(
                                    f"Legacy spreadsheet formula has no usable cached result at {sheet.name}!{coordinate}; "
                                    "recalculate and save a separate self-contained export. Python does not calculate formulas or run macros.")
                        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or cell.value == "":
                            continue
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, "#UNKNOWN_ERROR")
                            errors.append(coordinate)
                        populated.append(f"{coordinate}: {value}")
                        if cell.ctype not in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_TEXT):
                            continue
                        for number_format in ("dot_decimal", "comma_decimal"):
                            try:
                                parse_number(value, number_format)
                            except ValueError:
                                continue
                            format_key = workbook.xf_list[cell.xf_index].format_key
                            display_format = workbook.format_map[format_key].format_str
                            numbers.append({"index": len(numbers), "raw": value, "locator": coordinate,
                                            "display_format": display_format})
                            break
                    if not populated:
                        continue
                    unit = {"unit_id": f"{sheet.name}!R{row_index + 1}",
                            "text": " | ".join(populated), "numbers": numbers}
                    warnings = []
                    if formula_addresses:
                        warnings.append("Cached formula results are read as supplied, not recalculated or certified current: "
                                        + ", ".join(formula_addresses))
                    if errors:
                        warnings.append("Spreadsheet error cells are not observations: " + ", ".join(errors))
                    messages = parser_log.getvalue().strip()
                    if messages:
                        warnings.append("Legacy spreadsheet parser: " + messages)
                        parser_log.seek(0)
                        parser_log.truncate()
                    if warnings:
                        unit["warnings"] = warnings
                    yield unit
                workbook.unload_sheet(sheet_index)
        finally:
            workbook.release_resources()
    elif file_format == "xlsx":
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if len(archive.infolist()) > 5000 or sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Spreadsheet exceeds the uncompressed archive limit.")
        workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
        formulas = None
        try:
            formulas = load_workbook(io.BytesIO(payload), data_only=False, read_only=True)
            for sheet, formula_sheet in zip(workbook, formulas, strict=True):
                if sheet.max_row is None or sheet.max_column is None:
                    raise ValueError("Spreadsheet omits worksheet dimensions; bounded extraction cannot verify its extent.")
                if sheet.max_row > 10000 or sheet.max_column > 300:
                    raise ValueError("Spreadsheet dimensions exceed the bounded extraction limit.")
                for row, formula_row in zip(sheet.iter_rows(), formula_sheet.iter_rows(), strict=True):
                    formula_cells = [cell for cell in formula_row if cell.data_type == "f"]
                    for cell, formula_cell in zip(row, formula_row, strict=True):
                        if formula_cell.data_type == "f" and cell.value is None:
                            raise ValueError(
                                f"Spreadsheet formula has no cached result at {sheet.title}!{formula_cell.coordinate}; "
                                "recalculate and save the export in Excel or export self-contained values. "
                                "Python does not run vendor add-ins or calculate imported formulas.")
                    populated = [cell for cell in row if cell.value is not None]
                    if not populated:
                        continue
                    text = " | ".join(f"{cell.coordinate}: {cell.value}" for cell in populated)
                    numbers = []
                    for cell in populated:
                        value = cell.value
                        for number_format in ("dot_decimal", "comma_decimal"):
                            try:
                                parse_number(value, number_format)
                            except ValueError:
                                continue
                            numbers.append({"index": len(numbers), "raw": value, "locator": cell.coordinate,
                                            "display_format": cell.number_format})
                            break
                    unit = {"unit_id": f"{sheet.title}!R{populated[0].row}", "text": text, "numbers": numbers}
                    if formula_cells:
                        unit["warnings"] = [
                            "Cached formula results are read as supplied, not recalculated or certified current: "
                            + ", ".join(cell.coordinate for cell in formula_cells)]
                    errors = [cell.coordinate for cell in populated if cell.data_type == "e"]
                    if errors:
                        unit.setdefault("warnings", []).append("Spreadsheet error cells are not observations: " + ", ".join(errors))
                    yield unit
        finally:
            workbook.close()
            if formulas is not None:
                formulas.close()
    elif file_format == "csv":
        for index, row in enumerate(csv.reader(io.StringIO(payload.decode("utf-8-sig"))), 1):
            if len(row) > 300 or index > 10000:
                raise ValueError("CSV dimensions exceed the bounded spreadsheet extraction limit.")
            populated = [(f"{get_column_letter(column)}{index}", value)
                         for column, value in enumerate(row, 1) if value.strip()]
            if not populated:
                continue
            text = " | ".join(f"{cell}: {value}" for cell, value in populated)
            numbers = []
            for cell, value in populated:
                for number_format in ("dot_decimal", "comma_decimal"):
                    try:
                        parse_number(value, number_format)
                    except ValueError:
                        continue
                    numbers.append({"index": len(numbers), "raw": value, "locator": cell})
                    break
            yield {"unit_id": f"row:{index}", "text": text, "numbers": numbers}
    elif file_format in {"html", "txt"}:
        # Financial websites normally use UTF-8; unsupported encodings must not
        # silently corrupt symbols and digits into a supposedly verified source.
        text = payload.decode("utf-8-sig")
        if file_format == "html":
            parser = _HTMLText()
            parser.feed(text)
            text = "".join(parser.parts)
        lines = [line.strip(" \r") for line in text.splitlines() if line.strip()]
        chunk, length, index = [], 0, 1
        for line in lines:
            # These units provide stable citation locations, not limits on the
            # document text supplied to a research or review call.
            if chunk and length + len(line) > 5000:
                text = "\n".join(chunk)
                yield {"unit_id": f"chunk:{index}", "text": text, "numbers": _tokens(text)}
                chunk, length, index = [], 0, index + 1
            chunk.append(line)
            length += len(line) + 1
        if chunk:
            text = "\n".join(chunk)
            yield {"unit_id": f"chunk:{index}", "text": text, "numbers": _tokens(text)}
    else:
        raise ValueError("Unsupported evidence format.")


def _vendor_period(value):
    """Read a vendor's column label without inventing fiscal calendar dates."""
    if not isinstance(value, str):
        return None
    match = re.fullmatch(r"(?:(?P<n>[1-4])(?P<part>[QH])|(?P<ytd>9M))?FY[- ]?(?P<year>\d{4})(?P<estimate>E)?",
                         value.strip(), re.IGNORECASE)
    if not match or (match["part"] and match["part"].upper() == "H" and int(match["n"]) > 2):
        return None
    slot = (match["part"].upper() + match["n"] if match["part"] else
            "9M" if match["ytd"] else "FY")
    return {"year": int(match["year"]), "slot": slot, "estimate": bool(match["estimate"])}


def _vendor_profile(payload, file_format, units, dataset_role=None):
    """Classify explicit dataset/section roles after validating vendor schemas.

    A file role designates its dataset, not the economic meaning of every metric.
    Combined exports require named section labels; row position alone is invalid.
    """
    summary = {"status": "unsupported", "version": 2, "origin": "explicit_dataset_or_section_labels",
               "reason": "Unrecognised vendor layout; use company originals for historical numbers.",
               "exclusions": "Historical AI views omit standardised tables, estimates and analytics helper values; originals remain complete."}
    profile = {"header_sets": {}, "sheets": []}
    if file_format != "xlsx":
        return summary, profile
    workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    try:
        for sheet in workbook:
            rows = {n: {c.column: c.value for c in row if c.value is not None}
                    for n, row in enumerate(sheet.iter_rows(), 1)}
            headers = []
            for n, values in rows.items():
                unit_columns = [c for c, value in values.items() if str(value).strip().casefold() == "units"]
                if len(unit_columns) != 1:
                    continue
                unit_col = unit_columns[0]
                series_values = {c: str(v).strip() for c, v in values.items()
                                 if c > unit_col and str(v).strip().casefold() in {"va actuals", "filings", "consensus"}}
                if not series_values or not any(v.casefold() == "va actuals" for v in series_values.values()):
                    continue
                columns = {}
                for c, label in rows.get(n - 1, {}).items():
                    parsed = _vendor_period(label)
                    if parsed is None or c <= unit_col:
                        continue
                    series = str(values.get(c, "")).strip()
                    columns[get_column_letter(c)] = dict(parsed, period_label=str(label), series=series,
                        eligible=not parsed["estimate"] and series.casefold() in {"va actuals", "filings"})
                if not columns:
                    continue
                headers.append({"row": n, "period_row": n - 1, "unit_col": unit_col,
                                "label_col": unit_col - 1, "analytics_col": unit_col + 1,
                                "columns": columns})
            if not headers:continue
            markers=[]
            aliases={'company reported':'company_reported','company-reported':'company_reported',
                     'company reported data':'company_reported','standardised':'standardized','standardized':'standardized',
                     'standardised data':'standardized','standardized data':'standardized'}
            for n,values in rows.items():
                found={aliases[str(v).strip().casefold()] for v in values.values() if str(v).strip().casefold() in aliases}
                if len(found)>1:raise ValueError('Conflicting vendor section labels')
                if found:markers.append((n,next(iter(found))))
            role=dataset_role or aliases.get(sheet.title.casefold())
            if role not in (None,'company_reported','standardized','combined'):
                raise ValueError('Unknown vendor dataset_role')
            if role in ('company_reported','standardized'):
                if any(label!=role for n,label in markers):raise ValueError('Dataset role contradicts workbook section labels')
                markers=[(0,role)]
            elif not markers or {label for n,label in markers}!={'company_reported','standardized'}:
                continue
            sheet_info={'sheet':sheet.title,'dataset_role':role or 'combined',
                        'section_labels':[dict(row=n,role=label) for n,label in markers],
                        'boundary_references':[f'{sheet.title}!R{n}' for n,label in markers]}
            profile['sheets'].append(sheet_info)
            active = None
            for n, values in rows.items():
                if n in {h["period_row"] for h in headers}:
                    active = next(h for h in headers if h["period_row"] == n)
                    active["id"] = f"{sheet.title}!R{n}"
                    profile["header_sets"][active["id"]] = {
                        "period_reference": active["id"], "series_reference": f"{sheet.title}!R{n + 1}",
                        "columns": active["columns"]}
                key = f"{sheet.title}!R{n}"
                if key not in units:
                    continue
                if active is None:
                    units[key]["vendor_row"] = {"role": "metadata", "section": "metadata"}
                    continue
                label = str(values.get(active["label_col"], "")).strip()
                unit = str(values.get(active["unit_col"], "")).strip()
                analytics = str(values.get(active["analytics_col"], "")).strip()
                section = next((label for start,label in reversed(markers) if n>=start),'metadata')
                if n in (active["row"], active["period_row"]):
                    role = "header"
                elif label and unit and analytics.casefold() == "analytics":
                    role = "metric"
                elif analytics and analytics.casefold() != "analytics":
                    role = "helper"
                else:
                    role = "heading"
                # All nonnumeric text preceding the first period column is kept,
                # including subheadings and definition notes, even on blank rows.
                first_column = min(c for c in rows[active["period_row"]] if _vendor_period(rows[active["period_row"]][c]))
                context = " | ".join(f"{get_column_letter(c)}{n}: {v}" for c, v in values.items()
                                     if c < first_column and str(v).strip())
                if role == "heading" and not units[key]["numbers"]:
                    context = units[key]["text"]
                numeric_cells = {number["locator"] for number in units[key]["numbers"]}
                missing = {get_column_letter(c): str(v) for c, v in values.items()
                           if role == "metric" and c >= first_column and isinstance(v, str)
                           and (f"{get_column_letter(c)}{n}" not in numeric_cells
                                or v.strip() in {"-", "\u2013", "\u2014"})}
                units[key]["vendor_row"] = {"role": role, "section": section,
                    "header_id": active["id"], "label": label, "unit": unit, "context": context,
                    "missing_markers": missing}
        if profile["sheets"]:
            summary.update(status="recognized", reason="Explicit dataset or labelled section role, validated period/series/unit schema; no positional company mapping.",
                           sheets=profile["sheets"],
                           caveat="VA Actuals/Filings are vendor channels, not independent proof of company reporting or historical availability.")
            unclassified = set()
            for unit in units.values():
                if "vendor_row" in unit:
                    continue
                sheet_name = unit["unit_id"].rsplit("!R", 1)[0]
                if sheet_name.strip().casefold() in {"key", "legend", "definitions", "notes"}:
                    unit["vendor_row"] = {"role": "metadata", "section": "metadata"}
                else:
                    unit["vendor_row"] = {"role": "unclassified", "section": "unclassified"}
                    unclassified.add(sheet_name)
            if unclassified:
                summary["excluded_unclassified_sheets"] = sorted(unclassified)
        return summary, profile
    finally:
        workbook.close()


def _evidence_preview(unit):
    """Vendor navigation reveals row labels, not excluded numeric inputs."""
    row = unit.get("vendor_row")
    if not row or row["role"] == "metadata":
        return unit["text"]
    classification = f"{row['section']} {row['role']}"
    return classification + ": " + row.get("context", "Unclassified vendor worksheet; use company originals.")


class SourceStore:
    """Immutable originals with stable citations and complete document reads."""

    def __init__(self, directory: Path, fetcher=None):
        self.directory = Path(directory)
        self.fetcher = fetcher or download
        self.sources = []
        self._units = {}

    @staticmethod
    def _validate_spec(spec: dict):
        if not isinstance(spec, dict) or not all(key in spec for key in ("id", "url", "title", "published_at", "reason", "kind")):
            raise ValueError("Source metadata must include identity, URL, title, publication time, selection reason and kind.")
        if any(not isinstance(spec[key], str) or not spec[key].strip() for key in ("url", "title", "reason", "kind")):
            raise ValueError("Source URL, title, selection reason and kind must be nonempty text.")
        if spec["published_at"] is not None and not isinstance(spec["published_at"], str):
            raise ValueError("Publication time must be source-supplied text or null.")
        source_id = spec["id"]
        if not isinstance(source_id, str) or not re.fullmatch(r"[a-zA-Z0-9_-]{1,60}", source_id):
            raise ValueError("Invalid source ID.")

    @staticmethod
    def _index(payload: bytes, file_format: str, source_id: str, category=None, dataset_role=None):
        """Use the same bounded parser for new and restored originals."""
        units, text_length = {}, 0
        for unit in _units(payload, file_format):
            text_length += len(unit["text"])
            if len(units) >= MAX_UNITS:
                raise ValueError("Source exceeds 3000 evidence units.")
            unit["source_id"] = source_id
            units[unit["unit_id"]] = unit
        if not units or not text_length:
            raise ValueError("No readable source text; scanned reports require an OCR-capable adapter.")
        metadata = {"unit_count": len(units),
                    "empty_units": [key for key, unit in units.items() if not unit["text"].strip()],
                    "extraction_warnings": [{"unit_id": key, "messages": unit["warnings"]}
                                            for key, unit in units.items() if unit.get("warnings")]}
        if category == "visible_alpha":
            metadata["vendor_layout"], metadata["vendor_profile"] = _vendor_profile(payload, file_format, units, dataset_role)
        return units, metadata

    def add(self, spec: dict) -> dict:
        self._validate_spec(spec)
        source_id = spec["id"]
        for existing in self.sources:
            if existing["id"] == source_id:
                if existing["url"] != spec["url"]:
                    raise ValueError("A source ID cannot refer to two URLs.")
                if existing["status"] != "ready":
                    raise ValueError("Previously archived source could not be read.")
                return existing
        if len(self.sources) >= 12:
            raise ValueError("The run reached its 12-source limit.")
        payload, final_url, content_type = self.fetcher(spec["url"])
        if not isinstance(payload, bytes) or not payload or len(payload) > MAX_BYTES:
            raise ValueError("Source is empty, invalid or exceeds 25 MB.")
        if payload.startswith(b"%PDF"):
            file_format = "pdf"
        elif payload.startswith(b"PK\x03\x04"):
            file_format = "xlsx"
        elif "html" in content_type or payload.lstrip().lower().startswith((b"<!doctype html", b"<html")):
            file_format = "html"
        else:
            raise ValueError("Unsupported disclosure format; PDF, XLSX or UTF-8 HTML is required.")
        self.directory.mkdir(parents=True, exist_ok=True)
        path = self.directory / f"{source_id}.{file_format}"
        with path.open("xb") as output:
            output.write(payload)
        descriptor = {key: spec.get(key) for key in ("id", "url", "title", "published_at", "reason", "kind")}
        descriptor.update({"fetched_url": final_url, "fetched_at": datetime.now(timezone.utc).isoformat(),
                           "sha256": hashlib.sha256(payload).hexdigest(), "path": str(path),
                           "format": file_format, "status": "unreadable"})
        self.sources.append(descriptor)
        try:
            units, metadata = self._index(payload, file_format, source_id)
        except Exception as exc:
            descriptor["error"] = str(exc)
            raise
        self._units[source_id] = units
        descriptor.update(status="ready", **metadata)
        return descriptor

    def add_local(self, spec: dict, *, root: Path) -> dict:
        """Freeze and index a manifest member without any network operation.

        Check the actual bytes against the preflight inventory immediately before
        archiving. Originals remain in the run even when extraction fails, so a
        failure cannot silently become approved evidence on continuation.
        """
        self._validate_spec(spec)
        if (not isinstance(spec.get("local_path"), str) or spec.get("category") not in LOCAL_CATEGORIES
                or spec.get("format") not in set(LOCAL_FORMATS.values())
                or not re.fullmatch(r"[0-9a-f]{64}", spec.get("sha256", ""))):
            raise ValueError("Local source metadata is incomplete or invalid.")
        root = Path(root).resolve(strict=True)
        path = _local_file(root, spec["local_path"])
        if Path(spec["local_path"]).parts[0] != spec["category"]:
            raise ValueError("Local source category disagrees with its archive path.")
        if LOCAL_FORMATS.get(path.suffix.lower()) != spec["format"]:
            raise ValueError("Local source extension disagrees with the manifest format.")
        payload = _local_payload(path)
        if hashlib.sha256(payload).hexdigest() != spec["sha256"]:
            raise ValueError("Local source changed after inventory; start again with the new archive manifest.")
        file_format, source_id = spec["format"], spec["id"]
        signature = _binary_format(payload)
        if (file_format in {"pdf", "xlsx", "xls"} and signature != file_format) or (signature and signature != file_format):
            raise ValueError("Local source contents do not match their filename format.")
        for existing in self.sources:
            if existing["id"] == source_id:
                if existing.get("sha256") != spec["sha256"] or existing.get("local_path") != spec["local_path"]:
                    raise ValueError("A local source ID cannot refer to different files.")
                if existing["status"] != "ready":
                    raise ValueError("Previously archived source could not be read.")
                return existing
        self.directory.mkdir(parents=True, exist_ok=True)
        archived_path = self.directory / f"{source_id}.{file_format}"
        with archived_path.open("xb") as output:
            output.write(payload)
        descriptor = copy.deepcopy(spec)
        descriptor.update({"fetched_url": spec["url"], "fetched_at": datetime.now(timezone.utc).isoformat(),
                           "original_path": str(path), "path": str(archived_path),
                           "acquisition": "local_copy", "status": "unreadable"})
        self.sources.append(descriptor)
        try:
            units, metadata = self._index(payload, file_format, source_id, spec.get("category"), spec.get("dataset_role"))
        except Exception as exc:
            descriptor["error"] = str(exc)
            raise
        self._units[source_id] = units
        descriptor.update(status="ready", **metadata)
        return descriptor

    def restore(self, descriptor: dict, *, archive_root: Path) -> dict:
        """Reopen a verified original without fetching it or rewriting its history.

        The descriptor continues to point at its original run. Only ready sources
        regain evidence units; an earlier parsing failure cannot become approval
        merely because a run is resumed.
        """
        self._validate_spec(descriptor)
        required = ("fetched_url", "fetched_at", "sha256", "path", "format", "status")
        if any(not isinstance(descriptor.get(key), str) or not descriptor[key].strip() for key in required):
            raise ValueError("Archived source metadata is incomplete.")
        source_id, file_format = descriptor["id"], descriptor["format"]
        if file_format not in set(LOCAL_FORMATS.values()) or descriptor["status"] not in {"ready", "unreadable"}:
            raise ValueError("Archived source format or status is invalid.")
        if not re.fullmatch(r"[0-9a-f]{64}", descriptor["sha256"]):
            raise ValueError("Archived source hash is invalid.")
        for existing in self.sources:
            if existing["id"] == source_id:
                if existing != descriptor:
                    raise ValueError("A restored source ID cannot refer to different archived metadata.")
                return existing
        try:
            archive_root = Path(archive_root).resolve(strict=True)
            path = Path(descriptor["path"]).resolve(strict=True)
            path.relative_to(archive_root)
        except (OSError, ValueError) as exc:
            raise ValueError("Archived source is missing or outside the allowed output directory.") from exc
        if not path.is_file() or path.name != f"{source_id}.{file_format}":
            raise ValueError("Archived source path does not match its ID and format.")
        if not 0 < path.stat().st_size <= MAX_BYTES:
            raise ValueError("Archived source is empty or exceeds 25 MB.")
        with path.open("rb") as original:
            payload = original.read(MAX_BYTES + 1)
        if not payload or len(payload) > MAX_BYTES or hashlib.sha256(payload).hexdigest() != descriptor["sha256"]:
            raise ValueError("Archived source content does not match its saved hash or size limits.")
        signature = _binary_format(payload)
        if (file_format in {"pdf", "xlsx", "xls"} and signature != file_format) or (signature and signature != file_format):
            raise ValueError("Archived source content does not match its saved format.")

        units = None
        if descriptor["status"] == "ready":
            if isinstance(descriptor.get("unit_count"), bool) or not isinstance(descriptor.get("unit_count"), int):
                raise ValueError("Archived source unit count is invalid.")
            units, metadata = self._index(payload, file_format, source_id, descriptor.get("category"), descriptor.get("dataset_role"))
            if any(descriptor.get(key) != value for key, value in metadata.items()):
                raise ValueError("Archived source extraction metadata changed; saved evidence cannot be replayed.")
        elif not isinstance(descriptor.get("error"), str) or not descriptor["error"]:
            raise ValueError("An unreadable archived source must retain its parsing error.")
        restored = copy.deepcopy(descriptor)
        self.sources.append(restored)
        if units is not None:
            self._units[source_id] = units
        return restored

    def catalog(self, max_units=None) -> list:
        """List stable unit IDs, explicitly marking any requested preview limit."""
        if max_units is not None and (isinstance(max_units, bool) or not isinstance(max_units, int) or max_units < 1):
            raise ValueError("Catalog limit must be a positive integer or None.")
        result = []
        for source in self.sources:
            units = list(self._units.get(source["id"], {}).values())
            selected = units if max_units is None else units[:max_units]
            public_source = {key: value for key, value in source.items() if key != "vendor_profile"}
            result.append(dict(public_source, unit_count=len(units), catalog_truncated=len(selected) < len(units),
                               units=[{"id": unit["unit_id"], "preview": " ".join(_evidence_preview(unit).split())[:100],
                                       "number_count": len(unit["numbers"])} for unit in selected]))
        return result

    def search(self, query: str, limit=30, *, source_ids=None) -> dict:
        """Find literal text anywhere in archived units without enlarging catalogs."""
        if not isinstance(query, str) or not query.strip() or len(query) > 200:
            raise ValueError("Evidence search requires nonempty text of at most 200 characters.")
        if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 30:
            raise ValueError("Evidence search limit must be between 1 and 30.")
        if source_ids is not None:
            if (not isinstance(source_ids, (list, tuple, set, frozenset))
                    or any(not isinstance(source_id, str) for source_id in source_ids)):
                raise ValueError("Evidence search source IDs must be a collection of known IDs.")
            source_ids = set(source_ids)
            if not source_ids <= {source["id"] for source in self.sources}:
                raise ValueError("Evidence search refers to an unknown source ID.")
        query, matches = query.strip(), []
        needle = query.casefold()
        for source_id, units in self._units.items():
            # Restrict the search before applying the match limit. Filtering the
            # first 30 global matches afterward can hide all primary evidence.
            if source_ids is not None and source_id not in source_ids:
                continue
            for unit_id, unit in units.items():
                text = _evidence_preview(unit)
                folded = text.casefold()
                position = folded.find(needle)
                if position < 0:
                    continue
                if len(matches) == limit:
                    return {"query": query, "matches": matches, "truncated": True}
                if len(folded) != len(text):
                    # Unicode case folding can expand a character (e.g. sharp s).
                    # Recover the original offset so the preview still surrounds it.
                    offset = 0
                    for index, character in enumerate(text):
                        offset += len(character.casefold())
                        if offset > position:
                            position = index
                            break
                preview = " ".join(text[max(0, position - 60):position + 180].split())[:180]
                matches.append({"source_id": source_id, "unit_id": unit_id, "preview": preview})
        return {"query": query, "matches": matches, "truncated": False}

    def read(self, requests: list) -> list:
        """Read exact citation units for deterministic validation, without truncation."""
        result = []
        for reference in requests:
            try:
                unit = self._units[reference["source_id"]][reference["unit_id"]]
            except KeyError as exc:
                raise ValueError("Unknown source evidence unit.") from exc
            result.append(copy.deepcopy(unit))
        return result

    def _vendor_history(self, source, years):
        """Create an explicit historical input view, never change source evidence."""
        if source.get("vendor_layout", {}).get("status") != "recognized":
            raise ValueError("Unrecognised Visible Alpha layout; historical input cannot be classified safely. Use company financial reports instead.")
        header_sets = source["vendor_profile"]["header_sets"]
        result = []
        for original in self._units[source["id"]].values():
            row = original["vendor_row"]
            if row["section"] in {"standardized", "unclassified"} or row["role"] == "helper":
                continue
            unit = copy.deepcopy(original)
            if row["role"] == "metadata":
                # Company identity and the Key sheet's legend remain visible.
                result.append(unit)
                continue
            columns = header_sets[row["header_id"]]["columns"]
            selected = {column: value for column, value in columns.items()
                        if value["year"] in years and value["eligible"]}
            if row["role"] == "header":
                unit["text"] = "Historical columns: " + json.dumps(selected, ensure_ascii=False, separators=(",", ":"))
                unit["numbers"] = []
            else:
                unit["text"] = row["context"]
                unit["numbers"] = [number for number in unit["numbers"]
                                   if row["role"] == "metric"
                                   and re.sub(r"\d+$", "", number["locator"]) in selected
                                   and str(number["raw"]).strip() not in {"-", "\u2013", "\u2014"}]
                missing = {column: marker for column, marker in row.get("missing_markers", {}).items()
                           if column in selected}
                unit["vendor_row"]["missing_markers"] = missing
                if missing:
                    row_number = unit["unit_id"].rsplit("!R", 1)[1]
                    notes = {column + row_number: marker for column, marker in missing.items()}
                    unit["text"] += " | Missing/not-observation notes: " + json.dumps(notes, separators=(",", ":"))
            result.append(unit)
        if result:
            result[0]["vendor_view"] = {"historical_years": sorted(years),
                "selection": "User-designated company-reported base metrics, historical VA Actuals and Filings columns; labels and missing rows retained.",
                "excluded": "Standardized section, derived analytics helper values, consensus/estimate columns, unrequested years and unclassified worksheets. Complete original and numeric index remain archived.",
                "citation_indexes": "Original number_index values are preserved; gaps are intentional.",
                "basis_warning": "Vendor actual labels are not proof of directly reported quarters, comparable definitions or publication-time availability."}
        return result

    def read_documents(self, requests: list, *, historical_years=None) -> list:
        """Open every indexed unit of each document selected by a valid citation.

        Validate every requested location before returning any document. Repeated
        selections of one source do not duplicate its text or numeric index;
        sources follow request order and their units retain original index order.
        """
        if historical_years is not None:
            if (not isinstance(historical_years, (list, tuple, set, frozenset)) or not historical_years
                    or any(isinstance(year, bool) or not isinstance(year, int) or not 1900 <= year <= 2200 for year in historical_years)):
                raise ValueError("Historical input years must be a nonempty collection of fiscal years.")
            historical_years = set(historical_years)
        selected = {}
        for reference in requests:
            try:
                source_id = reference["source_id"]
                self._units[source_id][reference["unit_id"]]
            except KeyError as exc:
                raise ValueError("Unknown source evidence unit.") from exc
            selected[source_id] = None
        result = []
        by_id = {source["id"]: source for source in self.sources}
        for source_id in selected:
            source = by_id[source_id]
            if historical_years is not None and source.get("category") == "visible_alpha":
                result.extend(self._vendor_history(source, historical_years))
            else:
                result.extend(copy.deepcopy(unit) for unit in self._units[source_id].values())
        return result

    def validate_historical_reference(self, reference, period):
        """Reject vendor estimates, standardized/helper rows and wrong periods.

        This is eligibility, not source acceptance. Independent company-source
        reconstruction still establishes whether the company's reporting supports
        the observation. Interim slots require the discovered annual fiscal dates.
        """
        source = next((source for source in self.sources if source["id"] == reference.get("source_id")), None)
        if source is None:
            raise ValueError("Historical reference has an unknown source.")
        if source.get("category") != "visible_alpha":
            return None
        if source.get("vendor_layout", {}).get("status") != "recognized":
            raise ValueError("Unrecognised Visible Alpha layout cannot supply historical inputs; use company financial reports.")
        try:
            unit = self._units[source["id"]][reference["unit_id"]]
            index = reference["number_index"]
            if isinstance(index, bool) or not isinstance(index, int) or index < 0:
                raise ValueError("Invalid historical numeric token index.")
            number = unit["numbers"][index]
            row = unit["vendor_row"]
            if row["section"] != "company_reported" or row["role"] != "metric":
                raise ValueError("Visible Alpha historical inputs must use company-reported base metric rows, not standardized tables or analytics helpers.")
            column = re.sub(r"\d+$", "", number["locator"])
            header = source["vendor_profile"]["header_sets"][row["header_id"]]["columns"][column]
        except (KeyError, IndexError, TypeError) as exc:
            raise ValueError("Historical vendor cell has no verified row/column classification.") from exc
        if not header["eligible"]:
            raise ValueError("Consensus, estimate or unknown vendor columns cannot be historical observations.")
        if str(number["raw"]).strip() in {"-", "\u2013", "\u2014"}:
            raise ValueError("A vendor dash is missing/ambiguous evidence, not a verified historical zero.")
        if period.get("year") != header["year"]:
            raise ValueError("Vendor cell fiscal year does not match the historical period.")
        slot = header["slot"]
        expected_id = str(header["year"]) + ("" if slot == "FY" else "-" + slot)
        if "id" in period and period["id"] != expected_id:
            raise ValueError("Vendor cell fiscal slot does not match the historical period ID.")
        expected_kind = "annual" if slot == "FY" else "quarter" if slot.startswith("Q") else "half_year" if slot.startswith("H") else "year_to_date"
        if period.get("kind") != expected_kind:
            raise ValueError("Vendor cell period kind does not match the historical period; periods are not relabelled automatically.")
        try:
            start, end = date.fromisoformat(period["start"]), date.fromisoformat(period["end"])
            if slot == "FY":
                fiscal_start = date.fromisoformat(period.get("fiscal_year_start", period["start"]))
                fiscal_end = date.fromisoformat(period.get("fiscal_year_end", period["end"]))
            else:
                fiscal_start = date.fromisoformat(period["fiscal_year_start"])
                fiscal_end = date.fromisoformat(period["fiscal_year_end"])
        except (ValueError, KeyError, TypeError) as exc:
            raise ValueError("Vendor period validation requires explicit dates and annual fiscal-year anchors for interims.") from exc
        days = (fiscal_end - fiscal_start).days + 1
        fractions = {"FY": (0, 1), "Q1": (0, .25), "Q2": (.25, .5), "Q3": (.5, .75),
                     "Q4": (.75, 1), "H1": (0, .5), "H2": (.5, 1), "9M": (0, .75)}
        left, right = fractions[slot]
        if (not 1 < days <= 551 or start > end or start < fiscal_start or end > fiscal_end
                or abs((start - fiscal_start).days - days * left) > 14
                or abs((end - fiscal_start).days + 1 - days * right) > 14
                or (left == 0 and start != fiscal_start)
                or (right == 1 and end != fiscal_end)):
            raise ValueError("Vendor period dates do not match its fiscal slot and annual reporting calendar.")
        return {"section": row["section"], "series": header["series"], "period_label": header["period_label"],
                "locator": number["locator"], "year": header["year"], "slot": slot,
                "classification_origin": source["vendor_layout"]["origin"]}

    def resolve(self, reference: dict) -> float:
        try:
            unit = self._units[reference["source_id"]][reference["unit_id"]]
            index = reference["number_index"]
            if isinstance(index, bool) or not isinstance(index, int) or index < 0:
                raise ValueError("Invalid numeric token index.")
            token = unit["numbers"][index]
        except (KeyError, IndexError) as exc:
            raise ValueError("Numeric reference does not exist in the archived source.") from exc
        return parse_number(token["raw"], reference["number_format"])


def source_selection_audit(data,spec,frozen_sources,values=None):
    """Trace active historical cells to exact originals, never infer from Excel output."""
    from collections import Counter
    sources={s['id']:s for s in frozen_sources};books={};records=[];counts=Counter()
    observations={o['id']:o for o in spec.get('observations',[])}
    try:
        for row in spec['rows']:
            if row['section']=='fx' or row.get('evidence_role')=='shared_factor_diagnostic':continue
            for pid,cell in row.get('cells',{}).items():
                if cell.get('status') not in ('actual','reported','derived_actual'):continue
                if cell.get('value') is None and not cell.get('expression'):continue
                observation=observations.get(cell.get('observation_id'),{})
                refs=cell.get('source_refs',observation.get('source_refs',[]));bindings=[]
                for ref in refs:
                    source=sources.get(ref.get('source_id'),{});entry=dict(ref)
                    entry.update(file=source.get('local_path'),sha256=source.get('sha256'),publication=source.get('published_at'),
                        dataset_role=source.get('dataset_role'),category=source.get('category'))
                    if source.get('category')=='visible_alpha' and ref.get('cell') and source.get('path'):
                        if source['id'] not in books:books[source['id']]=load_workbook(source['path'],data_only=True)
                        sheet=books[source['id']][ref['sheet']];number=sheet[ref['cell']].row
                        entry['row_labels']=[str(c.value) for c in sheet[number][:6] if c.value is not None]
                    bindings.append(entry)
                origin='calculation' if cell.get('expression') else '+'.join(sorted({
                    ('VA '+(b.get('dataset_role') or 'legacy combined; role requires saved designation')) if b.get('category')=='visible_alpha' else
                    'company document' if b.get('category')=='financial_reports' else str(b.get('category')) for b in bindings})) or 'unresolved'
                if row['id']=='period_year_fraction':
                    origin='fiscal calendar calculation'
                    bindings=[dict(period=next(p for p in spec['periods'] if p['id']==pid),method='Inclusive period days divided by fiscal-year days')]
                counts[origin]+=1
                records.append(dict(row_id=row['id'],metric=row['label'],period=pid,status=cell['status'],
                    unit=row['unit'],currency=data['currency'],reporting_basis=cell.get('basis',observation.get('basis')),
                    definition=observation.get('definition',row.get('definition',row['label'])),
                    value=(values or {}).get(row['id'],{}).get(pid,cell.get('value')),origin=origin,bindings=bindings,
                    derivation=dict(expression=cell.get('expression'),bindings=cell.get('bindings',{}),observation_derivation=cell.get('derivation'))))
    finally:
        for book in books.values():book.close()
    return dict(method='Active cell graph and preserved original input bindings; calculations retain dependency IDs.',
        counts=dict(counts),records=records,standardized_policy='Secondary comparison only; never a missing primary substitute',
        measure_policy='Underlying/trading/statutory profits have distinct IDs and definitions; accounting identities retain explicit adjustment bridges')


def migrate_vendor_bindings(data,frozen_sources,new_sources):
    """Move unchanged active vendor facts to an explicitly company-reported file.

    Exact label, unit, fiscal slot and cached number must agree. Missing or changed
    data is an exception for research, never a standardized fallback or a zero.
    """
    books={};profiles={};audit=[];sources={s['id']:s for s in frozen_sources}
    try:
        candidates=[]
        for source in new_sources:
            if source.get('dataset_role')!='company_reported':continue
            payload=Path(source['path']).read_bytes();book=load_workbook(io.BytesIO(payload),data_only=True)
            books[source['id']]=book
            units,metadata=SourceStore._index(payload,'xlsx',source['id'],'visible_alpha','company_reported')
            if metadata['vendor_layout']['status']!='recognized':raise ValueError('Unsupported separated company-reported VA file')
            profiles[source['id']]=metadata
            for unit in units.values():
                row=unit['vendor_row']
                if row['role']!='metric' or row['section']!='company_reported':continue
                sheet,number=unit['unit_id'].rsplit('!R',1)
                candidates.append((source,sheet,int(number),row,metadata['vendor_profile']['header_sets'][row['header_id']]['columns']))
        def replace(ref,pid):
            source=sources.get(ref.get('source_id'))
            if not source or source.get('category')!='visible_alpha' or not ref.get('cell'):return ref
            if source.get('dataset_role')=='company_reported':return ref
            if source['id'] not in books:books[source['id']]=load_workbook(source['path'],data_only=True)
            sheet=books[source['id']][ref['sheet']];oldrow=sheet[ref['cell']].row
            # Unit/metric columns are located from labelled Units headers.
            headers=[c for row in sheet.iter_rows(max_row=oldrow) for c in row[:10] if str(c.value).strip().casefold()=='units']
            if not headers:raise ValueError('Original vendor binding has no Units schema')
            header=headers[-1];label=sheet.cell(oldrow,header.column-1).value;unit=sheet.cell(oldrow,header.column).value
            year,kind=(pid.split('-',1)+['FY'])[:2] if '-' in pid else (pid,'FY')
            options=[]
            for new,sname,n,row,columns in candidates:
                if row['label']!=label or row['unit']!=unit:continue
                for col,info in columns.items():
                    if info['year']!=int(year) or info['slot']!=kind or not info['eligible']:continue
                    value=books[new['id']][sname][col+str(n)].value
                    if type(value) not in (int,float):continue
                    options.append((info['series'].casefold()!='filings',new,sname,col+str(n),value,info))
            if not options:raise ValueError('Missing company-reported equivalent for '+str(label)+' '+pid)
            options.sort(key=lambda x:x[0]);_,new,sname,cell,value,info=options[0]
            if value!=ref['raw']:raise ValueError('Changed vendor observation needs source reconciliation: '+str(label)+' '+pid)
            revised=dict(ref,source_id=new['id'],sheet=sname,cell=cell,series=info['series'],row_label=label,
                dataset_role='company_reported',basis_note='Exact-label/unit/period/cache migration from preserved combined company-reported binding; original financial definition unchanged')
            audit.append(dict(period=pid,label=label,before=ref,after=revised));return revised
        for collection in ('historical','schedule_controls'):
            for role in data.get(collection,{}).values():
                for year,refs in role.get('sources',{}).items():role['sources'][year]=[replace(ref,str(year)+'-FY') for ref in refs]
        for observation in data.get('observations',[]):
            if observation.get('status') not in ('actual','reported','derived_actual'):continue
            observation['source_refs']=[replace(ref,observation['period_id']) for ref in observation.get('source_refs',[])]
        return audit
    finally:
        for book in books.values():book.close()


def verify_model_input_sources(data: dict, archive: Path) -> list[dict]:
    """Audit a frozen company input specification against unchanged originals.

    Spreadsheet checks read cached values, retain signs and apply the recorded
    multiplier once. PDF checks establish only that the signed numeric token
    occurs on the specified complete page, not that the concept/period mapping
    is correct. Derived observations must retain weighted original references.
    Workbook/PDF objects and page tokens are cached by content hash for this audit.
    """
    checks, payloads, books, pages = [], {}, {}, {}
    sources = data.get("sources", {})
    if isinstance(sources, list):
        sources = {source["id"]: source for source in sources}
    tolerance = 1e-8

    def check(name, status, detail, difference=None, area="source_data"):
        checks.append(dict(area=area, name=name, status=status, detail=detail,
                           difference=difference, tolerance=tolerance))

    for source_id, source in sources.items():
        try:
            path = _local_file(Path(archive), source["local_path"])
            payload = path.read_bytes()
            digest = hashlib.sha256(payload).hexdigest()
            if not source.get("sha256") or digest != source["sha256"]:
                raise ValueError("Original source bytes do not match the frozen SHA-256.")
            payloads[source_id] = (digest, payload, source.get("format", path.suffix.lstrip(".")))
            check("source_hash:" + source_id, "PASS", source["local_path"], area="source_integrity")
        except (OSError, ValueError, KeyError) as exc:
            check("source_hash:" + source_id, "FAIL", str(exc), area="source_integrity")

    def reference_value(reference):
        if "source_refs" in reference and "source_id" not in reference:
            return sum(reference_value(item) * parse_number(item.get("contribution_weight", 1))
                       for item in reference["source_refs"])
        if "raw" not in reference or reference["raw"] is None:
            raise ValueError("Numeric source reference is missing its raw value; blank is not zero.")
        source_id = reference.get("source_id")
        if source_id not in payloads:
            raise ValueError("Source is unknown, unreadable or failed its content hash check.")
        digest, payload, format_name = payloads[source_id]
        raw = reference["raw"]
        raw_value = parse_number(raw, reference.get("number_format", "dot_decimal"))
        multiplier = parse_number(reference.get("multiplier", 1))
        if format_name in {"xlsx", "xls"}:
            if digest not in books:
                books[digest] = (load_workbook(io.BytesIO(payload), data_only=True)
                                 if format_name == "xlsx" else xlrd.open_workbook(file_contents=payload))
            book = books[digest]
            sheet, cell = reference.get("sheet"), reference.get("cell")
            if not sheet or not isinstance(cell, str) or not re.fullmatch(r"[A-Z]+[1-9]\d*", cell):
                raise ValueError("Spreadsheet numeric reference requires an exact sheet and cell.")
            if format_name == "xlsx":
                saved = book[sheet][cell].value
            else:
                from openpyxl.utils.cell import coordinate_to_tuple
                row, column = coordinate_to_tuple(cell)
                saved = book.sheet_by_name(sheet).cell_value(row - 1, column - 1)
            if saved is None or saved == "":
                raise ValueError("Referenced source cell is blank or lacks a saved formula value.")
            saved_value = parse_number(saved, reference.get("number_format", "dot_decimal"))
            if abs(saved_value - raw_value) > tolerance:
                raise ValueError(f"Source cell differs from recorded raw value: {saved_value} versus {raw_value}.")
        elif format_name == "pdf":
            match = re.fullmatch(r"page:([1-9]\d*)(?::plain)?", reference.get("unit_id", ""))
            if not match:
                raise ValueError("PDF numeric reference requires a complete numbered page.")
            page_number = int(match[1]) - 1
            key = digest, page_number
            if key not in pages:
                if digest not in books:
                    books[digest] = PdfReader(io.BytesIO(payload))
                page = books[digest].pages[page_number]
                # Retain both layouts: either parser representation may preserve a
                # table's signed numeric token better. No page/text truncation.
                texts = [page.extract_text(extraction_mode="layout"), page.extract_text()]
                values = []
                for text in texts:
                    for token in _tokens(text or ""):
                        try:
                            values.append(parse_number(token["raw"], reference.get("number_format", "dot_decimal")))
                        except ValueError:
                            continue
                pages[key] = values
            if not any(abs(number - raw_value) <= tolerance for number in pages[key]):
                raise ValueError("The signed numeric value does not occur in the specified PDF page text.")
        else:
            raise ValueError("Unsupported numeric source format; no implicit parsing fallback.")
        return raw_value * multiplier

    def observation(name, expected, references, optional=False, missing_reason=None):
        if expected is None:
            detail = missing_reason or "Observation is explicitly missing; it has not been converted to zero."
            check(name, "WARNING" if optional else "MISSING", detail)
            return
        try:
            if not references:
                raise ValueError("Populated observation has no original numeric references.")
            expected_number = parse_number(expected)
            resolved = sum(reference_value(reference) * parse_number(reference.get("contribution_weight", 1))
                           for reference in references)
            difference = resolved - expected_number
            pdf = any(reference.get("unit_id", "").startswith("page:") for reference in references)
            detail = ("Signed numeric occurrence on complete PDF page and recorded conversion agree; semantic mapping is unverified."
                      if pdf else "Original cached cell values, signs, multipliers and weighted contributions agree.")
            check(name, "PASS" if abs(difference) <= tolerance else "FAIL", detail, difference)
        except (KeyError, IndexError, TypeError, ValueError, OSError) as exc:
            check(name, "FAIL", str(exc))

    try:
        for section in ("historical", "schedule_controls"):
            for role, metric in data.get(section, {}).items():
                for period, value in metric.get("values", {}).items():
                    observation(f"{section}:{role}:{period}", value, metric.get("sources", {}).get(period, []),
                                optional=metric.get("optional") is True, missing_reason=metric.get("missing_reason"))
        for section in ("segments", "historical_segments"):
            for index, segment in enumerate(data.get(section, [])):
                for metric, reference in segment.get("source_refs", {}).items():
                    observation(f"{section}:{index}:{metric}", segment.get(metric), [reference])
        for metric, reference in data.get("corporate", {}).get("source_refs", {}).items():
            observation("corporate:" + metric, data["corporate"].get(metric), [reference])
        for index, interim in enumerate(data.get("interims", [])):
            observation(f"interims:{index}:{interim.get('period')}", interim.get("value"),
                        [interim["source_ref"]] if interim.get("source_ref") else [])
        for item in data.get("observations", []):
            if item.get("status") in ("unverified", "conflict", "excluded", "vendor_unverified", "primary_conflict") or not item.get("eligible", True):
                check("observation_excluded:" + item["id"], "WARNING",
                      item.get("reason", "Candidate retained for audit but excluded from model inputs."))
                continue
            observation("observation:" + item["id"], item.get("value"), item.get("source_refs", []),
                        optional=item.get("optional", False), missing_reason=item.get("reason"))
    finally:
        for book in books.values():
            if hasattr(book, "close"):
                book.close()
            elif hasattr(book, "release_resources"):
                book.release_resources()
    return checks
