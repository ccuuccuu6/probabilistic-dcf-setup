"""Present an integrated company specification as an auditable financial model.

Numbers already use their display units. This module owns presentation only: the
same restricted financial AST used by Python creates every calculation formula.
"""
import json
import ast
import copy
import hashlib
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import textwrap

from openpyxl.cell.cell import MergedCell
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from .construction import dependencies, expression_tree, formula_excel, formula_value


STYLE_VERSION = "financial-model-5"
NAVY = "183B56"
INK = "111827"
PINK = "D63384"
BLUE = "0000FF"
GREEN = "008000"
YELLOW = "FFFFCC"
PALE = "E4EFF3"
GREY = "64748B"
NUMBER = '#,##0.0;(#,##0.0);"–"'
RATE = '0.0%;(0.0%);"–"'
BOX = Border(*(Side(style="thin", color="B2B2B2") for _ in range(4)))
TOTAL = Border(top=Side(style="thin", color="AABAC6"), bottom=Side(style="thin", color="AABAC6"))


def _text(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else value


def _lookup(mapping, year):
    return mapping.get(str(year), mapping.get(year))


def _number(value):
    return type(value) in (int, float) and math.isfinite(value)


def _format(unit):
    if unit == "exchange_rate":
        return '0.000000;(0.000000);"–"'
    if unit in ("ratio", "%", "rate"):
        return RATE
    if unit == "bps":
        return '0;[Red](0);"–"'
    return NUMBER


def _model_format(row):
    """Suppress display-only roundoff; neither formulas nor check gates change."""
    if "check_tolerance" not in row:
        return row.get("number_format", _format(row["unit"]))
    tolerance = row["check_tolerance"]
    if not _number(tolerance) or not 0 < tolerance < 1e12:
        raise ValueError("Display check tolerance must be finite and positive")
    threshold = format(tolerance, ".15f").rstrip("0").rstrip(".")
    if not threshold or threshold == "0":
        raise ValueError("Display check tolerance is smaller than supported precision")
    places = max(1, min(12, math.ceil(-math.log10(tolerance)) + 1))
    numeric = "#,##0." + "0" * places
    return f'{numeric};({numeric});"–"'


def _display_unit(unit):
    if unit.endswith("m/shares_m"):
        return unit.replace("m/shares_m", "/share")
    return {"ratio": "%", "shares_m": "m shares", "exchange_rate": "quote / unit"}.get(unit, unit)


def _style_sheet(sheet, freeze="D6"):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 70
    sheet.freeze_panes = freeze
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = False
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.print_title_rows = "1:5"
    sheet.sheet_format.defaultRowHeight = 15
    sheet.sheet_format.defaultColWidth = 14
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.page_margins.bottom = 0.5
    sheet.page_margins.footer = 1 / 6


def _cell(sheet, row, column, value=None, *, color=INK, bold=False, fill=None, number_format=None):
    cell = sheet.cell(row, column, value)
    cell.font = Font(name="Aptos", size=10, color=color, bold=bold)
    cell.alignment = Alignment(vertical="center")
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if number_format:
        cell.number_format = number_format
    return cell


def _bar(sheet, row, label, last_column):
    for column in range(2, last_column + 1):
        cell = _cell(sheet, row, column, label if column == 2 else None,
                     color="FFFFFF", bold=True, fill=NAVY)
        cell.font = Font(name="Aptos", size=11, color="FFFFFF", bold=True)
    sheet.row_dimensions[row].height = 27


def _table(workbook, name, columns, records, widths=None):
    sheet = workbook.create_sheet(name)
    _style_sheet(sheet, "B6")
    _cell(sheet, 2, 2, name, bold=True)
    for column, label in enumerate(columns, 1):
        _cell(sheet, 5, column, label, bold=True, color="FFFFFF", fill=NAVY)
        sheet.column_dimensions[get_column_letter(column)].width = (widths or {}).get(column, 24)
    for row, record in enumerate(records, 6):
        for column, value in enumerate(record, 1):
            c = _cell(sheet, row, column, _text(value))
            c.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = 30
    sheet.auto_filter.ref = f"A5:{get_column_letter(len(columns))}{max(5, sheet.max_row)}"
    return sheet


def _supporting(workbook, name, data):
    if isinstance(data, dict) and "columns" in data:
        columns, rows = data["columns"], data.get("rows", [])
        if rows and isinstance(rows[0], dict):
            rows = [[item.get(key) for key in columns] for item in rows]
    elif isinstance(data, list) and data and isinstance(data[0], dict):
        columns = list(dict.fromkeys(key for item in data for key in item))
        rows = [[item.get(key) for key in columns] for item in data]
    else:
        raise ValueError(f"{name} must be a table or list of records")
    return _table(workbook, name, columns, rows, {1: 32, 2: 44})


def _missing_value_audit(workbook, audit):
    """Expose every count and explanation as a readable row, retaining scope."""
    records = []

    def label(value):
        return str(value).replace("_", " ").capitalize()

    def visit(value, path):
        if isinstance(value, dict):
            for key, child in value.items():
                visit(child, path + [str(key)])
        elif isinstance(value, list):
            for index, child in enumerate(value, 1):
                visit(child, path + [str(index)])
        else:
            group = " / ".join(label(part) for part in path[:-1]) or "Audit context"
            field = label(path[-1]) if path else "Value"
            if _number(value):
                records.append([group, field, value, None])
            else:
                # Excel caps row height at409points. Long evidence text is split
                # into continuation rows without truncating any characters.
                content = "Not supplied" if value is None else str(value)
                chunks = textwrap.wrap(content, 700, replace_whitespace=False,
                                       drop_whitespace=False) or [""]
                records.extend([group, field if i == 0 else field + " (continued)", None, chunk]
                               for i, chunk in enumerate(chunks))

    visit(audit, [])
    widths = {1: 34, 2: 85, 3: 24, 4: 110}
    sheet = _table(workbook, "Missing Data Audit",
                   ["Audit group", "Status / cause / field", "Count or value", "Explanation / source"],
                   records, widths)
    _cell(sheet, 3, 2, "Before and after counts retain their stated scope; prior-workbook comparison matches the same row IDs and fiscal periods.", color=GREY)
    for row, record in enumerate(records, 6):
        lines = max(sum(max(1, math.ceil(len(line) / (widths[column] * .8)))
                        for line in str(value or "").splitlines()) or 1
                    for column, value in enumerate(record, 1))
        sheet.row_dimensions[row].height = max(30, 15 * lines + 6)
    return sheet


def export_integrated(path, spec, values, checks, *, valuation=None):
    """Export a reusable integrated specification and return Excel check targets.

    ``values`` is {row_id: {period: number_or_none}} in displayed units.
    Saved backend inputs own forecast controls; Excel projects a review copy.
    Identity expressions are residuals, with explicit source-precision tolerance.
    """
    if spec.get("periods"):
        if spec.get('assumption_overrides'):
            import copy
            spec=copy.deepcopy(spec)
            by_id={r['id']:r for r in spec['rows']}
            for rid,period_values in spec['assumption_overrides'].items():
                for pid,value in period_values.items():
                    cell=by_id[rid]['cells'][pid]
                    cell.pop('expression',None);cell.pop('bindings',None)
                    cell.update(value=value,reason='Validated saved backend scenario override; Excel edits remain local to this export.')
        return _export_periodic(path, spec, values, checks, valuation=valuation)
    path = Path(path)
    historical = list(spec["historical_years"])
    future = list(spec["forecast_years"])
    years = historical + future
    if not years or years != sorted(set(years)):
        raise ValueError("Periods must be distinct increasing annual years")
    rows = spec["rows"]
    if len({row["id"] for row in rows}) != len(rows):
        raise ValueError("Duplicate model row IDs")
    row_by_id = {row["id"]: row for row in rows}
    sections = spec["sections"]
    if len({section["id"] for section in sections}) != len(sections):
        raise ValueError("Duplicate section IDs")
    if any(row["section"] not in {section["id"] for section in sections} for row in rows):
        raise ValueError("Every row must belong to a declared section")
    workbook = Workbook()
    model = workbook.active
    model.title = "Model"
    _style_sheet(model)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    model.column_dimensions["A"].hidden = True
    model.column_dimensions["A"].width = 35
    model.column_dimensions["B"].width = 55
    model.column_dimensions["C"].width = 14
    _cell(model, 2, 2, str(spec["company"]), bold=True)
    model["B2"].font = Font(name="Aptos", size=18, bold=True, color=NAVY)
    model.row_dimensions[2].height = 30
    _cell(model, 3, 2, "Temporary assumptions · edit Assumptions; yellow Model boxes link to inputs · financial model, not a validated forecast", color=GREY)
    _cell(model, 4, 2, "Pink: historical pulls  |  Green/yellow: linked assumption  |  Black: formula  |  N/C: unavailable or non-comparable  |  Assumption history: initialization values, not reported assumptions", color=GREY)
    model.row_dimensions[3].height = 20
    _cell(model, 5, 2, "Financial model", bold=True, fill=NAVY, color="FFFFFF")
    _cell(model, 5, 3, "Units", bold=True, fill=NAVY, color="FFFFFF")
    columns = {year: i + 4 for i, year in enumerate(years)}
    last_column = max(columns.values())
    model.merge_cells(start_row=4, start_column=2, end_row=4, end_column=last_column)
    model["B4"].alignment = Alignment(vertical="center", wrap_text=True)
    model.row_dimensions[4].height = 30 if len(years) < 8 else 15
    for year, column in columns.items():
        cell = _cell(model, 5, column, f"{year}{'E' if year in future else 'A'}", bold=True,
                     fill=NAVY if year in historical else "376B7D", color="FFFFFF")
        cell.alignment = Alignment(horizontal="right")
        model.column_dimensions[get_column_letter(column)].width = 14
    model.row_dimensions[5].height = 24

    positions, section_positions = {}, {}
    current = 7
    for section in sections:
        members = [row for row in rows if row["section"] == section["id"]]
        if not members:
            continue
        section_positions[section["id"]] = current
        _bar(model, current, section["label"], last_column)
        current += 2
        previous_assumption = False
        for row in members:
            if previous_assumption and row["kind"] != "assumption":
                current += 1
            positions[row["id"]] = current
            _cell(model, current, 1, row["id"])
            c = _cell(model, current, 2, row["label"], bold=row.get("style") == "subtotal")
            c.alignment = Alignment(indent=1 if row["kind"] == "assumption" else 0, vertical="center")
            _cell(model, current, 3, row.get("display_unit", _display_unit(row["unit"])), color=GREY)
            note = "\n".join(str(row.get(key, "")) for key in ("definition", "basis", "note") if row.get(key))
            if note:
                c.comment = Comment(note, "Refractive")
            if row.get("supporting_input") or row.get("fx_only"):
                model.row_dimensions[current].hidden = True
                model.row_dimensions[current].outlineLevel = 1
            previous_assumption = row["kind"] == "assumption"
            current += 1
        current += 2

    history = workbook.create_sheet("Historics")
    _style_sheet(history)
    _cell(history, 2, 2, "Historical observations in model display units", bold=True)
    _cell(history, 3, 2, "Source bindings, sign conversions and reporting bases are retained in Sources. Blank means unavailable.", color=GREY)
    for column, label in enumerate(["Metric ID", "Metric", "Units"] + historical, 1):
        _cell(history, 5, column, label, bold=True, color="FFFFFF", fill=NAVY)
    history.column_dimensions["A"].width = 38
    history.column_dimensions["B"].width = 55
    history.column_dimensions["C"].width = 14
    for index, row in enumerate(rows, 6):
        for column, value in enumerate([row["id"], row["label"], row["unit"]], 1):
            _cell(history, index, column, value)
        for year, column in ((year, i + 4) for i, year in enumerate(historical)):
            value = _lookup(row.get("historical", {}), year)
            if value is not None and not _number(value):
                raise ValueError(f"Non-numeric historical input: {row['id']} {year}")
            _cell(history, index, column, value, color=BLUE, number_format=_format(row["unit"]))
    history.auto_filter.ref = f"A5:{get_column_letter(max(3, len(historical) + 3))}{max(6, len(rows) + 5)}"

    assumptions = [row for row in rows if row["kind"] == "assumption"]
    assumption_cells = {row["id"]: {str(year): {"sheet": "Assumptions", "cell": f"{get_column_letter(c)}{r}"}
                        for c, year in enumerate(future, 4)} for r, row in enumerate(assumptions, 6)}
    expected_model, expected_checks, formula_count = {}, {}, 0
    def reference(key, year, *, qualified=False):
        if key not in positions:
            raise ValueError("Unknown expression dependency: " + key)
        if year not in columns:
            return "NA()"
        return ("'Model'!" if qualified else "") + f"{get_column_letter(columns[year])}{positions[key]}"

    def expression(expression, year, *, qualified=False, historical_formula=False):
        tree = expression_tree(expression)
        resolved = list(dict.fromkeys(reference(key, year - lag, qualified=qualified) for key, lag in dependencies(tree)))
        calculation = formula_excel(tree, lambda key, lag: reference(key, year - lag, qualified=qualified))
        missing = '"N/C"' if historical_formula else "NA()"
        guarded = f"IFERROR({calculation},{missing})"
        return "=" + (f"IF(IFERROR(COUNT({','.join(resolved)}),0)={len(resolved)},{guarded},{missing})" if resolved else guarded)

    def expected(key, year):
        if key not in row_by_id:
            return None
        row = row_by_id[key]
        supplied = _lookup(values.get(key, {}), year)
        if supplied is not None:
            return supplied
        return _lookup(row.get("historical" if year in historical else "inputs", {}), year)

    for row in rows:
        r = positions[row["id"]]
        for year, column in columns.items():
            historical_formula = row.get("historical_expression")
            if isinstance(historical_formula, dict):
                historical_formula = _lookup(historical_formula, year)
            if "historical_expression_periods" in row and str(year) not in {str(p) for p in row["historical_expression_periods"]}:
                historical_formula = None
            is_input = year in future and row["kind"] == "assumption"
            if year in historical and historical_formula and _lookup(row.get("historical", {}), year) is None:
                content = expression(historical_formula, year, historical_formula=True)
                target = expected(row["id"], year)
                target = "N/C" if target is None else target
                color = INK
            elif year in historical:
                end = max(6, len(rows) + 5)
                col_end = get_column_letter(len(historical) + 3)
                lookup = (f"INDEX(Historics!$D$6:${col_end}${end},MATCH($A{r},Historics!$A$6:$A${end},0),"
                          f"MATCH({year},Historics!$D$5:${col_end}$5,0))")
                content = f'=IF(COUNT({lookup})=1,{lookup},"N/C")'
                target = _lookup(row.get("historical", {}), year)
                target = "N/C" if target is None else target
                color = PINK
            elif is_input:
                target = _lookup(row.get("inputs", {}), year)
                if target is not None and not _number(target):
                    raise ValueError(f"Non-numeric assumption: {row['id']} {year}")
                control = assumption_cells[row["id"]][str(year)]["cell"]
                content = f'=IF(ISNUMBER(Assumptions!{control}),Assumptions!{control},NA())'
                color = GREEN
            elif row.get("expression"):
                content = expression(row["expression"], year)
                target = expected(row["id"], year)
                color = INK
            else:
                content, target, color = "N/C", "N/C", GREY
            cell = _cell(model, r, column, content, color=color, bold=row.get("style") == "subtotal",
                         fill=YELLOW if is_input else (PALE if row.get("style") == "subtotal" else None),
                         number_format=_model_format(row))
            if is_input:
                cell.border = BOX
                cell.hyperlink = f"#'Assumptions'!{assumption_cells[row['id']][str(year)]['cell']}"
                cell.comment = Comment("Linked display. Edit the authoritative input on Assumptions. " + str(row.get("note", row.get("definition", ""))), "Refractive")
            elif row.get("style") == "subtotal":
                cell.border = TOTAL
            if future and year == future[0]:
                cell.border = Border(left=Side(style="medium", color=NAVY), right=cell.border.right,
                                     top=cell.border.top, bottom=cell.border.bottom)
            if not is_input or content is not None:
                expected_model[cell.coordinate] = target
            formula_count += isinstance(content, str) and content.startswith("=")
        if "check_tolerance" in row:
            # Explicit sign sections prevent Excel adding a minus to the zero-band dash.
            first = f"{get_column_letter(min(columns.values()))}{r}"
            area = f"{first}:{get_column_letter(last_column)}{r}"
            zero_band = Rule(type="expression", stopIfTrue=True,
                formula=[f"AND(ISNUMBER({first}),ABS({first})<={row['check_tolerance']:.15g})"],
                dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=164, formatCode='"–";"–";"–"')))
            model.conditional_formatting.add(area, zero_band)
        if row.get("style") == "subtotal":
            for column in (2, 3):
                model.cell(r, column).border = TOTAL
            model.row_dimensions[r].height = 20
    model.print_area = f"B2:{get_column_letter(last_column)}{current - 1}"

    assumption_records = []
    for row in assumptions:
        assumption_records.append([row["id"], row["label"], row["unit"]] +
            [_lookup(row.get("inputs", {}), year) for year in future] +
            [row.get("note", "Temporary editable input on Assumptions"), row["section"]])
    assumption_sheet = _table(workbook, "Assumptions", ["Driver ID", "Authoritative forecast input", "Unit"] + future + ["Initialisation / rationale", "Block"],
                              assumption_records, {1: 35, 2: 52, len(future) + 4: 80})
    _cell(assumption_sheet, 3, 2, "Edit yellow input cells here. Model boxes link to these controls; blank inputs remain missing, not zero.", color=GREY)
    for r, row in enumerate(assumptions, 6):
        assumption_sheet.cell(r, 2).hyperlink = f"#'Model'!B{positions[row['id']]}"
        for c in range(4, len(future) + 4):
            assumption_sheet.cell(r, c).font = Font(name="Aptos", size=10, color=BLUE)
            assumption_sheet.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)
            assumption_sheet.cell(r, c).border = BOX
            assumption_sheet.cell(r, c).number_format = _format(row["unit"])

    dependency_rows = []
    for row in assumptions:
        for year in future:
            control = assumption_cells[row["id"]][str(year)]["cell"]
            dependency_rows.append([row["id"], row["label"], row["id"], 0, "Assumptions → Model", year,
                                    f"Assumptions!{control} → Model!{reference(row['id'], year)}"])
    affects = {row["id"]: [] for row in assumptions}
    for row in rows:
        for label, expr in (("Forecast", row.get("expression")), ("Historical", row.get("historical_expression"))):
            if not expr:
                continue
            expressions = expr.items() if isinstance(expr, dict) else [("all", expr)]
            for period, item in expressions:
                for key, lag in dependencies(expression_tree(item)):
                    dependency_rows.append([row["id"], row["label"], key, lag, label, period, item])
                    if key in affects:
                        affects[key].append(f"{row['label']} (lag {lag}y)")
    driver_rows = [[row["id"], row["label"], row["section"], row["unit"], row.get("measure", "rate" if row["unit"] == "ratio" else "level"),
                    ", ".join(map(str, future)), "temporary_baseline", row.get("definition", row["label"]),
                    "; ".join(dict.fromkeys(affects[row["id"]])), row.get("note", ""), row.get("source_refs", {}),
                    f"Assumptions!{assumption_cells[row['id']][str(future[0])]['cell']}" if future else ""] for row in assumptions]
    _table(workbook, "Drivers", ["ID", "Driver", "Entity / block", "Unit", "Measure", "Periods", "Scenario", "Definition", "Direct dependencies / timing", "Initialisation / rationale", "Evidence", "Control"],
           driver_rows, {1: 35, 2: 50, 8: 65, 9: 70, 10: 80, 11: 80})
    _table(workbook, "Dependencies", ["Metric ID", "Financial line", "Dependency ID", "Annual lag", "Scope", "Periods", "Expression"], dependency_rows,
           {1: 35, 2: 52, 3: 35, 7: 95})

    source_rows = []
    for row in rows:
        for year, refs in row.get("source_refs", {}).items():
            for source in refs if isinstance(refs, list) else [refs]:
                source_rows.append([row["id"], row["label"], year, row["unit"], row.get("basis", ""), source])
    for source in spec.get("sources", []):
        source_rows.append(["SOURCE", "Document metadata", "", "", "", source])
    _table(workbook, "Sources", ["Metric / source ID", "Financial line", "Period", "Display units", "Reporting basis", "Exact reference / document metadata"], source_rows,
           {1: 35, 2: 50, 5: 60, 6: 120})

    check_sheet = _table(workbook, "Checks", ["Check ID", "Check", "Period", "Residual", "Tolerance", "Live result", "Control type", "Notes"], [],
                         {1: 35, 2: 65, 3: 14, 4: 18, 5: 16, 6: 18, 7: 28, 8: 90})
    _cell(check_sheet, 3, 2, "Live arithmetic responds to Excel edits. Frozen source and financial reviews are shown separately below.", color=GREY)
    r = 6
    for identity in spec.get("identities", []):
        tolerance = identity["tolerance"]
        if not _number(tolerance) or tolerance < 0:
            raise ValueError("Identity tolerance must be finite and nonnegative")
        for year in identity.get("periods") or years:
            year = int(year)
            if year not in columns:
                continue
            residual = formula_value(expression_tree(identity["expression"]), lambda key, lag: expected(key, year - lag))
            content = [identity["id"], identity["label"], year, expression(identity["expression"], year, qualified=True),
                       tolerance, f'=IF(ISNUMBER(D{r}),IF(ABS(D{r})<=E{r},"PASS","FAIL"),"MISSING")',
                       "Independent source control" if identity.get("independent") else "Arithmetic identity",
                       identity.get("note", "")]
            for col, value in enumerate(content, 1):
                _cell(check_sheet, r, col, value, number_format=NUMBER if col in (4, 5) else None)
            expected_checks[f"D{r}"] = residual
            expected_checks[f"F{r}"] = "MISSING" if residual is None else ("PASS" if abs(residual) <= tolerance else "FAIL")
            r += 1
    if checks:
        r += 2
        _bar(check_sheet, r, "FROZEN VALIDATION — rerun Python/source review after model changes", 8)
        r += 1
        for item in checks:
            if not isinstance(item, dict):
                item = {"detail": str(item)}
            content = [item.get("id", item.get("check", "")), item.get("label", item.get("name", item.get("detail", ""))), item.get("period", ""),
                       item.get("residual", item.get("difference")), item.get("tolerance"), item.get("status", ""), item.get("category", item.get("kind", "Frozen review")),
                       item.get("detail", item.get("message", ""))]
            for col, value in enumerate(content, 1):
                _cell(check_sheet, r, col, _text(value))
            r += 1
    check_sheet.auto_filter.ref = f"A5:H{max(6, r - 1)}"
    for status, color in (("PASS", "D9EAD3"), ("FAIL", "F4CCCC"), ("MISSING", "FFF2CC")):
        check_sheet.conditional_formatting.add(f"F6:F{max(6, r)}", CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)))
    if spec.get("gaps"):
        _table(workbook, "Gaps", ["Unresolved issue / limitation"], [[gap] for gap in spec["gaps"]], {1: 145})
    for key, name in (("historical_segments", "Historical Segments"), ("interims", "Interims"),
                      ("reporting_changes", "Reporting Changes")):
        if spec.get(key):
            _supporting(workbook, name, spec[key])
    valuation_inspection = _valuation_sheet(workbook, spec, values, valuation,
        lambda key, pid: reference(key, int(str(pid).split('-')[0]), qualified=True)) if valuation is not None else {}
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    # Verify the saved structure, not cached calculation values (Excel owns those).
    with path.open("rb") as stream:
        saved = load_workbook(stream, data_only=False)
        for sheet in saved:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f" and ("#REF!" in cell.value or "[" in cell.value):
                        raise ValueError(f"Broken/external formula: {sheet.title}!{cell.coordinate}")
        saved.close()
    return dict(status="structure_and_formula_syntax_checked", recalculation="not_run", formula_count=formula_count,
                expected_model_cells=expected_model, expected_check_cells=expected_checks,
                model_rows=positions, year_columns={str(year): get_column_letter(column) for year, column in columns.items()},
                section_rows=section_positions, assumption_count=len(assumptions), assumption_cells=assumption_cells, **valuation_inspection)


def _valuation_formula(name, references):
    """Render the valuation backend's restricted expression, without new rules."""
    from .valuation import VALUATION_EQUATIONS
    def render(node):
        if isinstance(node, ast.Expression): return render(node.body)
        if isinstance(node, ast.Name): return references[node.id]
        if isinstance(node, ast.Constant) and type(node.value) in (int, float): return str(node.value)
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            return '('+('+' if isinstance(node.op, ast.UAdd) else '-')+render(node.operand)+')'
        if isinstance(node, ast.BinOp):
            signs={ast.Add:'+',ast.Sub:'-',ast.Mult:'*',ast.Div:'/',ast.Pow:'^'}
            if type(node.op) in signs:return '('+render(node.left)+signs[type(node.op)]+render(node.right)+')'
        raise ValueError('Unsupported valuation display arithmetic')
    return '='+render(ast.parse(VALUATION_EQUATIONS[name], mode='eval'))


def _valuation_sheet(workbook, spec, values, valuation, source_reference):
    """Audit projection: live ordinary equations; dated timing stays a snapshot."""
    from .valuation import valuation_input_contract
    sheet=workbook.create_sheet('Valuation');_style_sheet(sheet, 'D8')
    detail=valuation.get('detailed_cashflows', []);transition=valuation.get('transition_cashflows', [])
    last=max(10, 3+len(detail), 3+len(transition));last_letter=get_column_letter(last)
    sheet.column_dimensions['A'].width=3;sheet.column_dimensions['B'].width=54;sheet.column_dimensions['C'].width=13
    for column in range(4,last+1):sheet.column_dimensions[get_column_letter(column)].width=16
    expected={};locations={};assumption_cells={};formula_count=0;section_starts=[]
    def section(row, title):
        section_starts.append(row);_bar(sheet,row,title,last)
    def scaled_unit(unit, scale=1):
        if unit=='ratio':return '%'
        if unit=='multiple':return 'x'
        suffix={1:'',1e3:'k',1e6:'m',1e9:'bn'}.get(scale)
        if suffix is None:return str(unit)+' × '+format(scale,',g')
        return (suffix+' shares').strip() if unit=='shares' else str(unit)+suffix
    def status(row, column, value):
        text=str(value or '').replace('_',' ')
        cell=_cell(sheet,row,column,text,color=GREY)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        cell.comment=Comment(str(value or ''),'Refractive')
        return 15*max(1,len(textwrap.wrap(text,width=18)))
    def note(row, text):
        sheet.merge_cells(start_row=row,start_column=2,end_row=row,end_column=last)
        cell=_cell(sheet,row,2,str(text),color=GREY);cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[row].height=min(72,max(30,15*math.ceil(len(str(text))/130)))
    def put(row,column,value,target=None,unit=None,comment=None,input_cell=False):
        nonlocal formula_count
        cell=_cell(sheet,row,column,value,color=BLUE if input_cell else GREEN if isinstance(value,str) and "'Model'!" in value else INK,
                   fill=YELLOW if input_cell else None,number_format=RATE if unit=='ratio' else NUMBER.replace('0.0','0.00') if unit=='per_share' else NUMBER)
        if isinstance(value,str) and value.startswith('='):formula_count+=1
        wanted=value if target is None and _number(value) else target
        if _number(wanted):expected[cell.coordinate]=wanted
        if comment:cell.comment=Comment(_text(comment),'Refractive')
        if input_cell:cell.border=BOX
        return cell.coordinate
    def label(row,text,unit=None):
        cell=_cell(sheet,row,2,text);cell.alignment=Alignment(wrap_text=True,vertical='center')
        if len(text)>60:sheet.row_dimensions[row].height=30
        _cell(sheet,row,3,unit or '',color=GREY)
    def equation(row,column,name,refs,target):
        return put(row,column,_valuation_formula(name,refs),target)
    _bar(sheet,2,str(spec.get('company',''))+' — discounted cash flow',last)
    note(3,'Fair value as of '+str(valuation.get('valuation_date','not specified'))+' | Information cutoff '+str(valuation.get('information_cutoff','not specified'))+' | '+str(valuation.get('status','incomplete')).upper())
    note(4,'Saved backend inputs own this valuation. Yellow cells are local review controls; save a typed backend scenario to preserve changes. Source financial inputs link to Model.')
    note(5,('INCOMPLETE: conditional arithmetic only; unresolved valuation inputs block acceptance. ' if valuation.get('status')=='incomplete' else '')+
        'Operating cash leases; lease debt excluded. SBC remains expensed. Timing, opening-balance bridges, transition paths and sensitivity tables are dated backend snapshots; rerun the backend after changing operating scenarios.')
    sheet.print_title_rows='2:5'
    inputs=valuation.get('input_snapshot') or {}
    if not detail or not inputs:
        _bar(sheet,7,'VALUATION INCOMPLETE',last)
        diagnostics=valuation.get('diagnostics') or [dict(message=valuation.get('detail') or valuation.get('reason') or 'Required dated valuation inputs or complete operating cash-flow components are missing.')]
        for row,item in enumerate(diagnostics,9):note(row,item.get('message',item.get('detail',str(item))) if isinstance(item,dict) else item)
        sheet.print_area=f'B2:{last_letter}{sheet.max_row}'
        return dict(expected_valuation_cells=expected,valuation_cells=locations,valuation_assumption_cells={},
                    valuation_print_area=f'B2:{last_letter}{sheet.max_row}',valuation_status='incomplete',valuation_formula_count=0)
    roles=inputs['row_roles'];unit=scaled_unit(valuation['currency'],valuation.get('money_scale',1))
    share_unit=scaled_unit('shares',valuation.get('share_scale',1))
    # Allocate scalar controls before rendering formulas that reference them.
    row=43;section(row,'DATED VALUATION ASSUMPTIONS');row+=1
    for column,text in ((2,'Assumption'),(3,'Unit'),(4,'Value'),(5,'As of'),(6,'Evidence status'),(7,'Method')):_cell(sheet,row,column,text,bold=True)
    refs={};contract=valuation_input_contract(spec,inputs)
    for field in [f for f in contract['fields'] if f['id'].startswith('assumptions.')]:
        row+=1;key=field['id'].split('.')[1]
        label(row,key.replace('_',' ').capitalize(),scaled_unit(field['unit'],field.get('scale',1)))
        refs[key]=put(row,4,field.get('value'),unit='ratio' if field['unit']=='ratio' else None,comment=field,input_cell=True)
        assumption_cells[field['id']]=refs[key]
        _cell(sheet,row,5,field.get('as_of'));status_height=status(row,6,field.get('status',''))
        sheet.merge_cells(start_row=row,start_column=7,end_row=row,end_column=last)
        method=_cell(sheet,row,7,field.get('method',''),color=GREY);method.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[row].height=min(72,max(30,status_height,15*math.ceil(len(str(field.get('method','')))/65)))
    row+=1;label(row,'Discount-rate sensitivity shift (baseline)','%');refs['wacc_shift']=put(row,4,0.,unit='ratio')
    for key,title in (('cost_of_equity','Cost of equity'),('pre_tax_cost_of_debt','Pre-tax cost of debt'),('wacc','Weighted average cost of capital')):
        row+=1;label(row,title,'%');target=valuation['wacc']['value'] if key=='wacc' else valuation['wacc'][key]
        refs[key]=equation(row,4,key,refs,target);sheet[refs[key]].number_format=RATE
    wacc_ref=refs['wacc'];locations['wacc']=wacc_ref
    # Detailed operating cash-flow bridge, generated from the same expressions.
    section(7,'OPERATING FORECAST TO UNLEVERED FREE CASH FLOW')
    for column,item in enumerate(detail,4):_cell(sheet,8,column,item['period_id'],bold=True,color='FFFFFF',fill=NAVY)
    row_keys=[('revenue','Revenue'),('operating_profit','Statutory operating profit'),
        ('noncash_portfolio_reversal','Reverse noncash portfolio gains / losses'),('normalized_operating_profit','Operating profit excluding portfolio gains / losses'),
        ('rou_depreciation_reversal','Reverse right-of-use depreciation'),('lease_principal','Cash lease principal'),('lease_interest','Lease interest'),
        ('cash_lease_cost','Cash lease cost'),('operating_profit_after_cash_leases','Operating profit after cash leases'),
        ('tax_rate','Unlevered tax rate'),('unlevered_tax','Unlevered operating tax'),('nopat','Net operating profit after tax'),
        ('owned_depreciation','Owned-asset depreciation'),('amortization','Amortization'),('eligible_da','Eligible depreciation and amortization'),
        ('capex_ppe','Cash investment in owned assets'),('capex_intangibles','Cash investment in intangible assets'),
        ('operating_working_capital_cash','Operating working-capital cash movement'),('other_cash_adjustments','Other evidenced operating cash adjustments'),
        ('sbc_addback','SBC addback — retained as an economic cost'),('full_period_fcff','Full-period unlevered free cash flow'),
        ('remaining_fraction','Proportion after valuation date'),('scope_timing_adjustment','Dated operating-scope timing adjustment'),
        ('remaining_fcff','Remaining unlevered free cash flow'),('discount_years','Mid-period discount duration'),
        ('discount_factor','Discount factor'),('present_value','Present value of remaining cash flow')]
    detail_rows={key:9+index for index,(key,_) in enumerate(row_keys)}
    ratios={'tax_rate','remaining_fraction','discount_factor'}
    sources=dict(revenue='revenue',operating_profit='operating_profit',noncash_portfolio_reversal='portfolio_noncash',
        rou_depreciation_reversal='rou_depreciation',lease_principal='lease_principal',lease_interest='lease_interest',
        owned_depreciation='owned_depreciation',amortization='amortization',capex_ppe='capex_ppe',capex_intangibles='capex_intangibles',
        operating_working_capital_cash='working_capital_cash')
    calc={'normalized_operating_profit','cash_lease_cost','operating_profit_after_cash_leases','unlevered_tax','nopat','eligible_da','full_period_fcff','remaining_fcff','discount_factor','present_value'}
    detail_pv=[]
    for key,title in row_keys:label(detail_rows[key],title,'%' if key in ratios else 'years' if key=='discount_years' else unit)
    for column,item in enumerate(detail,4):
        local={key:f'{get_column_letter(column)}{detail_rows[key]}' for key,_ in row_keys};local['wacc']=wacc_ref
        for key,_ in row_keys:
            target=item.get(key,0.);rr=detail_rows[key]
            if key in calc:equation(rr,column,key,local,target)
            elif key=='tax_rate':put(rr,column,'='+refs['tax_rate'],target)
            elif key in sources and roles.get(sources[key]) in values and values[roles[sources[key]]].get(item['period_id']) is not None:
                put(rr,column,'='+source_reference(roles[sources[key]],item['period_id']),target)
            else:put(rr,column,target,comment=dict(status='backend_snapshot',method=item.get('scope_timing_details') if key=='scope_timing_adjustment' else 'Dated backend input or allocation; rerun evaluation to refresh.'))
            if key in ratios:sheet[local[key]].number_format=RATE
            if key in ('nopat','full_period_fcff','remaining_fcff','present_value'):
                sheet[local[key]].border=TOTAL;sheet[local[key]].font=Font(name='Aptos',size=10,bold=True,color=INK)
        detail_pv.append(local['present_value'])
        locations[item['period_id']]=local
    note(37,'The scope timing adjustment uses existing event economics and control dates. Undisclosed event-specific lease, capex, working-capital and FX timing remains provisional. No elapsed operating cash flow is discounted again.')
    # Transition paths are snapshots; their NOPAT/reinvestment/discount equations remain inspectable.
    row+=3;section(row,'ANNUAL TRANSITION TO SUSTAINABLE ECONOMICS');row+=1
    for column,item in enumerate(transition,4):_cell(sheet,row,column,item['end'][:4],bold=True,color='FFFFFF',fill=NAVY)
    tkeys=[('revenue','Revenue'),('revenue_growth','Revenue growth'),('operating_margin','Operating margin after cash leases'),('tax_rate','Unlevered tax rate'),
           ('nopat','Net operating profit after tax'),('reinvestment_rate','Reinvestment / NOPAT'),('economic_reinvestment','Economic reinvestment'),
           ('full_period_fcff','Unlevered free cash flow'),('discount_years','Discount duration'),('discount_factor','Discount factor'),('present_value','Present value')]
    trows={key:row+1+index for index,(key,_) in enumerate(tkeys)};transition_pv=[];transition_refs=[]
    for key,title in tkeys:label(trows[key],title,'%' if key in ('revenue_growth','operating_margin','tax_rate','reinvestment_rate','discount_factor') else 'years' if key=='discount_years' else unit)
    for column,item in enumerate(transition,4):
        local={key:f'{get_column_letter(column)}{trows[key]}' for key,_ in tkeys};local.update(wacc=wacc_ref,next_year_revenue=local['revenue'],reinvestment=local['economic_reinvestment'],remaining_fcff=local['full_period_fcff'])
        for key,_ in tkeys:
            target=valuation['wacc']['tax_rate'] if key=='tax_rate' else item[key]
            if key=='tax_rate':put(trows[key],column,'='+refs['tax_rate'],target)
            elif key in ('nopat','economic_reinvestment','full_period_fcff','discount_factor','present_value'):
                equation(trows[key],column,dict(nopat='terminal_nopat',economic_reinvestment='reinvestment',full_period_fcff='terminal_fcff',discount_factor='discount_factor',present_value='present_value')[key],local,target)
            else:put(trows[key],column,target,comment='Backend transition snapshot: growth/margin/reinvestment fade is rebuilt by evaluation, not by editing this cell.')
            if key in ('revenue_growth','operating_margin','tax_rate','reinvestment_rate','discount_factor'):sheet[local[key]].number_format=RATE
        transition_pv.append(local['present_value']);transition_refs.append(local)
    row+=len(tkeys)+2;note(row,'Transition length and paths come from typed inputs. Changing terminal assumptions locally updates terminal equations; rerun the backend to regenerate the intervening transition path.');row+=2
    section(row,'TERMINAL VALUE');terminal_rows={};terminal=valuation['terminal'];row+=1
    terminal_items=[('base_revenue','Revenue entering terminal year'),('growth','Perpetual growth'),('next_year_revenue','Next-year revenue'),
        ('operating_margin','Sustainable operating margin'),('tax_rate','Unlevered tax rate'),('nopat','Sustainable next-year NOPAT'),
        ('roic','Sustainable return on invested capital'),('reinvestment_rate','Reinvestment / NOPAT'),('reinvestment','Required reinvestment'),
        ('fcff','Sustainable next-year free cash flow'),('value','Terminal value at final year end'),('discount_years','Discount years to final year end'),('present_value','Present value of terminal value')]
    for key,title in terminal_items:terminal_rows[key]=row;label(row,title,'%' if key in ('growth','operating_margin','tax_rate','roic','reinvestment_rate') else 'years' if key=='discount_years' else unit);row+=1
    tr={key:'D'+str(rr) for key,rr in terminal_rows.items()};tr['wacc']=wacc_ref
    for key,_ in terminal_items:
        target=valuation['wacc']['tax_rate'] if key=='tax_rate' else terminal[key]
        if key in ('growth','operating_margin','tax_rate','roic'):
            ar={'growth':'terminal_growth','operating_margin':'terminal_margin','tax_rate':'tax_rate','roic':'terminal_roic'}[key]
            put(terminal_rows[key],4,'='+refs[ar],target,unit='ratio')
        elif key=='base_revenue' and transition_refs:put(terminal_rows[key],4,'='+transition_refs[-1]['revenue'],target)
        elif key=='base_revenue':
            annual=next((p['id'] for p in spec.get('periods',[]) if p.get('kind')=='FY' and p['end']==detail[-1]['end']),detail[-1]['period_id'])
            put(terminal_rows[key],4,'='+source_reference(roles['revenue'],annual),target)
        elif key=='discount_years':put(terminal_rows[key],4,target,comment='Dated backend snapshot; valuation date and fiscal endpoint determine this duration.')
        else:equation(terminal_rows[key],4,{'next_year_revenue':'next_year_revenue','nopat':'terminal_nopat','reinvestment_rate':'reinvestment_rate','reinvestment':'reinvestment','fcff':'terminal_fcff','value':'terminal_value','present_value':'terminal_present_value'}[key],tr,target)
        if key in ('growth','operating_margin','tax_rate','roic','reinvestment_rate'):sheet[tr[key]].number_format=RATE
    note(row,'Terminal growth requires reinvestment = NOPAT × growth / ROIC. The mid-period perpetuity includes (1 + WACC)^0.5 and is discounted to the valuation date. WACC must exceed growth; backend evaluation rejects invalid cases.');row+=2
    # Event dates and economic claims are separate from operating cash flows.
    event_pv=[]
    if valuation.get('portfolio_events'):
        section(row,'DATED PORTFOLIO CASH AND RETAINED OWNERSHIP');row+=1
        for event in valuation['portfolio_events']:
            note(row,str(event['id'])+' | Closing '+event['date']);row+=1;event_refs={'wacc':wacc_ref}
            event_keys=[('net_cash','Net acquisition / disposal cash'),('event_cash_tax','Cash tax on transaction'),('retained_interest_value','Retained interest created at closing'),
                ('incremental_nci_claim','Incremental minority claim'),('existing_interest_transfer','Previously held associate stake transferred'),('nonlease_debt_transfer','Nonlease debt assumed / transferred'),
                ('discount_years','Years to closing'),('discount_factor','Discount factor'),('present_value','Present value of event claims')]
            for key,title in event_keys:
                label(row,title,'years' if key=='discount_years' else '%' if key=='discount_factor' else unit);event_refs[key]='D'+str(row)
                if key=='discount_factor':
                    equation(row,4,'discount_factor',event_refs,event[key]);sheet[event_refs[key]].number_format=RATE
                elif key=='present_value':equation(row,4,'event_present_value',event_refs,event[key])
                else:put(row,4,event.get(key,0.),comment='Dated backend event input. Ownership-adjusted amounts; applied once.');row+=1;continue
                row+=1
            event_pv.append(event_refs['present_value'])
    section(row,'ENTERPRISE VALUE TO COMMON EQUITY');row+=1
    er={};label(row,'Present value of operating cash flows',unit)
    er['operating_cashflow_present_value']=put(row,4,'=SUM('+','.join(detail_pv+transition_pv)+')',sum(c['present_value'] for c in detail+transition));row+=1
    label(row,'Present value of terminal value',unit);er['terminal_present_value']=put(row,4,'='+tr['present_value'],terminal['present_value']);row+=1
    for key,target in valuation['equity_bridge'].items():
        label(row,key.replace('_',' ').capitalize(),unit);er[key]='D'+str(row)
        if key=='enterprise_value':equation(row,4,'enterprise_value',er,target)
        elif key=='future_portfolio_events':put(row,4,'=SUM('+','.join(event_pv)+')' if event_pv else 0.,target)
        elif key in ('restricted_cash','pension_claim','deferred_consideration','other_claims'):put(row,4,'=-'+refs[key],target)
        elif key=='other_assets':put(row,4,'='+refs[key],target)
        else:put(row,4,target,comment=valuation.get('opening_bridge',{}).get({'nonlease_debt':'debt'}.get(key,key),'Dated valuation bridge snapshot; no forecast terminal cash is added.'))
        row+=1
    label(row,'Common equity value',unit);er['equity_value']=equation(row,4,'equity_value',er,valuation['equity_value']);locations['equity_value']=er['equity_value'];row+=1
    label(row,'Valuation-date basic outstanding shares',share_unit);er['basic_shares']=put(row,4,valuation['opening_bridge']['shares']['value'],comment=valuation['opening_bridge']['shares']);row+=1
    er['dilution_shares']=refs['dilution_shares']
    label(row,'Valuation-date shares including explicit dilution',share_unit);er['valuation_date_shares']=equation(row,4,'valuation_date_shares',er,valuation['valuation_date_shares']);row+=1
    for key in ('money_scale','share_scale'):
        label(row,key.replace('_',' ').capitalize(),'units');er[key]=put(row,4,valuation[key]);row+=1
    label(row,'Fair value per share — valuation date',valuation['currency']);locations['fair_value_per_share']=equation(row,4,'fair_value_per_share',er,valuation['fair_value_per_share']);row+=1
    sheet[locations['fair_value_per_share']].border=TOTAL;sheet[locations['fair_value_per_share']].font=Font(name='Aptos',size=12,bold=True,color=INK)
    sheet[locations['fair_value_per_share']].number_format=NUMBER.replace('0.0','0.00')
    label(row,'Terminal present value / enterprise value','%');put(row,4,'='+tr['present_value']+'/'+er['enterprise_value'],terminal['share_of_enterprise_value'],unit='ratio');row+=2
    section(row,'DATED MARKET COMPARISON');row+=1
    quote=inputs.get('market',{}).get('price',{});label(row,'Quoted price',inputs.get('market',{}).get('currency'))
    market_price=put(row,4,quote.get('value'),unit='per_share',comment=quote,input_cell=True);assumption_cells['market.price']=market_price
    _cell(sheet,row,5,quote.get('as_of'));sheet.row_dimensions[row].height=max(30,status(row,6,quote.get('status')));row+=1
    label(row,'Quote unit conversion to reporting currency','multiple');scale=put(row,4,inputs.get('market',{}).get('price_unit_scale',1.));row+=1
    if valuation.get('market_comparison',{}).get('status')=='available':
        label(row,'Current market price in reporting currency',valuation['currency']);market_normalized=put(row,4,'='+market_price+'*'+scale,valuation['market_comparison']['price_in_reporting_currency'],unit='per_share');row+=1
        label(row,'Implied upside / downside','%');put(row,4,'='+locations['fair_value_per_share']+'/'+market_normalized+'-1',valuation['market_comparison']['upside'],unit='ratio');row+=1
    else:note(row,'No reliable dated quote is available. Missing is not a zero share price; upside and current multiples remain unavailable.');row+=1
    for item in valuation.get('implied_multiples',[]):
        for key,title in (('implied_ev_ebit','DCF EV / operating profit after leases'),('implied_ev_ebitda','DCF EV / EBITDA after leases'),('market_ev_ebit','Market EV / operating profit after leases'),('market_ev_ebitda','Market EV / EBITDA after leases')):
            label(row,item['period_id']+' — '+title,'x');put(row,4,item.get(key),comment='Backend comparison snapshot; rerun valuation after operating or market changes.');row+=1
    row+=1;section(row,'VALUATION SENSITIVITIES — BACKEND SCENARIOS');row+=1
    for column,title in ((2,'WACC'),(3,'Growth'),(4,'Fair value / share'),(5,'Terminal reinvestment'),(6,'Status')):
        _cell(sheet,row,column,title,bold=True).alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[row].height=30
    for point in valuation.get('sensitivities',{}).get('wacc_growth',[]):
        row+=1;put(row,2,point.get('wacc'),unit='ratio');put(row,3,point.get('growth'),unit='ratio');put(row,4,point.get('fair_value_per_share'),unit='per_share');put(row,5,point.get('terminal_reinvestment'));_cell(sheet,row,6,point.get('status'))
    row+=2;note(row,'Every sensitivity calls the same Python valuation function. Growth changes recompute terminal reinvestment. These are saved scenarios, not Excel data tables.');row+=2
    for point in valuation.get('sensitivities',{}).get('event_consideration',[]):
        label(row,str(point['event_id'])+' — consideration × '+str(point['net_cash_multiplier']),valuation['currency']+'/share');put(row,4,point['fair_value_per_share'],unit='per_share');row+=1
    row+=1;section(row,'PROVENANCE, CASH-FLOW RECONCILIATION AND LIMITATIONS');row+=1
    for item in valuation.get('statement_cf_reconciliation',[]):
        note(row,item['period_id']+' statement CFO to FCFF: '+item['status']+(' | residual '+format(item['residual'],'.8g') if item.get('residual') is not None else ''))
        sheet.cell(row,2).comment=Comment(_text(item),'Refractive');row+=1
    for item in valuation.get('diagnostics',[]):note(row,str(item.get('message',item)));row+=1
    note(row,'Sources and assumption methods are preserved in valuation.json / valuation_inputs.json. Workbook edits do not update these saved inputs or renew evidence validation.');row+=1
    # A3 landscape pages reserve room for the repeated title and methods. Pack
    # complete semantic blocks where possible; long blocks break between rows.
    # This keeps headings with their data without company-specific coordinates.
    page_capacity=600.;used=0.
    starts=sorted(section_starts)
    for index,start in enumerate(starts):
        end=starts[index+1] if index+1<len(starts) else row+1
        heights=[sheet.row_dimensions[r].height or 15 for r in range(start,end)]
        block=sum(heights)
        minimum=sum(heights[:min(4,len(heights))])
        if used and used+(block if block<=page_capacity else minimum)>page_capacity:
            sheet.row_breaks.append(Break(id=start-1));used=0.
        for current,height in zip(range(start,end),heights):
            if used and used+height>page_capacity:
                sheet.row_breaks.append(Break(id=current-1));used=0.
            used+=height
    sheet.print_area=f'B2:{last_letter}{row}'
    return dict(expected_valuation_cells=expected,valuation_cells=locations,valuation_assumption_cells=assumption_cells,
                valuation_print_area=f'B2:{last_letter}{row}',valuation_status=valuation.get('status'),valuation_formula_count=formula_count,
                valuation_section_rows=starts,valuation_page_breaks=[item.id for item in sheet.row_breaks.brk])


def _legacy_history(workbook, spec):
    """Readable numerical histories keep each reported scope and vintage separate."""
    observations = [item for item in spec.get("observations", []) if item.get("entity_id") not in (None, "group")]
    if not observations:
        if spec.get("historical_segments"):
            rows = []
            for item in spec["historical_segments"]:
                rows.append([item.get("period", item.get("year")), item.get("label"), item.get("basis"), item.get("unit"),
                             item.get("sales"), item.get("utop"), item.get("top"), item.get("da")])
            _table(workbook, "Historical Segments", ["Period", "Reported segment", "Reporting basis", "Units", "Sales", "Underlying trading profit", "Trading profit", "D&A"], rows,
                   {1: 18, 2: 35, 3: 48, 6: 25})
        return {}
    sheet = _table(workbook, "Historical Segments", ["Period", "Financial metric", "Units", "Value", "Observation status", "Reporting basis", "Valid dates", "Observation ID"], [],
                   {1: 18, 2: 53, 3: 15, 4: 21, 5: 24, 6: 50, 7: 30, 8: 32})
    _cell(sheet, 3, 2, "Original, restated and derived observations remain separate. Scope and vintage determine comparability; these blocks are not added together.", color=GREY)
    groups = {}
    for item in observations:
        groups.setdefault((item.get("entity_id"), item.get("basis", "")), []).append(item)
    links, r = {}, 6
    for (entity, basis), members in groups.items():
        label = next((v.get("entity_label") for v in members if v.get("entity_label")), str(entity))
        _bar(sheet, r, f"{label}  |  {basis}", 8)
        for member in members:
            links.setdefault(member.get("model_entity_id", entity), r)
        r += 1
        for item in sorted(members, key=lambda v: (v.get("start", ""), v.get("end", ""), v.get("metric", ""), v.get("id", ""))):
            data = [item.get("period_id"), item.get("label", item.get("metric")), item.get("unit"), item.get("value"),
                    item.get("status", item.get("restatement")), basis, f"{item.get('start', '')} to {item.get('end', '')}", item.get("id")]
            for c, value in enumerate(data, 1):
                target = _cell(sheet, r, c, value, color=PINK if c == 4 else INK, number_format=_format(item.get("unit", "")) if c == 4 else None)
                target.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(r, 4).comment = Comment(_text({k: item.get(k) for k in ("source_refs", "restatement", "vintage", "derivation", "scope")}), "Refractive")
            sheet.row_dimensions[r].height = 30
            r += 1
        r += 1
    sheet.auto_filter.ref = None
    sheet.print_area = f"A2:H{r - 1}"
    return links


def _reporting_bridges(workbook, bridges):
    if not bridges:
        return
    sheet = workbook.create_sheet("Reporting Bridges")
    _style_sheet(sheet, "C6")
    sheet.column_dimensions["A"].width = 3
    sheet.column_dimensions["B"].width = 34
    for column in ("C", "D", "E"):
        sheet.column_dimensions[column].width = 20
    sheet.column_dimensions["F"].width = 55
    _cell(sheet, 2, 2, "Reporting Bridges", bold=True)
    _cell(sheet, 3, 2, "Exact reconstructions and analytical proxies remain distinct. Only the selected, non-overlapping basis enters consolidation.", color=GREY)
    sheet.merge_cells("B3:F3")
    sheet["B3"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[3].height = 30
    _cell(sheet, 5, 2, "Transition field", color="FFFFFF", bold=True, fill=NAVY)
    for c in range(3,7):
        _cell(sheet, 5, c, "Evidence and supported relationship" if c == 3 else None, color="FFFFFF", bold=True, fill=NAVY)
    sheet.merge_cells("C5:F5")
    r = 6
    for item in bridges:
        predecessor = item.get("predecessor_ids", item.get("predecessors", []))
        successor = item.get("successor_ids", item.get("successors", []))
        def names(values):
            return ", ".join(str(v.get("entity_id", v.get("label", ""))) if isinstance(v,dict) else str(v) for v in values) if isinstance(values,list) else str(values)
        title = f"{item.get('predecessor_basis', names(predecessor))} → {item.get('successor_basis', names(successor))}"
        if item.get("period_id") or item.get("metric"):
            title += f"  |  {item.get('period_id', '')}  |  {item.get('metric', '')}"
        _bar(sheet, r, title, 6)
        sheet.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        sheet.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[r].height = max(27,15*math.ceil(len(title)/110))
        sheet.cell(r, 2).comment = Comment(_text(item), "Refractive")
        r += 1
        fields = [("Effective / publication dates", {k: item.get(k) for k in ("effective_date", "effective", "publication_date", "published_at", "publication_status") if item.get(k)}),
                  ("Scope change", item.get("scope_change", item.get("description", item.get("summary", item.get("explanation"))))),
                  ("Affected metrics / periods", {k: item.get(k) for k in ("metric", "period_id", "metrics", "periods", "affected_periods", "unit") if item.get(k)}),
                  ("Relationship / exactness", item.get("relationship", item.get("status", item.get("exactness")))),
                  ("Supported reconstruction", {k: item.get(k) for k in ("amounts", "adjustments", "weights", "equations") if item.get(k)})]
        def field(label,value):
            nonlocal r
            content = _text(value) if value is not None and value != {} and value != [] else "Not supplied / not applicable"
            _cell(sheet, r, 2, label, bold=True)
            _cell(sheet, r, 3, content).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.merge_cells(start_row=r,start_column=3,end_row=r,end_column=6)
            sheet.row_dimensions[r].height = min(409, max(30, 15 * math.ceil(len(str(content)) / 100)))
            r += 1
        for label, value in fields:
            if label != "Supported reconstruction" or value:
                field(label,value)
        relationships = item.get("relationships", [])
        if relationships:
            for c,label in enumerate(["Predecessor → successor","Before","After","Change","Interpretation"],2):
                _cell(sheet,r,c,label,bold=True,fill=PALE)
            r += 1
            for relationship in relationships:
                _cell(sheet,r,2,f"{relationship.get('predecessor','')} → {relationship.get('successor','')}").alignment=Alignment(wrap_text=True,vertical="top")
                for c,key in enumerate(("before","after","delta"),3):
                    _cell(sheet,r,c,relationship.get(key),number_format=_format(item.get("unit", "")))
                _cell(sheet,r,6,relationship.get("interpretation",relationship.get("status",""))).alignment=Alignment(wrap_text=True,vertical="top")
                sheet.cell(r,3).comment=Comment(_text(relationship.get("source_refs",[])),"Refractive")
                sheet.row_dimensions[r].height=45
                r += 1
        if any(item.get(k) is not None for k in ("group_before","group_after","group_difference")):
            _cell(sheet,r,2,"Group control",bold=True,fill=PALE)
            for c,key in enumerate(("group_before","group_after","group_difference"),3):
                _cell(sheet,r,c,item.get(key),bold=True,fill=PALE,number_format=_format(item.get("unit", "")))
            r += 1
        field("Residual / limitations", {k:item.get(k) for k in ("residual","limitations","reason","attribution_limit") if item.get(k)})
        field("Source evidence", item.get("source_refs",item.get("source_ref",item.get("change_source_ref",item.get("primary_source_id")))))
        r += 1
    sheet.print_area = f"B2:F{r - 1}"


def _export_periodic(path, spec, values, checks, *, valuation=None):
    """Render one explicit cell graph at fiscal quarter, half-year and year frequency."""
    from .periods import cell_dependencies, cell_formula_excel, iter_identity_cells
    path = Path(path)
    all_periods, rows = spec["periods"], spec["rows"]
    # Cumulative source periods remain in the calculation graph and Historics,
    # but are not company presentation columns, even when outlines expand.
    periods = [p for p in all_periods if p['kind'] != '9M']
    internal_periods = [p for p in all_periods if p['kind'] == '9M']
    diagnostic = spec.get("model_configuration", {}).get("diagnostic_export", False)
    period_by_id = {str(p["id"]): p for p in all_periods}
    row_by_id = {row["id"]: row for row in rows}
    if len(period_by_id) != len(all_periods) or len(row_by_id) != len(rows):
        raise ValueError("Duplicate row or period identifier")
    columns = {str(p["id"]): i + 4 for i, p in enumerate(periods)}
    internal_columns={str(p['id']):len(columns)+i+4 for i,p in enumerate(internal_periods)}
    history_columns=columns | internal_columns
    if not columns:
        raise ValueError("The periodic model requires periods")
    workbook = Workbook()
    model = workbook.active
    model.title = "Model"
    _style_sheet(model, "D7")
    model.print_title_rows = "2:6"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    model.column_dimensions["A"].hidden = True
    model.column_dimensions["B"].width = 55
    model.column_dimensions["C"].width = 14
    last = max(columns.values())
    _cell(model, 2, 2, spec["company"], bold=True)
    model["B2"].font = Font(name="Aptos", size=18, bold=True, color=NAVY)
    model.row_dimensions[2].height = 30
    opening_note = "Temporary scenario · backend inputs are authoritative; Excel edits stay local · expand year groups for interim detail"
    if spec.get('financial_completeness') == 'incomplete_scope_accounting':
        opening_note = "Provisional model — material portfolio closing accounting is incomplete · edit Assumptions; yellow Model boxes are linked displays"
    _cell(model, 3, 2, opening_note, color=GREY)
    _cell(model, 4, 2, "A: disclosed actuals  D/italic: derived  F: forecast  E: estimate  M: mixed coverage. Blank: unavailable; details in cell comments and audit files.", color=GREY)
    model.merge_cells(start_row=4, start_column=2, end_row=4, end_column=last)
    model["B4"].alignment = Alignment(wrap_text=True, vertical="center")
    model.row_dimensions[4].height = 30 if len(periods) < 8 else 18
    _cell(model, 6, 2, "Financial model", color="FFFFFF", bold=True, fill=NAVY)
    _cell(model, 6, 3, "Units", color="FFFFFF", bold=True, fill=NAVY)
    forecast_years = sorted({p["fiscal_year"] for p in periods if p.get("status") == "forecast"})
    expanded = set(spec.get("expanded_fiscal_years", ([forecast_years[0] - 1, forecast_years[0]] if forecast_years else [max(p["fiscal_year"] for p in periods)])))
    suffix = {"actual": "A", "reported": "A", "derived_actual": "D", "forecast": "F", "estimate": "E", "mixed": "M"}
    year_groups = {}
    for period in periods:
        pid, c = str(period["id"]), columns[str(period["id"])]
        year_groups.setdefault(period["fiscal_year"], []).append(c)
        letter = get_column_letter(c)
        model.column_dimensions[letter].width = 13
        model.column_dimensions[letter].min=c
        model.column_dimensions[letter].max=c
        detail = period["kind"] != "FY"
        model.column_dimensions[letter].outlineLevel = 1 if detail else 0
        model.column_dimensions[letter].hidden = bool(period.get("hidden")) or not period.get("display", True) or (detail and period["fiscal_year"] not in expanded)
        if not detail:
            model.column_dimensions[letter].collapsed = period["fiscal_year"] not in expanded
        header = _cell(model, 6, c, period["kind"] + suffix.get(period.get("status"), ""), color="FFFFFF", bold=True,
                       fill=("28505E" if period.get("status") in ("forecast", "estimate") else "142C42") if period["kind"]=="FY" else ("376B7D" if period.get("status") in ("forecast", "estimate") else NAVY))
        header.comment = Comment(_text(period), "Refractive")
        header.alignment = Alignment(horizontal="right")
    for year, group in year_groups.items():
        for c in group:
            _cell(model, 5, c, year if c == group[0] else None, color="FFFFFF", bold=True, fill=NAVY)
        if len(group) > 1:
            model.merge_cells(start_row=5, start_column=group[0], end_row=5, end_column=group[-1])
        model.cell(5, group[0]).alignment = Alignment(horizontal="center")
    model.row_dimensions[5].height = 20
    model.row_dimensions[6].height = 22
    positions, section_positions, current = {}, {}, 8
    for section in spec["sections"]:
        members = [row for row in rows if row["section"] == section["id"] and not row.get('fx_only') and not row.get('support_matrix')]
        if not members:
            continue
        section_positions[section["id"]] = current
        _bar(model, current, section["label"], last)
        current += 2
        previous_assumption = False
        for row in members:
            if previous_assumption and row["kind"] != "assumption":
                current += 1
            positions[row["id"]] = current
            _cell(model, current, 1, row["id"])
            label = _cell(model, current, 2, row["label"], bold=row.get("style") == "subtotal")
            if len(row['label'])>62:
                label.alignment=Alignment(wrap_text=True,vertical='center')
                model.row_dimensions[current].height=15*math.ceil(len(row['label'])/62)
            label.comment = Comment("\n".join(str(row[k]) for k in ("definition", "basis", "note") if row.get(k)), "Refractive")
            _cell(model, current, 3, row.get("display_unit", _display_unit(row["unit"])), color=GREY)
            if row.get("supporting_input") or row.get("fx_only"):
                model.row_dimensions[current].hidden = True
                model.row_dimensions[current].outlineLevel = 1
            previous_assumption = row["kind"] == "assumption"
            current += 1
        current += 2
    if len(positions) != sum(not r.get('fx_only') and not r.get('support_matrix') for r in rows):
        raise ValueError("Every model row must belong to a declared section")
    model.print_area = f"B2:{get_column_letter(last)}{current - 1}"
    assumptions = [row for row in rows if row["kind"] == "assumption"]
    support_rows=[row for row in rows if row.get('support_matrix') and row['kind']!='assumption']
    control_rows=assumptions+support_rows
    assumption_positions, assumption_headers, next_assumption = {}, {}, 6
    for section in spec['sections']:
        members=[row for row in control_rows if row['section']==section['id']]
        if not members:continue
        assumption_headers[next_assumption]=section['label'];next_assumption+=1
        for row in members:
            assumption_positions[row['id']]=next_assumption;next_assumption+=1
        next_assumption+=1
    assumption_cells = {row["id"]: {pid: {"sheet": "Assumptions", "cell": f"{get_column_letter(c)}{assumption_positions[row['id']]}"} for pid, c in columns.items()} for row in control_rows}
    assumption_sheet = _table(workbook, "Assumptions", ["Driver ID", "Authoritative input / calculation", "Units"] + list(columns), [], {1: 35, 2: 55, 3: 14})
    _cell(assumption_sheet, 3, 2, "Backend input projection. Blue/yellow: workbook-local input. Black: calculation. Pink: calibration. Excel edits do not update the saved backend model.", color=GREY)
    for index,label in assumption_headers.items():_bar(assumption_sheet,index,label,last)
    fx_sheet = _table(workbook, "FX", ["Rate ID", "Imported period average", "Quote units"] + list(columns), [], {1: 35, 2: 55, 3: 20})
    _cell(fx_sheet, 3, 2, "Shared database snapshot: actual daily crosses plus spot held over future dates. Currency weights link to authoritative Assumptions; provisional exposures are estimates.", color=GREY)
    metadata=spec.get('fx_metadata',{})
    _cell(fx_sheet, 4, 2, "Latest observation across sources: " + str(metadata.get('daily_end','unavailable')) + "; individual currency coverage varies. See provider coverage below and fx_import.json.", color=GREY)
    fx_rows = [row for row in rows if row["section"] == "fx" and (row["kind"] != "assumption" or row.get("fx_only"))]
    fx_positions = {row["id"]: i for i, row in enumerate(fx_rows, 6)}
    if not fx_rows:
        _cell(fx_sheet, 6, 2, "No evidenced currency exposure selected. Configure relevant currencies before importing rates into the model.")
    for row in fx_rows:
        fr = fx_positions[row["id"]]
        for c, value in enumerate((row["id"], row["label"], _display_unit(row["unit"])), 1):
            _cell(fx_sheet, fr, c, value)
        for pid, c in columns.items():
            descriptor = row["cells"].get(pid, {})
            cell = _cell(fx_sheet, fr, c, descriptor.get("value"), color=BLUE, number_format=_format(row["unit"]))
            if descriptor.get('status')=='estimate':
                cell.fill=PatternFill('solid',fgColor=YELLOW)
                cell.number_format=';'.join(part+'" E"' for part in cell.number_format.split(';')[:3])
            if descriptor.get('status')!='not_applicable':cell.comment = Comment(_text(descriptor), "Refractive")
    history = _table(workbook, "Historics", ["Metric ID", "Historical inputs / supporting parameters", "Units"] + list(history_columns), [], {1: 35, 2: 55, 3: 14})
    _cell(history, 3, 2, "Source observations; internal cumulative periods retain live derivations in the hidden source columns. Original references and vintages remain in audit files.", color=GREY)
    for sheet in (assumption_sheet, history, fx_sheet):
        sheet.freeze_panes='D6'
        sheet.column_dimensions['A'].hidden=True
        for pid, c in columns.items():
            sheet.column_dimensions[get_column_letter(c)].width = 15
            sheet.column_dimensions[get_column_letter(c)].min=c
            sheet.column_dimensions[get_column_letter(c)].max=c
            sheet.column_dimensions[get_column_letter(c)].hidden = model.column_dimensions[get_column_letter(c)].hidden
            sheet.column_dimensions[get_column_letter(c)].outlineLevel = model.column_dimensions[get_column_letter(c)].outlineLevel
    for pid,c in internal_columns.items():history.column_dimensions[get_column_letter(c)].hidden=True
    history_rows=rows
    history_positions,next_history={},6
    for section in spec['sections']:
        members=[row for row in history_rows if row['section']==section['id']]
        if not members:continue
        _bar(history,next_history,section['label'],last);next_history+=1
        for row in members:
            history_positions[row['id']]=next_history;next_history+=1
        next_history+=1
    for row in history_rows:
        hr = history_positions[row["id"]]
        for c, value in enumerate((row["id"], row["label"], _display_unit(row["unit"])), 1):
            _cell(history, hr, c, value)
        for pid, c in columns.items():
            descriptor = row.get("cells", {}).get(pid, {})
            value = descriptor.get("value")
            if value is not None and not _number(value):
                raise ValueError(f"Non-numeric source/input value: {row['id']} {pid}")
            if descriptor.get('data_role')=='estimated_historical_baseline' or descriptor.get('calibration_eligible') is False:
                continue
            if row["kind"] != "assumption" or descriptor.get("status") not in ("forecast",) and not descriptor.get("editable"):
                _cell(history, hr, c, value, color=BLUE, number_format=_format(row["unit"]))
    def resolve(key, pid):
        if key not in row_by_id or pid not in history_columns:
            raise ValueError(f"Unknown financial cell reference: {key} {pid}")
        if pid in internal_columns:return "'Historics'!"+get_column_letter(internal_columns[pid])+str(history_positions[key])
        if key in assumption_cells:
            return "'Assumptions'!" + assumption_cells[key][pid]["cell"]
        if key in fx_positions:
            return "'FX'!" + get_column_letter(columns[pid]) + str(fx_positions[key])
        return "'Model'!" + get_column_letter(columns[pid]) + str(positions[key])
    def lookup(key, pid):
        if key in fx_positions:
            ref = "'FX'!" + get_column_letter(columns[pid]) + str(fx_positions[key])
            return f'=IF(ISNUMBER({ref}),{ref},IF(ISERROR({ref}),{ref},""))'
        end = max([6,*history_positions.values()])
        ref = (f'INDEX(Historics!$D$6:${get_column_letter(last)}${end},MATCH("{key}",Historics!$A$6:$A${end},0),'
               f'MATCH("{pid}",Historics!$D$5:${get_column_letter(last)}$5,0))')
        return f'=IF(COUNT({ref})=1,{ref},NA())'
    missing = {"not_disclosed": None, "not_applicable": None, "not_comparable": None, "error": "ERR"}
    expected_model, expected_assumptions, expected_checks = {}, {}, {}
    expected_history={}
    for row in rows:
        for pid,c in internal_columns.items():
            descriptor=row.get('cells',{}).get(pid,{})
            content=cell_formula_excel(descriptor,pid,resolve) if descriptor.get('expression') else descriptor.get('value')
            target=_cell(history,history_positions[row['id']],c,content,color=INK if descriptor.get('expression') else BLUE,number_format=_model_format(row))
            expected_history[target.coordinate]=values.get(row['id'],{}).get(pid)
    dependency_rows, driver_rows, source_rows, formula_count = [], [], [], 0
    for row in rows:
        r = positions.get(row["id"])
        ar = assumption_positions.get(row["id"])
        if r is None:
            # Currency support matrices have no duplicate hidden Model rows.
            if ar:
                for c,value in enumerate((row['id'],row['label'],row.get('display_unit',_display_unit(row['unit']))),1):_cell(assumption_sheet,ar,c,value)
                for pid,c in columns.items():
                    descriptor=row.get('cells',{}).get(pid,{})
                    content=cell_formula_excel(descriptor,pid,resolve) if descriptor.get('expression') else descriptor.get('value')
                    control=_cell(assumption_sheet,ar,c,content,color=BLUE if descriptor.get('editable') else INK,
                        fill=YELLOW if descriptor.get('editable') else None,number_format=_model_format(row))
                    if descriptor.get('editable'):control.border=BOX
                    expected_assumptions[control.coordinate]=values[row['id']].get(pid)
            continue
        if ar:
            for c, value in enumerate((row["id"], row["label"], row.get("display_unit", _display_unit(row["unit"]))), 1):
                _cell(assumption_sheet, ar, c, value)
            assumption_sheet.cell(ar, 2).hyperlink = f"#'Model'!B{r}"
            model.cell(r,2).hyperlink=f"#'Assumptions'!B{ar}"
        for pid, c in columns.items():
            descriptor = row.get("cells", {}).get(pid, {"status": "not_disclosed", "reason": "No mapped observation or formula"})
            status = descriptor.get("status", "forecast" if period_by_id[pid].get("status") == "forecast" else "actual")
            value = values.get(row["id"], {}).get(pid)
            formula = descriptor.get("expression")
            if ar and descriptor.get("editable") and not formula and descriptor.get("value") is None:
                content, target = None, None
            elif status in missing and not formula and descriptor.get("value") is None:
                content, target = missing[status], missing[status]
            elif formula:
                content, target = cell_formula_excel(descriptor, pid, resolve), value
            elif descriptor.get("value") is not None:
                estimated=descriptor.get('data_role')=='estimated_historical_baseline' or descriptor.get('calibration_eligible') is False
                content = descriptor["value"] if estimated or ar and (descriptor.get("editable") or status == "forecast") else lookup(row["id"], pid)
                target = descriptor["value"]
            else:
                content, target = (None, None)
            source = content
            color = INK if formula else (PINK if status in ("actual", "reported") else GREY)
            if ar:
                control = _cell(assumption_sheet, ar, c, content, color=BLUE if descriptor.get("editable") else color,
                                fill=YELLOW if descriptor.get("editable") else None, number_format=_model_format(row))
                if status!='not_applicable' or not row.get('fx_only'):
                    control.comment = Comment(_text({"status": status, "reason": descriptor.get("reason"), "source_refs": descriptor.get("source_refs"), "definition": row.get("definition")}), "Refractive")
                if descriptor.get("editable"):
                    control.border = BOX
                expected_assumptions[control.coordinate] = target
                ref = f"'Assumptions'!{control.coordinate}"
                if target is not None or formula or descriptor.get("editable"):
                    content = f'=IF(ISNUMBER({ref}),{ref},IF({ref}="","",{ref}))' if not descriptor.get("editable") else f'=IF(ISBLANK({ref}),"",{ref})'
                    if descriptor.get("editable") and target is None:
                        target = None
                color = GREEN if status in ("forecast", "estimate") else color
                dependency_rows.append([row["id"], pid, "Model", f"{get_column_letter(c)}{r}", row["id"], pid, "Assumptions", control.coordinate, "Linked display"])
                driver_rows.append([row["id"], row["label"], row["section"], pid, period_by_id[pid].get("start"), period_by_id[pid].get("end"), row["unit"], status, bool(descriptor.get("editable")), ref, descriptor.get("reason", row.get("note", ""))])
            target_cell = _cell(model, r, c, content, color=color, bold=row.get("style") == "subtotal",
                                fill=YELLOW if ar and status in ("forecast", "estimate") else PALE if row.get("style") == "subtotal" else None,
                                number_format=_model_format(row))
            if not row.get('fx_only'):
                target_cell.comment = Comment(_text({"status": status, "classification":descriptor.get('data_role'),"calibration_eligible":descriptor.get('calibration_eligible'),"reason": descriptor.get("reason"), "period": period_by_id[pid], "source_refs": descriptor.get("source_refs")}), "Refractive")
            if ar:
                # Excel can replace an empty formula's displayed cache with its
                # hyperlink destination. Row labels provide navigation instead.
                if status in ("forecast", "estimate"):
                    target_cell.border = BOX
            if status == "derived_actual":
                derived_font = copy.copy(target_cell.font)
                derived_font.italic = True
                target_cell.font = derived_font
            if not ar and row.get("style") == "subtotal":
                target_cell.border = TOTAL
            if status == "estimate" and row['section']=='fx':
                target_cell.fill = PatternFill("solid", fgColor="FFF2CC")
                target_cell.number_format = ";".join(part + '" E"' for part in target_cell.number_format.split(";")[:3])
            expected_model[target_cell.coordinate] = target
            formula_count += isinstance(content, str) and content.startswith("=")
            formula_count += ar is not None and isinstance(source, str) and source.startswith("=")
            for dependency in cell_dependencies(descriptor, pid) if formula else []:
                dep_ref = resolve(dependency["row_id"], dependency["period_id"])
                dependency_rows.append([row["id"], pid, "Assumptions" if ar else "Model", assumption_cells[row["id"]][pid]["cell"] if ar else target_cell.coordinate,
                                        dependency["row_id"], dependency["period_id"], dep_ref.split("!")[0].strip("'"), dep_ref.split("!")[1], descriptor["expression"]])
            if descriptor.get("source_refs") or formula or status in missing:
                source_rows.append([row["id"], row["label"], pid, row["unit"], status, descriptor.get("reason"), descriptor.get("source_refs"), descriptor.get("expression"), descriptor.get("bindings")])
        if "check_tolerance" in row:
            first = f"D{r}"
            model.conditional_formatting.add(f"D{r}:{get_column_letter(last)}{r}", Rule(type="expression", stopIfTrue=True,
                formula=[f"AND(ISNUMBER({first}),ABS({first})<={row['check_tolerance']:.15g})"],
                dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=164, formatCode='"–";"–";"–"'))))
        if row.get("style") == "subtotal":
            for c in (2, 3):
                model.cell(r, c).border = TOTAL
            model.row_dimensions[r].height = 20
    # FX owns the detailed rate/exposure displays; their Model rows stay hidden.
    # All references use stable row IDs and generated positions after reordering.
    expected_fx={}
    for row in fx_rows:
        for pid,c in columns.items():
            descriptor=row['cells'].get(pid,{})
            if descriptor.get('status')=='not_applicable':content=None
            elif row['id'] in assumption_cells:
                ref="'Assumptions'!"+assumption_cells[row['id']][pid]['cell']
                content=f'=IF(ISNUMBER({ref}),{ref},IF(ISERROR({ref}),{ref},""))'
            elif descriptor.get('expression'):
                content=cell_formula_excel(descriptor,pid,resolve)
            else:content=descriptor.get('value')
            cell=fx_sheet.cell(fx_positions[row['id']],c);cell.value=content
            cell.font=Font(name='Aptos',size=10,color=GREEN if row['id'] in assumption_cells else INK if descriptor.get('expression') else BLUE)
            expected_fx[cell.coordinate]=values[row['id']].get(pid)
    coverage_row=None
    if spec.get('fx_currency_coverage'):
        coverage_row=fx_sheet.max_row+3
        _bar(fx_sheet,coverage_row,'CURRENCY COVERAGE — genuine currency crosses and explicit proxies are different',min(last,12))
        headings=['Currency','Selected provider / series','Last genuine history','Spot timestamp','Spot type','Observation definition','Gaps','Transition approval']
        for c,label in enumerate(headings,2):
            _cell(fx_sheet,coverage_row+1,c,label,bold=True,fill=PALE).alignment=Alignment(wrap_text=True,vertical='top')
        fx_sheet.row_dimensions[coverage_row+1].height=36
        for offset,item in enumerate(spec['fx_currency_coverage'],2):
            transitions=item.get('provider_transitions',[])
            transitions=[t for t in transitions if item.get('currency') in (t.get('base'),t.get('quote'))]
            transition_summary='; '.join(str(t.get('provider',''))+': '+str(t.get('approval','unspecified')) for t in transitions) or 'No cutover approved'
            selected=set(item.get('selected_series',[]))
            held=[s for s in item.get('available_native_series',[]) if f"{s['provider']}:{s['base']}/{s['quote']}:{s['definition']}" not in selected]
            if held and not item.get('last_genuine_observation') and item.get('currency')!=item.get('reporting_currency'):
                transition_summary += '; native history held: '+str(sum(s.get('observations',0) for s in held))
            spot_time=item.get('spot_timestamp') or (', '.join(item.get('spot_source_dates',[]))+' (fixing)' if item.get('spot_source_dates') else None)
            spot_type=' / '.join(str(item.get(k)) for k in ('spot_rate_type','spot_status') if item.get(k))
            details=[item.get('currency'),'; '.join(item.get('selected_providers',[])) or 'No selected history',item.get('last_genuine_observation'),spot_time,spot_type,
                '; '.join(item.get('observation_definitions',[])),item.get('coverage_gap_count'),transition_summary]
            for c,value in enumerate(details,2):
                cell=_cell(fx_sheet,coverage_row+offset,c,_text(value));cell.alignment=Alignment(wrap_text=True,vertical='top')
            fx_sheet.row_dimensions[coverage_row+offset].height=72
            fx_sheet.cell(coverage_row+offset,2).comment=Comment(_text(item),'Refractive')
    hidden_assumption_rows={}
    for row in control_rows:
        if row.get('support_matrix') and not any(cell.get('editable') for cell in row.get('cells',{}).values()):
            index=assumption_positions[row['id']]
            assumption_sheet.row_dimensions[index].hidden=True
            assumption_sheet.row_dimensions[index].outlineLevel=1
            hidden_assumption_rows[row['id']]=index
    for index in assumption_positions.values():
        label=assumption_sheet.cell(index,2)
        if len(str(label.value or ''))>62:
            label.alignment=Alignment(wrap_text=True,vertical='center')
            assumption_sheet.row_dimensions[index].height=15*math.ceil(len(str(label.value))/62)
    # Darker FY body and header fills, retaining input boxes and subtotal hierarchy.
    for pid,c in columns.items():
        if period_by_id[pid]['kind']!='FY':continue
        for sheet in (model,assumption_sheet,fx_sheet,history):
            header=sheet.cell(6 if sheet is model else 5,c)
            header.fill=PatternFill('solid',fgColor='23465C')
            header.font=Font(name='Aptos',size=10,bold=True,color='FFFFFF')
            for index in range(8 if sheet is model else 6,sheet.max_row+1):
                cell=sheet.cell(index,c)
                if isinstance(cell,MergedCell):continue
                fill=cell.fill.fgColor.rgb if cell.fill.patternType=='solid' else None
                if fill in ('00'+NAVY,NAVY):continue
                if fill in ('00'+YELLOW,YELLOW):shade='F2DF9B'
                elif fill in ('00'+PALE,PALE):shade='D5DFE7'
                else:shade='E8EDF2'
                cell.fill=PatternFill('solid',fgColor=shade)
    if diagnostic:
        _table(workbook, "Drivers", ["Driver ID", "Driver", "Entity / block", "Period", "Start", "End", "Units", "Status", "Editable", "Authoritative cell", "Basis / rationale"], driver_rows,
               {1: 35, 2: 50, 3: 25, 10: 32, 11: 90})
    if diagnostic:
        _table(workbook, "Dependencies", ["Target ID", "Target period", "Target sheet", "Target cell", "Dependency ID", "Dependency period", "Dependency sheet", "Dependency cell", "Expression / link"], dependency_rows,
               {1: 35, 5: 35, 9: 90})
    for item in spec.get("observations", []):
        source_rows.append([item.get("id"), item.get("label", item.get("metric")), item.get("period_id"), item.get("unit"), item.get("status"), item.get("basis"), item.get("source_refs"), item.get("derivation"), {k: item.get(k) for k in ("entity_id", "scope", "restatement", "vintage")}])
    for source in spec.get("sources", []):
        source_rows.append(["SOURCE", "Document metadata", "", "", "", "", source, "", ""])
    if diagnostic:
        _table(workbook, "Sources", ["Metric / observation ID", "Metric", "Period", "Units", "Status", "Basis / reason", "Source references", "Derivation", "Bindings / vintage"], source_rows,
               {1: 36, 2: 52, 6: 65, 7: 95, 8: 75, 9: 85})
    if diagnostic:_supporting(workbook, "Periods", periods)
    check_sheet = _table(workbook, "Checks", ["Check ID", "Check", "Period", "Residual", "Tolerance", "Live result", "Control type", "Notes"], [], {1: 35, 2: 65, 3: 18, 8: 90})
    r = 6
    for identity, pid, descriptor in iter_identity_cells(spec):
        formula = cell_formula_excel(descriptor, pid, resolve)
        deps = {(dep["symbol"], dep["lag"]): dep for dep in cell_dependencies(descriptor, pid)}
        residual = formula_value(expression_tree(descriptor["expression"]), lambda symbol, lag: values.get(deps[(symbol, lag)]["row_id"], {}).get(deps[(symbol, lag)]["period_id"]))
        tolerance = identity.get("tolerance", 1e-6)
        content = [identity["id"], identity.get("label", identity["id"]), pid, formula, tolerance,
                   f'=IF(ISNUMBER(D{r}),IF(ABS(D{r})<=E{r},"PASS","FAIL"),"MISSING")',
                   "Independent source control" if identity.get("independent") else "Arithmetic identity", identity.get("note", "")]
        for c, value in enumerate(content, 1):
            _cell(check_sheet, r, c, value, number_format=NUMBER if c in (4, 5) else None)
        expected_checks[f"D{r}"] = residual
        expected_checks[f"F{r}"] = "MISSING" if residual is None else "PASS" if abs(residual) <= tolerance else "FAIL"
        r += 1
    if checks and diagnostic:
        r += 2
        _bar(check_sheet, r, "FROZEN VALIDATION — source and financial reviews require a new checked run", 8)
        r += 1
        for check in checks:
            for c, value in enumerate([check.get("id"), check.get("label", check.get("check")), check.get("period"), check.get("residual"), check.get("tolerance"), check.get("status"), check.get("category"), check.get("detail")], 1):
                _cell(check_sheet, r, c, _text(value))
            r += 1
    legacy_links = _legacy_history(workbook, spec)
    for section, r in section_positions.items():
        entity = section.removeprefix("segment_")
        if entity in legacy_links:
            model.cell(r, 2).hyperlink = f"#'Historical Segments'!B{legacy_links[entity]}"
            model.cell(r, 2).comment = Comment("Open the linked legacy operating histories, retaining scope and restatement variants.", "Refractive")
    if diagnostic:_reporting_bridges(workbook, spec.get("reporting_bridges"))
    for key, title in (("reporting_changes", "Reporting Changes"), ("interims", "Interims"), ("fx_observations", "FX Data")):
        if spec.get(key) and diagnostic:
            _supporting(workbook, title, spec[key])
    if spec.get("fx_interfaces") and diagnostic:
        interface_sheet = _table(workbook, "FX Interfaces", ["Interface", "Definition / current state"],
            [[key.replace("_", " ").capitalize(), value] for key, value in spec["fx_interfaces"].items()], {1: 38, 2: 125})
        _cell(interface_sheet, 3, 2, "Imported rates and company FX effects are separate. Unknown exposures and closing rates remain explicit; no double translation is applied.", color=GREY)
        for r in range(6, interface_sheet.max_row + 1):
            interface_sheet.row_dimensions[r].height = 60
    if spec.get("missing_value_audit") and diagnostic:
        _missing_value_audit(workbook, spec["missing_value_audit"])
    if spec.get("gaps") and diagnostic:
        _table(workbook, "Gaps", ["Disclosure / modelling limitation"], [[item] for item in spec["gaps"]], {1: 145})
    if not diagnostic:
        check_sheet.sheet_state = "hidden"  # Needed by the live compact summary.
    summary_row = 7
    live_end = max(6, 5 + len(expected_checks) // 2)
    frozen='Saved checks: '+str(sum(c.get('status') in ('FAIL','MISSING') for c in checks))+' FAIL/MISSING; '+str(sum(c.get('status')=='WARNING' for c in checks))+' warnings. Material limitations remain in the banner and audit.'
    _cell(model,summary_row,2,f'="Live arithmetic: "&(COUNTIF(Checks!F6:F{live_end},"FAIL")+COUNTIF(Checks!F6:F{live_end},"MISSING"))&" exceptions. {frozen}"',color=INK)
    model.merge_cells(start_row=summary_row,start_column=2,end_row=summary_row,end_column=last)
    finding_counts={}
    for item in checks:
        if item.get('status') in ('FAIL','MISSING','WARNING'):
            key=item.get('area','review')+': '+item['status']
            finding_counts[key]=finding_counts.get(key,0)+1
    model.cell(summary_row,2).comment=Comment(_text(dict(gaps=spec.get('gaps',[]),finding_counts=finding_counts,
        full_findings='Every finding and its evidence is preserved in model.audit.json under checks.',
        fx='Rates and exposure matrices on FX; provisional weights and adjustment factors on Assumptions. Historical reported FX is preserved.')),'Refractive')
    valuation_inspection = _valuation_sheet(workbook, spec, values, valuation, resolve) if valuation is not None else {}
    audit = dict(style_version=STYLE_VERSION, dependencies=dependency_rows, drivers=driver_rows, sources=source_rows, checks=checks, reporting_bridges=spec.get("reporting_bridges", []), missing_value_audit=spec.get("missing_value_audit", {}), fx_metadata=spec.get("fx_metadata", {}),fx_currency_coverage=spec.get('fx_currency_coverage',[]),fx_exposure_quality=spec.get('fx_exposure_quality',[]),financial_contracts=spec.get('financial_contracts',{}), valuation=valuation)
    for sheet in workbook:
        for cells in sheet:
            for cell in cells:
                if cell.comment and len(cell.comment.text)>32767:
                    reference=sheet.title+'!'+cell.coordinate
                    audit.setdefault('extended_cell_comments',{})[reference]=cell.comment.text
                    cell.comment=Comment('This commentary exceeds the Excel comment size supported by this exporter. The complete text is preserved in model.audit.json under extended_cell_comments['+reference+']. No source evidence has been discarded.','Refractive')
    path.parent.mkdir(parents=True, exist_ok=True)
    path.with_suffix(".audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    workbook._sheets.sort(key=lambda sh: {"Model": 0, "Valuation": 1, "Assumptions": 2, "Historics": 3, "FX": 4, "Historical Segments": 5}.get(sh.title, 6))
    for sheet in workbook:
        sheet.sheet_view.tabSelected = False
        sheet.sheet_view.topLeftCell='A1'
        sheet.sheet_view.selection=list({selection.pane:selection for selection in sheet.sheet_view.selection}.values())
        for selection in sheet.sheet_view.selection:
            selection.activeCell='D7' if sheet.title=='Model' else 'D6'
            selection.sqref=selection.activeCell
    workbook.active = 0
    model.sheet_view.tabSelected = True
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return dict(style_version=STYLE_VERSION, status="structure_and_formula_syntax_checked", recalculation="not_run", formula_count=formula_count,
                expected_model_cells=expected_model, expected_check_cells=expected_checks, expected_assumption_cells=expected_assumptions, expected_fx_cells=expected_fx,expected_history_cells=expected_history,hidden_assumption_rows=hidden_assumption_rows,
                model_rows=positions, fx_rows=fx_positions, period_columns={pid: get_column_letter(c) for pid, c in columns.items()},
                internal_period_columns={pid:get_column_letter(c) for pid,c in internal_columns.items()},history_rows=history_positions,
                year_columns={str(p["fiscal_year"]): get_column_letter(columns[str(p["id"])]) for p in periods if p["kind"] == "FY"},
                periods=periods, forecast_years=spec.get('forecast_years',[]),forecast_horizon=spec.get('forecast_horizon',{}),internal_periods=internal_periods,section_rows=section_positions, fx_coverage_row=coverage_row,fx_currency_coverage=spec.get('fx_currency_coverage',[]),assumption_count=len(assumptions), assumption_cells=assumption_cells, header_rows="2:6", **valuation_inspection)


def validate_workbook_layout(path, inspection):
    """Inspect saved columns and opening views, including Excel's outline spans."""
    from openpyxl.utils.cell import column_index_from_string
    book=load_workbook(path,read_only=False,data_only=False)
    findings=[]
    try:
        periods={p['id']:p for p in inspection.get('periods',[])}
        columns=inspection.get('period_columns',{})
        def hidden(sheet,index):
            for key,dim in sheet.column_dimensions.items():
                begin=dim.min or column_index_from_string(key);end=dim.max or begin
                if begin<=index<=end:return bool(dim.hidden)
            return False
        for title in ('Model','Assumptions'):
            sheet=book[title];header=6 if title=='Model' else 5
            if any('9M' in str(cell.value) for cell in sheet[header]):findings.append(title+': cumulative 9M presentation column')
            for pid,column in columns.items():
                p=periods[pid];index=column_index_from_string(column)
                if p['kind']=='FY' and hidden(sheet,index):findings.append(title+': FY column hidden '+pid)
                if p.get('annual_only') and p['kind']!='FY':findings.append(title+': unnecessary interim forecast '+pid)
                if p['kind']!='FY' and p['fiscal_year'] in inspection.get('forecast_years',[]) and p.get('status') in ('forecast','mixed') and hidden(sheet,index):findings.append(title+': forecast detail hidden '+pid)
            for key,dim in sheet.column_dimensions.items():
                if not dim.hidden:continue
                begin=dim.min or column_index_from_string(key);end=dim.max or begin
                years={periods[pid]['fiscal_year'] for pid,col in columns.items() if begin<=column_index_from_string(col)<=end}
                if len(years)>1:findings.append(title+': hidden outline spans fiscal years '+key)
        if book.active.title!='Model':findings.append('Opening sheet is not Model')
        pane=book['Model'].sheet_view.pane
        if pane and pane.topLeftCell and int(re.search(r'\d+',pane.topLeftCell)[0])>10:findings.append('Opening pane is deep in model')
        return dict(status='passed' if not findings else 'failed',findings=findings,
            columns=list(columns),internal_source_periods=list(inspection.get('internal_period_columns',{})),
            visible_model_columns=[pid for pid,col in columns.items() if not hidden(book['Model'],column_index_from_string(col))],
            active_sheet=book.active.title)
    finally:book.close()


def _default_sensitivity_cases(spec, baseline):
    """Direct policy controls for the standard integrated compiler, not AST replay."""
    years = spec["forecast_years"]
    if not years:
        raise ValueError("Sensitivity verification requires forecast periods")
    year = years[0]
    following = years[1] if len(years) > 1 else None
    rows = {row["id"]: row for row in spec["rows"]}
    prefix = next((key[:-8] for key, row in rows.items() if key.endswith("_pricing") and row["kind"] == "assumption"), None)
    cases = []
    def value(key, period=year):
        return baseline.get(key, {}).get(str(period))
    def add(key, driver, bump, expectations, directions):
        cases.append(dict(id=key, driver=driver, year=year, bump=bump,
                          expected_deltas=[dict(row=key, year=period, delta=delta) for key, period, delta in expectations],
                          directions=[dict(row=key, year=period, direction=direction) for key, period, direction in directions]))
    if prefix:
        sales_effect = value(prefix + "_sales", year - 1) * .01
        add("segment_pricing_plus_1pt", prefix + "_pricing", .01,
            [(prefix + "_sales", year, sales_effect), ("sales", year, sales_effect), (prefix + "_utop", year, sales_effect),
             (prefix + "_cash_cost", year, 0)], [("utop", year, 1), ("cfo", year, 1)])
        cost_effect = value(prefix + "_cash_cost", year - 1) * (1 + value(prefix + "_rig") + value(prefix + "_fx")) * .01
        add("operating_cost_plus_1pt", prefix + "_cost_inflation", .01,
            [(prefix + "_cash_cost", year, cost_effect), (prefix + "_utop", year, -cost_effect), ("sales", year, 0)],
            [("utop", year, -1), ("cfo", year, -1)])
    if all(key in rows for key in ("inventory_ratio", "inventories", "cogs", "cfo", "cash")):
        inventory_effect = -value("cogs") * .01
        add("inventory_plus_1pt", "inventory_ratio", .01,
            [("inventories", year, inventory_effect), ("cfo", year, -inventory_effect), ("cash", year, -inventory_effect)], [])
    if all(key in rows for key in ("capex_ppe_ratio", "sales", "capex_ppe", "owned_ppe", "owned_ppe_depreciation", "cash")):
        investment = value("sales") * .01
        expectations = [("capex_ppe", year, -investment), ("owned_ppe", year, investment),
                        ("cash", year, -investment), ("owned_ppe_depreciation", year, 0)]
        if following is not None:
            expectations.append(("owned_ppe_depreciation", following, investment * value("owned_depreciation_rate", following)))
        add("capex_plus_1pt", "capex_ppe_ratio", .01, expectations, [])
    if all(key in rows for key in ("borrowing_input", "nonlease_debt", "debt_interest", "cash")):
        expectations = [("nonlease_debt", year, 100), ("cash", year, 100), ("debt_interest", year, 0)]
        if following is not None:
            expectations.append(("debt_interest", following, 100 * value("borrowing_rate", following)))
        add("borrowing_plus_100", "borrowing_input", 100, expectations, [])
    return cases


def _excel_equal(actual, expected):
    if expected is None:
        return actual in (None, "")
    if isinstance(expected, str):
        return actual == expected
    return _number(actual) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-7)


def verify_excel_sensitivities(path, spec, inspection, *, cases=None, output_directory=None):
    """Edit controls in real Excel on an isolated copy and restore between cases.

    Cases accept {id, driver, year, bump, expected_deltas:[{row,year,delta}],
    directions:[{row,year,direction}]}. All populated forecast cells and live check
    cells are compared with Python, plus explicit independent delta/timing checks.
    The original workbook remains byte-for-byte unchanged. No provider calls occur.
    """
    from .construction import calculate_integrated
    periodic = bool(spec.get("periods"))
    if periodic:
        from .periods import calculate_periodic, cell_dependencies, iter_identity_cells
    calculate = calculate_periodic if periodic else calculate_integrated
    path = Path(path).resolve()
    initial_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    baseline = calculate(spec)
    cases = copy.deepcopy((spec.get("sensitivity_cases", []) if periodic else _default_sensitivity_cases(spec, baseline)) if cases is None else cases)
    if not cases:
        raise ValueError("No sensitivity cases apply to this specification")
    rows = {row["id"]: row for row in spec["rows"]}
    years = [p["id"] for p in spec["periods"]] if periodic else spec["historical_years"] + spec["forecast_years"]
    def coordinate(key, year):
        return inspection["period_columns" if periodic else "year_columns"][str(year)] + str(inspection["model_rows"][key])
    def expected_outputs(values):
        result = {}
        if periodic:
            missing = {"not_disclosed": None, "not_applicable": None, "not_comparable": None, "error": "ERR"}
            for key, row in rows.items():
                for pid, descriptor in row.get("cells", {}).items():
                    if descriptor.get("expression") or row["kind"] == "assumption" or descriptor.get("status") in ("forecast", "estimate"):
                        value = values.get(key, {}).get(pid)
                        control_value = value
                        if value is None and not descriptor.get("expression"):
                            value = missing.get(descriptor.get("status"))
                            if descriptor.get("editable"):
                                value = None
                            else:
                                control_value = value
                        if pid in inspection.get('internal_period_columns',{}):
                            result['Historics!'+inspection['internal_period_columns'][pid]+str(inspection['history_rows'][key])]=value
                            continue
                        if key in inspection['model_rows']:result["Model!" + coordinate(key, pid)] = value
                        elif key in inspection.get('fx_rows',{}):result['FX!'+inspection['period_columns'][pid]+str(inspection['fx_rows'][key])]=value
                        if key in inspection["assumption_cells"] and pid in inspection['assumption_cells'][key]:
                            control = inspection["assumption_cells"][key][pid]
                            result[control["sheet"] + "!" + control["cell"]] = control_value
            for line, (identity, pid, descriptor) in enumerate(iter_identity_cells(spec), 6):
                deps = {(dep["symbol"], dep["lag"]): dep for dep in cell_dependencies(descriptor, pid)}
                residual = formula_value(expression_tree(descriptor["expression"]), lambda symbol, lag: values.get(deps[(symbol, lag)]["row_id"], {}).get(deps[(symbol, lag)]["period_id"]))
                result[f"Checks!D{line}"] = residual
                result[f"Checks!F{line}"] = "MISSING" if residual is None else "PASS" if abs(residual) <= identity.get("tolerance", 1e-6) else "FAIL"
            return result
        for key, row in rows.items():
            if row["kind"] == "assumption" or row.get("expression"):
                for year in spec["forecast_years"]:
                    result["Model!" + coordinate(key, year)] = values[key].get(str(year))
        line = 6
        for identity in spec.get("identities", []):
            for year in identity.get("periods") or years:
                year = int(year)
                if year not in years:
                    continue
                residual = formula_value(expression_tree(identity["expression"]), lambda key, lag: values[key].get(str(year - lag)))
                result[f"Checks!D{line}"] = residual
                result[f"Checks!F{line}"] = "MISSING" if residual is None else ("PASS" if abs(residual) <= identity["tolerance"] else "FAIL")
                line += 1
        return result
    predictions = {"baseline": expected_outputs(baseline)}
    changed_values = {}
    prepared = []
    for case in cases:
        driver = case["driver"]
        year = str(case.get("period_id", case.get("year"))) if periodic else int(case["year"])
        eligible = (rows.get(driver, {}).get("cells", {}).get(year, {}).get("editable", False) if periodic else year in spec["forecast_years"])
        if driver not in rows or rows[driver]["kind"] != "assumption" or not eligible:
            raise ValueError("Sensitivity must edit a declared forecast assumption")
        old = baseline[driver].get(str(year))
        if not _number(old) or not _number(case["bump"]):
            raise ValueError("Sensitivity input and bump must be finite")
        changed = old + case["bump"]
        calculated = calculate(spec, {driver: {str(year): changed}})
        predictions[case["id"]] = expected_outputs(calculated)
        changed_values[case["id"]] = calculated
        control = inspection["assumption_cells"][driver][str(year)]
        prepared.append(dict(id=case["id"], sheet=control["sheet"], cell=control["cell"], value=changed))
    targets = []
    for key in predictions["baseline"]:
        sheet, cell = key.split("!")
        match = re.fullmatch(r"([A-Z]+)([0-9]+)", cell)
        from openpyxl.utils.cell import column_index_from_string
        targets.append(dict(key=key, sheet=sheet, row=int(match[2]), column=column_index_from_string(match[1])))
    result = dict(status="unavailable", engine="Microsoft Excel COM", cases=[], original_sha256=initial_hash,
                  limitation="Python/Excel agreement is implementation consistency; explicit delta checks test the stated financial policy, not forecasting accuracy.")
    if os.name != "nt":
        result["detail"] = "Excel COM is available on Windows only"
        return result
    script = r'''$ErrorActionPreference='Stop'
$excel=$null
$book=$null
try {
 $payload=Get-Content -LiteralPath $env:RESEARCH_EXCEL_TASK -Raw | ConvertFrom-Json
 $excel=New-Object -ComObject Excel.Application
 $excel.Visible=$false
 $excel.DisplayAlerts=$false
 $excel.AutomationSecurity=3
 $book=$excel.Workbooks.Open($payload.workbook,0,$false)
 function Read-Values {
  $result=@{}
  foreach($sheetName in @('Model','Checks','Assumptions','FX','Historics')) {
   $sheet=$book.Worksheets.Item($sheetName)
   $used=$sheet.UsedRange
   $lastRow=$used.Row+$used.Rows.Count-1
   $lastColumn=$used.Column+$used.Columns.Count-1
   $matrix=$sheet.Range($sheet.Cells.Item(1,1),$sheet.Cells.Item($lastRow,$lastColumn)).Value2
   foreach($target in $payload.targets) {
    if($target.sheet -ne $sheetName) { continue }
    $value=$matrix[$target.row,$target.column]
    if($value -is [int] -and $value -lt -2146800000) { $value=$sheet.Cells.Item($target.row,$target.column).Text }
    $result[$target.key]=$value
   }
  }
  return $result
 }
 $excel.CalculateFullRebuild()
 $out=@{baseline=(Read-Values);cases=@{};restored=@{}}
 foreach($case in $payload.cases) {
  $cell=$book.Worksheets.Item($case.sheet).Range($case.cell)
  $original=$cell.Formula
  try {
   $cell.Value2=[double]$case.value
   $excel.CalculateFullRebuild()
   $out.cases[$case.id]=(Read-Values)
  } finally {
   $cell.Formula=$original
   $excel.CalculateFullRebuild()
  }
 }
 $out.restored=(Read-Values)
 $out | ConvertTo-Json -Depth 10 | Set-Content -LiteralPath $payload.result -Encoding UTF8
 $book.Close($false)
 $book=$null
 Write-Output 'SENSITIVITIES_COMPLETE'
} finally {
 if($null -ne $book) { $book.Close($false) }
 if($null -ne $excel) { $excel.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) }
}'''
    with tempfile.TemporaryDirectory(prefix="refractive_excel_sensitivity_") as directory:
        directory = Path(directory)
        copied = directory / "model.xlsx"
        shutil.copy2(path, copied)
        task_path, response_path = directory / "task.json", directory / "results.json"
        task_path.write_text(json.dumps(dict(workbook=str(copied), result=str(response_path), targets=targets, cases=prepared)), encoding="utf-8")
        try:
            completed = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                env=dict(os.environ, RESEARCH_EXCEL_TASK=str(task_path)), capture_output=True, text=True, timeout=360 if periodic else 180,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if completed.returncode or not response_path.exists():
                result["detail"] = completed.stderr[-1500:]
            else:
                actuals = json.loads(response_path.read_text(encoding="utf-8-sig"))
                def differences(actual, expected):
                    return [dict(cell=key, expected=value, actual=actual.get(key)) for key, value in expected.items()
                            if not _excel_equal(actual.get(key), value)]
                result["baseline_differences"] = differences(actuals["baseline"], predictions["baseline"])
                result["restored_differences"] = differences(actuals["restored"], predictions["baseline"])
                for case in cases:
                    actual = actuals["cases"][case["id"]]
                    comparison = differences(actual, predictions[case["id"]])
                    effects = []
                    for expectation in case.get("expected_deltas", []) + case.get("directions", []):
                        key = "Model!" + coordinate(expectation["row"], expectation.get("period_id", expectation.get("year")))
                        old, new = actuals["baseline"].get(key), actual.get(key)
                        delta = new - old if _number(new) and _number(old) else None
                        if "delta" in expectation:
                            passed = _excel_equal(delta, expectation["delta"])
                        else:
                            direction = expectation["direction"]
                            passed = _number(delta) and (abs(delta) <= 1e-7 if direction == 0 else delta * direction > 1e-7)
                        effects.append(dict(expectation, before=old, after=new, observed_delta=delta, status="PASS" if passed else "FAIL"))
                    result["cases"].append(dict(id=case["id"], driver=case["driver"], period_id=case.get("period_id", str(case.get("year"))), bump=case["bump"],
                        compared_cells=len(predictions[case["id"]]), differences=comparison, independent_effects=effects,
                        status="PASS" if not comparison and all(item["status"] == "PASS" for item in effects) else "FAIL"))
                result["status"] = "passed" if not result["baseline_differences"] and not result["restored_differences"] and all(case["status"] == "PASS" for case in result["cases"]) else "failed"
                result["restored"] = not result["restored_differences"]
        except (OSError, subprocess.TimeoutExpired) as exc:
            result["detail"] = type(exc).__name__
    result["original_unchanged"] = hashlib.sha256(path.read_bytes()).hexdigest() == initial_hash
    if not result["original_unchanged"]:
        result["status"] = "failed"
    if output_directory is not None:
        output = Path(output_directory) / "excel_sensitivity_validation.json"
        if output.exists():
            raise FileExistsError("Preserve earlier sensitivity evidence: " + str(output))
        output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


def export_excel_review(path, inspection, output_directory, *, sections=None, exposure_matrices=True):
    """Render the real Excel Model as a full PDF plus selected financial blocks.

    The PDF keeps every annual column. A temporary workbook copy supplies print
    settings; the delivered workbook and its baseline inputs remain unchanged.
    """
    path, output_directory = Path(path).resolve(), Path(output_directory).resolve()
    if os.name != "nt":
        return dict(status="unavailable", detail="Excel PDF export requires Windows COM")
    output_directory.mkdir(parents=True, exist_ok=True)
    section_rows = inspection["section_rows"]
    selected = list(section_rows) if sections is None else list(sections)
    maximum = max(inspection["model_rows"].values()) + 1
    ordered = sorted(section_rows.items(), key=lambda item: item[1])
    bounds = {key: (start, ordered[i + 1][1] - 2 if i + 1 < len(ordered) else maximum)
              for i, (key, start) in enumerate(ordered)}
    periodic = bool(inspection.get("period_columns"))
    all_columns = inspection.get("period_columns", inspection["year_columns"])
    last_column = list(all_columns.values())[-1]
    items = [dict(name="complete_model", area=f"B2:{last_column}{maximum}", path=str(output_directory / "complete_model.pdf"))]
    operating=next((start for key,start in ordered if key.startswith('segment_')),ordered[0][1])
    items.append(dict(name='opening',area=f'B2:{last_column}{min(maximum,operating+12)}',path=str(output_directory/'opening.pdf')))
    for key in selected:
        start, end = bounds[key]
        items.append(dict(name=key, area=f"B{start}:{last_column}{end}", path=str(output_directory / (key + ".pdf"))))
    if inspection.get('valuation_print_area'):
        items.append(dict(name='valuation',sheet='Valuation',area=inspection['valuation_print_area'],
            path=str(output_directory/'valuation.pdf'),print_title_rows='$2:$5'))
    if periodic:
        periods = inspection["periods"]
        forecast_years = sorted({p["fiscal_year"] for p in periods if p.get("status") == "forecast"})
        horizon = inspection.get('forecast_horizon', {})
        current = horizon.get('current_fiscal_year')
        through = horizon.get('detailed_through')
        if current is not None:
            detailed = horizon.get('detailed_forecast_years', 2)
            through = through if through is not None else current + detailed - 1
            focus = set(range(current, through + 2))
        else:
            focus = set([forecast_years[0] - 1, forecast_years[0]] if forecast_years else [max(p["fiscal_year"] for p in periods)])
        annual = [p for p in periods if p["kind"] == "FY"]
        detail = [p for p in periods if p["fiscal_year"] in focus and p.get("display", True) and not p.get("hidden") and p["kind"] != "9M"]
        def view_metadata(selected_periods):
            groups = {}
            for period in selected_periods:
                groups.setdefault(period["fiscal_year"], []).append(all_columns[period["id"]])
            return dict(visible_columns=[all_columns[p["id"]] for p in selected_periods],
                        year_headers=[dict(year=y, first=columns[0], last=columns[-1]) for y, columns in groups.items()])
        for item in items:
            if item.get('sheet','Model')=='Model':item.update(view_metadata(annual))
        opening=next(item for item in items if item['name']=='opening')
        items.append(dict(name='opening_periods',area=opening['area'],path=str(output_directory/'opening_periods.pdf'),**view_metadata(detail)))
        for key in selected:
            start, end = bounds[key]
            items.append(dict(name=key + "_periods", area=f"B{start}:{last_column}{end}", path=str(output_directory / (key + "_periods.pdf")), **view_metadata(detail)))
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=False)
            for title in ("Assumptions", "Historics", "FX", "Historical Segments", "Reporting Bridges", "FX Interfaces", "Missing Data Audit"):
                if title in book:
                    sheet = book[title]
                    name = 'support_'+title.lower().replace(" ", "_")
                    end_column = last_column if title in ("Assumptions", "Historics", "FX") else get_column_letter(sheet.max_column)
                    preview_end=min(sheet.max_row,48) if title in ('Assumptions','Historics','Historical Segments') else sheet.max_row
                    item = dict(name=name, sheet=title, area=f"A2:{end_column}{preview_end}", sampled_rows=[2,preview_end],total_sheet_rows=sheet.max_row,path=str(output_directory / (name + ".pdf")),
                                visible_columns=[all_columns[p["id"]] for p in detail] if title in ("Assumptions", "Historics", "FX") else [])
                    if title not in ("Missing Data Audit", "FX Interfaces"):
                        item["max_pages"] = 3
                    items.append(item)
            book.close()
        if inspection.get('fx_coverage_row'):
            start = inspection['fx_coverage_row']
            items.append(dict(name='fx_provider_coverage',sheet='FX',
                area=f'B{start}:K{start + len(inspection.get("fx_currency_coverage", [])) + 4}',
                path=str(output_directory/'fx_provider_coverage.pdf'),
                print_title_rows=f'${start}:${start + 1}',
                visible_columns=list('DEFGHIJK'),max_pages=3))
        # Exposure weights have one authoritative FY cell. Render their complete
        # currency lists separately, since a first-page rate preview omits them.
        baskets={}
        for rid,index in inspection.get('fx_rows',{}).items():
            if exposure_matrices and '_fx_weight_' in rid and not rid.endswith('_total'):
                baskets.setdefault(rid.split('_weight_',1)[0],[]).append((rid,index))
        matrix_columns=[all_columns[p['id']] for p in annual if p['fiscal_year'] in focus]
        for basket,members in baskets.items():
            for title in ('FX','Assumptions'):
                indices=[index if title=='FX' else int(re.search(r'\d+$',next(iter(inspection['assumption_cells'][rid].values()))['cell'])[0]) for rid,index in members]
                first,last_row=min(indices),max(indices)
                name=title.lower()+'_'+basket+'_weights'
                items.append(dict(name=name,sheet=title,area=f'B{first}:{last_column}{last_row}',sampled_rows=[first,last_row],
                    path=str(output_directory/(name+'.pdf')),visible_columns=matrix_columns,max_pages=3))
    if any(Path(item["path"]).exists() for item in items):
        raise FileExistsError("Review exports must preserve earlier artifacts")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    script = r'''$ErrorActionPreference='Stop'
$excel=$null
$book=$null
try {
 $payload=Get-Content -LiteralPath $env:RESEARCH_EXCEL_TASK -Raw | ConvertFrom-Json
 $excel=New-Object -ComObject Excel.Application
 $excel.Visible=$false
 $excel.DisplayAlerts=$false
 $excel.AutomationSecurity=3
 $book=$excel.Workbooks.Open($payload.workbook,0,$false)
 $excel.CalculateFullRebuild()
 foreach($item in $payload.items) {
  $name=if($item.sheet){$item.sheet}else{'Model'}
  $sheet=$book.Worksheets.Item($name)
  $sheet.Activate()
  $excel.ActiveWindow.Zoom=70
  $sheet.PageSetup.Orientation=2
  $sheet.PageSetup.PaperSize=8
  $sheet.PageSetup.Zoom=$false
  $sheet.PageSetup.FitToPagesWide=1
  $sheet.PageSetup.FitToPagesTall=$false
  $sheet.PageSetup.PrintTitleRows=if($item.print_title_rows){$item.print_title_rows}elseif($payload.periodic -and $name -eq 'Model'){'$2:$6'}else{'$2:$5'}
  $sheet.PageSetup.LeftMargin=18
  $sheet.PageSetup.RightMargin=18
  $sheet.PageSetup.TopMargin=18
  $sheet.PageSetup.BottomMargin=36
  $sheet.PageSetup.FooterMargin=12
  if($item.visible_columns.Count -gt 0) {
   $sheet.Range('D:'+$payload.last_column).EntireColumn.Hidden=$true
   foreach($column in $item.visible_columns) {$sheet.Range($column+':'+$column).EntireColumn.Hidden=$false}
  }
  if($payload.periodic -and $name -eq 'Model') {
   $sheet.Range('D5:'+$payload.last_column+'5').UnMerge()
   $sheet.Range('D5:'+$payload.last_column+'5').ClearContents()
   foreach($header in $item.year_headers) {
    $range=$sheet.Range($header.first+'5:'+$header.last+'5')
    if($header.first -ne $header.last) {$range.Merge()}
    $sheet.Range($header.first+'5').Value2=[double]$header.year
    $sheet.Range($header.first+'5').HorizontalAlignment=-4108
   }
  }
  $sheet.PageSetup.PrintArea=$item.area
  $sheet.ExportAsFixedFormat(0,$item.path,0,$true,$false)
 }
 $book.Close($false)
 $book=$null
 Write-Output 'REVIEW_EXPORTED'
} finally {
 if($null -ne $book) { $book.Close($false) }
 if($null -ne $excel) { $excel.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) }
}'''
    with tempfile.TemporaryDirectory(prefix="refractive_excel_review_") as directory:
        copied = Path(directory) / "model.xlsx"
        shutil.copy2(path, copied)
        task_path = Path(directory) / "task.json"
        task_path.write_text(json.dumps(dict(workbook=str(copied), items=items, periodic=periodic, last_column=last_column)), encoding="utf-8")
        completed = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            env=dict(os.environ, RESEARCH_EXCEL_TASK=str(task_path)), capture_output=True, text=True, timeout=300 if periodic else 180,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    unchanged = hashlib.sha256(path.read_bytes()).hexdigest() == before
    return dict(status="exported" if completed.returncode == 0 and unchanged and all(Path(item["path"]).exists() for item in items) else "failed",
                original_unchanged=unchanged, periods=list(all_columns), zoom=70,
                exports=items, detail=completed.stderr[-1000:])


def render_review_images(review, *, scale=2):
    """Rasterise actual Excel PDF pages for visual inspection, with optional tools.

    pypdfium2/Pillow are validation-only installations, not package dependencies.
    Rasterisation preserves Excel's rendering; it does not redraw a model mock-up.
    """
    if not 1 <= scale <= 4:
        raise ValueError("Review render scale must be between 1 and 4")
    try:
        import pypdfium2
        from PIL import Image, ImageChops
    except ImportError as exc:
        raise RuntimeError("Install optional pypdfium2 and Pillow for review rendering") from exc
    images = []
    for exported in review["exports"]:
        if exported["name"] == "complete_model":
            continue
        pdf = pypdfium2.PdfDocument(exported["path"])
        try:
            for page_index in range(min(len(pdf), exported.get("max_pages", len(pdf)))):
                page = pdf[page_index]
                bitmap = page.render(scale=scale)
                pil = bitmap.to_pil().convert("RGB")
                difference = ImageChops.difference(pil, Image.new("RGB", pil.size, "white"))
                # A separated page-number footer otherwise retains an entire
                # empty A3 page beneath a short financial block. Keep the PDF
                # intact; trim that isolated footer only from the review image.
                occupied = difference.getprojection()[1]
                bands = []
                for y, has_ink in enumerate(occupied):
                    if has_ink:
                        if not bands or bands[-1][1] < y - 1:
                            bands.append([y, y])
                        else:
                            bands[-1][1] = y
                if (len(bands) > 1 and bands[-1][0] > pil.height * .85
                        and bands[-1][0] - bands[-2][1] > 40 * scale
                        and bands[-1][1] - bands[-1][0] < 30 * scale):
                    pil = pil.crop((0, 0, pil.width, bands[-2][1] + int(12 * scale)))
                    difference = ImageChops.difference(pil, Image.new("RGB", pil.size, "white"))
                bounds = difference.getbbox()
                if bounds:
                    margin = int(12 * scale)
                    bounds = (max(0, bounds[0] - margin), max(0, bounds[1] - margin),
                              min(pil.width, bounds[2] + margin), min(pil.height, bounds[3] + margin))
                    pil = pil.crop(bounds)
                output = Path(exported["path"]).with_name(Path(exported["path"]).stem + f"_page{page_index + 1}.png")
                if output.exists():
                    raise FileExistsError("Preserve prior review image: " + str(output))
                pil.save(output)
                images.append(dict(section=exported["name"], page=page_index + 1, path=str(output), width=pil.width, height=pil.height))
                bitmap.close()
                page.close()
        finally:
            pdf.close()
    return images
