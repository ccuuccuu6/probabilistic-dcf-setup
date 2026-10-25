"""Read a labelled FX workbook without changing it or inventing company exposures.

The adapter recognises currency codes, quarter labels, dates and quote headings;
it does not rely on sheet coordinates. Values are reporting units per one source
unit. Daily crosses are formed before averaging, on the source publication dates.
"""
from __future__ import annotations

import calendar
import copy
from datetime import date, datetime, timedelta, timezone
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re

import openpyxl


def publication_dates(start,end,policy=None):
    """Declared observation calendar; explicit dates take precedence.

    TARGET excludes weekends, New Year, Good Friday, Easter Monday, May Day,
    Christmas and Boxing Day. Other providers may specify their own holidays.
    This is a sampling calendar, not a claim that every currency trades then.
    """
    policy=policy or {};start,end=_date(start),_date(end)
    if 'dates' in policy:return sorted(_date(d).isoformat() for d in policy['dates'] if start<=_date(d)<=end)
    holidays=set(policy.get('holidays',[]))
    if policy.get('name')=='TARGET':
        for year in range(start.year,end.year+1):
            a=year%19;b=year//100;c=year%100;d=b//4;e=b%4;f=(b+8)//25;g=(b-f+1)//3
            h=(19*a+b-d-g+15)%30;i=c//4;k=c%4;l=(32+2*e+2*i-h-k)%7;m=(a+11*h+22*l)//451
            easter=date(year,(h+l-7*m+114)//31,(h+l-7*m+114)%31+1)
            holidays.update([f'{year}-01-01',f'{year}-05-01',f'{year}-12-25',f'{year}-12-26',
                (easter-timedelta(days=2)).isoformat(),(easter+timedelta(days=1)).isoformat()])
    return [(start+timedelta(days=n)).isoformat() for n in range((end-start).days+1)
            if (start+timedelta(days=n)).weekday() in policy.get('weekdays',[0,1,2,3,4])
            and (start+timedelta(days=n)).isoformat() not in holidays]


def prepare_snapshot(quotes,as_of,policy):
    """Select dated native/normalized evidence without replacing provider history."""
    daily={};spots={};findings=[];groups={};available_series={};overlaps={}
    for q in quotes:
        if q['status']=='projected':continue
        if q['definition']=='bid_ask_midpoint':
            spots.setdefault(q['base']+'/'+q['quote'],[]).append(q);continue
        if q['status']=='actual' and q['source_date']==q['day']:
            key=(q['provider'],q['base'],q['quote'],q['definition'])
            series=available_series.setdefault(key,dict(provider=q['provider'],base=q['base'],quote=q['quote'],definition=q['definition'],first=q['day'],last=q['day'],observations=0))
            series['first']=min(series['first'],q['day']);series['last']=max(series['last'],q['day']);series['observations']+=1
        groups.setdefault((q['day'],q['base'],q['quote']),[]).append(q)
    for (day,base,quote),items in sorted(groups.items()):
        by_provider={q['provider']:q for q in items if q['status']=='actual' and q['source_date']==day}
        providers=sorted(by_provider)
        for i,left in enumerate(providers):
            for right in providers[i+1:]:
                a,b=by_provider[left],by_provider[right];key=(base,quote,left,right)
                audit=overlaps.setdefault(key,dict(base=base,quote=quote,providers=[left,right],overlap_days=0,max_relative_difference=0.,first=day,last=day))
                audit['overlap_days']+=1;audit['last']=day
                audit['max_relative_difference']=max(audit['max_relative_difference'],abs(a['value']/b['value']-1))
        transition=next((t for t in policy.get('cutovers',[]) if t['base']==base and t['quote']==quote and day>=t['from']),None)
        if transition:items=[q for q in items if q['provider']==transition['provider']]
        else:items=[q for q in items if (q['provider']!='IBKR' or q.get('extension_approved')) and q.get('transition_status') not in ('insufficient_overlap','quarantined')]
        if not items:continue
        best=sorted(items,key=lambda q:(q.get('selection_priority',10),q['status']!='actual',q['provider']))[0]
        daily.setdefault(day,[]).append(best)
    for pair,items in spots.items():spots[pair]=max(items,key=lambda q:q.get('timestamp',q['day']))
    # Retain validated direct observations. A spanning tree can discard a
    # perfectly valid direct pair and silently force an unnecessary cross.
    daily={day:_selected_quotes(items,policy) for day,items in daily.items()}
    for day,items in daily.items():
        bad=set()
        for q in items:
            other=[x for x in items if {x['base'],x['quote']}!={q['base'],q['quote']}]
            cross,path=_cross_quotes(other,q['quote'],q['base'])
            if cross is not None and abs(q['value']/cross-1)>policy.get('cross_consistency_tolerance',.05):
                findings.append(dict(day=day,pair=q['base']+'/'+q['quote'],status='inconsistent_direct_cross',
                    direct_id=q['id'],cross_ids=[x['id'] for x in path],relative_difference=q['value']/cross-1))
                bad.add(q['id'])
        if bad:daily[day]=[q for q in items if q['id'] not in bad]
    dates=sorted(day for day,items in daily.items() if items)
    latest={}
    for day in dates:
        for q in daily[day]:latest[q['base'],q['quote']]=q
    spot_candidates=list(spots.values())+list(latest.values())
    spot_state=_selected_quotes(spot_candidates,policy)
    for audit in overlaps.values():
        audit['approval']='passed_overlap_check' if audit['overlap_days']>=3 and audit['max_relative_difference']<=policy.get('cross_consistency_tolerance',.05) else 'insufficient_or_inconsistent_overlap'
    return dict(schema='shared_fx_snapshot_v1',as_of=as_of,policy=copy.deepcopy(policy),daily_quotes=daily,spots=spots,spot_state=spot_state,spot_candidates=spot_candidates,findings=findings,
        metadata=dict(daily_start=dates[0] if dates else None,daily_end=dates[-1] if dates else None,
            actual_through=dates[-1] if dates else None,quote_direction='reporting-currency units per one local-currency unit',
            calendar=policy.get('calendar',{'name':'weekdays','weekdays':[0,1,2,3,4]}),
            forecast_method='Observed elapsed publication dates plus latest eligible spot on remaining dates; daily cross before average.',
            as_of=as_of,closing_policy='Last scheduled observation date at or before period end; future closing uses eligible spot.',
            cross_consistency_tolerance=policy.get('cross_consistency_tolerance',.05),
            cross_consistency_note='Direct observations are retained within the configured discrepancy tolerance. Different provider sampling times can leave residual cross discrepancies; this is not an exactly arbitrage-free currency state.',
            source_currencies=sorted({q['base'] for q in quotes}|{q['quote'] for q in quotes}),
            available_native_series=list(available_series.values()),provider_overlap_audits=list(overlaps.values())))


def _selected_quotes(quotes,policy):
    """One observation per native pair; preserve direct routes and provenance."""
    priority=policy.get('provider_priority',['IBKR','ECB','ECB Daily'])
    selected={}
    for q in sorted(quotes,key=lambda q:(q.get('definition')!='bid_ask_midpoint',
            q.get('native') is False,
            priority.index(q['provider']) if q['provider'] in priority else len(priority),
            q.get('selection_priority',10),q['base'],q['quote'])):
        selected.setdefault(tuple(sorted((q['base'],q['quote']))),q)
    return list(selected.values())


def _cross_quotes(quotes,reporting,currency,max_timestamp_spread=None):
    """Direct/inverse pair first; otherwise genuine currency legs, never proxies."""
    if reporting==currency:return 1.0,[]
    adjacency={}
    for q in quotes:
        if _number(q.get('value')) is None:continue
        if q.get('proxy_currency') or q.get('is_proxy') or q.get('definition')=='currency_proxy':continue
        adjacency.setdefault(q['base'],[]).append((q['quote'],q['value'],q))
        adjacency.setdefault(q['quote'],[]).append((q['base'],1/q['value'],q))
    queue=[(currency,1.0,[])];seen={currency}
    for node,value,path in queue:
        for target,rate,q in adjacency.get(node,[]):
            if target in seen:continue
            route=path+[q]
            # Daily crosses must align observation dates. Carried source dates
            # remain visible and are checked separately against the carry policy.
            if len({leg['day'] for leg in route})>1 and not all(leg.get('timestamp') for leg in route):continue
            stamps=[datetime.fromisoformat(leg['timestamp'].replace('Z','+00:00')) for leg in route if leg.get('timestamp')]
            if stamps and max_timestamp_spread is not None and (max(stamps)-min(stamps)).total_seconds()>max_timestamp_spread:continue
            if target==reporting:return value*rate,route
            seen.add(target);queue.append((target,value*rate,route))
    return None,[]


def _route_metadata(path):
    return dict(rate_type='identity' if not path else 'direct' if len(path)==1 else 'cross_rate',
        providers=sorted({q['provider'] for q in path}),
        selected_series=sorted({q['provider']+':'+q['base']+'/'+q['quote']+':'+q['definition'] for q in path}),
        observation_definitions=sorted({q['definition'] for q in path}),
        lineage=[dict(id=q.get('id'),provider=q['provider'],base=q['base'],quote=q['quote'],
            value=q['value'],day=q['day'],source_date=q['source_date'],definition=q['definition'],
            timestamp=q.get('timestamp')) for q in path])


def spot_quote(snapshot,reporting,currency):
    asof=datetime.fromisoformat(snapshot['as_of'].replace('Z','+00:00'))
    policy=snapshot['policy'];calendar=policy.get('calendar',{})
    if reporting==currency:return dict(value=1.0,status='identity',date=asof.date().isoformat(),ids=[],**_route_metadata([]))
    quotes=snapshot.get('spot_candidates',snapshot.get('spot_state',list(snapshot.get('spots',{}).values())))
    closed=asof.weekday()==5 or asof.weekday()==6 and asof.hour<21 or asof.weekday()==4 and asof.hour>=21
    def eligible(q):
        if q.get('timestamp'):
            age=(asof-datetime.fromisoformat(q['timestamp'].replace('Z','+00:00'))).total_seconds()
            return 0<=age<=(72*3600 if closed else policy.get('spot_max_age_seconds',900)) and q.get('market_data_type') in (('live','frozen') if closed else ('live',))
        return len(publication_dates(_date(q['source_date'])+timedelta(days=1),asof.date(),calendar))<=policy.get('fixing_max_age_business_days',1)
    chosen=_selected_quotes([q for q in quotes if eligible(q)],policy)
    value,path=_cross_quotes(chosen,reporting,currency,policy.get('spot_alignment_seconds',900))
    if value is not None:
        for leg in path:
            cross,alternate=_cross_quotes([q for q in chosen if {q['base'],q['quote']}!={leg['base'],leg['quote']}],leg['quote'],leg['base'],policy.get('spot_alignment_seconds',900))
            if cross is not None and abs(leg['value']/cross-1)>policy.get('cross_consistency_tolerance',.05):
                return dict(value=None,status='inconsistent_direct_cross',ids=[q['id'] for q in path],
                    relative_difference=leg['value']/cross-1,cross_ids=[q['id'] for q in alternate],**_route_metadata(path))
    if value is None:value,path=_cross_quotes(_selected_quotes(quotes,policy),reporting,currency)
    if value is not None:
        timed=[q for q in path if q.get('timestamp')];fixings=[q for q in path if not q.get('timestamp')]
        timestamps=[datetime.fromisoformat(q['timestamp'].replace('Z','+00:00')) for q in timed]
        age=(asof-min(timestamps)).total_seconds() if timestamps else None
        spread=(max(timestamps)-min(timestamps)).total_seconds() if timestamps else 0
        lag=max([len(publication_dates(_date(q['source_date'])+timedelta(days=1),asof.date(),calendar)) for q in fixings] or [0])
        valid_ticks=not timed or (age>=0 and age<=(72*3600 if closed else policy.get('spot_max_age_seconds',900)) and
                    all(q.get('market_data_type') in (('live','frozen') if closed else ('live',)) for q in timed))
        valid_fixings=lag<=policy.get('fixing_max_age_business_days',1)
        valid=valid_ticks and valid_fixings and spread<=policy.get('spot_alignment_seconds',900)
        return dict(value=value if valid else None,status=('mixed_fixing_proxy' if timed else 'dated_fixing_proxy') if valid and fixings else 'closed_market_frozen' if valid and closed else 'live' if valid else 'stale_or_delayed',
            timestamp=min(timestamps).isoformat() if timestamps else None,age_seconds=age,age_business_days=lag,
            source_dates=sorted({q['source_date'] for q in path}),ids=[q['id'] for q in path],
            timestamp_spread_seconds=spread,**_route_metadata(path),
            convention='Direct observations first; genuine same-date currency crosses. Dated fixing proxies are not synchronous market spot.' if fixings else 'Timestamp-aligned bid/ask currency observations')
    # A source fixing is an explicit dated spot proxy, never labelled a live quote.
    for day in sorted(snapshot['daily_quotes'],reverse=True):
        value,path=_cross_quotes(snapshot['daily_quotes'][day],reporting,currency)
        if value is None:continue
        oldest=min(q['source_date'] for q in path)
        elapsed=len(publication_dates(_date(oldest)+timedelta(days=1),asof.date(),calendar))
        valid=elapsed<=policy.get('fixing_max_age_business_days',1)
        allowed=valid or policy.get('allow_stale_spot_scenario',False)
        return dict(value=value if allowed else None,status='dated_fixing_proxy' if valid else 'stale_scenario' if allowed else 'stale',
            date=day,source_date=oldest,age_business_days=elapsed,ids=[q['id'] for q in path],**_route_metadata(path))
    return dict(value=None,status='unavailable_currency',ids=[],**dict(_route_metadata([]),rate_type='unavailable'))


def forecast_rate(snapshot,reporting,currency,start,end,kind='average'):
    start,end=_date(start),_date(end)
    if start>end or kind not in ('average','closing'):raise ValueError('Invalid FX period')
    asof=datetime.fromisoformat(snapshot['as_of'].replace('Z','+00:00'))
    cutoff=asof.date();calendar=snapshot['policy'].get('calendar',{})
    days=publication_dates(start,end,calendar)
    result=dict(value=None,status='missing',kind=kind,reporting_currency=reporting,source_currency=currency,
        start=start.isoformat(),end=end.isoformat(),quote_direction=snapshot['metadata']['quote_direction'],
        provenance=[snapshot.get('source',{})],qualifications=[],coverage={},as_of=snapshot['as_of'])
    if not days:
        result['qualifications']=['No expected observations under the declared calendar'];return result
    if kind=='closing':days=days[-1:]
    spot=spot_quote(snapshot,reporting,currency);values=[];missing=[];actual=projected=carried=0
    source_ids=set();providers=set();definitions=set();series=set();route_types=set()
    for day in days:
        if _date(day)<=cutoff:
            value,path=_cross_quotes(snapshot['daily_quotes'].get(day,[]),reporting,currency)
            if value is not None:
                lag=max([len(publication_dates(_date(q['source_date'])+timedelta(days=1),day,calendar)) for q in path] or [0])
                if lag>snapshot['policy'].get('max_carry_business_days',0):value=None
                elif any(q['status']=='carried' for q in path):carried+=1
                else:actual+=1
                source_ids.update(q['id'] for q in path)
                metadata=_route_metadata(path)
                providers.update(metadata['providers']);definitions.update(metadata['observation_definitions'])
                series.update(metadata['selected_series']);route_types.add(metadata['rate_type'])
            # Today's unfinished publication slot may use spot as a projection.
            if value is None and _date(day)==cutoff and end>=cutoff and asof.hour<snapshot['policy'].get('publication_hour_utc',17):
                value=spot['value'];projected+=int(value is not None)
        else:
            value=spot['value'];projected+=int(value is not None)
        if value is None:missing.append(day)
        else:values.append(value)
    if projected:
        providers.update(spot.get('providers',[]));definitions.update(spot.get('observation_definitions',[]))
        series.update(spot.get('selected_series',[]));route_types.add(spot.get('rate_type','unavailable'))
    result.update(spot=spot,selected_providers=sorted(providers),selected_series=sorted(series),
        observation_definitions=sorted(definitions),rate_types=sorted(route_types),
        coverage=dict(expected_days=len(days),actual_days=actual,carried_days=carried,
        projected_days=projected,available_days=len(values),complete=not missing,missing_dates=missing),
        observation_ids_sha256=hashlib.sha256('|'.join(sorted(source_ids)).encode()).hexdigest())
    if missing:result['qualifications'].append('Missing elapsed observations or eligible spot; no present-day spot backfill for past dates.')
    else:result.update(value=sum(values)/len(values),status='estimate' if projected else 'actual')
    if carried:result['qualifications'].append('Carried provider observations retain original dates; authorized bounded carry in snapshot policy.')
    if spot['status'] in ('stale','stale_scenario','stale_or_delayed','unavailable_currency'):
        result['qualifications'].append('Spot freshness: '+spot['status'])
    return result


def currency_coverage(snapshot,reporting,currencies,start=None,end=None):
    """Compact pair-specific evidence; a global latest date is not coverage."""
    output=[];days=sorted(snapshot['daily_quotes']);calendar=snapshot['policy'].get('calendar',{})
    cutoff=_date(snapshot['as_of'][:10])
    expected=publication_dates(start or (days[0] if days else cutoff),min(_date(end or cutoff),cutoff),calendar)
    for currency in sorted(set(currencies)):
        observed=[];gaps=[];providers=set();definitions=set();series=set();types=set();latest_path=[]
        for day in expected:
            value,path=_cross_quotes(snapshot['daily_quotes'].get(day,[]),reporting,currency)
            genuine=value is not None and all(q['status']=='actual' and q['source_date']==day for q in path)
            if not genuine:gaps.append(day);continue
            observed.append(day);latest_path=path;meta=_route_metadata(path)
            providers.update(meta['providers']);definitions.update(meta['observation_definitions'])
            series.update(meta['selected_series']);types.add(meta['rate_type'])
        spot=spot_quote(snapshot,reporting,currency)
        transitions=[dict(base=t['base'],quote=t['quote'],provider=t['provider'],effective_date=t['from'],
            approval=t['overlap_evidence']['status']) for t in snapshot['policy'].get('cutovers',[])
            if any({q['base'],q['quote']}=={t['base'],t['quote']} for q in latest_path+spot.get('lineage',[]))]
        output.append(dict(currency=currency,reporting_currency=reporting,
            selected_providers=sorted(providers),selected_series=sorted(series),observation_definitions=sorted(definitions),
            rate_types=sorted(types),first_genuine_observation=observed[0] if observed else None,
            last_genuine_observation=observed[-1] if observed else None,spot_timestamp=spot.get('timestamp'),
            spot_source_dates=spot.get('source_dates',[spot.get('source_date')] if spot.get('source_date') else []),
            spot_status=spot['status'],spot_value=spot['value'],spot_rate_type=spot.get('rate_type'),
            spot_selected_series=spot.get('selected_series',[]),coverage_gap_count=len(gaps),
            coverage_first_gap=gaps[0] if gaps else None,coverage_last_gap=gaps[-1] if gaps else None,
            provider_transitions=transitions,provider_transition_rule='Only approved dated cutovers; unapproved broker daily series excluded',
            provider_overlap_audits=[a for a in snapshot.get('metadata',{}).get('provider_overlap_audits',[]) if currency in (a['base'],a['quote'])],
            available_native_series=[s for s in snapshot.get('metadata',{}).get('available_native_series',[]) if currency in (s['base'],s['quote'])],
            selected_rate_is_currency_proxy=False))
    return output


def exposure_quality(basket,reporting):
    """Keep uncertain exposure amounts visible; never renormalise missing mix."""
    weights=basket.get('weights',[]);threshold=basket.get('proxy_materiality_threshold',.05)
    proxy=weak=missing_evidence=0.;details=[]
    status_weights={status:0. for status in ('reported','estimated','scenario','unresolved')}
    weak_bases={'customer_geography_proxy','geographic_sales_proxy','production_geography_proxy',
                'factory_footprint_proxy','factory_count_proxy','unmeasured','unspecified'}
    for item in weights:
        status=item.get('exposure_status',basket.get('exposure_status','estimated'))
        neutral=item.get('scenario_neutral') and item['currency']=='UNALLOCATED'
        if neutral and (item.get('rate_currency')!=reporting or status!='unresolved'):
            raise ValueError('Unallocated neutral exposure must be unresolved and use the reporting-currency scenario')
        is_proxy=item.get('rate_currency',item['currency'])!=item['currency'] and not neutral
        if is_proxy and item.get('rate_type') in ('cross','cross_rate','triangulated'):
            raise ValueError('A different currency is a proxy, never a cross-rate')
        if is_proxy and not item.get('proxy_reason'):raise ValueError('A proxy rate requires a documented currency-specific reason')
        if status not in ('reported','estimated','scenario','unresolved'):raise ValueError('Invalid exposure evidence status')
        refs=item.get('source_refs',basket.get('source_refs',[]));effective=item.get('effective_date',basket.get('effective_date'))
        basis=item.get('evidence_basis',basket.get('evidence_basis','unspecified'))
        low=status in ('scenario','unresolved') or not refs or basis in weak_bases
        status_weights[status]+=item['weight']
        proxy+=item['weight'] if is_proxy else 0.;weak+=item['weight'] if low else 0.
        missing_evidence+=item['weight'] if not refs or not effective else 0.
        details.append(dict(currency=item['currency'],rate_currency=item.get('rate_currency',item['currency']),weight=item['weight'],
            exposure_status=status,evidence_basis=basis,effective_date=effective,source_refs=refs,rate_type='unresolved_neutral_scenario' if neutral else 'currency_proxy' if is_proxy else 'actual_currency',
            scenario_neutral=bool(item.get('scenario_neutral')),weak_evidence=low,label=item.get('label',item['currency'])))
    return dict(proxy_weight=proxy,weak_evidence_weight=weak,missing_source_or_effective_date_weight=missing_evidence,
        **{status+'_weight':weight for status,weight in status_weights.items()},
        materiality_threshold=threshold,material_weak_exposure=max(proxy,weak,missing_evidence)>=threshold,
        weights=details,historical_use='Current exposure version on historical rates is a diagnostic, not historically reported currency exposure')


def _shift_months(day,months):
    day=_date(day);year,month=divmod(day.year*12+day.month-1+months,12);month+=1
    return date(year,month,min(day.day,calendar.monthrange(year,month)[1]))


def transaction_hedge_terms(component,start,end,*,prior=False):
    """Resolve current and comparison purchase windows independently."""
    start,end=_date(start),_date(end)
    matches=[t for t in component.get('hedge_terms',[]) if _date(t['start'])<=end and _date(t['end'])>=start]
    if matches:
        if len(matches)!=1 or _date(matches[0]['start'])>start or _date(matches[0]['end'])<end:
            raise ValueError('Hedge terms must cover the complete purchase window without overlap; split the scenario period explicitly')
        terms=matches[0];coverage=terms.get('hedge_coverage',0.);hedge=terms.get('hedge_rate')
    else:
        prefix='prior_' if prior else '';coverage=component.get(prefix+'hedge_coverage',0.);hedge=component.get(prefix+'hedge_rate')
    if type(coverage) not in (int,float) or not math.isfinite(coverage) or not 0<=coverage<=1:
        raise ValueError('Hedge coverage must be between zero and one')
    if coverage and _number(hedge) is None:raise ValueError('Positive hedge coverage requires a positive functional/invoice hedge rate')
    return dict(hedge_coverage=coverage,hedge_rate=hedge)


def transaction_fx_for_period(snapshot,component,start,end,prior_start=None,prior_end=None):
    """Functional-currency costing route, before consolidation translation.

    Lag shifts the purchase-rate window back from consumption. An explicitly
    entered fixed hedge rate blends with market rates; no hedge is inferred.
    Current and comparable coverage/rates are independent. Undocumented prior
    hedging is an explicitly unhedged scenario, never inherited from today.
    """
    functional=component['functional_currency'];invoice=component['invoice_currency']
    lag=component.get('lag_months',0)
    if type(lag) is not int or lag<0:raise ValueError('Procurement lag must be nonnegative whole months')
    if functional==invoice:return dict(value=0.,status='identity',current_rate=1.,prior_rate=1.,functional_currency=functional,invoice_currency=invoice,lag_months=lag)
    start,end=_shift_months(start,-lag),_shift_months(end,-lag)
    previous_start=_shift_months(prior_start,-lag) if prior_start else _shift_months(start,-12)
    previous_end=_shift_months(prior_end,-lag) if prior_end else _shift_months(end,-12)
    current_terms=transaction_hedge_terms(component,start,end)
    prior_terms=transaction_hedge_terms(component,previous_start,previous_end,prior=True)
    coverage=current_terms['hedge_coverage'];hedge=current_terms['hedge_rate']
    prior_coverage=prior_terms['hedge_coverage'];prior_hedge=prior_terms['hedge_rate']
    current=forecast_rate(snapshot,functional,invoice,start,end);prior=forecast_rate(snapshot,functional,invoice,previous_start,previous_end)
    value=None
    if current['value'] is not None and prior['value'] is not None:
        effective=(1-coverage)*current['value']+coverage*(hedge or 0)
        baseline=component.get('baseline_rate')
        if baseline is None:baseline=(1-prior_coverage)*prior['value']+prior_coverage*(prior_hedge or 0)
        if _number(baseline) is None:raise ValueError('Transaction FX baseline rate must be positive')
        value=effective/baseline-1
    return dict(value=value,status='estimate' if value is not None else 'missing',current_rate=current['value'],prior_rate=prior['value'],
        functional_currency=functional,invoice_currency=invoice,lag_months=lag,hedge_coverage=coverage,hedge_rate=hedge,
        prior_hedge_coverage=prior_coverage,prior_hedge_rate=prior_hedge,
        current=current,prior=prior,costing_route='Functional-currency cost x transaction factor, then reporting translation once')


def _attach_cost_pool_fx(spec,data,snapshot,pools,add,formula,empty,rows):
    """Independent translation baskets and procurement FX for causal cost pools."""
    periods=spec['periods'];pmap={p['id']:p for p in periods};reporting=data['currency']
    by_kind={(p['fiscal_year'],p['kind']):p['id'] for p in periods}
    frequency=spec.get('frequency_policy',{}).get('statements','half_year')
    annual=set(spec.get('frequency_policy',{}).get('annual_only_years',[]))
    for entity in [s['id'] for s in data['segments']]+['corporate']:
        for pool in pools:
            if entity=='corporate' and entity not in pool.get('fx_by_entity',{}):continue
            prefix=f'{entity}_pool_{pool["id"]}'
            cfg=pool.get('fx_by_entity',{}).get(entity,pool.get('fx',{}))
            translation=cfg.get('translation',{});transaction=cfg.get('transaction',{})
            active=[p for p in periods if p['fiscal_year'] in spec['forecast_years'] and
                (p['kind']=='FY' if p['fiscal_year'] in annual else p['kind'] in (('H1','H2') if frequency=='half_year' else ('Q1','Q2','Q3','Q4'))) and
                rows['utop']['cells'][p['id']].get('status') not in ('actual','reported','derived_actual')]
            if not active:continue
            anchor=by_kind[active[0]['fiscal_year'],'FY']
            trans=add(prefix+'_translation_fx',entity+' '+pool['id']+' translation FX',fx_only=True)
            txn=add(prefix+'_transaction_fx',entity+' '+pool['id']+' transaction FX',fx_only=True)
            trans.update(cost_pool=pool['id'],entity=entity,currency_basis='functional_to_reporting')
            txn.update(cost_pool=pool['id'],entity=entity,currency_basis='invoice_to_functional',
                costing_route='Functional-currency costs; apply transaction effect once, then translation once')
            weights=translation.get('weights',[])
            neutral_unresolved=bool(weights) and all(w.get('scenario_neutral') and
                w.get('exposure_status',translation.get('exposure_status')) in ('unresolved','scenario') and
                w.get('rate_currency',w['currency'])==reporting for w in weights if w['weight']>0)
            if weights:
                if not math.isclose(sum(w['weight'] for w in weights),1,abs_tol=1e-9) or any(w['weight']<0 for w in weights):
                    raise ValueError('Cost pool translation weights must sum to one without renormalisation')
                if reporting not in {w['currency'] for w in weights}:raise ValueError('Cost basket must include reporting currency explicitly')
                quality=exposure_quality(translation,reporting)
                spec['fx_exposure_quality'].append(dict(entity=entity,role='cost_pool',cost_pool=pool['id'],**quality))
            if weights and not neutral_unresolved:
                controls=[]
                for i,w in enumerate(weights):
                    control=add(prefix+'_translation_weight_'+w['currency'].lower(),entity+' '+pool['id']+' functional currency: '+w.get('label',w['currency']),
                        kind='assumption',fx_only=True,note=translation.get('method',''))
                    control.update(effective_date=w.get('effective_date',translation.get('effective_date')),
                        exposure_status=w.get('exposure_status','estimated'))
                    control['cells'][anchor]=dict(value=w['weight'],status='estimate',editable=True,
                        source_refs=w.get('source_refs',translation.get('source_refs',[])),reason=translation.get('method',''))
                    controls.append((control,w))
                total=add(prefix+'_translation_weight_total',entity+' '+pool['id']+' currency mix total',fx_only=True)
                factor=add(prefix+'_translation_multiplier',entity+' '+pool['id']+' translation multiplier',kind='assumption',fx_only=True)
                factor.update(display_unit='x',number_format='0.00"x"')
                factor['cells'][anchor]=dict(value=translation.get('multiplier',1.),status='estimate',editable=True,reason='Uncalibrated multiplier; does not remedy missing exposure evidence')
                for p in active:
                    pid=p['id'];bindings={};terms=[]
                    for i,(control,w) in enumerate(controls):
                        bindings[f'w{i}']=dict(row_id=control['id'],period_id=anchor)
                        bindings[f'r{i}']=dict(row_id='fx_yoy_'+w.get('rate_currency',w['currency']).lower(),period_id=pid)
                        terms.append(f'w{i} * r{i}')
                    total['cells'][pid]=formula(' + '.join(f'w{i}' for i in range(len(controls))),
                        {k:v for k,v in bindings.items() if k.startswith('w')})
                    bindings['factor']=dict(row_id=factor['id'],period_id=anchor)
                    trans['cells'][pid]=formula('('+' + '.join(terms)+') * factor',bindings,
                        reason='Functional-currency pool cost translated once using its own explicit currency mix')
                    spec['identities'].append(dict(id=prefix+'_fx_weights_'+pid,label='Cost pool currency weights sum to one',expression=total['id']+' - 1',periods=[pid],tolerance=1e-9,independent=False))
                if quality['material_weak_exposure']:
                    spec['fx_validation'].append(dict(area='fx_exposure',name=prefix,status='WARNING',
                        detail=f"Cost-pool exposure remains provisional: weak {quality['weak_evidence_weight']:.1%}, proxy {quality['proxy_weight']:.1%}; no revenue-basket inheritance."))
            else:
                trans['kind']='assumption'
                trans['sensitivity_range']=copy.deepcopy(translation.get('sensitivity_range',{}))
                trans['note']='Explicit unmodelled translation scenario. The reporting-currency neutral placeholder is not a measured functional-currency mix.'
                for p in active:trans['cells'][p['id']]=dict(value=translation.get('scenario_effect',0.),status='estimate',editable=True,
                    source_refs=translation.get('source_refs',[]),
                    reason='Independent editable translation stress; zero effect is not evidence of zero exposure. '+str(trans['sensitivity_range']))
                spec['fx_validation'].append(dict(area='fx_exposure',name=prefix+'_translation',status='WARNING',detail='No supported pool functional-currency map; independent provisional zero translation scenario'))
            components=transaction.get('components',[])
            if pool.get('input_cost_includes_transaction_fx',False):
                if any(c.get('weight',0) and c.get('functional_currency')!=c.get('invoice_currency') for c in components) or transaction.get('scenario_effect',0):
                    raise ValueError('Input-cost forecast already includes transaction FX; separate transaction exposure must be disabled: '+prefix)
                txn['note']='Transaction FX is included in the unit-cost input; no second transaction adjustment is applied.'
                txn['input_cost_includes_transaction_fx']=True
                for p in active:txn['cells'][p['id']]=dict(value=0.,status='forecast',editable=False,reason=txn['note'])
                continue
            if any(not isinstance(c.get('weight'),(int,float)) or not math.isfinite(c['weight']) or c['weight']<0 for c in components) or sum(c['weight'] for c in components)>1+1e-9:
                raise ValueError('Transaction exposure fractions must be finite, nonnegative and at most one')
            if weights and not any(w.get('exposure_status')=='unresolved' for w in weights):
                for currency in {c['functional_currency'] for c in components}:
                    share=sum(c['weight'] for c in components if c['functional_currency']==currency)
                    functional=sum(w['weight'] for w in weights if w['currency']==currency)
                    if share>functional+1e-9:raise ValueError('Transaction component exceeds its functional-currency share of the cost pool')
            if not components:
                txn['kind']='assumption'
                for p in active:txn['cells'][p['id']]=dict(value=0.,status='estimate',editable=True,
                    reason='Explicit unmodelled procurement FX scenario; no invoice currency, hedge or lag inferred')
                spec['fx_validation'].append(dict(area='transaction_fx',name=prefix,status='WARNING',detail='Transaction currency exposures are unmodelled; separate zero scenario input, not inferred hedging'))
                continue
            component_rows=[]
            for i,component in enumerate(components):
                part=prefix+'_transaction_'+str(i+1);effect=add(part+'_effect',part+' effect',fx_only=True)
                weight=add(part+'_weight',part+' fraction of pool',kind='assumption',fx_only=True)
                weight['cells'][anchor]=dict(value=component['weight'],status='estimate',editable=True,
                    source_refs=component.get('source_refs',transaction.get('source_refs',[])),reason=transaction.get('method','Explicit transaction exposure scenario'))
                component_rows.append((effect,weight))
                if component['functional_currency']==component['invoice_currency']:
                    for p in active:effect['cells'][p['id']]=dict(value=0.,status='derived_actual',editable=False,reason='Invoice and functional currency are identical; no transaction FX')
                    continue
                current=add(part+'_market_current',part+' purchase-period market rate',unit='exchange_rate',fx_only=True)
                prior=add(part+'_market_prior',part+' comparable purchase-period market rate',unit='exchange_rate',fx_only=True)
                cover=add(part+'_hedge_coverage',part+' explicitly assumed hedge coverage',kind='assumption',fx_only=True)
                hedge=add(part+'_hedge_rate',part+' functional units per invoice unit hedge rate',unit='exchange_rate',kind='assumption',fx_only=True)
                prior_cover=add(part+'_prior_hedge_coverage',part+' comparable-period hedge coverage',kind='assumption',fx_only=True)
                prior_hedge=add(part+'_prior_hedge_rate',part+' comparable-period functional per invoice hedge rate',unit='exchange_rate',kind='assumption',fx_only=True)
                hedged_current=add(part+'_hedged_current',part+' current blended rate',unit='exchange_rate',fx_only=True)
                hedged_prior=add(part+'_hedged_prior',part+' comparable blended rate',unit='exchange_rate',fx_only=True)
                effective_current=add(part+'_effective_current',part+' current effective rate',unit='exchange_rate',fx_only=True)
                effective_prior=add(part+'_effective_prior',part+' prior effective rate',unit='exchange_rate',fx_only=True)
                fixed_baseline=component.get('baseline_rate') is not None
                if fixed_baseline:effective_prior['kind']='assumption'
                for row in (effect,current,prior,cover,hedge,prior_cover,prior_hedge,weight,effective_prior):row.update(lag_months=component.get('lag_months',0),functional_currency=component['functional_currency'],invoice_currency=component['invoice_currency'])
                for p in active:
                    pid=p['id'];previous=pmap.get(by_kind.get((p['fiscal_year']-1,p['kind'])))
                    result=transaction_fx_for_period(snapshot,component,p['start'],p['end'],previous['start'] if previous else None,previous['end'] if previous else None)
                    for row,key in ((current,'current_rate'),(prior,'prior_rate')):
                        row['cells'][pid]=dict(value=result[key],status='estimate',editable=False,required=False,source_refs=[snapshot['source']],reason='Market invoice-to-functional average; purchase window shifted by explicit lag')
                    for terms_prefix,coverage_row,rate_row in (('',cover,hedge),('prior_',prior_cover,prior_hedge)):
                        coverage_row['cells'][pid]=dict(value=result[terms_prefix+'hedge_coverage'],status='estimate',editable=not (terms_prefix and fixed_baseline),
                            reason='Independent '+('comparable-period' if terms_prefix else 'current-period')+' hedge scenario; zero is provisional unhedged, not a company hedge disclosure',source_refs=component.get('source_refs',[]))
                        rate_row['cells'][pid]=dict(value=result[terms_prefix+'hedge_rate'],status='estimate',editable=not (terms_prefix and fixed_baseline),required=False,
                            reason='Missing is permitted only while this period hedge coverage is zero; never a fabricated execution rate')
                    for market,blended,coverage_row,rate_row in ((current,hedged_current,cover,hedge),(prior,hedged_prior,prior_cover,prior_hedge)):
                        blended['cells'][pid]=formula(f'(1 - {coverage_row["id"]}) * {market["id"]} + {coverage_row["id"]} * {rate_row["id"]}')
                        blended['cells'][pid]['required']=False
                    effective_current['cells'][pid]=formula(f'if_zero({cover["id"]}, {current["id"]}, {hedged_current["id"]})')
                    effective_prior['cells'][pid]=dict(value=component['baseline_rate'],status='estimate',editable=True,reason='Explicit effective historical procurement rate; replaces the dormant comparable hedge-term controls') if fixed_baseline else formula(f'if_zero({prior_cover["id"]}, {prior["id"]}, {hedged_prior["id"]})')
                    effect['cells'][pid]=formula(f'{effective_current["id"]} / {effective_prior["id"]} - 1')
            transaction_total=add(prefix+'_transaction_weight_total',prefix+' total transaction exposure fraction',fx_only=True)
            for p in active:
                bindings={};terms=[]
                for i,(effect,weight) in enumerate(component_rows):
                    bindings[f'w{i}']=dict(row_id=weight['id'],period_id=anchor);bindings[f'x{i}']=dict(row_id=effect['id'],period_id=p['id']);terms.append(f'w{i} * x{i}')
                txn['cells'][p['id']]=formula(' + '.join(terms),bindings,reason='Invoice-to-functional transaction changes, applied once before pool translation')
                transaction_total['cells'][p['id']]=formula(' + '.join(f'w{i}' for i in range(len(component_rows))),
                    {k:v for k,v in bindings.items() if k.startswith('w')},reason='Remaining fraction is outside the configured transaction exposures; no renormalisation')


def attach_company_fx(spec,data,snapshot):
    """Compile explicit baskets into the shared financial expression graph.

    Weights are company configuration, not geographic guesses in the engine.
    Revenue growth contributions use the declared bridge basis. Costs translate
    the constant-currency cost base once, at the statement calculation cadence.
    """
    config=data.get('model_configuration',{});exposures=config.get('fx_exposures',{})
    from .integrated import financial_cost_pools
    pools=financial_cost_pools(data)
    periods=spec['periods'];pmap={p['id']:p for p in periods};by_kind={(p['fiscal_year'],p['kind']):p['id'] for p in periods}
    rows={r['id']:r for r in spec['rows']};reporting=data['currency'];base=max(spec['historical_years'])
    spec['sections'].append(dict(id='fx',label='Foreign exchange — shared rates and company exposure'))
    spec.update(fx_source=copy.deepcopy(snapshot['source']),fx_metadata=copy.deepcopy(snapshot['metadata']),fx_observations=[],fx_validation=[])
    def empty(reason='No applicable exposure control at this frequency'):
        return dict(value=None,status='not_applicable',editable=False,reason=reason)
    def add(rid,label,section='fx',unit='ratio',kind='calculated',fx_only=False,note=''):
        row=dict(id=rid,label=label,section=section,unit=unit,kind=kind,metric_type='ratio' if unit in ('ratio','exchange_rate') else 'flow',
            cells={p['id']:empty() for p in periods},historical={},source_refs={},fx_only=fx_only,note=note,optional=fx_only,
            assumption_support=unit in ('ratio','exchange_rate'),evidence_role='shared_factor_diagnostic')
        rows[rid]=row;spec['rows'].append(row);return row
    def formula(expr,bindings=None,editable=False,reason='Company basket applied to a frozen shared rate snapshot'):
        return dict(expression=expr,bindings=bindings or {},status='forecast',editable=editable,reason=reason)
    currencies={reporting}
    for entity,exposure in exposures.items():
        for role in ('revenue','cost'):
            basket=exposure.get(role,{})
            for item in basket.get('weights',[]):currencies.add(item.get('rate_currency',item['currency']))
    for pool in pools:
        for pool_fx in [pool.get('fx',{}),*pool.get('fx_by_entity',{}).values()]:
            for item in pool_fx.get('translation',{}).get('weights',[]):currencies.add(item.get('rate_currency',item['currency']))
            for item in pool_fx.get('transaction',{}).get('components',[]):
                currencies.update((item['functional_currency'],item['invoice_currency']))
    accounting=data.get('financial_contracts',config.get('financial_contracts',{})).get('fx_accounting',{})
    currencies.update(item['currency'] for item in accounting.get('exposures',[]))
    currencies.update(config.get('fx_currencies',[]))
    spec['fx_currency_coverage']=currency_coverage(snapshot,reporting,currencies,
        min(p['start'] for p in periods),max(p['end'] for p in periods))
    spec['fx_exposure_quality']=[]
    for currency in sorted(currencies):
        code=currency.lower();avg=add('fx_average_'+code,f'{reporting} per {currency} — average',unit='exchange_rate',fx_only=True)
        close=add('fx_closing_'+code,f'{reporting} per {currency} — closing',unit='exchange_rate',fx_only=True)
        change=add('fx_yoy_'+code,f'{currency} translation change versus same prior-year period',fx_only=True)
        for p in periods:
            pid=p['id']
            for kind,row in (('average',avg),('closing',close)):
                rate=forecast_rate(snapshot,reporting,currency,p['start'],p['end'],kind)
                spec['fx_observations'].append(dict(rate,row_id=row['id'],period_id=pid))
                row['cells'][pid]=dict(value=rate['value'],status=rate['status'] if rate['value'] is not None else 'not_disclosed',
                    editable=False,required=False,reason='; '.join(rate['qualifications']),source_refs=rate['provenance'])
            previous=by_kind.get((p['fiscal_year']-1,p['kind']))
            if previous:
                change['cells'][pid]=formula('current / prior - 1',{
                    'current':dict(row_id=avg['id'],period_id=pid),'prior':dict(row_id=avg['id'],period_id=previous)},
                    reason='Average of daily reporting/local crosses compared with the same prior-year fiscal period')
                rate_states={avg['cells'][pid]['status'],avg['cells'][previous]['status']}
                change['cells'][pid]['status']='derived_actual' if rate_states=={'actual'} else 'estimate'
    for seg in data['segments']:
        key=seg['id'];exposure=exposures.get(key)
        if not exposure:
            spec['fx_validation'].append(dict(area='fx_exposure',name=key,status='MISSING',detail='No company currency baskets; no neutral exposure is inferred'))
            for suffix in ('_fx','_cost_fx'):
                for p in periods:
                    cell=rows[key+suffix]['cells'][p['id']]
                    if cell.get('status') in ('forecast','estimate'):
                        rows[key+suffix]['cells'][p['id']]=dict(value=None,status='not_disclosed',editable=False,reason='Required company exposure is missing')
            continue
        bridge=exposure.get('revenue_bridge')
        if bridge not in ('translation_of_constant_currency','additive_reported_growth_contribution'):
            raise ValueError('Each basket requires an explicit revenue_bridge convention: '+key)
        for role,target_suffix in (('revenue','fx'),('cost','cost_fx')):
            if role=='cost' and pools:continue  # Pool-specific factors own causal costs.
            basket=exposure.get(role,{})
            weights=basket.get('weights',[])
            if not weights or not basket.get('method') or not basket.get('source_refs'):
                raise ValueError(f'{key} {role} basket needs weights, method and source evidence')
            if len({item['currency'] for item in weights})!=len(weights):raise ValueError('Duplicate basket currency')
            if any(not isinstance(item.get('weight'),(float,int)) or not math.isfinite(item['weight']) or item['weight']<0 for item in weights):
                raise ValueError('Currency weights must be finite and nonnegative')
            if not math.isclose(sum(i['weight'] for i in weights),1,abs_tol=1e-9):raise ValueError(f'{key} {role} weights must sum to one, including reporting currency and explicit proxies')
            if reporting not in {i['currency'] for i in weights}:raise ValueError('Basket must explicitly include reporting currency, even at zero weight')
            factor=basket.get('multiplier',1.0)
            if type(factor) not in (int,float) or not math.isfinite(factor) or factor<0:raise ValueError('FX calibration multiplier must be finite and nonnegative')
            for item in weights:
                if item.get('rate_currency',item['currency'])!=item['currency'] and not item.get('proxy_reason') and not (item.get('scenario_neutral') and item['currency']=='UNALLOCATED'):
                    raise ValueError('A proxy rate requires a documented currency-specific reason')
            quality=exposure_quality(basket,reporting)
            spec['fx_exposure_quality'].append(dict(entity=key,role=role,**quality))
            rid=f'{key}_{role}_fx';target=rows[key+'_'+target_suffix]
            target['label']='Applied '+role+' FX '+('growth contribution' if role=='revenue' else 'translation factor')
            raw=add(rid+'_raw',seg['label']+' '+role+' FX — raw basket','segment_'+key,note=basket['method'])
            multiplier=add(rid+'_multiplier',role.capitalize()+' FX adjustment multiplier','segment_'+key,kind='assumption',
                note='Provisional 1.0 unless separately evidenced; never fitted by division through a near-zero historical impact.')
            multiplier.update(display_unit='x',number_format='0.00"x";(0.00"x");0.00"x"')
            adjusted=add(rid+'_adjusted',role.capitalize()+' FX — adjusted translation factor','segment_'+key)
            monetary=add(rid+'_amount',role.capitalize()+' FX — monetary impact','segment_'+key,unit=reporting+'m')
            residual=add(rid+'_residual',seg['label']+' reported FX less adjusted estimate',fx_only=True)
            weight_rows=[]
            forecast_cells=[p['id'] for p in periods if target['cells'][p['id']].get('editable') and target['cells'][p['id']].get('status') in ('forecast','estimate')]
            if role=='revenue':
                reported_sales=[pid for pid in forecast_cells if rows[key+'_sales']['cells'][pid].get('status') in ('actual','reported','derived_actual')]
                for pid in reported_sales:
                    target['cells'][pid]=empty('Sales are reported; undisclosed historical FX is not replaced by a forecast contribution. Raw basket remains a separate diagnostic.')
                    target['cells'][pid]['status']='not_disclosed'
                # A missing reported FX component leaves these growth-bridge
                # diagnostics unavailable. It must not invalidate independently
                # reported sales or invent a historical translation contribution.
                diagnostics={key+suffix for suffix in ('_fx','_growth','_fx_effect')}
                diagnostics|={rid+suffix for rid in tuple(diagnostics) for suffix in ('_weighted_amount','_comparable_revenue') if rid+suffix in rows}
                affected={(target['id'],pid) for pid in reported_sales}
                changed=True
                while changed:
                    changed=False
                    for row_id in sorted(diagnostics):
                        if row_id not in rows:continue
                        for pid,cell in rows[row_id]['cells'].items():
                            if not cell.get('expression') or (row_id,pid) in affected:continue
                            from .periods import cell_dependencies
                            deps=cell_dependencies(cell,pid)
                            if any((d['row_id'],d['period_id']) in affected for d in deps):
                                cell['required']=False
                                cell['reason']='Reported historical FX component is undisclosed; this bridge diagnostic is unavailable, while reported sales remain unchanged.'
                                affected.add((row_id,pid));changed=True
                forecast_cells=[pid for pid in forecast_cells if pid not in reported_sales]
            if not forecast_cells:continue
            anchor=forecast_cells[0]
            weight_anchor=by_kind[pmap[anchor]['fiscal_year'],'FY']
            for item in weights:
                proxy_label=' (neutral unresolved scenario)' if item.get('scenario_neutral') and item['currency']=='UNALLOCATED' else ' (proxy '+item['rate_currency']+')' if item.get('rate_currency',item['currency'])!=item['currency'] else ''
                w=add(rid+'_weight_'+item['currency'].lower(),seg['label']+' '+role+' exposure: '+item.get('label',item['currency'])+proxy_label,kind='assumption',fx_only=True,
                    note=basket['method']+' '+item.get('proxy_reason','')+' '+basket.get('limitations','')+
                        ' Exposure evidence: '+item.get('exposure_status','estimated')+'; effective '+str(item.get('effective_date',basket.get('effective_date','undated'))))
                w['effective_date']=item.get('effective_date',basket.get('effective_date'))
                weight_rows.append((w,item))
                # One editable weight per basket/version. Every period links the
                # same authority, preventing unnoticed per-quarter weight drift.
                w['cells'][weight_anchor]=dict(value=item['weight'],status='estimate',editable=True,source_refs=basket['source_refs'],reason=w['note']+' One exposure version applies to every calculation period.')
            total=add(rid+'_weight_total',seg['label']+' '+role+' basket sum',fx_only=True)
            total['check_tolerance']=1e-9
            for p in periods:
                pid=p['id'];terms=[];bindings={}
                for i,(w,item) in enumerate(weight_rows):
                    bindings[f'w{i}']=dict(row_id=w['id'],period_id=weight_anchor)
                    bindings[f'f{i}']=dict(row_id='fx_yoy_'+item.get('rate_currency',item['currency']).lower(),period_id=pid)
                    # Retain dependencies for editable zero weights as well.
                    terms.append(f'w{i} * f{i}')
                total['cells'][pid]=formula(' + '.join(f'w{i}' for i in range(len(weight_rows))),{k:v for k,v in bindings.items() if k.startswith('w')})
                raw['cells'][pid]=formula(' + '.join(terms) if terms else '0',bindings)
                raw['cells'][pid]['required']=False
                if p['fiscal_year']<=base or target['cells'][pid].get('status') in ('actual','reported','derived_actual'):
                    raw['cells'][pid]['status']='estimate'
                    raw['cells'][pid]['reason']='Current estimated exposure on historical market rates; diagnostic only, not a historically reported exposure or validated backtest'
                historical=target['cells'][pid].get('status') in ('actual','reported','derived_actual')
                active=pid in forecast_cells
                if active:
                    multiplier['cells'][pid]=dict(value=basket.get('multiplier',1.0),status='estimate',editable=True,
                        reason=basket.get('calibration',{}).get('method','Uncalibrated provisional multiplier 1.0'),source_refs=basket.get('calibration',{}).get('source_refs',[]))
                    adjusted['cells'][pid]=formula('raw * factor',{'raw':dict(row_id=raw['id'],period_id=pid),'factor':dict(row_id=multiplier['id'],period_id=pid)})
                    if role=='revenue':
                        expr=f'(1 + {key}_pricing + {key}_rig) * '+adjusted['id'] if bridge=='translation_of_constant_currency' else adjusted['id']
                        target['cells'][pid]=formula(expr,editable=True,
                            reason='FX percentage-point contribution on prior-year reported sales; translates the constant-currency growth base once' if bridge=='translation_of_constant_currency' else 'Declared additive growth bridge; adjusted basket is a contribution on prior-year reported sales')
                        monetary['cells'][pid]=formula(f'{key}_sales_seed * {key}_fx')
                    else:
                        target['cells'][pid]=formula(adjusted['id'],editable=True,reason='Cost translation factor applied once after activity and local cost inflation')
                        if not pools:
                            rows[key+'_cash_cost']['cells'][pid]=formula(f'{key}_cash_cost_seed * (1 + {key}_rig) * (1 + {key}_cost_inflation) * (1 + {key}_cost_fx)')
                        monetary['cells'][pid]=formula(f'{key}_cash_cost_seed * (1 + {key}_rig) * (1 + {key}_cost_inflation) * {key}_cost_fx')
                elif historical and role=='revenue' and p['fiscal_year']>=exposure.get('comparison_start_year',base) and (exposure.get('comparison_periods') is None or pid in exposure['comparison_periods']):
                    multiplier['cells'][pid]=formula('factor',{'factor':dict(row_id=multiplier['id'],period_id=anchor)},reason='Current provisional calibration applied for a comparable-basis diagnostic; reported FX is preserved')
                    adjusted['cells'][pid]=formula('raw * factor',{'raw':dict(row_id=raw['id'],period_id=pid),'factor':dict(row_id=multiplier['id'],period_id=anchor)})
                    comparison=f'(1 + {key}_pricing + {key}_rig) * '+adjusted['id'] if bridge=='translation_of_constant_currency' else adjusted['id']
                    residual['cells'][pid]=formula(target['id']+' - ('+comparison+')',reason='Reported additive FX impact minus basket-implied contribution; not an exact-match calibration')
                    residual['cells'][pid]['status']='estimate'
                    adjusted['cells'][pid]['required']=False
                if role=='cost' and p['kind'].startswith('Q') and spec.get('frequency_policy',{}).get('statements')=='half_year':
                    for row in (raw,multiplier,adjusted,monetary):row['cells'][pid]=empty('Cost FX is calculated directly at the company statement frequency')
            for pid in forecast_cells:
                spec['identities'].append(dict(id=rid+'_weights_'+pid,label=seg['label']+' '+role+' currency weights sum to one',
                    expression=total['id']+' - 1',periods=[pid],tolerance=1e-9,independent=False))
            spec['fx_validation'].append(dict(area='fx_exposure',name=rid,status='WARNING' if basket.get('status')=='provisional' else 'PASS',
                detail=basket['method']+' '+basket.get('limitations','')+
                f" Proxy currency weight {quality['proxy_weight']:.1%}; weak evidence {quality['weak_evidence_weight']:.1%}; missing date/source {quality['missing_source_or_effective_date_weight']:.1%}."))
    # Amount summaries sum compatible modeled periods, never percentage rates.
    for seg in data['segments']:
        for role in ('revenue','cost'):
            row=rows.get(seg['id']+'_'+role+'_fx_amount')
            if not row:continue
            for year in spec['forecast_years']:
                for kind,parts in (('H1',('Q1','Q2')),('H2',('Q3','Q4')),('FY',('H1','H2'))):
                    if (year,kind) not in by_kind or any((year,k) not in by_kind for k in parts):continue
                    pid=by_kind[year,kind];pids=[by_kind[year,k] for k in parts]
                    if row['cells'][pid].get('expression'):continue
                    if all(row['cells'][q].get('expression') for q in pids):
                        row['cells'][pid]=formula('left + right',{'left':dict(row_id=row['id'],period_id=pids[0]),'right':dict(row_id=row['id'],period_id=pids[1])})
                        if role=='cost' and not pools:
                            # A financial impact factor is weighted by the
                            # actual cost base, never by half-year duration.
                            target=rows[seg['id']+'_cost_fx']
                            target['cells'][pid]=formula(f'{row["id"]} / ({seg["id"]}_cash_cost - {row["id"]})',
                                reason='Sum monetary translation effects divided by sum pre-translation cost bases')
                            adjusted=rows[seg['id']+'_cost_fx_adjusted'];adjusted['cells'][pid]=formula(target['id'])
    # Group FX sums monetary effects; it never averages percentage impacts across
    # differently sized regions. Existing company historical FX cells are intact.
    for role in ('revenue','cost'):
        group=add('group_'+role+'_fx_amount','Group '+role+' FX impact',section='corporate',unit=reporting+'m')
        for p in periods:
            keys=[seg['id']+'_'+role+'_fx_amount' for seg in data['segments']]
            if all(k in rows and rows[k]['cells'][p['id']].get('expression') for k in keys):
                group['cells'][p['id']]=formula(' + '.join(keys))
    if pools:
        _attach_cost_pool_fx(spec,data,snapshot,pools,add,formula,empty,rows)
    for row in rows.values():
        if row.get('evidence_role')!='shared_factor_diagnostic' or not any(c.get('editable') for c in row['cells'].values()):continue
        entity=next((s['id'] for s in data['segments'] if row['id'].startswith(s['id']+'_')),'corporate' if row['id'].startswith('corporate_') else 'group')
        pool=next((p['id'] for p in pools if '_pool_'+p['id']+'_' in row['id']),None)
        outputs=[entity+'_pool_'+pool] if pool else [entity+'_sales'] if '_revenue_' in row['id'] else [entity+'_cash_cost']
        row['driver_contract']=dict(id=row['id'],definition=row['label'],units=row['unit'],financial_basis=data.get('forecast_basis'),
            frequency='effective-dated exposure' if 'weight' in row['id'] else spec.get('frequency_policy',{}).get('statements','configured'),
            forecast_horizon=spec['forecast_years'],entity_id=entity,cost_pool=pool,
            currency_basis='invoice_to_functional' if '_transaction_' in row['id'] else 'functional_to_reporting',
            timing=dict(effective_date=row.get('effective_date'),lag_months=row.get('lag_months',0)),
            parameter_role='exposure_assumption' if 'weight' in row['id'] else 'economic_input',expected_downstream_outputs=outputs)
    spec['fx_interfaces']=dict(reporting_currency=reporting,snapshot=snapshot['source']['snapshot_id'],exposures=copy.deepcopy(exposures),
        revenue_translation='Regional basket, multiplier, declared growth bridge, then sum monetary impacts',
        cost_translation='Separate explicit functional-currency mix for each causal cost pool; no automatic revenue-basket inheritance; financial amounts aggregate before ratios',
        transaction_fx='Functional-currency costing route: invoice-to-functional purchase rate changes at explicit lag/hedge terms, then translation once. Unmodelled exposures use separate labelled zero scenarios.',
        historical='Reported company FX observations are immutable; raw/adjusted/residual diagnostics are separate',
        balance_translation='Closing market rates support separately configured balance-sheet exposures; no currency map or OCI adjustment is inferred from cash-flow FX')
    spec['policies']['fx']='Deterministic shared spot-held rates; company input baskets and editable adjustment multipliers. No second translation of reported history.'


def apply_fx_accounting_contracts(spec,data,snapshot):
    """Derive statement FX from dated exposures, with no balancing inputs.

    This deliberately bounded interface covers opening foreign-operation
    balances, explicitly supplied native-currency movements, and reporting-
    functional-currency cash/debt monetary positions. It does not infer legal
    entity functional currencies or a treasury hedge book from bond currency.
    """
    cfg=data.get('financial_contracts',data.get('model_configuration',{}).get('financial_contracts',{}))
    policy=cfg.get('fx_accounting')
    if policy is None:return spec
    if not snapshot:
        if policy.get('exposures'):raise ValueError('Exposure-derived FX accounting requires a saved FX snapshot; configured company exposure is distinct from missing market data')
        spec['fx_accounting_coverage']=dict(status='incomplete',market_snapshot='missing',company_exposures='unquantified')
        spec.setdefault('fx_validation',[]).append(dict(area='fx_accounting',name='empty_exposure_without_market_snapshot',status='WARNING',
            detail='Company accounting exposures are unquantified and no market snapshot was supplied. Existing neutral scenarios do not establish FX-accounting coverage.'))
        return spec
    legacy=cfg.get('balance_sheet_fx',{})
    if any(legacy.get(k,0) for k in ('annual_cash_translation_oci','annual_financing_remeasurement')) or any(legacy.get('annual_movements',{}).values()):
        raise ValueError('Exposure-derived FX accounting conflicts with legacy independent FX amounts; migrate the explicit override to an exposure scenario')
    exposures=copy.deepcopy(policy.get('exposures',[]));periods=spec['periods']
    rows={r['id']:r for r in spec['rows']};py={(p['fiscal_year'],p['kind']):p['id'] for p in periods}
    reporting=data['currency'];money=reporting+'m'
    assets={'cash','inventories','receivables','owned_ppe','right_of_use_assets','intangibles','goodwill'}
    liabilities={'trade_payables','nonlease_debt','lease_liabilities'}
    frequency=spec.get('frequency_policy',{}).get('statements','half_year')
    regular=('H1','H2') if frequency=='half_year' else ('Q1','Q2','Q3','Q4')
    def actual(c):return c.get('status') in ('actual','reported','derived_actual')
    def populated(c):return c.get('value') is not None or bool(c.get('expression'))
    def leaves(year):
        kinds=regular if any((year,k) in py for k in regular) else ('FY',)
        return [p for p in periods if p['fiscal_year']==year and p['kind'] in kinds]
    active=[p for year in spec['forecast_years'] for p in leaves(year)
        if populated(rows['utop']['cells'][p['id']]) and not actual(rows['utop']['cells'][p['id']])]
    if not active:return spec
    def add(rid,label,unit=money,kind='calculated'):
        if rid not in rows:
            row=dict(id=rid,label=label,section='fx',unit=unit,kind=kind,metric_type='ratio' if unit in ('ratio','exchange_rate') else 'flow',
                historical={},source_refs={},cells={p['id']:dict(value=None,status='not_applicable',editable=False) for p in periods},
                fx_only=True,assumption_support=True,support_matrix=True,optional=True,evidence_role='shared_factor_diagnostic')
            rows[rid]=row;spec['rows'].append(row)
        else:rows[rid]['kind']=kind
        return rows[rid]
    def setf(rid,pid,expr,bindings=None,reason='Derived from exposure and shared average/closing rates'):
        rows[rid]['cells'][pid]=dict(expression=expr,bindings=bindings or {},status='forecast',editable=False,reason=reason)
    def sumf(rid,pid,terms):
        if terms:setf(rid,pid,' + '.join(terms))
        else:rows[rid]['cells'][pid]=dict(value=0.,status='forecast',editable=False,reason='No configured exposure in this accounting category; coverage limitation is explicit')
    def aggregate(rid):
        for year in spec['forecast_years']:
            for p in periods:
                if p['fiscal_year']!=year or actual(rows[rid]['cells'][p['id']]):continue
                members=[q for q in leaves(year) if q['id']!=p['id'] and q['start']>=p['start'] and q['end']<=p['end']]
                if len(members)>1 and members[0]['start']==p['start'] and members[-1]['end']==p['end']:
                    bindings={f'x{i}':dict(row_id=rid,period_id=q['id']) for i,q in enumerate(members)}
                    setf(rid,p['id'],' + '.join(bindings),bindings,'Sum of separate financial-period FX amounts; no duration averaging')
                    if any(not populated(rows[rid]['cells'][q['id']]) for q in members):rows[rid]['cells'][p['id']]['required']=False
    for rid in ['cash_fx','cash_translation_oci','financing_fx_remeasurement','oci_translation']:
        add(rid,rows[rid]['label'])
    add('cash_fx_remeasurement','Cash monetary remeasurement gain / (loss)')
    role_targets={role:role+'_translation_movement' for role in (assets|liabilities)-{'cash'}}
    for role,rid in role_targets.items():add(rid,rows[rid]['label'])
    targets={p['id']:{r:[] for r in [*role_targets.values(),'cash_translation_oci','cash_fx_remeasurement','financing_fx_remeasurement']} for p in active}
    contracts=[];seen=set();share_totals={}
    for exposure in exposures:
        eid=exposure['id'];role=exposure['role'];kind=exposure['type'];currency=exposure['currency']
        if eid in seen or not re.fullmatch(r'[a-z][a-z0-9_]*',eid):raise ValueError('FX accounting exposure IDs must be unique stable business identifiers')
        seen.add(eid)
        if role not in assets|liabilities:raise ValueError('Unsupported FX balance role: '+role)
        if kind not in ('foreign_operation','monetary'):raise ValueError('FX exposure requires foreign_operation or monetary accounting treatment')
        if kind=='monetary' and (role not in ('cash','nonlease_debt') or exposure.get('functional_currency',reporting)!=reporting):
            raise ValueError('Monetary FX currently supports cash/debt in the reporting functional currency; other functional ledgers require explicit extension')
        share=exposure.get('opening_share')
        if type(share) not in (int,float) or not math.isfinite(share) or not 0<=share<=1:raise ValueError('FX opening exposure share must be between zero and one')
        share_totals[role]=share_totals.get(role,0)+share
        if share_totals[role]>1+1e-9:raise ValueError('Foreign-operation and monetary shares cannot overlap beyond the total balance: '+role)
        prefix='fx_accounting_'+eid
        share_row=add(prefix+'_weight',exposure.get('label',eid)+' opening balance exposure share','ratio','assumption')
        share_row.update(exposure_role=role,exposure_type=kind,effective_date=exposure.get('effective_from'),source_refs=exposure.get('source_refs',[]))
        rate_unit='ratio' if currency==reporting else money+'/'+currency+'m'
        for suffix,label,unit in [('_opening_rate','opening closing rate',rate_unit),('_closing_rate','current closing rate',rate_unit),
            ('_average_rate','current average rate',rate_unit),('_opening_native','opening native-currency amount',currency+'m'),
            ('_movement_native','net native-currency movement',currency+'m'),('_effect','currency movement',money)]:
            row=add(prefix+suffix,exposure.get('label',eid)+' '+label,unit,'assumption' if suffix=='_movement_native' else 'calculated')
            if suffix.endswith('_rate'):row.update(metric_type='ratio',number_format='0.000000')
        for p in active:
            pid=p['id'];start=p['start'];end=p['end']
            applies=(not exposure.get('effective_from') or exposure['effective_from']<=end) and (not exposure.get('effective_to') or exposure['effective_to']>=start)
            if applies and ((exposure.get('effective_from') or start)>start or (exposure.get('effective_to') or end)<end):
                raise ValueError('Balance exposure changes inside a period need a separate dated event amount; split the profile at a financial-period boundary: '+eid)
            prior=[q for q in periods if q['end']<start and q['kind'] in ('FY','H1','Q1','Q2','Q3','Q4') and populated(rows[role]['cells'][q['id']])]
            if not prior:raise ValueError('No supported opening balance for FX exposure '+eid+':'+pid)
            opening=max(prior,key=lambda q:(q['end'],q['kind']=='FY',q['kind']=='H1'))
            share_row['cells'][pid]=dict(value=share if applies else 0.,status='estimate',editable=applies,source_refs=exposure.get('source_refs',[]),reason=exposure.get('method','Explicit provisional opening balance exposure'))
            for suffix,a,b,method in [('_opening_rate',opening['start'],opening['end'],'closing'),('_closing_rate',start,end,'closing'),('_average_rate',start,end,'average')]:
                result=forecast_rate(snapshot,reporting,currency,a,b,method)
                rows[prefix+suffix]['cells'][pid]=dict(value=result['value'],status=result['status'],editable=False,source_refs=result['provenance'],reason='Shared '+method+' rate; missing observations remain missing')
            setf(prefix+'_opening_native',pid,'opening * share / rate',dict(opening=dict(row_id=role,period_id=opening['id']),
                share=dict(row_id=prefix+'_weight',period_id=pid),rate=dict(row_id=prefix+'_opening_rate',period_id=pid)))
            movement=exposure.get('net_movement_native_by_period',{}).get(pid,0.) if applies else 0.
            rows[prefix+'_movement_native']['cells'][pid]=dict(value=movement,status='estimate',editable=applies,source_refs=exposure.get('source_refs',[]),
                reason='Explicit native-currency balance movement already represented in underlying cash/asset/debt schedules; zero limits this scenario to opening exposures. Not a reported treasury flow.')
            setf(prefix+'_effect',pid,f'{prefix}_opening_native * ({prefix}_closing_rate - {prefix}_opening_rate) + {prefix}_movement_native * ({prefix}_closing_rate - {prefix}_average_rate)')
            target=('cash_translation_oci' if role=='cash' else role_targets[role]) if kind=='foreign_operation' else ('cash_fx_remeasurement' if role=='cash' else 'financing_fx_remeasurement')
            targets[pid][target].append(prefix+'_effect')
        for suffix in ('_weight','_movement_native'):
            rid=prefix+suffix
            contract=dict(id=rid,definition=rows[rid]['label'],units=rows[rid]['unit'],financial_basis=data.get('forecast_basis'),
                frequency=frequency,forecast_horizon=spec['forecast_years'],entity_id=exposure.get('entity_id','group'),cost_pool=None,
                currency_basis='foreign_operation' if kind=='foreign_operation' else 'foreign_monetary_in_reporting_functional_currency',
                timing=dict(effective_date=exposure.get('effective_from'),anchor='previous disclosed/modelled closing stock'),parameter_role='exposure_assumption',
                expected_downstream_outputs=[target],source_refs=exposure.get('source_refs',[]),provisional=exposure.get('status')!='reported')
            rows[rid]['driver_contract']=contract;contracts.append(contract)
    for p in active:
        pid=p['id']
        for target,terms in targets[pid].items():sumf(target,pid,terms)
        setf('cash_fx',pid,'cash_translation_oci + cash_fx_remeasurement',reason='Foreign-operation cash translation plus monetary cash remeasurement; both have explicit accounting counterparts')
        c=rows['financial_income']['cells'][pid];c['expression']='('+c['expression']+') + cash_fx_remeasurement'
        c=rows['cfo_fx_noncash']['cells'][pid];c['expression']='('+c['expression']+') - cash_fx_remeasurement'
    for rid in ['cash_translation_oci','cash_fx_remeasurement','financing_fx_remeasurement',*role_targets.values()]:aggregate(rid)
    # Retire the former scalar authority. Historical reported cash FX stays intact.
    if 'cash_fx_assumption' in rows:
        rows['cash_fx_assumption'].update(kind='calculated',support_matrix=True,note='Retired scalar; exposure/rate equations determine forecast cash FX.')
        for p in periods:
            if p['fiscal_year'] in spec['forecast_years'] and not actual(rows['cash_fx_assumption']['cells'][p['id']]):
                rows['cash_fx_assumption']['cells'][p['id']]=dict(value=None,status='not_applicable',editable=False)
    old=spec['financial_contracts'].get('drivers',[])
    derived=set(role_targets.values())|{'cash_fx_assumption','cash_translation_oci','financing_fx_remeasurement','oci_translation'}
    spec['financial_contracts']['drivers']=[c for c in old if c['id'] not in derived]+contracts
    limitations=['Only configured opening balances and explicit native movements are covered; undisclosed intra-period balances, functional currencies and hedge books remain unresolved. Newly earned foreign profits and uncovered assets are not inferred from a cash-FX plug.']
    spec['fx_accounting_contract']=dict(version=1,exposures=exposures,covered_opening_shares=share_totals,
        unresolved_opening_shares={role:1-share_totals.get(role,0.) for role in sorted(assets|liabilities)},limitations=limitations,
        cash_attribution='cash_fx = foreign-operation cash translation (OCI) + cash monetary remeasurement (financial income)',
        monetary_debt='Debt remeasurement expense increases debt and is added back in CFO; translation debt affects OCI only',
        rate_policy='Opening closing rate; current closing rate; current average for explicitly supplied native movements')
    obsolete='Foreign balance-sheet exposures remain incomplete unless explicit closing-rate movements are supplied; cash-flow FX and OCI are independent.'
    spec['financial_contracts']['limitations']=[s for s in spec['financial_contracts'].get('limitations',[]) if s!=obsolete]+limitations
    spec['gaps']=[s for s in spec.get('gaps',[]) if s!=obsolete]
    spec['policies']['fx_accounting']='Exposure-derived cash translation and monetary remeasurement have distinct OCI/earnings counterparts; net-asset translation is separate from cash FX. Uncovered opening balances and native flows are explicit limitations.'
    spec['gaps']+=limitations
    spec.setdefault('fx_validation',[]).append(dict(area='fx_accounting',name='exposure_coverage',status='WARNING',detail=limitations[0]))
    return spec


def validate_fx_contracts(spec,values):
    """Bounds and activation checks for editable currency/hedging controls."""
    checks=[];rows={r['id']:r for r in spec['rows']}
    for row in rows.values():
        rid=row['id']
        if row.get('evidence_role')!='shared_factor_diagnostic':continue
        for pid,cell in row['cells'].items():
            value=values.get(rid,{}).get(pid)
            if rid.endswith('_transaction_weight_total') and value is not None:
                checks.append(dict(area='fx_contract',name=rid+':'+pid,status='PASS' if 0<=value<=1+1e-9 else 'FAIL',detail='Configured transaction components cannot exceed the whole cost pool'))
            if value is None or not cell.get('editable'):continue
            valid=True;detail=None
            if 'weight' in rid or rid.endswith('_hedge_coverage'):
                valid=0<=value<=1;detail='Exposure weight / hedge coverage must be between zero and one'
            elif rid.endswith('_multiplier'):
                valid=value>=0;detail='FX multiplier must be nonnegative and cannot replace missing coverage'
            if detail:checks.append(dict(area='fx_contract',name=rid+':'+pid,status='PASS' if valid else 'FAIL',detail=detail))
            if rid.endswith('_hedge_coverage') and value>0:
                hedge=values.get(rid.removesuffix('_coverage')+'_rate',{}).get(pid)
                checks.append(dict(area='fx_contract',name=rid+':'+pid+':activation',status='PASS' if _number(hedge) is not None else 'FAIL',detail='Active hedge coverage requires a positive functional-currency-per-invoice-currency rate'))
    if spec.get('fx_accounting_contract'):
        for p in spec['periods']:
            pid=p['id'];cash=values.get('cash_fx',{}).get(pid)
            translation=values.get('cash_translation_oci',{}).get(pid);remeasurement=values.get('cash_fx_remeasurement',{}).get(pid)
            if all(v is not None for v in (cash,translation,remeasurement)):
                difference=cash-translation-remeasurement
                checks.append(dict(area='fx_contract',name='cash_accounting_counterpart:'+pid,status='PASS' if abs(difference)<1e-8 else 'FAIL',
                    detail='Cash FX equals its exposure-derived translation OCI and monetary earnings components',difference=difference))
            for role in spec['fx_accounting_contract']['covered_opening_shares']:
                weights=[values.get('fx_accounting_'+e['id']+'_weight',{}).get(pid) for e in spec['fx_accounting_contract']['exposures'] if e['role']==role]
                if weights and all(v is not None for v in weights):checks.append(dict(area='fx_contract',name='nonoverlapping_balance_exposures:'+role+':'+pid,
                    status='PASS' if sum(weights)<=1+1e-9 else 'FAIL',detail='Foreign-operation and monetary positions cannot exceed the same underlying balance'))
    return checks


_CODE = re.compile(r"^[A-Z]{3}$")
_QUARTER = re.compile(r"^Q([1-4])\s*[: /-]?\s*(\d{4}|\d{2})(E)?$", re.I)
_RANGE = re.compile(r"'([^']+)'!\$?([A-Z]+)\$?\d+(?::\$?[A-Z]+\$?\d+)?")
_MONTHS = {name.lower(): number for number, name in enumerate(calendar.month_name) if name}


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def _number(value):
    return float(value) if type(value) in (int, float) and math.isfinite(value) and value > 0 else None


def _quarter(value):
    match = _QUARTER.fullmatch(str(value or '').strip())
    if not match:
        return None
    quarter, year = int(match[1]), int(match[2])
    year = 2000 + year if year < 100 else year
    first = date(year, quarter * 3 - 2, 1)
    last = date(year, quarter * 3, calendar.monthrange(year, quarter * 3)[1])
    return first, last, bool(match[3])


def _header(sheet, label):
    found = [(cell.row, cell.column) for row in sheet for cell in row
             if isinstance(cell.value, str) and cell.value.strip().lower() == label.lower()]
    return found[0] if found else None


def _citation(sheet, cell, cached=None):
    result = dict(sheet=sheet.title, cell=cell.coordinate)
    if cell.data_type == 'f':
        result['formula'] = cell.value
        result['cached_value'] = cached
    return result


def load_fx_source(path, *, adapter=None):
    """Return JSON-safe source, observations, daily_series and metadata.

    Optional adapter keys are reporting_sheets (currency -> sheet name),
    daily_base_currencies (sheet name -> base code), and
    projected_publication_dates (YYYY-Qn -> explicit ISO date list), with optional
    projection_calendar_assumptions (YYYY-Qn -> auditable assumption metadata).
    Configured projection calendars are distinguished from workbook disclosures;
    none of these settings imply company mappings or exposure assumptions.
    Unsupported or contradictory quotation layouts fail rather than guess.
    """
    adapter = dict(adapter or {})
    path = Path(path).resolve()
    original = path.read_bytes()
    formulas = openpyxl.load_workbook(BytesIO(original), data_only=False)
    values = openpyxl.load_workbook(BytesIO(original), data_only=True)
    try:
        source = dict(path=str(path), filename=path.name, sha256=hashlib.sha256(original).hexdigest(),
                      size_bytes=len(original), format='labelled_quarterly_daily_fx_v1')
        observations, qualifications, reference_counts = [], [], {}
        reporting = dict(adapter.get('reporting_sheets', {}))
        if not reporting:
            for sheet in formulas:
                name = re.fullmatch(r'([A-Z]{3})\s+reporting', sheet.title, re.I)
                if name:
                    reporting[name[1].upper()] = sheet.title
        if not reporting:
            raise ValueError('No currency-labelled reporting sheets; configure reporting_sheets')
        for currency, sheet_name in reporting.items():
            if not _CODE.fullmatch(currency):
                raise ValueError('Reporting currencies must use three-letter codes')
            sheet, cached = formulas[sheet_name], values[sheet_name]
            heading = _header(sheet, 'Code')
            if not heading:
                raise ValueError(f'No currency Code header in {sheet_name}')
            header_row, code_column = heading
            periods = {c.column: _quarter(c.value) for c in sheet[header_row] if _quarter(c.value)}
            if not periods:
                raise ValueError(f'No validated quarter labels in {sheet_name}')
            if len({(p[0],p[1]) for p in periods.values()}) != len(periods):
                raise ValueError(f'Duplicate quarter columns in {sheet_name}')
            notes = [dict(sheet=sheet_name, cell=c.coordinate, text=c.value)
                     for row in sheet.iter_rows(max_row=header_row - 1) for c in row
                     if isinstance(c.value, str) and c.data_type != 'f']
            direction = ' '.join(n['text'] for n in notes)
            if not re.search(r'1\s+source.currency\s+unit\s+in\s+' + currency, direction, re.I):
                raise ValueError(f'Quote direction is not documented for {sheet_name}')
            qualifications.extend(notes)
            seen = set()
            for row in sheet.iter_rows(min_row=header_row + 1):
                code = row[code_column - 1].value
                if not isinstance(code, str) or not _CODE.fullmatch(code):
                    continue
                if code in seen:
                    raise ValueError(f'Duplicate source currency {code} in {sheet_name}')
                seen.add(code)
                for column, (start, end, estimate) in periods.items():
                    cell = sheet.cell(row[0].row, column)
                    value = _number(cached[cell.coordinate].value)
                    observation = dict(reporting_currency=currency, source_currency=code,
                        start=start.isoformat(), end=end.isoformat(), kind='average',
                        period=f'{start.year}-Q{(start.month + 2)//3}', value=value,
                        status='missing' if value is None else 'estimate' if estimate else 'actual',
                        source_status='estimate' if estimate else 'actual',
                        provenance=[_citation(sheet, cell, value)], period_cell=sheet.cell(header_row,column).coordinate)
                    observations.append(observation)
                    if cell.data_type == 'f':
                        for name, letters in _RANGE.findall(cell.value):
                            reference_counts[(name, letters)] = reference_counts.get((name, letters), 0) + 1

        candidates, spot_date, projected, assumption_refs, currency_sources = {}, None, {}, [], []
        for sheet in formulas:
            cached = values[sheet.title]
            if sheet.title in reporting.values():
                continue
            for row in sheet:
                table = {str(c.value).strip().lower():c.column for c in row if isinstance(c.value,str)}
                if not {'code','data url','source'} <= table.keys():
                    continue
                for r in range(row[0].row+1,sheet.max_row+1):
                    code=sheet.cell(r,table['code']).value
                    if not isinstance(code,str) or not _CODE.fullmatch(code):
                        if code:
                            for name,column in table.items():
                                value=sheet.cell(r,column).value
                                if name in ('source','data url','method') and isinstance(value,str):
                                    qualifications.append(dict(sheet=sheet.title,cell=sheet.cell(r,column).coordinate,text=value))
                        continue
                    currency_sources.append(dict(currency=code,sheet=sheet.title,
                        fields={name:dict(value=sheet.cell(r,column).value,cell=sheet.cell(r,column).coordinate)
                                for name,column in table.items() if name in ('country','source','data url','method')}))
            header = _header(sheet, 'Date')
            note_end = header[0] - 1 if header else sheet.max_row
            for row in sheet.iter_rows(max_row=note_end):
                for cell in row:
                    if not isinstance(cell.value, str) or cell.data_type == 'f':
                        continue
                    qualifications.append(dict(sheet=sheet.title, cell=cell.coordinate, text=cell.value))
                    if cell.value.strip().lower() == 'spot reference date':
                        spot_date = _date(cached.cell(cell.row,cell.column + 1).value).isoformat()
                        assumption_refs.append(_citation(sheet,sheet.cell(cell.row,cell.column + 1)))
            # Projection dates must be explicit source evidence, not today's calendar.
            for row in sheet:
                cells = [c for c in row if c.value is not None]
                label = next((c for c in cells if isinstance(c.value,str)
                              and re.fullmatch(r'Q[1-4] projected publication days',c.value,re.I)),None)
                if label:
                    count = cached.cell(label.row,label.column + 1).value
                    dates = []
                    for cell in cells:
                        if not isinstance(cell.value,str):
                            continue
                        match = re.fullmatch(r'([\d,\s]+(?:and\s+\d+)?)\s+([A-Za-z]+)\s+(\d{4})',cell.value.strip())
                        if match and match[2].lower() in _MONTHS:
                            dates = [date(int(match[3]),_MONTHS[match[2].lower()],int(d)).isoformat()
                                     for d in re.findall(r'\d+',match[1])]
                    if dates and len(dates) == count:
                        key = dates[0][:4] + '-' + label.value[:2].upper()
                        projected[key] = sorted(dates)
                        assumption_refs.append(_citation(sheet,label))
            if not header:
                continue
            header_row, date_column = header
            headings = {c.column:str(c.value).strip() for c in sheet[header_row] if c.value is not None}
            notes = ' '.join(q['text'] for q in qualifications if q['sheet']==sheet.title)
            base_match = re.search(r'currency units per\s+(?:1\s+)?([A-Z]{3})', notes)
            default_base = adapter.get('daily_base_currencies',{}).get(sheet.title)
            default_base = default_base or (base_match[1] if base_match else None)
            for column, title in headings.items():
                match = re.fullmatch(r'([A-Z]{3})\s+per\s+([A-Z]{3})',title)
                code, base = (match[1],match[2]) if match else (title,default_base)
                if not _CODE.fullmatch(code) or not base:
                    continue
                source_date_col = next((c for c,t in headings.items() if t==code+' source date'),None)
                provider_col = next((c for c,t in headings.items() if t==code+' provider'),None)
                points = []
                for row in sheet.iter_rows(min_row=header_row+1):
                    day = cached.cell(row[0].row,date_column).value
                    if not isinstance(day,(date,datetime)):
                        continue
                    cell = sheet.cell(row[0].row,column)
                    value = _number(cached[cell.coordinate].value)
                    original_date = cached.cell(cell.row,source_date_col).value if source_date_col else day
                    provider = cached.cell(cell.row,provider_col).value if provider_col else sheet.title
                    points.append(dict(date=_date(day).isoformat(),value=value,
                        source_date=_date(original_date).isoformat() if isinstance(original_date,(date,datetime)) else None,
                        provider=provider,provenance=_citation(sheet,cell,value)))
                if not points:
                    continue
                if len({p['date'] for p in points}) != len(points):
                    raise ValueError(f'Duplicate daily dates for {code} in {sheet.title}')
                points.sort(key=lambda p:p['date'])
                candidate=dict(base_currency=base,points=points,sheet=sheet.title,header_cell=sheet.cell(header_row,column).coordinate)
                score=reference_counts.get((sheet.title,openpyxl.utils.get_column_letter(column)),0)
                candidates.setdefault(code,[]).append((score,candidate))
        if not candidates:
            raise ValueError('No dated daily currency series with a documented quote basis')
        daily = {code:max(options,key=lambda p:(p[0],len(p[1]['points'])))[1] for code,options in candidates.items()}
        source_projected=copy.deepcopy(projected)
        configured_projected={key:sorted(_date(d).isoformat() for d in days)
                              for key,days in adapter.get('projected_publication_dates',{}).items()}
        if any(not days for days in configured_projected.values()):
            raise ValueError('Configured projected publication dates must not be empty')
        projected.update(configured_projected)
        for period, days in projected.items():
            if len(set(days)) != len(days):
                raise ValueError('Duplicate projected publication dates')
            for day in days:
                dt=_date(day)
                if f'{dt.year}-Q{(dt.month+2)//3}' != period:
                    raise ValueError('Projected publication date falls outside its declared quarter')
        dates = sorted({p['date'] for series in daily.values() for p in series['points']})
        actual = [o for o in observations if o['source_status']=='actual']
        estimates = sorted({(o['period'],o['start'],o['end']) for o in observations if o['source_status']=='estimate'})
        metadata = dict(reporting_currencies=sorted(reporting),source_currencies=sorted({o['source_currency'] for o in observations}),
            coverage_start=min(o['start'] for o in observations),coverage_end=max(o['end'] for o in observations),
            actual_through=max(o['end'] for o in actual) if actual else None,
            daily_start=dates[0],daily_end=dates[-1],publication_dates=dates,
            spot_reference_date=spot_date,projected_publication_dates=projected,
            source_projected_publication_dates=source_projected,
            configured_projected_publication_dates=configured_projected,
            projection_calendar_provenance={period:dict(
                origin='adapter_configuration' if period in configured_projected else 'source_workbook',
                dates=copy.deepcopy(days),
                configuration=copy.deepcopy(adapter.get('projection_calendar_assumptions',{}).get(period,{})),
                note='Explicit projected publication dates; repeated source spot quotes remain estimates.' if period in configured_projected else 'Projected dates disclosed in the preserved source workbook.')
                for period,days in projected.items()},
            currency_sources=currency_sources,
            estimate_periods=[dict(period=p,start=s,end=e) for p,s,e in estimates],
            quote_direction='reporting-currency units per one source-currency unit',
            cross_convention='For each common publication date: reporting units per base / source units per base; arithmetic mean of daily crosses, one equal weight per published or explicitly projected date.',
            qualifications=qualifications,assumption_provenance=assumption_refs,
            closing_policy='Only exact dated daily observations; average estimates are not closing-rate assumptions.')
        return dict(source=source,observations=observations,daily_series=daily,metadata=metadata,
            all_daily_series=[dict(item,currency=code,selection_score=score) for code,options in candidates.items() for score,item in options])
    finally:
        formulas.close();values.close()


def fx_for_period(source, reporting_currency, source_currency, start, end, kind='average'):
    """Look up a quote, retaining source estimate labels and missing coverage.

    Quarterly workbook observations are preserved, including partial-source
    coverage. Other averages require the original daily publication weights;
    aggregate periods crossing an estimate with no dated weights stay missing.
    Closing rates use exact dated historical data only and never quarter averages.
    """
    if source.get('schema')=='shared_fx_snapshot_v1':
        return forecast_rate(source,reporting_currency,source_currency,start,end,kind)
    start,end=_date(start).isoformat(),_date(end).isoformat()
    if start>end or kind not in ('average','closing'):
        raise ValueError('FX period must be ordered and kind must be average or closing')
    meta=source['metadata']
    result=dict(value=None,status='missing',kind=kind,reporting_currency=reporting_currency,
        source_currency=source_currency,start=start,end=end,quote_direction=meta['quote_direction'],
        coverage=dict(available_days=0,expected_days=None,complete=False),
        provenance=[copy.deepcopy(source['source'])],qualifications=[])
    def missing(reason):
        result['qualifications'].append(reason)
        return result
    if reporting_currency not in meta['reporting_currencies']:
        return missing('Reporting currency has no configured reporting sheet')
    if start<meta['coverage_start'] or end>meta['coverage_end']:
        return missing('Period is outside source coverage; an explicit future input is required')
    a=source['daily_series'].get(reporting_currency)
    b=source['daily_series'].get(source_currency)
    if not a or not b or a['base_currency']!=b['base_currency']:
        return missing('No compatible common-base daily series for this currency pair')
    left={p['date']:p for p in a['points']};right={p['date']:p for p in b['points']}
    def cross(day):
        x,y=left.get(day),right.get(day)
        if x and y and x['value'] is not None and y['value'] is not None:
            return x['value']/y['value'],[x,y]
        return None,[]
    if kind=='closing':
        value,points=cross(end)
        if value is None:
            return missing('No exact dated closing observation; no implicit carry or average-to-closing conversion')
        result.update(value=value,status='actual',coverage=dict(available_days=1,expected_days=1,complete=True))
        result['provenance'].extend(copy.deepcopy(p['provenance']) for p in points)
        return result
    exact=next((o for o in source['observations'] if o['reporting_currency']==reporting_currency and
                o['source_currency']==source_currency and o['start']==start and o['end']==end),None)
    days=[d for d in meta['publication_dates'] if start<=d<=end]
    projected=[];projection_periods=[]
    unweighted=[]
    for period in meta['estimate_periods']:
        if period['end']<start or period['start']>end:
            continue
        extra=meta['projected_publication_dates'].get(period['period'])
        if extra is None and period['end']>meta['daily_end']:
            unweighted.append(period['period'])
        else:
            projected.extend(d for d in extra or [] if start<=d<=end and d not in days)
            if any(start<=d<=end and d not in days for d in extra or []):projection_periods.append(period['period'])
    observations=[];used=[]
    for day in days:
        value,points=cross(day)
        if value is not None:
            observations.append(value);used.extend(points)
    spot,spot_points=cross(meta['spot_reference_date']) if meta['spot_reference_date'] else (None,[])
    if projected and spot is not None:
        observations.extend([spot]*len(projected));used.extend(spot_points)
    expected=len(days)+len(projected)
    result['coverage']=dict(available_days=len(observations),expected_days=None if unweighted else expected,
                            complete=not unweighted and len(observations)==expected)
    result['provenance'].extend(dict(sheet=s['sheet'],header_cell=s['header_cell'],base_currency=s['base_currency']) for s in (a,b))
    result['provenance'].append(dict(method=meta['cross_convention'],publication_dates=days,projected_dates=projected,
        spot_reference_date=meta['spot_reference_date'] if projected or unweighted else None,
        projection_calendars={period:copy.deepcopy(meta.get('projection_calendar_provenance',{}).get(period,{}))
                              for period in projection_periods}))
    if any(period in meta.get('configured_projected_publication_dates',{}) for period in projection_periods):
        result['qualifications'].append('Includes an explicitly configured projected publication calendar; future dates repeat the source spot estimate, not observed future quotes')
        for period in projection_periods:
            description=meta.get('projection_calendar_provenance',{}).get(period,{}).get('configuration',{}).get('description')
            if description:result['qualifications'].append(period+': '+str(description))
    carried=[p for p in used if p.get('source_date') and p['source_date']!=p['date']]
    if carried:
        result['qualifications'].append(f'{len(carried)} component observations retain prior source dates under the source carry convention')
    if exact:
        result.update(value=exact['value'],status=exact['status'])
        result['provenance'].extend(copy.deepcopy(exact['provenance']))
        if unweighted:
            result['qualifications'].append('Source provides a quarterly estimate but no projected daily calendar for combined-period weighting')
        elif len(observations)!=expected:
            result['qualifications'].append(f'Partial source coverage: {len(observations)} of {expected} publication dates')
        if not unweighted and exact['value'] is not None and observations and not math.isclose(exact['value'],sum(observations)/len(observations),rel_tol=1e-9,abs_tol=1e-12):
            result.update(value=None,status='missing')
            result['qualifications'].append('Cached quarter rate disagrees with documented daily-cross calculation')
        if exact['value'] is None:
            result['qualifications'].append('Quarterly source value is missing, invalid or lacks a saved formula cache')
        return result
    if unweighted:
        return missing('Projected publication-date weights are absent for '+', '.join(unweighted))
    if projected and spot is None:
        return missing('Projected dates lack the documented spot reference observation')
    if not observations:
        return missing('No valid paired daily observations in the requested period')
    result.update(value=sum(observations)/len(observations),status='estimate' if projected else 'actual')
    if len(observations)!=expected:
        result['qualifications'].append(f'Partial source coverage: {len(observations)} of {expected} publication dates')
    return result
