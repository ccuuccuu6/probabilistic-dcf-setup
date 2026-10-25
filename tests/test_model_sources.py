"""Original-byte and numeric provenance audits, independent of AI/network calls."""
import copy
import hashlib
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject, NumberObject

from research_platform.evidence import verify_model_input_sources


class DateEvidenceTests(unittest.TestCase):
    def test_declared_date_must_match_original_single_full_date(self):
        from research_platform.evidence import date_phrase_matches
        cases = [('2025-07-29T16:30:00Z', '2025-07-29'),
                 ('Paris, 29 juillet 2025', '2025-07-29'),
                 ('February 21st, 2024', '2024-02-21'),
                 ('21 février 2024', '2024-02-21'),
                 ('31 December 2025', '2025-12-31'),
                 ('29/07/2025', '2025-07-29'), ('07/29/2025', '2025-07-29')]
        for phrase, expected in cases:
            with self.subTest(phrase=phrase):
                self.assertTrue(date_phrase_matches(phrase, expected))
                self.assertFalse(date_phrase_matches(phrase, '2026-10-25'))

    def test_ambiguous_partial_invalid_and_conflicting_dates_fail(self):
        from research_platform.evidence import date_phrase_matches
        for phrase in ('2025', 'July 2025', '01/02/2025', '31 February 2025',
                       '31 December 2025 and 31 December 2024', 'retrieved today'):
            with self.subTest(phrase=phrase):
                self.assertFalse(date_phrase_matches(phrase, '2025-12-31'))
        self.assertFalse(date_phrase_matches('01/02/2025', '2025-01-02'))
        self.assertFalse(date_phrase_matches('01/02/2025', '2025-02-01'))


class ModelSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        book = Workbook()
        book.active.title = "Original"
        book.active["A1"] = -12
        book.active["A2"] = 2.5
        book.active["A4"] = "—"
        book.save(self.root / "original.xlsx")
        self.data = {"sources": {"src": self.source("original.xlsx", "xlsx")},
                     "historical": {}, "segments": [], "interims": []}

    def source(self, path, format_name):
        return dict(id="src", local_path=path, format=format_name,
                    sha256=hashlib.sha256((self.root / path).read_bytes()).hexdigest())

    def ref(self, cell="A1", raw=-12, multiplier=1, **kwargs):
        return dict(source_id="src", sheet="Original", cell=cell, raw=raw,
                    multiplier=multiplier, **kwargs)

    def point(self, value, refs):
        self.data["historical"]["measure"] = dict(label="Measure", unit="currency",
            values={"2025": value}, sources={"2025": refs})

    def failures(self, data=None):
        return [check for check in verify_model_input_sources(data or self.data, self.root)
                if check["status"] == "FAIL"]

    def test_signed_value_and_unit_conversion_are_both_checked(self):
        self.point(.025, [self.ref("A2", 2.5, .01)])
        self.assertEqual(self.failures(), [])
        self.data["historical"]["measure"]["sources"]["2025"][0]["multiplier"] = 1
        self.assertTrue(self.failures())
        self.point(12, [self.ref(raw=12)])
        self.assertTrue(self.failures())

    def test_derived_values_use_explicit_weighted_originals(self):
        self.point(-7, [self.ref(contribution_weight=1),
                        self.ref("A2", 2.5, contribution_weight=2)])
        self.assertEqual(self.failures(), [])
        self.data["historical"]["measure"]["values"]["2025"] = -6
        self.assertEqual(len(self.failures()), 1)

    def test_changed_original_and_missing_raw_are_rejected(self):
        self.point(-12, [self.ref()])
        bad = copy.deepcopy(self.data)
        del bad["historical"]["measure"]["sources"]["2025"][0]["raw"]
        self.assertTrue(self.failures(bad))
        with (self.root / "original.xlsx").open("ab") as stream:
            stream.write(b"changed")
        self.assertTrue(any(c["area"] == "source_integrity" for c in self.failures()))

    def test_blank_caches_and_none_do_not_become_zero(self):
        self.point(0, [self.ref("A3", 0)])
        self.assertTrue(self.failures())
        self.point(None, [])
        checks = verify_model_input_sources(self.data, self.root)
        self.assertEqual([c["status"] for c in checks if c["area"] == "source_data"], ["MISSING"])

    def test_optional_absence_stays_none_and_is_a_warning(self):
        self.point(None, [])
        metric = self.data["historical"]["measure"]
        metric.update(optional=True, missing_reason="Optional comparison unavailable; no forecast uses it.")
        checks = verify_model_input_sources(self.data, self.root)
        self.assertIsNone(metric["values"]["2025"])
        self.assertEqual([c["status"] for c in checks if c["area"] == "source_data"], ["WARNING"])

    def test_segments_and_interims_use_the_same_provenance_contract(self):
        self.data["segments"] = [dict(id="segment", sales2025=12,
            source_refs={"sales2025": self.ref(multiplier=-1)})]
        self.data["interims"] = [dict(period="2025-H1", value=-12, source_ref=self.ref())]
        self.assertEqual(self.failures(), [])
        self.data["interims"][0]["value"] = 12
        self.assertEqual(len(self.failures()), 1)

    def test_pdf_presence_is_signed_and_is_not_semantic_acceptance(self):
        writer = PdfWriter()
        page = writer.add_blank_page(600, 800)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 60 740 Td (Charge \\(12\\) and rate 2.5%) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(self.root / "original.pdf")
        self.data["sources"]["src"] = self.source("original.pdf", "pdf")
        self.point(-12, [dict(source_id="src", unit_id="page:1", raw=-12, multiplier=1)])
        checks = verify_model_input_sources(self.data, self.root)
        self.assertFalse([c for c in checks if c["status"] == "FAIL"])
        self.assertTrue(any("semantic mapping is unverified" in c["detail"] for c in checks))
        self.point(12, [dict(source_id="src", unit_id="page:1", raw=12, multiplier=1)])
        self.assertTrue(self.failures())
        self.data["historical"] = {}
        self.data["observations"] = [dict(id="rate", period="2025-H1", value=.025,
            source_refs=[dict(source_id="src", unit_id="page:1", raw="2.5%", multiplier=1)])]
        self.assertEqual(self.failures(), [])
        self.data["observations"][0]["source_refs"][0]["multiplier"] = .01
        self.assertTrue(self.failures(), "Literal percentages must not be divided by 100 twice")

    def test_archive_escape_is_rejected(self):
        self.data["sources"]["src"]["local_path"] = "../original.xlsx"
        self.assertTrue(self.failures())

    def test_real_blank_pdf_pages_preserve_empty_units_and_original_bytes(self):
        from research_platform.evidence import _units
        writer=PdfWriter()
        for _ in range(3):writer.add_blank_page(600,800)
        path=self.root/'blank.pdf';writer.write(path);original=path.read_bytes()
        units=list(_units(original,'pdf'))
        self.assertEqual([unit['unit_id'] for unit in units],['page:1','page:2','page:3'])
        for unit in units:
            self.assertEqual(unit['text'],'');self.assertEqual(unit['numbers'],[])
            self.assertEqual(unit['extraction_status'],'verified_empty_page')
        self.assertEqual(path.read_bytes(),original)

    def test_missing_or_malformed_content_with_resources_is_not_silently_blank(self):
        from research_platform.evidence import _units
        for malformed in (False,True):
            writer=PdfWriter();page=writer.add_blank_page(600,800)
            if malformed:page[NameObject('/Contents')]=NumberObject(42)
            else:page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):DictionaryObject()})})
            path=self.root/'malformed.pdf';writer.write(path)
            with self.subTest(malformed=malformed),self.assertRaisesRegex(ValueError,'PDF page 1'):
                list(_units(path.read_bytes(),'pdf'))


class ConstructionVendorPeriodTests(unittest.TestCase):
    """The source adapter accepts fiscal periods without changing their schema."""

    def setUp(self):
        from research_platform.evidence import SourceStore
        from research_platform.periods import build_periods
        import io
        book=Workbook();sheet=book.active;sheet.title='Company Reported'
        sheet['D1']='Synthetic manufacturing business'
        labels=('FY-2023','FY-2024E','1QFY-2023','2QFY-2023','3QFY-2023',
                '4QFY-2023','1HFY-2023','2HFY-2023','9MFY-2023')
        values=(240,999999,40,50,60,90,90,150,150)
        sheet['E3']='Units';sheet['D4']='Revenue';sheet['E4']='GBPM';sheet['F4']='Analytics'
        for column,(label,value) in enumerate(zip(labels,values),7):
            sheet.cell(2,column,label);sheet.cell(3,column,'Consensus' if label.endswith('E') else 'VA Actuals')
            sheet.cell(4,column,value)
        stream=io.BytesIO();book.save(stream);book.close()
        units,metadata=SourceStore._index(stream.getvalue(),'xlsx','vendor','visible_alpha')
        self.store=SourceStore(Path('.'));self.store._units={'vendor':units}
        self.store.sources=[dict(id='vendor',category='visible_alpha',sha256=hashlib.sha256(stream.getvalue()).hexdigest(),**metadata)]
        self.assertEqual(metadata['vendor_layout']['status'],'recognized')
        self.periods=build_periods([2023],dict(fiscal_year_end='06-30'))
        self.by_kind={p['kind']:p for p in self.periods}
        self.indexes=dict(FY=0,Q1=2,Q2=3,Q3=4,Q4=5,H1=6,H2=7,**{'9M':8})

    def reference(self,slot):
        return dict(source_id='vendor',unit_id='Company Reported!R4',number_index=self.indexes[slot],number_format='dot_decimal')

    def test_june_fiscal_year_all_slots_and_exact_original_values(self):
        before=copy.deepcopy((self.periods,self.store._units))
        for slot,expected in dict(FY=240,Q1=40,Q2=50,Q3=60,Q4=90,H1=90,H2=150,**{'9M':150}).items():
            with self.subTest(slot=slot):
                result=self.store.validate_historical_reference(self.reference(slot),self.by_kind[slot],fiscal_periods=self.periods)
                self.assertEqual(result['slot'],slot);self.assertEqual(self.store.resolve(self.reference(slot)),expected)
        self.assertEqual(self.by_kind['Q1']['start'],'2022-07-01')
        self.assertEqual(self.by_kind['Q1']['end'],'2022-09-30')
        self.assertEqual(self.by_kind['FY']['end'],'2023-06-30')
        self.assertEqual((self.periods,self.store._units),before)
        self.assertNotIn('year',self.by_kind['Q1']);self.assertNotIn('fiscal_year_start',self.by_kind['Q1'])

    def test_wrong_year_slot_cumulative_half_and_calendar_reject(self):
        cases=[dict(self.by_kind['Q1'],fiscal_year=2024),self.by_kind['Q2'],self.by_kind['H1'],self.by_kind['9M'],
               dict(self.by_kind['Q1'],start='2023-01-01',end='2023-03-31'),
               dict(self.by_kind['Q1'],id='2023-Q2'),dict(self.by_kind['Q1'],kind='Q2'),
               dict(self.by_kind['Q1'],year=2024)]
        for period in cases:
            with self.subTest(period=period),self.assertRaises(ValueError):
                self.store.validate_historical_reference(self.reference('Q1'),period,fiscal_periods=self.periods)
        for source_slot,period_slot in [('9M','Q3'),('H1','H2'),('FY','9M'),('Q3','9M')]:
            with self.subTest(source=source_slot,period=period_slot),self.assertRaises(ValueError):
                self.store.validate_historical_reference(self.reference(source_slot),self.by_kind[period_slot],fiscal_periods=self.periods)

    def test_missing_conflicting_or_foreign_fy_anchor_cannot_default_to_calendar_year(self):
        period=self.by_kind['H1'];reference=self.reference('H1')
        for anchors in (None,[],[dict(self.by_kind['FY'],fiscal_year=2022)],
                        self.periods+[dict(self.by_kind['FY'],start='2023-01-01',end='2023-12-31')]):
            with self.subTest(anchors=anchors),self.assertRaisesRegex(ValueError,'FY|fiscal-year'):
                self.store.validate_historical_reference(reference,period,fiscal_periods=anchors)
        with self.assertRaisesRegex(ValueError,'conflict'):
            self.store.validate_historical_reference(reference,dict(period,fiscal_year_start='2023-01-01'),fiscal_periods=self.periods)
        with self.assertRaisesRegex(ValueError,'period ID conflicts'):
            self.store.validate_historical_reference(reference,period,fiscal_periods=[dict(self.by_kind['FY'],id='2023-H2')])
        self.store.validate_historical_reference(self.reference('FY'),self.by_kind['FY'])

    def test_custom_model_ids_use_explicit_kind_and_dates_without_rewriting_input(self):
        custom=dict(self.by_kind['H1'],id='first_half_of_retail_reporting_year')
        before=copy.deepcopy(custom)
        result=self.store.validate_historical_reference(self.reference('H1'),custom,fiscal_periods=self.periods)
        self.assertEqual(result['slot'],'H1');self.assertEqual(custom,before)

    def test_legacy_initiation_contract_and_forecast_exclusion_are_unchanged(self):
        from research_platform.evidence import _vendor_reference_period
        legacy=dict(id='2023-H1',year=2023,kind='half_year',start='2022-07-01',end='2022-12-31',
                    fiscal_year_start='2022-07-01',fiscal_year_end='2023-06-30')
        self.assertIs(_vendor_reference_period(legacy),legacy)
        self.store.validate_historical_reference(self.reference('H1'),legacy)
        with self.assertRaises(ValueError):
            self.store.validate_historical_reference(self.reference('H1'),dict(legacy,id='2023-Q1'))
        reference=dict(self.reference('FY'),number_index=1)
        with self.assertRaisesRegex(ValueError,'Consensus, estimate'):
            self.store.validate_historical_reference(reference,self.by_kind['FY'],fiscal_periods=self.periods)

    def test_materializer_supplies_fy_anchor_even_for_single_quarter_cohort(self):
        from research_platform.integrated import integrated_role_contract
        from research_platform.workflow import _materialize_financial_observations
        ref=self.reference('Q1');period=self.by_kind['Q1']
        entry=dict(row_id='sales',label='Revenue',period_id=period['id'],unit='GBPm',metric_type='flow',entity_id='group',
            basis='consolidated',definition_id='revenue',scope='consolidated',rounding_increment=1,
            restatement='original',restatement_vintage=None,classification_basis='Synthetic company-reported table.',
            classification_source_refs=[dict(source_id='vendor',unit_id=ref['unit_id'])],
            terms=[dict(reference=ref,source_raw=40,multiplier=1,weight=1)])
        evidence=self.store.read_documents([ref],historical_years=[2023])
        answer=dict(observations=[entry],reporting_changes=[],gaps=[])
        before=copy.deepcopy((answer,self.periods,evidence))
        kwargs=dict(store=self.store,catalog={'vendor':self.store.sources[0]},periods=self.periods,currency='GBP',
                    roles=integrated_role_contract(),segments=[],evidence=evidence,requested_period_ids=[period['id']])
        result=_materialize_financial_observations(answer,**kwargs)
        self.assertEqual(result[0]['value'],40);self.assertEqual(result[0]['period_id'],'2023-Q1')
        self.assertEqual(result[0]['source_refs'][0]['number_index'],2)
        self.assertEqual(result[0]['source_refs'][0]['cell'],'I4')
        self.assertEqual((answer,self.periods,evidence),before)
        bad=copy.deepcopy(answer);bad['observations'][0]['terms'][0]['reference']=self.reference('Q2')
        with self.assertRaisesRegex(ValueError,'fiscal slot'):
            _materialize_financial_observations(bad,**kwargs)


if __name__ == "__main__":
    unittest.main()
