"""Known-answer construction cases; synthetic data never implies company accuracy."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from company_initiation.shared_ai import OpenAIProvider, AIError, DEFAULT_MODEL
from company_initiation.historic_construction.evidence import inventory_local_sources, SourceStore
from company_initiation.forecast_construction.construction import source_unit, calculate, validate_spec, initialise_assumptions, validate_results, normalise_history, export_model, construct_company, validate_assumptions
from company_initiation.shared_runtime import expression_tree, RunJournal, SELECT, digest
import test_setup as fixtures


def known_model():
    def row(key, value, unit="USD"):
        return dict(id=key, label=key, entity="Group", unit=unit, scale=1,
                    time_basis="flow", section="income_statement", values={"2025": value},
                    origins={}, basis="Synthetic", definition=key)
    history = dict(company="SYNTHETIC Known Answers", currency="USD", rows=[row("revenue",100),row("costs",60),
        row("marketing",10),row("profit",30),row("price_history",0,"ratio")],
        periods=[dict(id="2025",year=2025,kind="annual",start="2025-01-01",end="2025-12-31")], identities=[])
    def driver(key, target, component, measure="growth_rate", direction=1):
        return dict(id=key,label=key,definition=key,kind="business",measure=measure,unit="ratio",scale=1,
            frequency="annual",entity="Group",materiality="Synthetic mechanism with known answer",priority="high",
            relationship="hypothesis",references=[],historical_row="price_history",limitations=["SYNTHETIC"],
            effects=[dict(line=target,direction=direction,lag=0,explanation="Known formula")],components=[component],
            baseline=dict(rule="zero",numerator="",denominator="",explanation="Neutral temporary growth"))
    drivers=dict(drivers=[driver("price","revenue","price"),driver("volume","revenue","volume"),
                         driver("unit_cost","costs","commodity_cost")],gaps=[])
    def forecast(key, expression):
        return dict(id=key,label=key,entity="Group",section="income_statement",unit="USD",scale=1,time_basis="flow",
            expression=expression,rationale="Synthetic known answer",references=[],simplifications=["Synthetic"])
    spec=dict(rows=[forecast("revenue",'lag("revenue", 1) * (1 + price) * (1 + volume)'),
        forecast("costs",'lag("costs", 1) * (1 + volume) * (1 + unit_cost)'),
        forecast("marketing",'lag("marketing", 1)'),forecast("profit","revenue - costs - marketing")],
        identities=[dict(id="profit_bridge",target="profit",expression="revenue-costs-marketing",periods=[],tolerance=.0001,rationale="Identity",references=[])],gaps=[])
    return history, drivers, spec


class FormulaTests(unittest.TestCase):
    def test_empty_forecast_and_ratio_scale_cannot_appear_complete(self):
        h,d,s=known_model()
        empty=dict(s,rows=[])
        self.assertTrue(validate_spec(h,d,empty))
        d['drivers'][0]['scale']=100
        self.assertTrue(any('display scale 1' in e for e in validate_spec(h,d,s)))
    def test_known_answers_and_unrelated_price_unchanged_on_cost_shock(self):
        history,drivers,spec=known_model()
        assumptions=initialise_assumptions(history,drivers,[2026,2027])
        for a in assumptions:
            a["value"]={"price":.1,"volume":.2,"unit_cost":.05}[a["driver"]]
        values=calculate(history,drivers,spec,assumptions,[2026,2027])
        self.assertAlmostEqual(values["revenue"]["2026"],132)
        self.assertAlmostEqual(values["costs"]["2026"],75.6)
        self.assertAlmostEqual(values["profit"]["2026"],46.4)
        self.assertAlmostEqual(values["revenue"]["2027"],174.24)
        modified=copy.deepcopy(assumptions)
        next(a for a in modified if a["driver"]=="unit_cost")["value"] += .1
        changed=calculate(history,drivers,spec,modified,[2026,2027])
        self.assertEqual(changed["revenue"],values["revenue"])
        self.assertAlmostEqual(changed["costs"]["2026"],82.8)
        self.assertEqual(assumptions[0]["value"],.1)

    def test_missing_is_not_zero_and_no_forecast_from_interim(self):
        h,d,s=known_model()
        h["periods"].append(dict(id="2025-H1",year=2025,kind="half_year",start="2025-01-01",end="2025-06-30"))
        h["rows"][0]["values"]["2025-H1"]=900
        a=initialise_assumptions(h,d,[2026])
        self.assertEqual(calculate(h,d,s,a,[2026])["revenue"]["2026"],100)
        a[0]["value"]=None
        self.assertIsNone(calculate(h,d,s,a,[2026])["revenue"]["2026"])

    def test_invalid_and_circular_dependencies_are_rejected(self):
        h,d,s=known_model()
        for expr in ('__import__("os")',"revenue[0]","revenue**2","revenue * .07"):
            with self.subTest(expr=expr),self.assertRaises(ValueError):
                expression_tree(expr)
        s["rows"][0]["expression"]="profit"
        self.assertTrue(any("Circular" in e for e in validate_spec(h,d,s)))
        s["rows"][0]["expression"]="unknown"
        self.assertTrue(any("Unknown" in e for e in validate_spec(h,d,s)))

    def test_double_counted_organic_price_and_volume_rejected(self):
        h,d,s=known_model()
        organic=copy.deepcopy(d["drivers"][0]);organic.update(id="organic",components=["price","volume"])
        d["drivers"].append(organic)
        s["rows"][0]["expression"] += " * (1 + organic)"
        self.assertTrue(any("Double-counted" in e for e in validate_spec(h,d,s)))

    def test_incorrect_economic_relationship_and_lag_detected(self):
        h,d,s=known_model()
        s["rows"][1]["expression"]='lag("costs", 1) * (1 - unit_cost)'
        a=initialise_assumptions(h,d,[2026,2027])
        v=calculate(h,d,s,a,[2026,2027])
        checks=validate_results(h,d,s,a,[2026,2027],v)
        self.assertTrue(any(c["status"]=="FAIL" and c["name"].startswith("unit_cost") for c in checks))
        h,d,s=known_model()
        d["drivers"][2]["effects"][0]["lag"]=1
        v=calculate(h,d,s,a,[2026,2027])
        checks=validate_results(h,d,s,a,[2026,2027],v)
        self.assertTrue(any(c["status"]=="FAIL" and "timing" in c["name"] for c in checks))

    def test_real_lag_uses_historical_driver_interface(self):
        h,d,s=known_model()
        s["rows"][1]["expression"]='lag("costs", 1) * (1 + lag("unit_cost", 1))'
        d["drivers"][2]["effects"][0]["lag"]=1
        a=initialise_assumptions(h,d,[2026,2027])
        next(i for i in a if i["driver"]=="unit_cost" and i["period"]=="2026")["value"]=.1
        v=calculate(h,d,s,a,[2026,2027])
        self.assertEqual(v["costs"]["2026"],60)
        self.assertEqual(v["costs"]["2027"],66)

    def test_mismatched_units_and_scenario_interface_rejected(self):
        h,d,s=known_model()
        s["rows"][0]["expression"]="price + costs"
        self.assertTrue(any("incompatible units" in e for e in validate_spec(h,d,s)))
        a=initialise_assumptions(h,d,[2026]);a[0]["unit"]="CHF"
        with self.assertRaises(ValueError):validate_assumptions(a,d,[2026])

    def test_workbook_formulas_inputs_formatting_and_portability(self):
        h,d,s=known_model();a=initialise_assumptions(h,d,[2026]);v=calculate(h,d,s,a,[2026])
        h['rows'][3]['values']['2025']=31  # Planted reported profit discrepancy: 31 vs 100-60-10.
        h['identities']=[dict(s['identities'][0],periods=['2025'])]
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/"model.xlsx"
            report=export_model(path,h,d,s,a,[2026],v,[],[],["SYNTHETIC"])
            book=load_workbook(path)
            self.assertEqual(len(book.sheetnames),8)
            self.assertEqual(book["Model"].freeze_panes,"D5")
            self.assertIn("'Assumptions'!",book["Model"]["E5"].value)
            self.assertIn("COUNT",book["Model"]["E5"].value)
            self.assertEqual(book["Assumptions"]["F2"].fill.fgColor.rgb,"00FFF2CC")
            self.assertFalse(book._external_links)
            self.assertGreater(report["formula_count"],5)
            self.assertEqual(report["expected_model_cells"]["E5"],100)
            self.assertEqual(report['expected_check_cells']['G2'],1)
            self.assertEqual(report['expected_check_cells']['H2'],'FAIL')
            self.assertIn("'Model'!D8",book['Checks']['G2'].value)
            self.assertEqual(book['Historics']['D5'].value,31)
            self.assertEqual(book['Drivers']['B2'].font.name,'Aptos')
            book.close()


def vendor_mapping(store):
    source=next(s['id'] for s in store.sources if s.get('category')=='visible_alpha')
    def row(key, label, address):
        return dict(id=key,label=label,entity="Group",section="income_statement",definition=label,
            unit="USD",scale=1e6,sign=1,time_basis="flow",basis="SYNTHETIC unchanged basis",references=[],
            source_rows=[dict(source_id=source,unit_id="Vendor data!R"+str(address))],observations=[])
    return dict(company="SYNTHETIC Manufacturing Group",currency="USD",identity_references=[],
        periods=[dict(fixtures.VisibleAlphaTests.period(y),references=[]) for y in (2024,2025)] +
                [dict({k:v for k,v in fixtures.VisibleAlphaTests.period(2024,"H1").items() if not k.startswith("fiscal_")},references=[])],
        rows=[row("segment_sales","Manufacturing sales",19),row("revenue","Group sales",28)],
        identities=[dict(id="segment_total",target="revenue",expression="segment_sales",periods=[],
                         tolerance=0,rationale="Synthetic one segment",references=[])],gaps=["Synthetic; other statements unavailable"],reporting_changes=[])


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.archive=self.root/"archive"
        (self.archive/"visible_alpha").mkdir(parents=True)
        self.path=self.archive/"visible_alpha"/"synthetic.xlsx"
        fixtures.VisibleAlphaTests.make_workbook(self.path)
        self.store=SourceStore(self.root/"sources")
        for source in inventory_local_sources(self.archive)["sources"]:
            self.store.add_local(source,root=self.archive)

    def test_aliases_filings_precedence_estimates_missing_and_cumulative(self):
        history=normalise_history(vendor_mapping(self.store),self.store,[2024,2025])
        revenue=history["rows"][1]
        self.assertEqual(revenue["values"],{"2024":100e6,"2025":110e6,"2024-H1":45e6})
        self.assertEqual(revenue["origins"]["2024"]["series"],"Filings")
        self.assertTrue(any(x["series"]=="VA Actuals" and "2025E" in x["period_label"] for x in history["excluded_observations"]))
        self.assertNotIn("2024-Q1",revenue["values"])
        self.assertEqual(self.path.read_bytes(),Path(self.store.sources[0]["path"]).read_bytes())

    def test_segment_discrepancy_preserved(self):
        key=self.store.sources[0]["id"]
        unit=self.store._units[key]["Vendor data!R28"]
        next(n for n in unit["numbers"] if n["locator"]=="I28")["raw"]=111
        h=normalise_history(vendor_mapping(self.store),self.store,[2024,2025])
        check=next(c for c in h["checks"] if c["name"]=="segment_total:2024")
        self.assertEqual(check["status"],"FAIL");self.assertEqual(check["difference"],11e6)
        self.assertEqual(h["rows"][1]["values"]["2024"],111e6)

    def test_units_percentage_basis_points_and_missing(self):
        self.assertEqual(source_unit("USDM",{"raw":1}), ("USD",1e6))
        self.assertEqual(source_unit("%",{"raw":5}), ("ratio",.01))
        self.assertEqual(source_unit("%",{"raw":.05,"display_format":"0.0%"}), ("ratio",1))
        self.assertEqual(source_unit("bps",{"raw":50}), ("ratio",.0001))
        with self.assertRaises(ValueError):source_unit("unknown",{"raw":1})
        m=vendor_mapping(self.store);m["rows"][0]["source_rows"][0]["unit_id"]="Vendor data!R21"
        h=normalise_history(m,self.store,[2024,2025])
        self.assertFalse(h["rows"][0]["values"])

    def test_standardised_rows_cannot_be_mapped(self):
        m=vendor_mapping(self.store);m["rows"][0]["source_rows"][0]["unit_id"]="Vendor data!R7"
        with self.assertRaisesRegex(ValueError,"excluded"):
            normalise_history(m,self.store,[2024,2025])

    def test_requested_year_projection_is_explicit_and_preserves_raw_scope(self):
        from company_initiation.forecast_construction.construction import _valid_scope
        m=vendor_mapping(self.store)
        m['periods'].append(dict(fixtures.VisibleAlphaTests.period(2023),references=[]))
        scope={k:v for k,v in m.items() if k not in ('rows','identities')}
        selected=_valid_scope(scope,self.store,[2024,2025])
        self.assertNotIn(2023,{p['year'] for p in selected['periods']})
        self.assertIn(2023,{p['year'] for p in scope['periods']})
        self.assertTrue(any('2023' in g for g in selected['gaps']))

    def test_mocked_production_workflow_and_exact_resume(self):
        m=vendor_mapping(self.store)
        provider=FixtureProvider(m)
        directory=construct_company(self.root,"SYNTHETIC Manufacturing Group",self.archive,years=[2024,2025],provider=provider)
        record=json.loads((directory/"record.json").read_text())
        self.assertTrue((directory/"model.xlsx").exists(),record["issues"])
        self.assertEqual(record["status"],"incomplete_model")
        self.assertEqual(len(provider.requests),13)
        before=len(provider.requests)
        self.assertEqual(construct_company(self.root,"SYNTHETIC Manufacturing Group",self.archive,years=[2024,2025],provider=provider),directory)
        self.assertEqual(len(provider.requests),before)
        blind=next(r for r in provider.requests if r["payload"]["stage"]=="construction_blind")["payload"]
        self.assertNotIn("historical_model",blind)
        self.assertNotIn("values",json.dumps(blind))
        coverage=next(r for r in provider.requests if r["payload"]["stage"]=="construction_coverage")["payload"]
        self.assertNotIn("concepts",coverage)

    def test_explicit_code_migration_reuses_calls_and_preserves_parent_budget(self):
        provider=FixtureProvider(vendor_mapping(self.store))
        first=construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider)
        path=first/'record.json';record=json.loads(path.read_text())
        record['configuration']['code']['construction.py']='synthetic-old-code'
        path.write_text(json.dumps(record),encoding='utf-8')
        before=len(provider.requests)
        second=construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider,continue_from=first)
        saved=json.loads((second/'record.json').read_text())
        self.assertNotEqual(first,second)
        self.assertEqual(saved['parent_run'],str(first))
        self.assertEqual(saved['usage'],record['usage'])
        self.assertEqual(len(provider.requests),before)
        self.assertEqual(json.loads(path.read_text())['configuration']['code']['construction.py'],'synthetic-old-code')

    def test_fresh_or_changed_years_cannot_bypass_unknown_company_charge(self):
        provider=FixtureProvider(vendor_mapping(self.store))
        company='SYNTHETIC Manufacturing Group'
        directory=construct_company(self.root,company,self.archive,years=[2024,2025],provider=provider)
        path=directory/'record.json'
        record=json.loads(path.read_text())
        record['pending_call']={'client_request_id':'synthetic-lost-response'}
        path.write_text(json.dumps(record),encoding='utf-8')
        count=len(provider.requests)
        for years in ([2024,2025],[2024]):
            with self.assertRaisesRegex(AIError,'unknown call charges'):
                construct_company(self.root,company,self.archive,years=years,provider=provider,fresh=True)
        self.assertEqual(len(provider.requests),count)

    def test_changed_usage_cannot_be_recovered_as_request_alias_bug(self):
        provider=FixtureProvider(vendor_mapping(self.store))
        first=construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider)
        path=first/'record.json';record=json.loads(path.read_text())
        record['calls'][0]['usage']['estimated_cost_usd']=0
        path.write_text(json.dumps(record),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'journal disagrees'):
            construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider,continue_from=first)

    def test_completed_result_still_verifies_frozen_originals(self):
        provider=FixtureProvider(vendor_mapping(self.store))
        first=construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider)
        record=json.loads((first/'record.json').read_text())
        Path(record['sources'][0]['path']).write_bytes(b'DELIBERATELY DAMAGED TEST COPY')
        before=len(provider.requests)
        with self.assertRaisesRegex(ValueError,'Frozen source bytes changed'):
            construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider)
        self.assertEqual(len(provider.requests),before)

    def test_offline_review_recalculates_without_provider_or_credentials(self):
        from company_initiation.forecast_construction.construction import review_saved
        provider=FixtureProvider(vendor_mapping(self.store))
        first=construct_company(self.root,'SYNTHETIC Manufacturing Group',self.archive,years=[2024,2025],provider=provider)
        with patch('company_initiation.shared_ai.OpenAIProvider.from_env',side_effect=AssertionError('Must not read credentials')):
            second=review_saved(self.root,first)
        self.assertTrue((second/'model.xlsx').exists())
        self.assertEqual(json.loads((second/'record.json').read_text())['new_paid_calls'],0)


class FixtureProvider(OpenAIProvider):
    def __init__(self,mapping):
        super().__init__(api_key="SYNTHETIC-TEST-KEY")
        self.mapping=mapping;self.requests=[]
    def metadata(self):
        return dict(super().metadata(),provider='scripted_fixture_no_network',
                    usage_note='Synthetic token counts for offline guard tests; no real provider charges')
    def complete(self,**kwargs):
        self.requests.append(copy.deepcopy(kwargs))
        stage=kwargs["payload"]["stage"]
        self.last_response={"model":DEFAULT_MODEL,"usage":{"input_tokens":100,"output_tokens":50,"total_tokens":150},"output":[]}
        if stage=="construction_select":return dict(source_ids=[],reasons=["No company documents"],gaps=["Synthetic no primary"])
        if stage in ("construction_review","construction_coverage"):return dict(issues=[],missing_data=["No primary documents"])
        if stage=="construction_history":
            if "historical_rows" in kwargs["payload"]:
                return dict(identities=self.mapping["identities"],gaps=[])
            return {k:v for k,v in self.mapping.items() if k not in ("rows","identities")}
        if stage=="construction_map":
            return dict(rows=[r for r in self.mapping["rows"] if r["section"]==kwargs["payload"]["requested_section"]],gaps=[])
        if stage=="construction_blind":return dict(items=[],coverage=[],gaps=["No primary sources"])
        if stage=="construction_drivers":
            driver=known_model()[1]["drivers"][0]
            driver["historical_row"]=""
            return dict(drivers=[driver],gaps=["Synthetic growth framework"])
        if stage=="construction_formulas":
            row=known_model()[2]["rows"][0]
            row["scale"]=1e6;row["expression"]='lag("revenue",1)*(1+price)'
            return dict(rows=[row],identities=[],gaps=["Other statements unavailable"])
        raise AssertionError(stage)


class JournalTests(unittest.TestCase):
    def test_provider_incomplete_status_cannot_be_accepted_even_when_partial_json_is_parseable(self):
        provider=FixtureProvider({})
        def partial(**kwargs):
            provider.last_response=dict(status='incomplete',incomplete_details=dict(reason='max_output_tokens'),
                usage=dict(input_tokens=100,output_tokens=32000,total_tokens=32100))
            return dict(source_ids=[],reasons=['Syntactically valid but incomplete'],gaps=[])
        with tempfile.TemporaryDirectory() as temporary:
            journal=RunJournal(Path(temporary),provider,dict(limits=dict(usd=1,repairs=0,retries=2)))
            with patch.object(provider,'complete',side_effect=partial) as dispatch:
                with self.assertRaisesRegex(AIError,'max_output_tokens'):
                    journal.ask('construction_select','Complete original',{},SELECT)
                self.assertEqual(dispatch.call_count,1)
            self.assertNotIn('result',journal.entries[0]);self.assertTrue(journal.state['usage']['cost_is_complete'])

    def test_output_cap_identity_preserves_legacy_failures_but_distinguishes_changed_limits(self):
        from company_initiation.shared_runtime import saved_output_limit_matches
        response=dict(status='incomplete',incomplete_details=dict(reason='max_output_tokens'),
            usage=dict(output_tokens=32000))
        legacy=dict(response=response)
        self.assertTrue(saved_output_limit_matches(legacy,32000))
        self.assertFalse(saved_output_limit_matches(legacy,16000))
        self.assertTrue(saved_output_limit_matches(dict(legacy,max_output_tokens=30000),30000))
        self.assertFalse(saved_output_limit_matches(dict(legacy,max_output_tokens=30000),32000))
        self.assertFalse(saved_output_limit_matches(dict(response=dict(status='completed',usage=dict(output_tokens=32000))),32000))

    def test_output_cap_is_settled_once_and_never_technically_retried_or_redispatched(self):
        from company_initiation.shared_budget import BudgetLedger
        provider=FixtureProvider({})
        response=dict(status='incomplete',incomplete_details=dict(reason='max_output_tokens'),
            usage=dict(input_tokens=260306,output_tokens=32000,total_tokens=292306,
                output_tokens_details=dict(reasoning_tokens=4690)),
            output=[dict(type='message',content=[dict(type='output_text',text='{"observations":[{"row_id":"sales","terms":[')])])
        def capped(**kwargs):
            provider.last_response=copy.deepcopy(response)
            raise AIError('AI response was incomplete or failed; it cannot be accepted.')
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);first=root/'first';first.mkdir();second=root/'second';second.mkdir()
            ledger=BudgetLedger(root/'budget.json','Synthetic output-cap settlement',1)
            config=dict(limits=dict(usd=1,calls=10,repairs=1,retries=2,generated_tokens=100000))
            journal=RunJournal(first,provider,config,ledger=ledger)
            with patch.object(provider,'complete',side_effect=capped) as dispatch:
                with self.assertRaisesRegex(AIError,'max_output_tokens'):
                    journal.ask('company_year_mapping','Complete originals',{'periods':['FY','H1']},SELECT)
                self.assertEqual(dispatch.call_count,1)
            saved=json.loads((first/'record.json').read_text());before=ledger.summary()
            self.assertEqual(saved['calls'][0]['response'],response)
            self.assertNotIn('result',saved['calls'][0]);self.assertIsNone(saved['pending_call'])
            self.assertEqual(before['settled_requests'],1);self.assertEqual(before['unresolved_reservations_usd'],0)
            self.assertGreater(before['known_cost_usd'],0)
            resumed=RunJournal(second,provider,config,previous=saved,ledger=ledger)
            with patch.object(provider,'complete',side_effect=AssertionError('Must not redispatch capped request')) as dispatch:
                with self.assertRaisesRegex(AIError,'max_output_tokens: exact saved request'):
                    resumed.ask('company_year_mapping','Complete originals',{'periods':['FY','H1']},SELECT)
                dispatch.assert_not_called()
            self.assertEqual(ledger.summary(),before)

    def test_saved_context_failure_is_not_redispatched_or_reserved_again(self):
        from company_initiation.shared_budget import BudgetLedger
        provider=FixtureProvider({})
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);first=root/'first';first.mkdir();second=root/'second';second.mkdir()
            ledger=BudgetLedger(root/'budget.json','Synthetic oversized review',1)
            config=dict(limits=dict(usd=1,calls=10,repairs=1,retries=2,generated_tokens=50000))
            journal=RunJournal(first,provider,config,ledger=ledger)
            with patch.object(provider,'complete',side_effect=AIError('HTTP 400 context_length_exceeded')) as dispatch:
                with self.assertRaisesRegex(AIError,'context_length_exceeded'):
                    journal.ask('construction_select','Original complete evidence',{'original_sections':['complete']},SELECT)
                self.assertEqual(dispatch.call_count,1)
            saved=json.loads((first/'record.json').read_text());before=ledger.summary()
            self.assertGreater(before['unresolved_reservations_usd'],0)
            resumed=RunJournal(second,provider,config,previous=saved,ledger=ledger)
            with patch.object(provider,'complete',side_effect=AssertionError('Oversized request must not be sent again')) as dispatch:
                with self.assertRaisesRegex(AIError,'exact saved request will not be redispatched') as error:
                    resumed.ask('construction_select','Original complete evidence',{'original_sections':['complete']},SELECT)
                dispatch.assert_not_called()
            self.assertIn(saved['calls'][0]['call_id'],str(error.exception))
            self.assertIn(saved['calls'][0]['signature'],str(error.exception))
            self.assertEqual(ledger.summary(),before);self.assertEqual(len(resumed.entries),1)
            # Different payloads are permitted by the ordinary remaining budget.
            with patch.object(provider,'complete',wraps=provider.complete) as dispatch:
                resumed.ask('construction_select','Original complete evidence',{'original_sections':['complete relevant topic']},SELECT)
                self.assertEqual(dispatch.call_count,1)
            self.assertEqual(ledger.summary()['unresolved_reservations_usd'],before['unresolved_reservations_usd'])

    def test_context_failure_guard_requires_matching_durable_request_identity(self):
        from company_initiation.shared_runtime import digest
        from company_initiation.shared_ai import response_usage
        provider=FixtureProvider({})
        with tempfile.TemporaryDirectory() as t:
            journal=RunJournal(Path(t),provider,dict(limits=dict(usd=1,calls=10,repairs=0)))
            signature=dict(stage='construction_select',instructions='target',payload={},schema=SELECT,
                model=DEFAULT_MODEL,effort='low',attempt=0)
            journal.entries.append(dict(signature=digest(signature),call_id='synthetic_original',
                request=dict(signature,payload={'different':'actual failed request'}),error='context_length_exceeded',
                usage=response_usage(dict(usage=dict(input_tokens=0,output_tokens=0,total_tokens=0)),DEFAULT_MODEL)))
            with patch.object(provider,'complete',wraps=provider.complete) as dispatch:
                journal.ask('construction_select','target',{},SELECT)
                self.assertEqual(dispatch.call_count,1)

    def test_source_change_blocks_new_dispatch_and_cached_reuse(self):
        from company_initiation.shared_runtime import _LOADED_CODE_IDENTITY
        provider = FixtureProvider({})
        with tempfile.TemporaryDirectory() as t:
            journal = RunJournal(Path(t), provider, dict(limits=dict(usd=1, repairs=0)))
            journal.ask('construction_select', 'test', {}, SELECT)
            changed = dict(_LOADED_CODE_IDENTITY, integrated_py='changed')
            with patch('company_initiation.shared_runtime._current_code_identity', return_value=changed):
                for prompt in ('test', 'new request'):
                    with self.assertRaisesRegex(AIError, 'Application source changed'):
                        journal.ask('construction_select', prompt, {}, SELECT)
            self.assertEqual(len(provider.requests), 1)
            self.assertEqual(journal.state['status'], 'incomplete')
            self.assertIsNone(journal.state['pending_call'])

    def test_source_change_during_response_settles_once_and_allows_exact_resume(self):
        from company_initiation.shared_runtime import _LOADED_CODE_IDENTITY
        from company_initiation.shared_budget import BudgetLedger
        provider = FixtureProvider({})
        original_complete = provider.complete
        identity = dict(_LOADED_CODE_IDENTITY)
        def completed_then_changed(**kwargs):
            answer = original_complete(**kwargs)
            identity['integrated.py'] = 'changed'
            return answer
        with tempfile.TemporaryDirectory() as t:
            root = Path(t)
            first = root/'first'; first.mkdir()
            second = root/'second'; second.mkdir()
            ledger = BudgetLedger(root/'budget.json', 'Synthetic runtime guard', 1)
            config = dict(limits=dict(usd=1, repairs=0, retries=1))
            journal = RunJournal(first, provider, config, ledger=ledger)
            with patch('company_initiation.shared_runtime._current_code_identity', side_effect=lambda: dict(identity)), \
                    patch.object(provider, 'complete', side_effect=completed_then_changed):
                with self.assertRaisesRegex(AIError, 'Application source changed'):
                    journal.ask('construction_select', 'test', {}, SELECT)
            state = json.loads((first/'record.json').read_text())
            calls = [json.loads(line) for line in (first/'calls.jsonl').read_text().splitlines()]
            self.assertIsNone(state['pending_call'])
            self.assertEqual(calls, state['calls'])
            self.assertIn('result', calls[0])
            self.assertNotIn('error', calls[0])
            self.assertEqual(state['usage']['total_tokens'], 150)
            totals = ledger.summary()
            self.assertEqual(totals['settled_requests'], 1)
            self.assertEqual(totals['unresolved_reservations_usd'], 0)
            self.assertGreater(totals['incurred_estimates_usd'], 0)
            resumed = RunJournal(second, provider, config, previous=state, ledger=ledger)
            with patch.object(provider, 'complete', side_effect=AssertionError('Must reuse saved response')):
                answer = resumed.ask('construction_select', 'test', {}, SELECT)
            self.assertEqual(answer, calls[0]['result'])
            self.assertEqual(ledger.summary()['known_cost_usd'], totals['known_cost_usd'])

    def test_known_charge_retry_is_bounded_and_unknown_charge_never_retries(self):
        p=FixtureProvider({});original=p.complete;count=0
        def once(**kwargs):
            nonlocal count
            count+=1
            if count==1:
                p.last_response={'model':DEFAULT_MODEL,'usage':{'input_tokens':100,'output_tokens':50,'total_tokens':150},'output':[]}
                raise AIError('Synthetic malformed response with known usage')
            return original(**kwargs)
        with tempfile.TemporaryDirectory() as t:
            config=dict(limits=dict(calls=3,repairs=0,retries=1,generated_tokens=50000,usd=1))
            j=RunJournal(Path(t),p,config)
            with patch.object(p,'complete',side_effect=once):
                j.ask('construction_select','test',{},SELECT)
            self.assertEqual([e['retry_count'] for e in j.entries],[0,1])
            self.assertEqual(j.state['usage']['total_tokens'],300)
        with tempfile.TemporaryDirectory() as t:
            j=RunJournal(Path(t),p,config)
            with patch.object(p,'complete',side_effect=AIError('Unknown transport charge')) as call:
                with self.assertRaisesRegex(AIError,'unknown usage'):
                    j.ask('construction_select','test',{},SELECT)
                self.assertEqual(call.call_count,1)
            self.assertFalse(j.state['usage']['cost_is_complete'])

    def test_caller_mutation_cannot_rewrite_durable_requests_or_results(self):
        p=FixtureProvider({})
        with tempfile.TemporaryDirectory() as t:
            j=RunJournal(Path(t),p,dict(limits=dict(calls=2,repairs=1,generated_tokens=50000,usd=1)))
            payload={'scope':{'gaps':[]}}
            answer=j.ask('construction_select','test',payload,SELECT)
            payload['scope']['gaps'].append('Later scope update')
            answer['gaps'].append('Caller mutation')
            j.save()
            state=json.loads((Path(t)/'record.json').read_text())
            saved=[json.loads(line) for line in (Path(t)/'calls.jsonl').read_text().splitlines()]
            self.assertEqual(saved,state['calls'])
            self.assertEqual(saved[0]['request']['payload']['scope']['gaps'],[])
            self.assertNotIn('Caller mutation',saved[0]['result']['gaps'])
    def test_stage_routing_defaults_and_no_silent_sol(self):
        p=OpenAIProvider(api_key="SYNTHETIC")
        self.assertEqual(p.route("extractor","construction_map"),(DEFAULT_MODEL,"low"))
        self.assertEqual(p.route("reviewer","construction_review"),(DEFAULT_MODEL,"medium"))
        self.assertEqual(set(p.models.values()),{DEFAULT_MODEL})

    def test_unknown_charge_and_call_limits_block(self):
        p=FixtureProvider({})
        with tempfile.TemporaryDirectory() as t:
            config=dict(limits=dict(calls=1,repairs=0,generated_tokens=50000,usd=1),escalations={})
            j=RunJournal(Path(t),p,config)
            j.ask("construction_select","test",{},SELECT)
            with self.assertRaisesRegex(AIError,"limit"):
                j.ask("construction_select","different",{},SELECT)
            self.assertEqual(len(p.requests),1)
            previous=dict(j.state,pending_call={"id":"unknown"})
            with self.assertRaisesRegex(AIError,"unknown charge"):
                RunJournal(Path(t),p,config,previous=previous)

    def test_expensive_route_requires_documented_cheaper_failures(self):
        p=FixtureProvider({});p.stage_models["construction_select"]="gpt-6.1-sol"
        with tempfile.TemporaryDirectory() as t:
            j=RunJournal(Path(t),p,dict(limits=dict(calls=2,repairs=1,generated_tokens=50000,usd=1)))
            with self.assertRaisesRegex(AIError,"Non-Luna"):
                j.ask("construction_select","test",{},SELECT)
            self.assertFalse(p.requests)

    def test_mock_provider_response_usage_and_pending_are_saved(self):
        p=FixtureProvider({})
        with tempfile.TemporaryDirectory() as t:
            j=RunJournal(Path(t),p,dict(limits=dict(calls=2,repairs=1,generated_tokens=50000,usd=1)))
            j.ask("construction_select","test",{},SELECT)
            state=json.loads((Path(t)/"record.json").read_text())
            self.assertIsNone(state["pending_call"])
            self.assertEqual(state["usage"]["total_tokens"],150)
            self.assertEqual(state["calls"][0]["reasoning_effort"],"low")
            self.assertEqual(state["calls"][0]["repair_count"],0)


def evaluate_luna(root):
    """Explicit paid developer benchmark using production prompts, schemas and tools.

    Called manually after offline tests, never during unittest discovery. Fixtures
    and reports are generated under outputs and are visibly synthetic.
    """
    from datetime import datetime, timezone
    from company_initiation.forecast_construction.construction import FORMULA_PROMPT, FORMULAS, metric_catalog, _valid_spec, _initiation_lock
    from company_initiation.shared_runtime import HISTORY_PROMPT, MAPPING_BATCH, write_json
    root=Path(root).resolve()
    directory=root/"outputs"/("luna_known_answer_evaluation_"+datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"))
    directory.mkdir()
    archive=directory/"inputs"; (archive/"visible_alpha").mkdir(parents=True)
    (archive/'financial_reports').mkdir()
    (archive/'financial_reports'/'SYNTHETIC-disclosures.txt').write_text(
        'SYNTHETIC Manufacturing Group: calendar fiscal years 2024 and 2025; H1 runs January 1 to June 30. '
        'Net turnover means group revenue in USD millions: 2024 = 111; 2025 = 110; 2024 H1 = 45. '
        'Manufacturing sales is the sole segment: 2024 = 100; 2025 = 110; 2024 H1 = 45. '
        'The 2024 mismatch is intentional and must remain. Operating margin 2024 = 5 percent; 2025 = 6 percent. '
        'The discontinued division is missing, not zero. '
        'SEPARATE SYNTHETIC FORMULA CASE, BASE USD: Revenue 2025 = 100, costs = 60, marketing = 10, profit = 30. '
        'Revenue is prior annual revenue times one plus price growth times one plus volume growth. '
        'Costs are prior costs times one plus volume growth times one plus unit-cost growth. '
        'Marketing stays at its prior annual level. Profit is revenue minus costs minus marketing. '
        'Costs do not affect pricing. All lags are one year.',encoding='utf-8')
    fixtures.VisibleAlphaTests.make_workbook(archive/"visible_alpha"/"SYNTHETIC.xlsx")
    book=load_workbook(archive/"visible_alpha"/"SYNTHETIC.xlsx")
    sheet=book["Vendor data"]
    sheet['D28']='Net turnover (Group sales)'
    sheet['I28']=111  # Planted segment discrepancy; source must not be changed.
    sheet['D30']='Operating margin';sheet['E30']='%';sheet['F30']='Analytics'
    sheet['I30']=5;sheet['J30']=6
    book.save(archive/"visible_alpha"/"SYNTHETIC.xlsx");book.close()
    store=SourceStore(directory/"sources")
    for source in inventory_local_sources(archive)["sources"]:store.add_local(source,root=archive)
    evidence=store.read_documents([dict(source_id=s['id'],unit_id=next(iter(store._units[s['id']])))
                                  for s in store.sources if s.get('category')=='financial_reports'])
    provider=OpenAIProvider.from_env(root/".env");provider.max_output_tokens=6000
    from company_initiation.__main__ import TerminalProgress
    config=dict(limits=dict(usd=.5,calls=5,repairs=1,generated_tokens=30000),escalations={},
                purpose="SYNTHETIC known-answer evaluation; not a company forecast",provider=provider.metadata())
    with _initiation_lock(root/"outputs"):
        journal=RunJournal(directory,provider,config,TerminalProgress())
        try:
            mapping=vendor_mapping(store)
            payload=dict(company="SYNTHETIC Manufacturing Group",years=[2024,2025],
                reporting_scope={k:v for k,v in mapping.items() if k not in ('rows','identities')},
                vendor_catalog=metric_catalog(store,[2024,2025]),
                evidence=evidence,
                requested_section="income_statement", existing_rows=[],
                task="Map net turnover to revenue, Manufacturing sales to segment_sales, Operating margin to margin and the missing discontinued division to discontinued_sales. Use the corresponding vendor rows only; preserve all four lines. Return rows and gaps, no numerical tables.")
            batch=journal.stage('construction_map',HISTORY_PROMPT,payload,MAPPING_BATCH)
            mapped=dict(mapping,rows=batch['rows'])
            history=normalise_history(mapped,store,[2024,2025])
            by_id={r['id']:r for r in history['rows']}
            results=[dict(case='label_alias_and_source_value',passed=by_id.get('revenue',{}).get('values',{}).get('2024')==111e6),
                dict(case='units_and_percent',passed=by_id.get('margin',{}).get('values',{}).get('2024')==.05),
                dict(case='estimate_rejection_and_annual_vs_cumulative',passed=by_id.get('revenue',{}).get('values',{}).get('2024-H1')==45e6 and by_id.get('revenue',{}).get('values',{}).get('2025')==110e6),
                dict(case='missing_not_zero',passed=not by_id.get('discontinued_sales',{}).get('values',{'bad':0})),
                dict(case='segment_error_retained',passed=any(c['status']=='FAIL' and c['difference']==11e6 for c in history['checks']))]
            journal.artifact('mapped_history.json',history)
            h,d,s=known_model()
            broken=copy.deepcopy(s)
            broken['rows'][0]['expression']='revenue * (1 + price) * (1 + volume) * (1 + price)'
            broken['rows'][1]['expression']='lag("costs",1) * (1 - unit_cost)'
            repaired=journal.stage('construction_formulas',FORMULA_PROMPT,
                dict(company=h['company'],historical_model=h,drivers=d,previous_specification=broken,
                     economic_errors=['revenue refers to itself circularly and applies price twice','cost growth has wrong direction and omits activity growth'],
                     evidence=evidence),
                FORMULAS,lambda value: value if not validate_spec(h,d,value) else (_ for _ in ()).throw(ValueError('; '.join(validate_spec(h,d,value)))))
            a=initialise_assumptions(h,d,[2026,2027])
            for item in a:item['value']={'price':.1,'volume':.2,'unit_cost':.05}[item['driver']]
            values=calculate(h,d,repaired,a,[2026,2027])
            results.append(dict(case='repaired_economic_formula_and_double_count_cycle',passed=abs(values.get('revenue',{}).get('2026',0)-132)<1e-8 and abs(values.get('costs',{}).get('2026',0)-75.6)<1e-8 and abs(values.get('profit',{}).get('2026',0)-46.4)<1e-8))
            checks=validate_results(h,d,repaired,a,[2026,2027],values)
            results.append(dict(case='sensitivity_direction',passed=not any(c['status']=='FAIL' and c['area']=='behaviour' for c in checks)))
            journal.artifact('evaluation.json',dict(cases=results,checks=checks,limitations='Two focused tasks, single trial; no statistical accuracy estimate or real-company validation'))
            journal.state['status']='passed' if all(c['passed'] for c in results) else 'failed_cases'
            journal.save()
        except Exception as exc:
            journal.state['status']='incomplete';journal.state['issues'].append(str(exc));journal.save()
    return directory


if __name__=="__main__":unittest.main()
