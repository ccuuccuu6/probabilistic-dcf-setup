"""Independent FX arithmetic, persistence and synthetic company integration tests."""
import copy
from datetime import date,datetime,timedelta,timezone
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook,load_workbook
from research_platform.fx import forecast_rate,publication_dates,spot_quote,exposure_quality,currency_coverage,transaction_fx_for_period,apply_fx_accounting_contracts
from research_platform.fx_store import FXStore
from research_platform.ibkr import BrokerError,overlap_check,update_fx,validate_spot
from research_platform.integrated import build_periodic_spec,periodic_sensitivities,financial_cost_pools
from research_platform.periods import calculate_periodic,validate_periodic,cell_dependencies
from research_platform.model_excel import export_integrated
from research_platform.evidence import SourceStore,_vendor_profile,inventory_local_sources
from tests.test_fx import source_workbook
from tests.test_periodic_model import different_period_company


def quote(day,value,base='EUR',currency='GBP',**extra):
    return dict(provider='SYNTHETIC',base=base,quote=currency,day=day,source_date=day,value=value,
        definition='reference_fixing',status='actual',timezone='UTC',**extra)


def spot(stamp,value,base='EUR',currency='GBP'):
    return dict(provider='IBKR',base=base,quote=currency,day=stamp[:10],source_date=stamp[:10],value=value,
        definition='bid_ask_midpoint',status='actual',timezone='UTC',timestamp=stamp,
        bid=value-.001,ask=value+.001,market_data_type='live')


class SharedRateTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.store=FXStore(Path(self.temp.name)/'fx.sqlite');self.addCleanup(self.store.db.close)

    def test_empty_accounting_profile_without_rates_is_explicitly_incomplete(self):
        spec={};data=dict(financial_contracts=dict(fx_accounting=dict(exposures=[])))
        apply_fx_accounting_contracts(spec,data,None)
        self.assertEqual(spec['fx_accounting_coverage']['market_snapshot'],'missing')
        self.assertEqual(spec['fx_accounting_coverage']['company_exposures'],'unquantified')
        self.assertEqual(spec['fx_validation'][0]['status'],'WARNING')
        self.assertNotIn('fx_accounting_contract',spec)
        data['financial_contracts']['fx_accounting']['exposures']=[dict(id='disclosed_cash')]
        with self.assertRaisesRegex(ValueError,'company exposure is distinct from missing market data'):
            apply_fx_accounting_contracts({},data,None)

    def test_workbook_bootstrap_idempotence_and_projected_summary_separation(self):
        path=Path(self.temp.name)/'original.xlsx';source_workbook(path)
        first=self.store.import_workbook(path);second=self.store.import_workbook(path)
        self.assertGreater(first['added'],0);self.assertEqual(second['added'],0)
        self.assertEqual(second['duplicates'],first['added'])
        periods=[json.loads(r[0]) for r in self.store.db.execute('SELECT payload FROM period_snapshots')]
        self.assertTrue(any(p['source_status']=='estimate' for p in periods))
        self.store.import_quotes([dict(quote('2024-09-30',9),status='projected')],{'synthetic_projection':True})
        snapshot=self.store.snapshot('2024-10-01')
        self.assertNotIn('2024-09-30',snapshot['daily_quotes'])

    def test_partial_period_is_four_actual_plus_six_spot_not_spot_backfill(self):
        days=['2026-01-01','2026-01-02','2026-01-05','2026-01-06','2026-01-07','2026-01-08','2026-01-09','2026-01-12','2026-01-13','2026-01-14']
        self.store.import_quotes([quote(d,1) for d in days[:4]]+[spot('2026-01-06T12:00:00+00:00',1.10)],{'fixture':'partial'})
        snapshot=self.store.snapshot('2026-01-06T12:00:00+00:00',{'calendar':{'dates':days}})
        result=forecast_rate(snapshot,'GBP','EUR',days[0],days[-1])
        self.assertAlmostEqual(result['value'],1.06);self.assertEqual(result['coverage']['actual_days'],4)
        self.assertEqual(result['coverage']['projected_days'],6)
        later=self.store.snapshot('2026-01-08T12:00:00+00:00',{'calendar':{'dates':days},'spot_max_age_seconds':300000})
        missing=forecast_rate(later,'GBP','EUR',days[0],days[-1])
        self.assertIsNone(missing['value']);self.assertIn('2026-01-07',missing['coverage']['missing_dates'])

    def test_daily_crosses_then_average_and_inverse_not_inverse_of_mean(self):
        days=['2025-01-02','2025-01-03']
        self.store.import_quotes([quote(days[0],2),quote(days[1],4),quote(days[0],2,currency='JPY'),quote(days[1],8,currency='JPY')],{'fixture':'crosses'})
        s=self.store.snapshot('2025-01-03',{'calendar':{'dates':days}})
        self.assertAlmostEqual(forecast_rate(s,'GBP','JPY',days[0],days[-1])['value'],.75)
        self.assertAlmostEqual(forecast_rate(s,'JPY','GBP',days[0],days[-1])['value'],1.5)
        self.assertNotAlmostEqual(.75,3/5)

    def test_future_horizon_and_nonzero_yoy_under_flat_spot(self):
        self.store.import_quotes([quote('2024-01-02',1),quote('2025-01-02',1.1)],{'fixture':'flat'})
        policy={'calendar':{'dates':['2024-01-02','2025-01-02','2026-01-02','2030-01-02']}}
        s=self.store.snapshot('2025-01-02',policy)
        first=forecast_rate(s,'GBP','EUR','2025-01-01','2025-03-31')['value']
        prior=forecast_rate(s,'GBP','EUR','2024-01-01','2024-03-31')['value']
        self.assertAlmostEqual(first/prior-1,.1)
        self.assertAlmostEqual(forecast_rate(s,'GBP','EUR','2030-01-01','2030-12-31')['value'],1.1)

    def test_stale_unavailable_and_carried_observations_are_explicit(self):
        self.store.import_quotes([quote('2025-01-02',1.1),dict(quote('2025-01-03',1.1),source_date='2025-01-02',status='carried')],{'fixture':'carry'})
        s=self.store.snapshot('2025-02-03')
        self.assertIsNone(forecast_rate(s,'GBP','EUR','2025-02-04','2025-03-01')['value'])
        self.assertIsNone(forecast_rate(s,'GBP','XYZ','2025-02-04','2025-03-01')['value'])
        self.assertIsNone(forecast_rate(s,'GBP','EUR','2025-01-03','2025-01-03')['value'])
        with self.assertRaises(ValueError):self.store.import_quotes([dict(quote('2025-01-04',1),source_date='2025-01-02')],{'bad':True})

    def test_revision_quarantine_and_immutable_snapshot(self):
        self.store.import_quotes([quote('2025-01-02',1)],{'version':1});s=self.store.snapshot('2025-01-02')
        changed=self.store.import_quotes([quote('2025-01-02',100)],{'version':2})
        self.assertEqual(changed['quarantined'],1)
        self.assertEqual(self.store.get_snapshot(s['source']['snapshot_id']),s)
        self.assertEqual(self.store.snapshot('2025-01-02')['daily_quotes']['2025-01-02'][0]['value'],1)

    def test_repeated_frozen_quote_is_idempotent_despite_retrieval_age(self):
        q=spot('2025-01-03T20:59:59+00:00',1.1)
        q.update(market_data_type='frozen',age_seconds=100)
        self.assertEqual(self.store.import_quotes([q],{'retrieval':1})['added'],1)
        q['age_seconds']=10000
        self.assertEqual(self.store.import_quotes([q],{'retrieval':2})['added'],0)
        q['value']=9
        with self.assertRaisesRegex(ValueError,'midpoint'):self.store.import_quotes([q],{'bad':True})

    def test_provider_overlap_mismatch_and_bid_ask_validation(self):
        old=[dict(quote(f'2025-01-0{d}',1),id=str(d)) for d in (2,3,4)]
        new=[dict(quote(f'2025-01-0{d}',100),provider='IBKR',definition='MIDPOINT_daily_close') for d in (2,3,4)]
        self.assertEqual(overlap_check(old,new,'EUR','GBP')['status'],'quarantined')
        now=datetime(2025,1,6,12,tzinfo=timezone.utc)
        for changes in ({'bid':2,'ask':1},{'market_data_type':'delayed'},{'timestamp':'2024-01-06T12:00:00+00:00'}):
            q=dict(bid=1,ask=1.1,timestamp=now.isoformat(),market_data_type='live');q.update(changes)
            with self.assertRaises(BrokerError):validate_spot(q,now)

    def test_incremental_update_preserves_good_pairs_and_never_calls_llm(self):
        days=['2025-01-02','2025-01-03','2025-01-06','2025-01-07']
        self.store.import_quotes([quote(d,1) for d in days],{'fixture':'broker'})
        class Client:
            session_timezone='UTC'
            def resolve(self,a,b):
                if b=='XYZ':raise BrokerError('Unavailable contract')
                return dict(base=a,quote=b,con_id=1)
            def history(self,r,start,end):
                self.request=(start,end)
                return [dict(day=d,value=1) for d in days]+[dict(day='2025-01-08',value=1.01)]
            def spot(self,r):return dict(bid=1.009,ask=1.011,value=1.01,timestamp='2025-01-09T12:00:00+00:00',market_data_type='live')
        client=Client()
        with patch('research_platform.ai.OpenAIProvider.from_env',side_effect=AssertionError('No LLM')):
            report=update_fx(self.store,client,['EUR/GBP','EUR/XYZ'],as_of=datetime(2025,1,9,12,tzinfo=timezone.utc),pacing_seconds=0)
        self.assertEqual(report['ai_calls'],0);self.assertEqual(report['exit_code'],2)
        self.assertEqual(client.request,('2024-12-31','2025-01-08'))
        self.assertTrue(any(q['provider']=='IBKR' and q['day']=='2025-01-08' for q in self.store.quotes()))
        self.assertEqual(report['pairs'][1]['status'],'unavailable_contract')

    def test_unavailable_direct_contract_uses_native_legs_for_cross(self):
        days=['2025-01-02','2025-01-03','2025-01-06']
        self.store.import_quotes([quote(d,.8,base='USD') for d in days]+[quote(d,150,base='USD',currency='JPY') for d in days],{'fixture':'cross legs'})
        class Client:
            session_timezone='UTC'
            def resolve(self,a,b):
                if set((a,b))=={'GBP','JPY'}:raise BrokerError('No direct contract')
                currency=b if a=='USD' else a
                return dict(base='USD',quote=currency,con_id=currency)
            def history(self,r,start,end):return [dict(day=d,value=.8 if r['quote']=='GBP' else 150) for d in days]
            def spot(self,r):
                v=.8 if r['quote']=='GBP' else 150
                return dict(bid=v-.001,ask=v+.001,value=v,timestamp='2025-01-07T12:00:00+00:00',market_data_type='live')
        report=update_fx(self.store,Client(),['GBP/JPY'],as_of=datetime(2025,1,7,12,tzinfo=timezone.utc),pacing_seconds=0)
        self.assertEqual(report['pairs'][0]['status'],'derived_cross')
        self.assertEqual(report['exit_code'],0)
        s=self.store.snapshot('2025-01-07T12:00:00+00:00')
        self.assertAlmostEqual(forecast_rate(s,'GBP','JPY','2025-01-08','2025-01-08')['value'],.8/150)
        self.assertFalse(any(q['base']=='GBP' and q['quote']=='JPY' for q in self.store.quotes()))

    def test_validated_direct_pair_precedes_longer_cross_and_keeps_lineage(self):
        day='2025-01-02'
        self.store.import_quotes([quote(day,1.2),quote(day,1.3,currency='USD'),quote(day,.9,base='USD')],{'fixture':'direct preference'})
        snapshot=self.store.snapshot(day,{'calendar':{'dates':[day]}})
        direct=forecast_rate(snapshot,'GBP','EUR',day,day)
        inverse=forecast_rate(snapshot,'EUR','GBP',day,day)
        self.assertAlmostEqual(direct['value'],1.2)
        self.assertAlmostEqual(direct['value']*inverse['value'],1)
        self.assertEqual(direct['rate_types'],['direct'])
        self.assertEqual(direct['selected_series'],['SYNTHETIC:EUR/GBP:reference_fixing'])

    def test_inconsistent_direct_cross_is_flagged_not_silently_selected(self):
        day='2025-01-02'
        self.store.import_quotes([quote(day,1.2),quote(day,1.3,currency='USD'),quote(day,.5,base='USD')],{'fixture':'inconsistent'})
        snapshot=self.store.snapshot(day,{'calendar':{'dates':[day]}})
        self.assertTrue(snapshot['findings'])
        self.assertIsNone(forecast_rate(snapshot,'GBP','EUR',day,day)['value'])
        self.assertIsNone(spot_quote(snapshot,'GBP','EUR')['value'])

    def test_cross_spot_requires_aligned_timestamps_and_keeps_actual_currencies(self):
        self.store.import_quotes([spot('2025-01-07T11:30:00+00:00',1.2),spot('2025-01-07T12:00:00+00:00',150,currency='JPY')],{'fixture':'unaligned'})
        snapshot=self.store.snapshot('2025-01-07T12:00:00+00:00',{'spot_max_age_seconds':3600,'spot_alignment_seconds':900})
        result=spot_quote(snapshot,'GBP','JPY')
        self.assertIsNone(result['value'])
        self.assertEqual(result['timestamp_spread_seconds'],1800)
        self.assertEqual(result['rate_type'],'cross_rate')

    def test_stale_direct_spot_does_not_block_valid_same_currency_fixing(self):
        self.store.import_quotes([spot('2025-01-06T12:00:00+00:00',1.1),quote('2025-01-07',1.2)],{'fixture':'stale direct'})
        snapshot=self.store.snapshot('2025-01-07T12:00:00+00:00')
        result=spot_quote(snapshot,'GBP','EUR')
        self.assertAlmostEqual(result['value'],1.2)
        self.assertEqual(result['status'],'dated_fixing_proxy')
        self.assertEqual(result['rate_type'],'direct')

    def test_timestamp_aligned_cross_can_span_midnight_without_mixing_fixing_days(self):
        self.store.import_quotes([spot('2025-01-06T23:59:00+00:00',1.2),spot('2025-01-07T00:01:00+00:00',150,currency='JPY')],{'fixture':'midnight'})
        snapshot=self.store.snapshot('2025-01-07T00:02:00+00:00')
        result=spot_quote(snapshot,'GBP','JPY')
        self.assertAlmostEqual(result['value'],1.2/150)
        self.assertEqual(result['timestamp_spread_seconds'],120)

    def test_currency_coverage_distinguishes_held_native_history_from_selected_series(self):
        days=['2025-01-02','2025-01-03','2025-01-06']
        qs=[quote(d,1.2) for d in days]
        qs += [dict(quote(d,7,base='USD',currency='CNH'),provider='IBKR',definition='MIDPOINT_daily_close') for d in days]
        self.store.import_quotes(qs,{'fixture':'held transition'})
        snapshot=self.store.snapshot('2025-01-06',{'calendar':{'dates':days}})
        coverage={c['currency']:c for c in currency_coverage(snapshot,'GBP',['EUR','CNH'])}
        self.assertEqual(coverage['EUR']['last_genuine_observation'],'2025-01-06')
        self.assertIsNone(coverage['CNH']['last_genuine_observation'])
        self.assertEqual(coverage['CNH']['available_native_series'][0]['last'],'2025-01-06')
        self.assertEqual(coverage['CNH']['coverage_gap_count'],3)

    def test_exposure_quality_never_calls_proxy_a_cross_or_hides_residual(self):
        basket=dict(method='Fixture',source_refs=[{'source_id':'synthetic'}],effective_date='2025-01-01',
            weights=[dict(currency='USD',weight=.6,exposure_status='estimated',evidence_basis='functional_currency_estimate'),
                dict(currency='GBP',weight=0,exposure_status='reported'),
                dict(currency='UNALLOCATED',rate_currency='GBP',weight=.4,scenario_neutral=True,exposure_status='unresolved')])
        quality=exposure_quality(basket,'GBP')
        self.assertEqual(quality['proxy_weight'],0)
        self.assertEqual(quality['weak_evidence_weight'],.4)
        self.assertTrue(quality['material_weak_exposure'])
        self.assertEqual(quality['reported_weight'],0)
        self.assertEqual(quality['estimated_weight'],.6)
        self.assertEqual(quality['scenario_weight'],0)
        self.assertEqual(quality['unresolved_weight'],.4)
        # A cited customer-location split still does not establish functional
        # currency exposure; its whole weight is a weak geographic proxy.
        basket['weights'][0]['evidence_basis']='customer_geography_proxy'
        geographical=exposure_quality(basket,'GBP')
        self.assertEqual(geographical['weak_evidence_weight'],1)
        self.assertEqual(geographical['estimated_weight'],.6)
        # A separately supported functional-currency estimate remains distinct
        # from weak geography, scenario, and unresolved assumptions.
        mixed=dict(basket,weights=[dict(currency='USD',weight=.4,exposure_status='estimated',evidence_basis='customer_geography_proxy'),
            dict(currency='EUR',weight=.2,exposure_status='estimated',evidence_basis='functional_currency_estimate'),
            dict(currency='GBP',weight=.1,exposure_status='reported',evidence_basis='functional_currency_disclosure'),
            dict(currency='JPY',weight=.1,exposure_status='scenario',evidence_basis='scenario'),
            dict(currency='UNALLOCATED',rate_currency='GBP',weight=.2,scenario_neutral=True,exposure_status='unresolved')])
        measured=exposure_quality(mixed,'GBP')
        self.assertAlmostEqual(measured['weak_evidence_weight'],.7)
        self.assertAlmostEqual(measured['reported_weight'],.1)
        self.assertAlmostEqual(measured['estimated_weight'],.6)
        self.assertAlmostEqual(measured['scenario_weight'],.1)
        self.assertAlmostEqual(measured['unresolved_weight'],.2)
        basket['weights'][2]=dict(currency='ARS',rate_currency='BRL',weight=.4,proxy_reason='Scenario only',rate_type='cross_rate')
        with self.assertRaisesRegex(ValueError,'proxy, never a cross'):exposure_quality(basket,'GBP')

    def test_procurement_fx_known_answers_for_currency_identity_lag_and_hedge(self):
        rates={'2023-12-04':5.,'2024-01-02':5.,'2024-12-02':5.5,'2025-01-02':6.}
        self.store.import_quotes([quote(d,v,base='USD',currency='BRL') for d,v in rates.items()],{'fixture':'procurement'})
        snapshot=self.store.snapshot('2025-01-02',{'calendar':{'dates':list(rates)}})
        component=dict(weight=1.,functional_currency='BRL',invoice_currency='USD')
        baseline=transaction_fx_for_period(snapshot,component,'2025-01-01','2025-01-31')
        self.assertAlmostEqual(baseline['value'],.2)
        lagged=transaction_fx_for_period(snapshot,dict(component,lag_months=1),'2025-01-01','2025-01-31')
        self.assertAlmostEqual(lagged['value'],.1)
        half_hedged=transaction_fx_for_period(snapshot,dict(component,hedge_coverage=.5,hedge_rate=5),'2025-01-01','2025-01-31')
        self.assertAlmostEqual(half_hedged['value'],.1)
        hedged=transaction_fx_for_period(snapshot,dict(component,hedge_coverage=1,hedge_rate=5),'2025-01-01','2025-01-31')
        self.assertEqual(hedged['value'],0)
        independently_hedged=transaction_fx_for_period(snapshot,dict(component,hedge_coverage=.5,hedge_rate=4,
            prior_hedge_coverage=.5,prior_hedge_rate=6),'2025-01-01','2025-01-31')
        self.assertAlmostEqual(independently_hedged['value'],5/5.5-1)
        current_only=transaction_fx_for_period(snapshot,dict(component,hedge_coverage=1,hedge_rate=4),
            '2025-01-01','2025-01-31')
        self.assertAlmostEqual(current_only['value'],-.2)
        prior_only=transaction_fx_for_period(snapshot,dict(component,hedge_coverage=0,hedge_rate=None,
            prior_hedge_coverage=1,prior_hedge_rate=4),'2025-01-01','2025-01-31')
        self.assertAlmostEqual(prior_only['value'],.5)
        dated=transaction_fx_for_period(snapshot,dict(component,hedge_terms=[
            dict(start='2024-01-01',end='2024-01-31',hedge_coverage=1,hedge_rate=4),
            dict(start='2025-01-01',end='2025-01-31',hedge_coverage=1,hedge_rate=5)]),'2025-01-01','2025-01-31')
        self.assertAlmostEqual(dated['value'],.25)
        local=transaction_fx_for_period(snapshot,dict(component,invoice_currency='BRL'),'2025-01-01','2025-01-31')
        self.assertEqual(local['value'],0)
        with self.assertRaisesRegex(ValueError,'requires a positive'):
            transaction_fx_for_period(snapshot,dict(component,hedge_coverage=.5),'2025-01-01','2025-01-31')


class CurrencyCompanyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.addClassCleanup(cls.temp.cleanup)
        cls.data=different_period_company();cls.data['model_configuration'].update(full_statement_frequency='half_year',primary_sales_frequency='quarterly')
        # The different company retains GBP, June year-end, two components
        # businesses and a published predecessor split. This is fixture evidence.
        weights=[dict(currency='JPY',weight=.7),dict(currency='GBP',weight=.3)]
        basket=dict(weights=weights,method='SYNTHETIC disclosed currency sales mix',status='provisional',limitations='Fixture only',
            source_refs=[dict(source_id='synthetic_primary_segment_note',unit_id='sheet:mix')])
        cls.data['model_configuration']['fx_exposures']={s['id']:dict(revenue=copy.deepcopy(basket),cost=copy.deepcopy(basket),
            revenue_bridge='translation_of_constant_currency') for s in cls.data['segments']}
        cls.data['model_configuration']['fx_currencies']=[]
        pools=financial_cost_pools(cls.data)
        for pool in pools:
            own=copy.deepcopy(basket)
            if pool['id']=='marketing_admin':own['weights']=[dict(currency='GBP',weight=1.)]
            own['effective_date']='2025-07-01'
            pool['fx']=dict(translation=own,transaction=dict(components=[],status='unmodelled',method='Synthetic unmodelled procurement FX scenario'))
        cls.data['financial_contracts']=dict(cost_pools=pools)
        # Keep actual company FX as a separate reported fact.
        template=next(o for o in cls.data['observations'] if o['row_id']=='north_pricing' and o['kind']=='FY')
        cls.data['observations'].append(dict(template,id='synthetic_reported_fx',row_id='north_fx',definition_id='north_fx',value=-.03))
        with FXStore(Path(cls.temp.name)/'fx.sqlite') as store:
            qs=[]
            for day in publication_dates('2023-07-01','2025-06-30'):
                qs.extend([quote(day,1.2),quote(day,150 if day<'2025-01-01' else 180,currency='JPY')])
            store.import_quotes(qs,{'fixture':'different company'});cls.snapshot=store.snapshot('2025-06-30')
        cls.spec=build_periodic_spec(cls.data,2,cls.snapshot);cls.values=calculate_periodic(cls.spec)

    def test_mixed_frequency_has_no_quarter_statements_and_balances_integrate(self):
        rows={r['id']:r for r in self.spec['rows']}
        for rid in ('operating_profit','cfo','cash','north_cash_cost','capex_ppe_ratio','borrowing_input'):
            for q in ('Q1','Q2','Q3','Q4'):
                cell=rows[rid]['cells']['2026-'+q]
                self.assertIsNone(self.values[rid]['2026-'+q]);self.assertNotIn('expression',cell);self.assertFalse(cell.get('editable'))
        self.assertIsNone(self.values['cash']['2026-H2'])

    def test_reported_sales_without_fx_keeps_transitive_growth_diagnostics_optional_only(self):
        data=copy.deepcopy(self.data)
        for rid,amount in [('north_sales',16.),('specialty_sales',12.8),('sales',28.8)]:
            original=next(o for o in data['observations'] if o['row_id']==rid and o['period_id']=='2025-Q1')
            data['observations'].append(dict(original,id='current_'+rid,period_id='2026-Q1',fiscal_year=2026,
                start='2025-07-01',end='2025-09-30',value=amount,source_refs=[dict(source_id='synthetic_current_sales',raw=amount)]))
        spec=build_periodic_spec(data,2,self.snapshot);values=calculate_periodic(spec);rows={r['id']:r for r in spec['rows']}
        self.assertEqual(values['north_sales']['2026-Q1'],16.)
        self.assertIsNone(values['north_fx']['2026-Q1'])
        self.assertEqual(rows['north_fx']['cells']['2026-Q1']['status'],'not_disclosed')
        for rid in ('north_fx_weighted_amount','north_fx'):
            for pid in ('2026-H1','2026-FY'):
                self.assertIsNone(values[rid][pid])
                self.assertFalse(rows[rid]['cells'][pid].get('required',True))
        # A supported denominator in the same period is not made unavailable by
        # another row's missing FX; only actual dependency edges propagate it.
        self.assertIsNotNone(values['north_fx_comparable_revenue']['2026-H1'])
        self.assertIsNotNone(values['sales']['2026-FY'])
        self.assertIsNotNone(values['cfo']['2026-H1'])
        self.assertFalse([c for c in validate_periodic(spec,values) if c['status'] in ('FAIL','MISSING')])
        broken=copy.deepcopy(spec);broken_rows={r['id']:r for r in broken['rows']}
        broken_rows['cfo']['cells']['2026-H1']=dict(value=None,status='forecast',required=True,reason='Required financial output')
        failures=validate_periodic(broken,calculate_periodic(broken))
        self.assertTrue(any(c['name']=='cfo:2026-H1' and c['status']=='MISSING' for c in failures))

    def accounting_case(self,exposures,rate_change='2025-06-30'):
        data=copy.deepcopy(self.data);data['financial_contracts']['fx_accounting']=dict(exposures=exposures)
        quotes=[]
        for day in publication_dates('2023-07-01','2025-12-31'):
            quotes.extend([quote(day,1.2),quote(day,150 if day>rate_change or day<'2025-01-01' else 180,currency='JPY')])
        with FXStore(Path(self.temp.name)/('accounting_'+rate_change+'.sqlite')) as store:
            store.import_quotes(quotes,{'fixture':'independent accounting rates'});snapshot=store.snapshot('2025-12-31')
        spec=build_periodic_spec(data,2,snapshot)
        return spec,calculate_periodic(spec)

    def test_nonzero_cash_balance_oci_and_monetary_fx_have_distinct_counterparts(self):
        exposures=[dict(id='foreign_cash',type='foreign_operation',role='cash',currency='JPY',opening_share=.3),
            dict(id='monetary_cash',type='monetary',role='cash',currency='JPY',opening_share=.2),
            dict(id='foreign_inventory',type='foreign_operation',role='inventories',currency='JPY',opening_share=.4),
            dict(id='foreign_payables',type='foreign_operation',role='trade_payables',currency='JPY',opening_share=.5),
            dict(id='foreign_debt',type='foreign_operation',role='nonlease_debt',currency='JPY',opening_share=.1),
            dict(id='monetary_debt',type='monetary',role='nonlease_debt',currency='JPY',opening_share=.25)]
        spec,value=self.accounting_case(exposures);base_spec,base=self.accounting_case([]);pid='2026-H1'
        # GBP/JPY increases from 1.2/180 to 1.2/150: exactly +20%.
        cash_gain=29.5*.2*.2;cash_oci=29.5*.3*.2;inventory_fx=30*.4*.2;debt_loss=58*.25*.2
        payables_fx=20*.5*.2;debt_translation=58*.1*.2
        self.assertAlmostEqual(value['cash_fx'][pid],cash_gain+cash_oci)
        self.assertAlmostEqual(value['cash_translation_oci'][pid],cash_oci)
        self.assertAlmostEqual(value['cash_fx_remeasurement'][pid],cash_gain)
        self.assertAlmostEqual(value['financing_fx_remeasurement'][pid],debt_loss)
        self.assertAlmostEqual(value['cfo_fx_noncash'][pid],debt_loss-cash_gain)
        self.assertAlmostEqual(value['cfo'][pid],sum(value[r][pid] for r in spec['cashflow_component_contract']['cfo_components']))
        self.assertAlmostEqual(value['oci_translation'][pid],cash_oci+inventory_fx-payables_fx-debt_translation)
        self.assertAlmostEqual(value['financial_income'][pid]-base['financial_income'][pid],cash_gain)
        self.assertAlmostEqual(value['financial_expense'][pid]-base['financial_expense'][pid],-debt_loss)
        self.assertAlmostEqual(value['inventories'][pid]-base['inventories'][pid],inventory_fx)
        self.assertAlmostEqual(value['wc_inventory_cf'][pid],base['wc_inventory_cf'][pid])
        self.assertAlmostEqual(value['wc_payables_cf'][pid],base['wc_payables_cf'][pid])
        self.assertAlmostEqual(value['nonlease_debt'][pid]-base['nonlease_debt'][pid],debt_loss+debt_translation)
        self.assertAlmostEqual(value['balance_check'][pid],0)
        self.assertFalse([c for c in validate_periodic(spec,value) if c['status'] in ('FAIL','MISSING')])
        rows={r['id']:r for r in spec['rows']}
        self.assertFalse(rows['cash_translation_oci']['cells'][pid]['editable'])
        with self.assertRaisesRegex(ValueError,'authoritative editable'):
            calculate_periodic(spec,{('cash_translation_oci',pid):99})
        # Source observations and their cells remain identical.
        base_rows={r['id']:r for r in base_spec['rows']}
        for rid in ('cash','cash_fx','inventories','nonlease_debt'):
            self.assertEqual(rows[rid]['cells']['2025-FY'],base_rows[rid]['cells']['2025-FY'])

    def test_fx_accounting_exposure_overlap_rejected_and_live_weights_validated(self):
        exposures=[dict(id='one',type='foreign_operation',role='cash',currency='JPY',opening_share=.6),
            dict(id='two',type='monetary',role='cash',currency='JPY',opening_share=.3)]
        spec,_=self.accounting_case(exposures)
        changed=calculate_periodic(spec,{('fx_accounting_two_weight','2026-H1'):.5})
        self.assertTrue(any(c['status']=='FAIL' and 'nonoverlapping_balance_exposures:cash' in c['name'] for c in validate_periodic(spec,changed)))
        exposures[1]['opening_share']=.5
        with self.assertRaisesRegex(ValueError,'cannot overlap'):self.accounting_case(exposures)
        data=copy.deepcopy(self.data)
        data['financial_contracts'].update(fx_accounting=dict(exposures=[]),balance_sheet_fx=dict(annual_cash_translation_oci=5.))
        with self.assertRaisesRegex(ValueError,'conflicts with legacy independent FX amounts'):
            build_periodic_spec(data,1,self.snapshot)

    def test_native_balance_movement_uses_closing_minus_average_not_closing_change(self):
        exposure=dict(id='cash_flow',type='foreign_operation',role='cash',currency='JPY',opening_share=0.,
            net_movement_native_by_period={'2026-H1':100.})
        spec,values=self.accounting_case([exposure],rate_change='2025-09-30');pid='2026-H1'
        # Independent fixture calendar:66 weekdays in each of Jul-Sep2025
        #and Oct-Dec2025; this synthetic provider has no holiday exclusions.
        average=((1.2/180)+(1.2/150))/2
        expected=100*(.008-average)
        self.assertAlmostEqual(values['fx_accounting_cash_flow_average_rate'][pid],average)
        self.assertAlmostEqual(values['cash_fx'][pid],expected)
        self.assertAlmostEqual(values['oci_translation'][pid],expected)
        self.assertAlmostEqual(values['balance_check'][pid],0.)
        self.assertNotAlmostEqual(expected,100*(.008-1.2/180))

    def test_unit_cost_including_fx_disables_second_transaction_adjustment(self):
        data=copy.deepcopy(self.data);pool=data['financial_contracts']['cost_pools'][0]
        pool['input_cost_includes_transaction_fx']=True
        spec=build_periodic_spec(data,2,self.snapshot);rows={r['id']:r for r in spec['rows']}
        rid='north_pool_'+pool['id']+'_transaction_fx';pid='2026-H1'
        self.assertFalse(rows[rid]['cells'][pid]['editable'])
        self.assertEqual(calculate_periodic(spec)[rid][pid],0)
        pool['fx']['transaction']['components']=[dict(weight=.1,functional_currency='GBP',invoice_currency='JPY')]
        with self.assertRaisesRegex(ValueError,'already includes transaction FX'):build_periodic_spec(data,2,self.snapshot)
        self.assertIsNotNone(self.values['cash']['2026-FY'])
        self.assertAlmostEqual(self.values['sales']['2026-H1'],self.values['sales']['2026-Q1']+self.values['sales']['2026-Q2'])
        self.assertAlmostEqual(self.values['balance_check']['2026-FY'],0,places=8)
        self.assertFalse([c for c in validate_periodic(self.spec,self.values) if c['status'] in ('FAIL','MISSING')])

    def test_non_major_weight_and_adjustment_changes_have_expected_economic_effect(self):
        pid='2026-Q1';weight='north_revenue_fx_weight_jpy';rows={r['id']:r for r in self.spec['rows']}
        anchor=next(p for p,c in rows[weight]['cells'].items() if c.get('editable'))
        changes={weight:{anchor:.8},'north_revenue_fx_weight_gbp':{anchor:.2},'north_revenue_fx_multiplier':{pid:2}}
        changed=calculate_periodic(self.spec,changes)
        rate=self.values['fx_yoy_jpy'][pid]
        expected_raw=.8*rate;self.assertAlmostEqual(changed['north_revenue_fx_raw'][pid],expected_raw)
        constant_currency=self.values['north_sales_seed'][pid]*(1+self.values['north_pricing'][pid]+self.values['north_rig'][pid])
        self.assertAlmostEqual(changed['north_sales'][pid],constant_currency*(1+expected_raw*2))
        self.assertEqual(changed['north_fx']['2025-FY'],-.03)
        cost=calculate_periodic(self.spec,{'north_pool_cogs_translation_multiplier':{'2026-FY':2}})
        self.assertNotEqual(cost['north_cash_cost']['2026-H1'],self.values['north_cash_cost']['2026-H1'])
        self.assertNotEqual(cost['utop']['2026-H1'],self.values['utop']['2026-H1'])

    def test_currency_refresh_changes_revenue_and_cost_without_touching_reported_fx(self):
        snapshot=copy.deepcopy(self.snapshot)
        for q in snapshot['spot_state']:
            if q['quote']=='JPY':q['value']*=1.1
        with patch('research_platform.ai.OpenAIProvider.from_env',side_effect=AssertionError('No LLM')):
            spec=build_periodic_spec(self.data,2,snapshot);changed=calculate_periodic(spec)
        self.assertNotEqual(changed['north_sales']['2026-Q1'],self.values['north_sales']['2026-Q1'])
        self.assertNotEqual(changed['north_cash_cost']['2026-H1'],self.values['north_cash_cost']['2026-H1'])
        self.assertEqual(changed['north_fx']['2025-FY'],-.03)
        self.assertEqual(spec['reporting_bridges'],self.spec['reporting_bridges'])

    def test_pool_translation_mixes_differ_and_annual_money_reconciles_to_cost_base(self):
        pid='2026-H1';v=self.values
        self.assertNotEqual(v['north_pool_cogs_translation_fx'][pid],v['north_pool_marketing_admin_translation_fx'][pid])
        self.assertEqual(v['north_pool_marketing_admin_translation_fx'][pid],0)
        annual='2026-FY'
        self.assertAlmostEqual(v['north_cost_fx'][annual],v['north_cost_fx_amount'][annual]/v['north_cost_fx_base'][annual])
        self.assertAlmostEqual(v['north_cost_fx_amount'][annual],v['north_cost_fx_amount']['2026-H1']+v['north_cost_fx_amount']['2026-H2'])

    def test_unresolved_neutral_cost_mix_is_editable_stress_not_a_reported_domestic_mix(self):
        data=copy.deepcopy(self.data);pool=data['financial_contracts']['cost_pools'][0]
        basket=pool['fx']['translation']
        basket['weights']=[dict(currency='GBP',weight=1.,exposure_status='unresolved',scenario_neutral=True)]
        basket['sensitivity_range']=dict(translation_shock_low=-.1,translation_shock_high=.1,
            interpretation='Synthetic stress range, not a calibrated bound')
        spec=build_periodic_spec(data,1,self.snapshot);base=calculate_periodic(spec);pid='2026-H1'
        rows={r['id']:r for r in spec['rows']};rid='north_pool_cogs_translation_fx'
        self.assertEqual(rows[rid]['kind'],'assumption')
        self.assertTrue(rows[rid]['cells'][pid]['editable'])
        self.assertEqual(rows[rid]['sensitivity_range'],basket['sensitivity_range'])
        self.assertNotIn('north_pool_cogs_translation_weight_gbp',rows)
        quality=next(q for q in spec['fx_exposure_quality'] if q['entity']=='north' and q.get('cost_pool')=='cogs')
        self.assertEqual(quality['unresolved_weight'],1.)
        for shock in (-.1,.1):
            changed=calculate_periodic(spec,{rid:{pid:shock}})
            delta=base['north_pool_cogs'][pid]*shock
            self.assertAlmostEqual(changed['north_pool_cogs'][pid]-base['north_pool_cogs'][pid],delta)
            self.assertAlmostEqual(changed['north_utop'][pid]-base['north_utop'][pid],-delta)
        # A genuine reported domestic functional-currency mix remains a
        # calculated identity, not an unresolved economic scenario.
        basket['weights']=[dict(currency='GBP',weight=1.,exposure_status='reported',evidence_basis='functional_currency_disclosure')]
        domestic=build_periodic_spec(data,1,self.snapshot)
        domestic_row=next(r for r in domestic['rows'] if r['id']==rid)
        self.assertEqual(domestic_row['kind'],'calculated')
        self.assertFalse(domestic_row['cells'][pid]['editable'])
        self.assertEqual(calculate_periodic(domestic)[rid][pid],0.)

    def test_translation_only_preserves_functional_cash_operating_margin(self):
        data=copy.deepcopy(self.data)
        # A separately fixed reporting-currency SBC budget is not a foreign
        # cash cost. Exclude it from this deliberately pure translation case.
        data['financial_contracts']['sbc']=dict(annual_amount=0.)
        single=[dict(currency='JPY',weight=1.,exposure_status='reported',evidence_basis='functional_currency_disclosure'),
            dict(currency='GBP',weight=0.,exposure_status='reported',evidence_basis='functional_currency_disclosure')]
        for exposure in data['model_configuration']['fx_exposures'].values():exposure['revenue']['weights']=copy.deepcopy(single)
        for pool in data['financial_contracts']['cost_pools']:pool['fx']['translation']['weights']=copy.deepcopy(single)
        original_spec=build_periodic_spec(data,1,self.snapshot);original=calculate_periodic(original_spec)
        shifted_snapshot=copy.deepcopy(self.snapshot)
        for q in shifted_snapshot['spot_candidates']:
            if q['quote']=='JPY':q['value']/=1.1
        shifted_spec=build_periodic_spec(data,1,shifted_snapshot);shifted=calculate_periodic(shifted_spec)
        pid='2026-H1'
        self.assertAlmostEqual(shifted['fx_average_jpy'][pid]/original['fx_average_jpy'][pid],1.1)
        for rid in ('north_sales','north_cash_cost','north_ebitda'):
            self.assertAlmostEqual(shifted[rid][pid],original[rid][pid]*1.1)
            self.assertAlmostEqual(shifted[rid][pid]/shifted['fx_average_jpy'][pid],original[rid][pid]/original['fx_average_jpy'][pid])
        self.assertAlmostEqual(shifted['north_ebitda'][pid]/shifted['north_sales'][pid],original['north_ebitda'][pid]/original['north_sales'][pid])
        # This covers cash-operating margin. Full EBIT translation additionally
        # requires explicit D&A/asset currency exposures, which this fixture
        # deliberately does not invent.

    def test_procurement_currency_shock_hits_only_configured_cost_pool_once(self):
        data=copy.deepcopy(self.data);pool=data['financial_contracts']['cost_pools'][0]
        pool['fx']['translation']['weights']=[dict(currency='GBP',weight=1.)]
        plain_spec=build_periodic_spec(data,1,self.snapshot);plain=calculate_periodic(plain_spec)
        pool['fx']['transaction']=dict(method='SYNTHETIC JPY procurement by GBP-functional operation',
            components=[dict(weight=1.,functional_currency='GBP',invoice_currency='JPY',lag_months=0,hedge_coverage=0.,status='scenario')])
        spec=build_periodic_spec(data,1,self.snapshot);values=calculate_periodic(spec);pid='2026-H1'
        self.assertAlmostEqual(values['north_pool_cogs_transaction_fx'][pid],150/180-1)
        expected=plain['north_pool_cogs'][pid]*(150/180-1)
        self.assertAlmostEqual(values['north_pool_cogs'][pid]-plain['north_pool_cogs'][pid],expected)
        self.assertEqual(values['north_pool_marketing_admin'][pid],plain['north_pool_marketing_admin'][pid])
        self.assertEqual(values['north_sales'][pid],plain['north_sales'][pid])
        self.assertAlmostEqual(values['north_utop'][pid]-plain['north_utop'][pid],-expected)
        edited=calculate_periodic(spec,{'north_pool_cogs_transaction_1_hedge_coverage':{pid:1.},
            'north_pool_cogs_transaction_1_hedge_rate':{pid:.008}})
        self.assertAlmostEqual(edited['north_pool_cogs_transaction_fx'][pid],0)
        self.assertAlmostEqual(edited['north_pool_cogs'][pid],plain['north_pool_cogs'][pid])
        independent=calculate_periodic(spec,{'north_pool_cogs_transaction_1_hedge_coverage':{pid:1.},
            'north_pool_cogs_transaction_1_hedge_rate':{pid:.008},
            'north_pool_cogs_transaction_1_prior_hedge_coverage':{pid:1.},
            'north_pool_cogs_transaction_1_prior_hedge_rate':{pid:.010}})
        self.assertAlmostEqual(independent['north_pool_cogs_transaction_fx'][pid],-.2)
        self.assertAlmostEqual(independent['north_pool_cogs'][pid],plain['north_pool_cogs'][pid]*.8)
        invalid=calculate_periodic(spec,{'north_pool_cogs_transaction_1_hedge_coverage':{pid:1.}})
        checks=validate_periodic(spec,invalid)
        self.assertTrue(any(c['area']=='fx_contract' and c['status']=='FAIL' for c in checks))
        invalid_prior=calculate_periodic(spec,{'north_pool_cogs_transaction_1_prior_hedge_coverage':{pid:1.}})
        self.assertTrue(any(c['area']=='fx_contract' and c['status']=='FAIL' for c in validate_periodic(spec,invalid_prior)))

    def test_true_quarterly_statement_company_still_calculates_quarters(self):
        data=copy.deepcopy(self.data);data['model_configuration'].update(full_statement_frequency='quarterly',forecast_frequency='quarterly')
        spec=build_periodic_spec(data,1,self.snapshot);values=calculate_periodic(spec)
        for rid in ('operating_profit','cfo','cash'):self.assertIsNotNone(values[rid]['2026-Q1'])
        self.assertAlmostEqual(values['balance_check']['2026-Q1'],0,places=8)

    def test_fx_and_cost_pools_forecast_directly_in_annual_only_years(self):
        spec=build_periodic_spec(self.data,4,self.snapshot);values=calculate_periodic(spec)
        # October 2026 is FY2027 for this June-year-end company: detailed
        # FY2027/FY2028, then FY2029 is genuinely annual-only.
        self.assertEqual([p['kind'] for p in spec['periods'] if p['fiscal_year']==2029],['FY'])
        for rid in ('fx_average_jpy','north_pool_cogs_translation_fx','north_cost_fx_amount','north_cost_fx_base','sales','cfo','cash'):
            self.assertIsNotNone(values[rid]['2029-FY'],rid)
        self.assertAlmostEqual(values['north_cost_fx']['2029-FY'],values['north_cost_fx_amount']['2029-FY']/values['north_cost_fx_base']['2029-FY'])
        self.assertAlmostEqual(values['balance_check']['2029-FY'],0,places=7)

    def test_renderer_order_removes_nine_month_columns_and_keeps_internal_reference(self):
        path=Path(self.temp.name)/'model.xlsx';inspection=export_integrated(path,self.spec,self.values,[])
        book=load_workbook(path);self.addCleanup(book.close)
        columns=inspection['period_columns']
        for sheet in ('Model','Assumptions'):
            self.assertNotIn('2025-9M',columns)
            visible=[p['kind'] for p in self.spec['periods'] if p['fiscal_year']==2026 and p['id'] in columns and not book[sheet].column_dimensions[columns[p['id']]].hidden]
            self.assertEqual(visible,['Q1','Q2','Q3','Q4','H1','H2','FY'])
            header='6' if sheet=='Model' else '5'
            self.assertNotEqual(book[sheet][columns['2025-FY']+header].fill.fgColor.rgb,book[sheet][columns['2025-H2']+header].fill.fgColor.rgb)
        cell=next(r for r in self.spec['rows'] if r['id']=='sales')['cells']['2025-Q3']
        self.assertIn('2025-9M',{d['period_id'] for d in cell_dependencies(cell,'2025-Q3')})
        self.assertAlmostEqual(self.values['sales']['2025-Q3'],36)

    def test_weights_reject_incomplete_or_unexplained_proxy(self):
        data=copy.deepcopy(self.data);data['model_configuration']['fx_exposures']['north']['revenue']['weights'][0]['weight']=.6
        with self.assertRaisesRegex(ValueError,'sum to one'):build_periodic_spec(data,1,self.snapshot)
        data=copy.deepcopy(self.data);data['model_configuration']['fx_exposures']['north']['revenue']['weights'][0]['rate_currency']='EUR'
        with self.assertRaisesRegex(ValueError,'proxy'):build_periodic_spec(data,1,self.snapshot)


class SeparatedVendorTests(unittest.TestCase):
    def workbook(self,value,estimate=999):
        book=Workbook();s=book.active;s.title='Source metrics'
        s.append(['Metric','Units','Helper','FY-2025','FY-2025','FY-2026E'])
        s.append(['','Units','','VA Actuals','Filings','Consensus'])
        s.append(['Underlying trading operating profit','GBPM','Analytics',value,value,estimate])
        out=io.BytesIO();book.save(out);book.close();return out.getvalue()

    def test_explicit_company_role_excludes_different_standardized_and_consensus(self):
        company=self.workbook(12);standard=self.workbook(99)
        cu,cm=SourceStore._index(company,'xlsx','company','visible_alpha','company_reported')
        su,sm=SourceStore._index(standard,'xlsx','standard','visible_alpha','standardized')
        self.assertEqual(cm['vendor_layout']['status'],'recognized');self.assertEqual(sm['vendor_layout']['status'],'recognized')
        with tempfile.TemporaryDirectory() as tmp:
            store=SourceStore(Path(tmp));store._units={'company':cu,'standard':su}
            store.sources=[dict(id='company',category='visible_alpha',**cm),dict(id='standard',category='visible_alpha',**sm)]
            view=store._vendor_history(store.sources[0],{2025})
            numbers=[n['raw'] for u in view for n in u['numbers']]
            self.assertIn(12,numbers);self.assertNotIn(999,numbers)
            self.assertFalse([n for u in store._vendor_history(store.sources[1],{2025}) for n in u['numbers']])
            ref=dict(source_id='standard',unit_id='Source metrics!R3',number_index=0)
            with self.assertRaisesRegex(ValueError,'company-reported'):store.validate_historical_reference(ref,dict(year=2025,kind='annual'))

    def test_unlabelled_combined_or_missing_primary_cannot_supply_numbers(self):
        units,metadata=SourceStore._index(self.workbook(None),'xlsx','unlabelled','visible_alpha')
        self.assertEqual(metadata['vendor_layout']['status'],'unsupported')
        units,metadata=SourceStore._index(self.workbook(None),'xlsx','company','visible_alpha','company_reported')
        with tempfile.TemporaryDirectory() as tmp:
            store=SourceStore(Path(tmp));store._units={'company':units};source=dict(id='company',category='visible_alpha',**metadata)
            self.assertFalse([n for u in store._vendor_history(source,{2025}) for n in u['numbers']])


if __name__=='__main__':unittest.main()
