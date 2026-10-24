"""Financial parsing, arithmetic and reproducibility checks."""

import ast
import copy
import hashlib
from http.client import BadStatusLine, IncompleteRead
import io
import json
import os
import re
import socket
import ssl
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError, URLError
from uuid import UUID

from openpyxl import Workbook, load_workbook

from research_platform.project import find_project_root, read_json
from research_platform.__main__ import TerminalProgress
from research_platform.ai import AIError, OpenAIProvider, aggregate_usage, response_usage
from research_platform.evidence import SourceStore, download, parse_number
from research_platform.initiation import (REVIEW_TOPICS, SECTIONS, _rounding_bound,
                                          export_workbook, initiate_company, materialize_and_check)


PROJECT = Path(__file__).resolve().parents[1]


class HistoricalParsingTests(unittest.TestCase):
    def test_numbers_preserve_reported_sign_and_thousands(self):
        for value, expected in (
            (1234, 1234.0),
            (-1234.5, -1234.5),
            ("1 234", 1234.0),
            ("1\u00a0234", 1234.0),
            ("(1 234)", -1234.0),
            ("(3.51)", -3.51),
        ):
            with self.subTest(value=value):
                self.assertEqual(parse_number(value), expected)

    def test_only_explicit_reported_dashes_are_zero(self):
        for value in ("-", "\u2013", "\u2014", 0, "0"):
            with self.subTest(value=value):
                self.assertEqual(parse_number(value), 0.0)

    def test_missing_ambiguous_and_nonfinite_numbers_are_rejected(self):
        for value in (
            None, "", "   ", "not disclosed", "n/a", "1 2 3", True, False,
            float("nan"), float("inf"), float("-inf"), "NaN", "Infinity",
        ):
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    parse_number(value)

    def test_decimal_locale_is_explicit_and_percent_is_not_an_amount(self):
        self.assertEqual(parse_number("1.234,50", "comma_decimal"), 1234.5)
        self.assertEqual(parse_number("(1,234.50)", "dot_decimal"), -1234.5)
        self.assertEqual(parse_number("3,5%", "comma_decimal"), 0.035)
        self.assertEqual(parse_number("1,234", "dot_decimal"), 1234)
        self.assertEqual(parse_number("1,234", "comma_decimal"), 1.234)
        with self.assertRaises(ValueError):
            parse_number("1.234,50", "dot_decimal")


class ProjectTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.package = self.root / "src" / "research_platform"
        self.package.mkdir(parents=True)
        (self.package / "__init__.py").write_bytes(
            (PROJECT / "src" / "research_platform" / "__init__.py").read_bytes())
        (self.root / "pyproject.toml").write_bytes((PROJECT / "pyproject.toml").read_bytes())

    def tearDown(self):
        self.temporary.cleanup()

    def test_root_from_subdirectory_without_company_configs(self):
        self.assertFalse((self.root / "configs").exists())
        self.assertEqual(find_project_root(self.package), self.root)

    def test_unrelated_pyproject_is_not_a_project_root(self):
        (self.package / "__init__.py").unlink()
        with self.assertRaises(ValueError):
            find_project_root(self.root)

    def test_read_run_record(self):
        path = self.root / "record.json"
        path.write_text('{"status": "incomplete"}', encoding="utf-8")
        self.assertEqual(read_json(path), {"status": "incomplete"})

    def test_read_run_record_rejects_nonobjects(self):
        path = self.root / "record.json"
        for value in ("[]", "null", "42"):
            with self.subTest(value=value):
                path.write_text(value, encoding="utf-8")
                with self.assertRaises(ValueError):
                    read_json(path)


class TerminalProgressTests(unittest.TestCase):
    class InteractiveStream(io.StringIO):
        def isatty(self):
            return True

    def test_timer_counts_up_while_request_is_still_running(self):
        stream = self.InteractiveStream()
        with TerminalProgress(stream).call_timer("Call 1: extract"):
            self.assertIn("waiting 00:00", stream.getvalue())

            def elapsed_ticks():
                return [int(minutes) * 60 + int(seconds)
                        for minutes, seconds in re.findall(r"waiting (\d+):(\d+)", stream.getvalue())]

            deadline = time.monotonic() + 3
            while max(elapsed_ticks()) < 1 and time.monotonic() < deadline:
                time.sleep(0.02)
            self.assertGreaterEqual(max(elapsed_ticks()), 1)
            self.assertGreaterEqual(len(elapsed_ticks()), 2)
            self.assertNotIn("completed", stream.getvalue())
        self.assertIn("completed after", stream.getvalue())
        self.assertTrue(stream.getvalue().endswith("s\n"))

    def test_failure_stops_timer_thread_and_reports_elapsed_time(self):
        stream = self.InteractiveStream()
        before = {thread.ident for thread in threading.enumerate()}
        with self.assertRaisesRegex(RuntimeError, "synthetic transport failure"):
            with TerminalProgress(stream).call_timer("Call 2: extract"):
                self.assertTrue(any(thread.name == "ai-call-timer" and thread.ident not in before
                                    for thread in threading.enumerate()))
                raise RuntimeError("synthetic transport failure")
        self.assertFalse(any(thread.name == "ai-call-timer" and thread.ident not in before
                             for thread in threading.enumerate()))
        self.assertIn("failed after", stream.getvalue())

    def test_redirected_output_has_one_completion_line_without_carriage_returns(self):
        stream = io.StringIO()
        with TerminalProgress(stream).call_timer("Call 3: model_review"):
            self.assertEqual(stream.getvalue(), "")
        self.assertNotIn("\r", stream.getvalue())
        self.assertEqual(len(stream.getvalue().splitlines()), 1)
        self.assertIn("completed after", stream.getvalue())


class SyntheticCompany:
    """Invented disclosures only; deliberately different layouts, no company map."""

    def __init__(self, variant=0):
        self.name = ("Fictional Hearth Appliances plc", "Fictional Meadow Diagnostics SA")[variant]
        self.currency = ("USD", "EUR")[variant]
        self.sheet = ("Statements", "Group accounts")[variant]
        self.reverse_columns = bool(variant)
        self.column_periods = ["2024", "2025"]
        self.source = {
            "id": "annual", "url": f"https://example.invalid/synthetic-{variant}.xlsx",
            "title": "SYNTHETIC TEST DISCLOSURE", "published_at": "2026-02-10",
            "reason": "Invented primary disclosure for an offline software test.",
            "kind": "annual_report",
        }
        # ID, section, accounting role, entity, two annual observations.
        specifications = [
            ("segment_a", "segments", "segment_revenue", "Alpha", (600, 650)),
            ("segment_b", "segments", "segment_revenue", "Beta", (400, 450)),
            ("profit_a", "segments", "segment_profit", "Alpha", (60, 70)),
            ("profit_b", "segments", "segment_profit", "Beta", (40, 50)),
            ("segment_group_profit", "income_statement", "segment_group_profit", "group", (100, 120)),
            ("revenue", "income_statement", "revenue", "group", (1000, 1100)),
            ("expenses", "income_statement", "other", "group", (-900, -980)),
            ("net_income", "income_statement", "net_income", "group", (100, 120)),
            ("assets", "balance_sheet", "total_assets", "group", (1000, 1200)),
            ("liabilities", "balance_sheet", "total_liabilities", "group", (600, 720)),
            ("equity", "balance_sheet", "total_equity", "group", (400, 480)),
            ("cash_close", "cash_flow", "cash_close", "group", (100, 130)),
            ("balance_sheet_cash", "balance_sheet", "balance_sheet_cash", "group", (100, 130)),
            ("inventory", "balance_sheet", "other", "group", (100, 110)),
            ("receivables", "balance_sheet", "other", "group", (120, 130)),
            ("payables", "balance_sheet", "other", "group", (-80, -90)),
            ("debt", "balance_sheet", "other", "group", (300, 350)),
            ("cash_open", "cash_flow", "cash_open", "group", (70, 100)),
            ("cash_change", "cash_flow", "cash_change", "group", (30, 30)),
            ("operating", "cash_flow", "operating_cash", "group", (60, 70)),
            ("investing", "cash_flow", "investing_cash", "group", (-20, -25)),
            ("financing", "cash_flow", "financing_cash", "group", (-10, -15)),
            ("working_capital", "working_capital", "working_capital", "group", (140, 150)),
            ("net_debt", "net_debt", "net_debt", "group", (200, 220)),
            ("share_capital", "equity", "other", "group", (200, 220)),
            ("minority_interest", "equity", "other", "group", (20, 25)),
            ("shares", "shares", "other", "group", (100, 100)),
        ]
        if variant:
            specifications.reverse()
        self.rows = []
        self.display_values = {}
        self.units = []
        for number, (row_id, section, role, entity, values) in enumerate(specifications, 3):
            ref = {"source_id": "annual", "unit_id": f"{self.sheet}!R{number}"}
            self.units.append(ref)
            self.display_values[row_id] = list(values)
            self.rows.append({
                "id": row_id, "section": section, "label": row_id.replace("_", " "),
                "entity": entity, "definition": "Synthetic disclosed " + row_id + "; reported sign retained.",
                "unit": "shares" if row_id == "shares" else self.currency,
                "scale": 1, "sign": 1, "role": role,
                "time_basis": ("opening_balance" if row_id == "cash_open" else
                               "closing_balance" if section in ("balance_sheet", "equity", "shares", "net_debt", "working_capital")
                               or row_id == "cash_close" else "flow"),
                "unavailable": [],
                "observations": [
                    {"period": str(year), "reference": dict(ref, number_index=(1 - index if variant else index),
                                                      number_format="dot_decimal")}
                    for index, year in enumerate((2024, 2025))
                ],
                "terms": [], "definition_refs": [
                    {"source_id": "annual", "unit_id": f"{self.sheet}!R1"},
                    {"source_id": "annual", "unit_id": f"{self.sheet}!R2"}, ref],
            })
        self.discovery = {
            "company": {"legal_name": self.name, "identifier": "SYNTHETIC-" + str(variant),
                        "exchange": "NONE", "country": ("US", "FR")[variant], "currency": self.currency},
            "periods": [{"id": str(year), "year": year, "kind": "annual", "start": f"{year}-01-01", "end": f"{year}-12-31"}
                        for year in (2024, 2025)],
            "period_coverage": [
                {"year": year, "slot": slot, "status": "not_disclosed",
                 "explanation": "Synthetic issuer publishes annual accounts only; no interim disclosure.",
                 "source_ids": ["annual"]}
                for year in (2024, 2025) for slot in ("Q1", "Q2", "Q3", "Q4", "H1", "H2", "9M")],
            "sources": [self.source], "notes": [], "blocking_issues": [],
        }
        relations = {
            "revenue": [("segment_a", 1), ("segment_b", 1)],
            "segment_group_profit": [("profit_a", 1), ("profit_b", 1)],
            "net_income": [("revenue", 1), ("expenses", 1)],
            "balance_sheet_cash": [("cash_close", 1)],
            "cash_change": [("operating", 1), ("investing", 1), ("financing", 1)],
            "working_capital": [("inventory", 1), ("receivables", 1), ("payables", 1)],
            "net_debt": [("debt", 1), ("cash_close", -1)],
        }
        self.candidate = {
            "reporting_basis": "Synthetic FY2025 disclosure with FY2024 comparatives; no restatement.",
            "rows": self.rows,
            "equations": [
                {"target": target, "terms": [{"row_id": row, "coefficient": coefficient}
                                            for row, coefficient in terms],
                 "explanation": "Synthetic disclosed composition of " + target}
                for target, terms in relations.items()
            ],
            "coverage": [{"section": section, "period": str(year), "status": "complete",
                          "explanation": "All material lines in this deliberately small synthetic disclosure.",
                          "references": [{"source_id": "annual", "unit_id": f"{self.sheet}!R1"}]}
                         for section in SECTIONS for year in (2024, 2025)],
            "cash_continuity": [{"period": "2025", "previous_period": "2024", "terms": [],
                                 "explanation": "The synthetic adjacent annual periods use the same cash definition.",
                                 "references": [{"source_id": "annual", "unit_id": f"{self.sheet}!R1"}]}],
            "gaps": [],
        }
        self.fetches = []

    def payload(self):
        book = Workbook()
        sheet = book.active
        sheet.title = self.sheet
        scope = ("annual accounts only, no interim disclosure" if len(self.column_periods) == 2 else
                 "full annual and H1 accounts; Q1 and nine-month sales-only releases; no standalone Q2, Q3, Q4 or H2")
        sheet.append(["SYNTHETIC: " + self.name + "; reporting currency " + self.currency
                      + "; units 1; " + scope + "; no further schedules."])
        headings = [("FY" + value if "-" not in value else value) for value in self.column_periods]
        sheet.append(["Reported line", *(list(reversed(headings)) if self.reverse_columns else headings)])
        for row in self.rows:
            values = self.display_values[row["id"]]
            sheet.append([row["label"], *(list(reversed(values)) if self.reverse_columns else values)])
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    def fetch(self, url):
        if url != self.source["url"]:
            raise AssertionError("Unrecognized synthetic source URL")
        self.fetches.append(url)
        return self.payload(), url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def include_interims(self):
        """Invented complete H1 accounts and sales-only Q1/9M disclosures."""
        self.reverse_columns = False
        interim_periods = [
            {"id": "2025-Q1", "year": 2025, "kind": "quarter", "start": "2025-01-01", "end": "2025-03-31"},
            {"id": "2025-H1", "year": 2025, "kind": "half_year", "start": "2025-01-01", "end": "2025-06-30"},
            {"id": "2025-9M", "year": 2025, "kind": "year_to_date", "start": "2025-01-01", "end": "2025-09-30"},
        ]
        self.discovery["periods"] = [self.discovery["periods"][0], *interim_periods, self.discovery["periods"][1]]
        self.column_periods += [item["id"] for item in interim_periods]
        for item in self.discovery["period_coverage"]:
            if item["year"] == 2025 and item["slot"] in ("Q1", "H1", "9M"):
                item.update(status="reported", explanation="Synthetic H1 full accounts and Q1/9M sales disclosures.")
        half_values = {
            "assets": 1100, "liabilities": 660, "equity": 440, "cash_close": 115,
            "balance_sheet_cash": 115, "inventory": 105, "receivables": 125, "payables": -85,
            "debt": 325, "cash_open": 100, "working_capital": 145, "net_debt": 210,
            "share_capital": 210, "minority_interest": 22.5, "shares": 100,
        }
        context = {"source_id": "annual", "unit_id": f"{self.sheet}!R1"}
        for row in self.rows:
            values = self.display_values[row["id"]]
            sales = row["role"] in ("revenue", "segment_revenue")
            half = half_values.get(row["id"], values[1] * 0.5)
            extras = [values[1] * 0.2 if sales else None, half, values[1] * 0.75 if sales else None]
            numeric_index = 2
            for period, value in zip(interim_periods, extras):
                if value is None:
                    row["unavailable"].append({"period": period["id"],
                                               "explanation": "Sales-only release does not disclose this financial line.",
                                               "references": [copy.deepcopy(context)]})
                else:
                    reference = dict(row["observations"][0]["reference"], number_index=numeric_index)
                    row["observations"].append({"period": period["id"], "reference": reference})
                    numeric_index += 1
            values.extend(extras)
        for period in interim_periods:
            for section in SECTIONS:
                status = ("complete" if period["kind"] == "half_year" else
                          "partial" if section in ("segments", "income_statement") else "not_disclosed")
                self.candidate["coverage"].append({"section": section, "period": period["id"], "status": status,
                                                   "explanation": "H1 full accounts; Q1 and 9M disclose sales only.",
                                                   "references": [copy.deepcopy(context)]})
        self.candidate["cash_continuity"].append({"period": "2025-H1", "previous_period": "2024", "terms": [],
                                                  "explanation": "H1 opens at prior annual closing cash, on the same basis.",
                                                  "references": [context]})
        return self


class ScriptedProvider:
    """Scripted decisions test plumbing and gates, not actual AI research quality."""

    last_response = None

    def __init__(self, company, *, mutate=None, review_fault=None, repair=False, batch=False):
        self.company = company
        self.mutate = mutate
        self.review_fault = review_fault
        self.repair = repair
        self.batch = batch
        self.calls = []
        self.proposals = 0
        self.read_requested = False
        self.saved_batches = 0

    def validate_configuration(self):
        pass

    def metadata(self):
        return {"provider": "SYNTHETIC SCRIPTED TEST", "models": {
            role: "gpt-6.1-sol" for role in ("researcher", "extractor", "reviewer")}}

    @staticmethod
    def synthetic_response(stage):
        return {"synthetic": True, "stage": stage, "status": "completed", "model": "gpt-6.1-sol",
                "output": [], "usage": {"input_tokens": 1000, "output_tokens": 100, "total_tokens": 1100}}

    def complete(self, *, role, instructions, payload, schema, web=False):
        stage = payload["stage"]
        self.calls.append(copy.deepcopy(payload))
        self.last_response = self.synthetic_response(stage)
        if stage == "discover":
            if not web:
                raise AssertionError("Discovery must request web research")
            return copy.deepcopy(self.company.discovery)
        if stage == "extract":
            action = {"action": "propose", "read_requests": [], "sources": [], "rows": [],
                      "delete_row_ids": [], "search_query": "", "candidate": None}
            if not self.read_requested:
                self.read_requested = True
                action.update(action="read", read_requests=[
                    {"source_id": source["id"], "unit_id": unit["id"]}
                    for source in payload["catalog"] for unit in source["units"]])
                return action
            candidate = copy.deepcopy(self.company.candidate)
            if self.mutate and (not self.repair or self.proposals == 0):
                self.mutate(candidate)
            if self.batch and self.saved_batches * 12 < len(candidate["rows"]):
                start = self.saved_batches * 12
                action.update(action="save_rows", rows=candidate["rows"][start:start + 12])
                self.saved_batches += 1
                return action
            self.proposals += 1
            if self.batch:
                candidate["rows"] = []
            action["candidate"] = candidate
            return action
        if stage == "source_review":
            result = {"facts": [], "issues": [], "read_requests": [], "search_query": ""}
            if not payload["evidence"]:
                result["read_requests"] = [
                    {"source_id": source["id"], "unit_id": unit["id"]}
                    for source in payload["catalog"] for unit in source["units"]]
                return result
            disclosed = {row["id"]: row for row in self.company.rows}
            facts = []
            for concept in payload["concepts"]:
                row = disclosed[concept["id"]]
                fact = {key: copy.deepcopy(row[key])
                        for key in ("entity", "unit", "scale", "sign", "definition", "time_basis")}
                fact.update(row_id=concept["id"], reporting_basis=self.company.candidate["reporting_basis"],
                            observations=[dict(copy.deepcopy(observation),
                                               context_refs=copy.deepcopy(row["definition_refs"]))
                                          for observation in row["observations"]
                                          if observation["period"] in concept["periods"]])
                facts.append(fact)
            if self.review_fault == "missing_source":
                facts.pop()
            if self.review_fault == "duplicate_source":
                facts.append(copy.deepcopy(facts[0]))
            if self.review_fault == "missing_source_context":
                facts[0]["observations"][0]["context_refs"] = []
            result["facts"] = facts
            return result
        if stage == "source_alignment":
            result = {"checks": [], "issues": [], "read_requests": [], "search_query": ""}
            if not payload["evidence"]:
                result["read_requests"] = [
                    {"source_id": source["id"], "unit_id": unit["id"]}
                    for source in payload["catalog"] for unit in source["units"]]
                return result
            disclosed = {row["id"]: row for row in self.company.rows}
            for row in payload["rows"]:
                passed, reason = True, "Synthetic semantic source alignment verified."
                if self.review_fault == "definition_disagreement" and row["definition"] != disclosed[row["id"]]["definition"]:
                    passed, reason = False, "Equal numbers describe a different economic definition."
                if self.review_fault == "basis_disagreement" and payload["reporting_basis"] != self.company.candidate["reporting_basis"]:
                    passed, reason = False, "The candidate reporting vintage differs from the source basis."
                result["checks"].append({"row_id": row["id"], "passed": passed, "reason": reason,
                                         "references": copy.deepcopy(disclosed[row["id"]]["definition_refs"])})
            if self.review_fault == "missing_alignment":
                result["checks"].pop()
            if self.review_fault == "duplicate_alignment":
                result["checks"].append(copy.deepcopy(result["checks"][0]))
            if self.review_fault == "missing_alignment_context":
                result["checks"][0]["references"] = []
            if self.review_fault == "unread_alignment_context":
                result["checks"][0]["references"] = [{"source_id": "annual", "unit_id": "Never read"}]
            return result
        if stage == "coverage_inventory":
            section = payload["section"]
            result = {"search_query": "", "read_requests": [], "section": section,
                      "coverage": [{key: copy.deepcopy(value) for key, value in item.items() if key != "section"}
                                   for item in self.company.candidate["coverage"] if item["section"] == section],
                      "items": [], "issues": []}
            if not payload["evidence"]:
                result["read_requests"] = [
                    {"source_id": source["id"], "unit_id": unit["id"]}
                    for source in payload["catalog"] for unit in source["units"]]
                return result
            result["items"] = [{"id": section + "_" + row["id"], "label": row["label"],
                                "entity": row["entity"], "references": copy.deepcopy(row["definition_refs"]),
                                "periods": [observation["period"] for observation in row["observations"]]}
                               for row in self.company.rows if row["section"] == section]
            if self.review_fault == "duplicate_inventory":
                result["items"].append(copy.deepcopy(result["items"][0]))
            if self.review_fault == "empty_inventory":
                result["items"] = []
            if self.review_fault == "wrong_inventory_section":
                result["section"] = "shares" if section != "shares" else "equity"
            return result
        if stage == "model_review":
            checks = [{"topic": topic, "passed": True, "reason": "Synthetic documentary scope checked."}
                      for topic in REVIEW_TOPICS]
            if self.review_fault == "missing_model":
                checks.pop()
            if self.review_fault == "model_reject":
                checks[0].update(passed=False, reason="Synthetic legal identity remains unresolved.")
            if self.review_fault == "duplicate_model":
                checks.append(copy.deepcopy(checks[0]))
            matches = []
            for item in self.company.rows:
                row = next((row for row in payload["rows"] if row["label"] == item["label"]
                            and row["section"] == item["section"]), None)
                matches.append({"inventory_id": item["section"] + "_" + item["id"],
                                "row_id": row["id"] if row else "",
                                "reason": "The independently inventoried line is present." if row
                                          else "The independently inventoried disclosure is missing.",
                                "references": copy.deepcopy(item["definition_refs"])})
            if self.review_fault == "missing_match":
                matches.pop()
            if self.review_fault == "duplicate_match":
                matches.append(copy.deepcopy(matches[0]))
            if self.review_fault == "reused_match":
                next(match for match in matches if match["inventory_id"] == "equity_minority_interest")["row_id"] = "share_capital"
            result = {"checks": checks, "matches": matches, "issues": [], "read_requests": [], "search_query": ""}
            read_ids = {(item["source_id"], item["unit_id"]) for item in payload["evidence"]}
            required = {(ref["source_id"], ref["unit_id"])
                        for match in matches for ref in match["references"]}
            if not required <= read_ids:
                result.update(checks=[], matches=[], read_requests=[
                    {"source_id": source_id, "unit_id": unit_id} for source_id, unit_id in sorted(required - read_ids)])
            return result
        raise AssertionError("Unknown workflow stage: " + stage)


class InitiationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def run_case(self, company=None, provider=None, **kwargs):
        company = company or SyntheticCompany()
        provider = provider or ScriptedProvider(company)
        # These financial scenarios intentionally start independent runs. Resume
        # behavior has separate tests with the public default preserved.
        kwargs.setdefault("fresh", True)
        kwargs.setdefault("historical_years", 2)
        directory = initiate_company(self.root, company.name, provider=provider, fetcher=company.fetch,
                                     **kwargs)
        record = json.loads((directory / "record.json").read_text(encoding="utf-8"))
        return directory, record, provider

    def test_mixed_annual_half_year_and_sales_only_interims_are_preserved_and_checked(self):
        company = SyntheticCompany().include_interims()
        directory, record, provider = self.run_case(company, years=[2024, 2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertTrue(all(check["passed"] for check in record["checks"]), record["checks"])
        periods = {item["id"]: item for item in record["discovery"]["periods"]}
        self.assertEqual(set(periods), {"2024", "2025", "2025-Q1", "2025-H1", "2025-9M"})
        self.assertEqual(periods["2025-Q1"]["kind"], "quarter")
        self.assertEqual(periods["2025-H1"]["kind"], "half_year")
        self.assertEqual(periods["2025-9M"]["kind"], "year_to_date")
        rows = {row["id"]: row for row in record["rows"]}
        self.assertEqual(rows["revenue"]["values"]["2025-Q1"], 220)
        self.assertEqual(rows["revenue"]["values"]["2025-H1"], 550)
        self.assertEqual(rows["revenue"]["values"]["2025-9M"], 825)
        self.assertNotIn("2025-Q2", rows["revenue"]["values"])
        self.assertNotIn("2025-Q3", rows["revenue"]["values"])
        self.assertNotIn("2025-Q1", rows["net_debt"]["values"])
        self.assertEqual({item["period"] for item in rows["net_debt"]["unavailable"]}, {"2025-Q1", "2025-9M"})
        coverage = {(item["section"], item["period"]): item["status"] for item in record["candidate"]["coverage"]}
        self.assertEqual(coverage[("income_statement", "2025-Q1")], "partial")
        self.assertEqual(coverage[("balance_sheet", "2025-Q1")], "not_disclosed")
        self.assertEqual(coverage[("balance_sheet", "2025-H1")], "complete")
        self.assertTrue(any(check["name"] == "balance_sheet:2025-H1" for check in record["checks"]))
        self.assertTrue(any(check["name"] == "cash_rollforward:2025-H1" for check in record["checks"]))
        source_calls = [call for call in provider.calls if call["stage"] == "source_review"]
        revenue = [concept for call in source_calls for concept in call["concepts"] if concept["id"] == "revenue"]
        self.assertTrue(revenue and all(set(concept["periods"]) == set(periods) for concept in revenue))
        book = load_workbook(directory / "model.xlsx", data_only=False)
        try:
            sheet = book["net_debt"]
            header = {str(cell.value): cell.column for cell in sheet[4]}
            quarter_column = next(column for value, column in header.items() if "2025-Q1" in value)
            debt_row = next(row for row in sheet if row[0].value == "net debt")
            self.assertIsNone(debt_row[quarter_column - 1].value)
            self.assertTrue(debt_row[quarter_column - 1].comment)
        finally:
            book.close()

    def test_cash_continuity_uses_dated_predecessor_not_previous_interim_column(self):
        company = SyntheticCompany().include_interims()
        directory, record, _ = self.run_case(company)
        self.assertEqual(record["status"], "accepted", record["issues"])
        book = load_workbook(directory / "model.xlsx", data_only=False)
        try:
            sheet = book["cash_flow"]
            header = {str(cell.value): cell.column for cell in sheet[4]}
            annual_2024 = next(column for value, column in header.items() if "2024" in value)
            half_2025 = next(column for value, column in header.items() if "2025-H1" in value)
            close_row = next(row for row in sheet if row[0].value == "cash close")
            prior_annual_cell = close_row[annual_2024 - 1]
            half_cell = close_row[half_2025 - 1]
            for name in ("cash_continuity:2025-H1", "cash_continuity:2025"):
                check = next(row for row in book["Checks"] if row[0].value == name)
                self.assertIn(f"'cash_flow'!{prior_annual_cell.coordinate}", check[6].value)
                self.assertNotIn(f"'cash_flow'!{half_cell.coordinate}", check[6].value)
                self.assertEqual(WorkbookTests.arithmetic_value(book, check[6]), 0)
            same_date = next(row for row in book["Checks"] if str(row[0].value).startswith("stock_date_consistency:cash_open:"))
            self.assertEqual(same_date[6].data_type, "f")
            self.assertEqual(WorkbookTests.arithmetic_value(book, same_date[6]), 0)
            opening = next(row for row in sheet if row[0].value == "cash open")
            opening[half_2025 - 1].value += 1
            self.assertEqual(abs(WorkbookTests.arithmetic_value(book, same_date[6])), 1)
        finally:
            book.close()

    def test_discovery_cannot_silently_skip_interim_cadence_or_accept_wrong_dates(self):
        for fault in ("missing_slot", "duplicate_slot", "missing_reported_period", "wrong_quarter_dates", "wrong_half_dates"):
            with self.subTest(fault=fault):
                company = SyntheticCompany().include_interims()
                if fault == "missing_slot":
                    company.discovery["period_coverage"].pop()
                elif fault == "duplicate_slot":
                    company.discovery["period_coverage"].append(copy.deepcopy(company.discovery["period_coverage"][0]))
                elif fault == "missing_reported_period":
                    company.discovery["periods"] = [item for item in company.discovery["periods"] if item["id"] != "2025-Q1"]
                elif fault == "wrong_quarter_dates":
                    next(item for item in company.discovery["periods"] if item["id"] == "2025-Q1")["end"] = "2025-06-30"
                else:
                    next(item for item in company.discovery["periods"] if item["id"] == "2025-H1")["end"] = "2025-09-30"
                directory, record, provider = self.run_case(company)
                self.assertEqual(record["status"], "incomplete")
                self.assertEqual([call["stage"] for call in provider.calls], ["discover"])
                self.assertEqual(company.fetches, [])
                self.assertFalse((directory / "model.xlsx").exists())

    def test_missing_interim_values_require_read_citations_and_explicit_absence(self):
        def missing_observation(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "shares")
            row["observations"] = [item for item in row["observations"] if item["period"] != "2025-H1"]

        def missing_citations(candidate):
            next(row for row in candidate["rows"] if row["id"] == "net_debt")["unavailable"][0]["references"] = []

        def unread_citations(candidate):
            next(row for row in candidate["rows"] if row["id"] == "net_debt")["unavailable"][0]["references"][0]["unit_id"] = "Never read"

        for mutation in (missing_observation, missing_citations, unread_citations):
            with self.subTest(mutation=mutation.__name__):
                company = SyntheticCompany().include_interims()
                directory, record, _ = self.run_case(company, ScriptedProvider(company, mutate=mutation))
                self.assertEqual(record["status"], "incomplete")
                self.assertFalse((directory / "model.xlsx").exists())

    def test_independent_inventory_requires_every_disclosed_period_of_a_matched_row(self):
        def omit_half_year(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "minority_interest")
            row["observations"] = [item for item in row["observations"] if item["period"] != "2025-H1"]
            row["unavailable"].append({"period": "2025-H1", "explanation": "Incorrect draft omission despite source disclosure.",
                                       "references": copy.deepcopy(row["definition_refs"])})

        company = SyntheticCompany().include_interims()
        directory, record, _ = self.run_case(company, ScriptedProvider(company, mutate=omit_half_year))
        self.assertEqual(record["status"], "incomplete")
        self.assertFalse((directory / "model.xlsx").exists())
        failed = [check for check in record["checks"] if not check["passed"]]
        self.assertTrue(any(check["name"] == "inventory_match:equity_minority_interest" for check in failed), record["issues"])

    def test_shared_stock_date_can_reuse_evidence_but_different_durations_or_dates_cannot(self):
        def share_opening_cash(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "cash_open")
            annual = next(item for item in row["observations"] if item["period"] == "2025")
            half = next(item for item in row["observations"] if item["period"] == "2025-H1")
            half["reference"] = copy.deepcopy(annual["reference"])

        company = SyntheticCompany().include_interims()
        _, checks = self.deterministic_case(company, share_opening_cash)
        self.assertTrue(all(check["passed"] for check in checks), checks)
        for row_id in ("revenue", "cash_close"):
            def reuse_different_period(candidate):
                row = next(row for row in candidate["rows"] if row["id"] == row_id)
                annual = next(item for item in row["observations"] if item["period"] == "2025")
                half = next(item for item in row["observations"] if item["period"] == "2025-H1")
                half["reference"] = copy.deepcopy(annual["reference"])

            with self.subTest(row_id=row_id), self.assertRaisesRegex(ValueError, "different duration periods or stock dates"):
                self.deterministic_case(company, reuse_different_period)

    def test_stock_balances_for_the_same_date_cannot_disagree_between_period_columns(self):
        company = SyntheticCompany().include_interims()
        for row_id, half_value in (("cash_open", 101), ("cash_close", 116),
                                   ("balance_sheet_cash", 116), ("net_debt", 209)):
            company.display_values[row_id][3] = half_value
        _, checks = self.deterministic_case(company)
        duplicate_date = [check for check in checks if check["name"].startswith("stock_date_consistency:cash_open:")]
        self.assertTrue(duplicate_date)
        self.assertTrue(any(not check["passed"] and abs(check["difference"]) == 1 for check in duplicate_date))
        self.assertTrue(next(check for check in checks if check["name"] == "cash_rollforward:2025-H1")["passed"])

    def test_optional_annual_subtotal_absence_does_not_invalidate_complete_half_year_accounts(self):
        company = SyntheticCompany().include_interims()
        subtotal = copy.deepcopy(next(row for row in company.rows if row["id"] == "minority_interest"))
        reference = {"source_id": "annual", "unit_id": f"{company.sheet}!R{len(company.rows) + 3}"}
        subtotal.update(id="optional_equity_subtotal", label="Optional annual equity subtotal",
                        definition="Separately disclosed annual subtotal, not a mandatory equity role.",
                        observations=[{"period": str(year), "reference": dict(reference, number_index=index,
                                                                              number_format="dot_decimal")}
                                      for index, year in enumerate((2024, 2025))],
                        definition_refs=[*subtotal["definition_refs"][:2], reference])
        subtotal["unavailable"].append({"period": "2025-H1", "explanation": "H1 accounts omit this optional annual subtotal.",
                                        "references": [reference]})
        company.rows.append(subtotal)
        company.display_values[subtotal["id"]] = [220, 245, None, None, None]
        company.candidate["equations"].append({"target": subtotal["id"],
                                                "terms": [{"row_id": "share_capital", "coefficient": 1},
                                                          {"row_id": "minority_interest", "coefficient": 1}],
                                                "explanation": "Synthetic optional subtotal of separately reported components."})
        _, record, _ = self.run_case(company)
        self.assertEqual(record["status"], "accepted", record["issues"])
        half_check = next(check for check in record["checks"] if check["name"] == "equation:optional_equity_subtotal:2025-H1")
        self.assertTrue(half_check["passed"])
        self.assertEqual(half_check["basis"], "not_applicable")
        row = next(row for row in record["rows"] if row["id"] == subtotal["id"])
        self.assertNotIn("2025-H1", row["values"])

    def test_discovery_notes_continue_and_preserve_independent_review_contexts(self):
        company = SyntheticCompany()
        notes = [
            "Use consolidated group accounts, not the separate parent-company accounts.",
            "Use FY2025 accounts and their FY2024 comparatives for a consistent basis.",
            "Do not mix later reorganized segments with the earlier comparable series.",
            "An optional spreadsheet mirror is unavailable; the primary disclosure is usable.",
            "Future periods are blank workbook columns, not forecasts to research.",
        ]
        company.discovery["notes"] = notes
        directory, record, provider = self.run_case(company, years=[2024, 2025], forecast_years=3)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(record["discovery"]["notes"], notes)
        self.assertEqual(record["issues"], [])
        discovery_call = provider.calls[0]
        self.assertEqual(set(discovery_call),
                         {"stage", "company_query", "historical_years", "years", "as_of"})
        self.assertEqual(discovery_call["years"], [2024, 2025])
        self.assertEqual(record["request"]["forecast_years"], 3)
        self.assertEqual(record["request"]["max_cost_usd"], 20)
        log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(log[0]["result"]["notes"], notes)
        for stage in ("extract", "model_review"):
            entries = [entry for entry in log if entry["stage"] == stage]
            self.assertTrue(entries)
            self.assertTrue(all(entry["payload"]["discovery"]["notes"] == notes for entry in entries))
        for stage in ("source_review", "coverage_inventory"):
            calls = [call for call in provider.calls if call["stage"] == stage]
            self.assertTrue(calls)
            for call in calls:
                self.assertEqual(set(call["discovery"]), {"company", "periods"})
                self.assertTrue(all(note not in json.dumps(call) for note in notes))
        book = load_workbook(directory / "model.xlsx", data_only=False)
        try:
            for section in SECTIONS:
                sheet = book[section]
                columns = [cell.column for cell in sheet[4] if "forecast (blank)" in str(cell.value)]
                self.assertEqual(len(columns), 3)
                self.assertTrue(all(sheet.cell(row, column).value is None
                                    for row in range(5, sheet.max_row + 1) for column in columns))
        finally:
            book.close()

    def test_discovery_blockers_stop_before_downloads_and_preserve_usage(self):
        for blocker in ("Company identity is ambiguous between two unrelated issuers.",
                        "Required annual accounts could not be located."):
            with self.subTest(blocker=blocker):
                company = SyntheticCompany()
                company.discovery["notes"] = ["The workflow should continue if evidence is sufficient."]
                company.discovery["blocking_issues"] = [blocker]
                directory, record, provider = self.run_case(company)
                self.assertEqual(record["status"], "incomplete")
                self.assertIn(blocker, " ".join(record["issues"]))
                self.assertEqual([call["stage"] for call in provider.calls], ["discover"])
                self.assertEqual(company.fetches, [])
                self.assertEqual(record["source_attempts"], [])
                self.assertEqual(record["checks"], [])
                self.assertFalse((directory / "model.xlsx").exists())
                self.assertEqual(record["usage"]["calls"], 1)
                self.assertEqual(record["usage"]["by_stage"]["discover"]["total_tokens"], 1100)
                self.assertGreater(record["usage"]["estimated_cost_usd"], 0)
                log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
                self.assertEqual(log[0]["result"]["blocking_issues"], [blocker])

    def test_discovery_notes_cannot_replace_required_identity_periods_or_sources(self):
        for missing in ("identity", "periods", "sources"):
            with self.subTest(missing=missing):
                company = SyntheticCompany()
                company.discovery["notes"] = ["Everything needed for extraction is available."]
                if missing == "identity":
                    company.discovery["company"]["legal_name"] = ""
                else:
                    company.discovery[missing] = []
                directory, record, provider = self.run_case(company)
                self.assertEqual(record["status"], "incomplete")
                self.assertTrue(record["issues"])
                self.assertEqual([call["stage"] for call in provider.calls], ["discover"])
                self.assertEqual(company.fetches, [])
                self.assertFalse((directory / "model.xlsx").exists())

    def test_optional_source_download_failure_does_not_discard_complete_evidence(self):
        company = SyntheticCompany()
        optional = dict(company.source, id="optional", url="https://example.invalid/optional.xlsx",
                        title="SYNTHETIC OPTIONAL SPREADSHEET MIRROR")
        company.discovery["sources"].append(optional)
        company.discovery["notes"] = ["Use the primary disclosure if the optional spreadsheet mirror fails."]
        primary_fetch = company.fetch
        attempted = []

        def fetch(url):
            attempted.append(url)
            if url == optional["url"]:
                raise HTTPError(url, 503, "Synthetic unavailable mirror", {}, None)
            return primary_fetch(url)

        company.fetch = fetch
        _, record, provider = self.run_case(company)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(attempted, [company.source["url"], optional["url"]])
        self.assertEqual([source["id"] for source in record["sources"]], ["annual"])
        failure = next(item for item in record["source_attempts"] if item["source"]["id"] == "optional")
        self.assertEqual(failure["status"], "failed")
        self.assertIn("503", failure["error"])
        self.assertTrue(failure["started_at"] and failure["finished_at"])
        extraction = next(call for call in provider.calls if call["stage"] == "extract")
        self.assertIn("optional", " ".join(extraction["feedback"]))
        self.assertTrue(record["checks"] and all(check["passed"] for check in record["checks"]))

    def test_unavailable_required_source_cannot_be_approved_by_reassuring_notes(self):
        company = SyntheticCompany()
        company.discovery["notes"] = ["The primary disclosure is accessible and sufficient for every statement."]

        def unavailable(url):
            raise HTTPError(url, 503, "Synthetic unavailable primary disclosure", {}, None)

        company.fetch = unavailable
        directory, record, _ = self.run_case(company)
        self.assertEqual(record["status"], "incomplete")
        self.assertFalse((directory / "model.xlsx").exists())
        self.assertEqual(record["sources"], [])
        self.assertEqual(record["rows"], [])
        self.assertTrue(record["issues"])
        self.assertTrue(record["source_attempts"])
        self.assertTrue(all(item["status"] == "failed" and "503" in item["error"]
                            for item in record["source_attempts"]))

    def test_two_companies_discover_different_layouts_without_configuration(self):
        first = None
        before = None
        for variant in (0, 1):
            company = SyntheticCompany(variant)
            directory, record, provider = self.run_case(company)
            self.assertEqual(record["status"], "accepted", record["issues"])
            self.assertTrue(all(item["passed"] for item in record["checks"]))
            rows = {row["id"]: row for row in record["rows"]}
            self.assertEqual(rows["revenue"]["values"], {"2024": 1000.0, "2025": 1100.0})
            self.assertEqual(rows["net_debt"]["values"], {"2024": 200.0, "2025": 220.0})
            self.assertEqual(rows["revenue"]["unit"], company.currency)
            self.assertFalse((self.root / "configs").exists())
            archived = record["sources"][0]
            self.assertEqual(archived["sha256"], hashlib.sha256(Path(archived["path"]).read_bytes()).hexdigest())
            self.assertTrue(archived["fetched_at"])
            self.assertEqual(record["workbook_sha256"],
                             hashlib.sha256((directory / "model.xlsx").read_bytes()).hexdigest())
            book = load_workbook(directory / "model.xlsx", data_only=False)
            try:
                revenue_row = next(row for row in book["income_statement"] if row[0].value == "revenue")
                self.assertEqual(revenue_row[3].value, 1000)
                for sheet in book:
                    if sheet.title in SECTIONS:
                        forecast_columns = [cell.column for cell in sheet[4]
                                            if "forecast (blank)" in str(cell.value)]
                        self.assertEqual(len(forecast_columns), 5)
                        self.assertTrue(all(sheet.cell(row, column).value is None
                                            for row in range(5, sheet.max_row + 1)
                                            for column in forecast_columns))
                self.assertTrue(any(row[6].data_type == "f" for row in book["Checks"].iter_rows(min_row=2)))
            finally:
                book.close()
            log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
            self.assertEqual(len(log), record["call_count"])
            self.assertTrue(all(entry["started_at"] and entry["finished_at"] for entry in log))
            self.assertTrue(all(isinstance(entry["elapsed_seconds"], (int, float))
                                and entry["elapsed_seconds"] >= 0 for entry in log))
            self.assertIn("source_review", [entry["stage"] for entry in log])
            self.assertIn("model_review", [entry["stage"] for entry in log])
            if first is None:
                first = directory
                before = {file.relative_to(first): file.read_bytes() for file in first.rglob("*") if file.is_file()}
            else:
                self.assertNotEqual(first, directory)
                self.assertEqual(before, {file.relative_to(first): file.read_bytes()
                                          for file in first.rglob("*") if file.is_file()})
        # Repeating the identical request must also preserve its earlier run.
        self.run_case(SyntheticCompany())
        self.assertEqual(before, {file.relative_to(first): file.read_bytes()
                                  for file in first.rglob("*") if file.is_file()})

    def test_large_candidates_can_be_saved_in_batches_before_review(self):
        company = SyntheticCompany()
        provider = ScriptedProvider(company, batch=True)
        _, record, _ = self.run_case(company, provider)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(len(record["rows"]), len(company.rows))
        self.assertEqual(provider.saved_batches, (len(company.rows) + 11) // 12)

    def test_wrong_source_arithmetic_cannot_be_approved_by_agreement(self):
        for row_id in ("assets", "cash_change"):
            with self.subTest(row_id=row_id):
                company = SyntheticCompany()
                company.discovery["notes"] = ["The source figures appear suitable for historical modelling."]
                company.display_values[row_id][1] += 25
                directory, record, provider = self.run_case(company)
                self.assertEqual(record["status"], "incomplete")
                self.assertFalse((directory / "model.xlsx").exists())
                self.assertTrue(any(not item["passed"] for item in record["checks"]))
                self.assertNotIn("model_review", [call["stage"] for call in provider.calls])

    def test_decimal_precision_limits_tolerance_for_balance_sheet_checks(self):
        company = SyntheticCompany()
        for row_id, amount in (("assets", "1200.25"), ("liabilities", "720.00"), ("equity", "480.00")):
            company.display_values[row_id][1] = amount
        directory, record, _ = self.run_case(company)
        self.assertEqual(record["status"], "incomplete")
        check = next(item for item in record["checks"] if item["name"] == "balance_sheet:2025")
        self.assertAlmostEqual(check["difference"], 0.25)
        self.assertLess(check["tolerance"], 0.02)
        self.assertFalse(check["passed"])
        self.assertFalse((directory / "model.xlsx").exists())

    def test_missing_duplicate_and_unread_references_or_coverage_cannot_pass(self):
        def missing(candidate):
            candidate["rows"][0]["observations"].pop()

        def duplicate(candidate):
            candidate["rows"][0]["observations"][1] = copy.deepcopy(candidate["rows"][0]["observations"][0])

        def unread(candidate):
            candidate["rows"][0]["observations"][0]["reference"]["unit_id"] = "Never read"

        def incomplete_coverage(candidate):
            candidate["coverage"].pop()

        for mutation in (missing, duplicate, unread, incomplete_coverage):
            with self.subTest(mutation=mutation.__name__):
                company = SyntheticCompany()
                company.discovery["notes"] = ["All necessary evidence is believed to be available."]
                directory, record, _ = self.run_case(company, ScriptedProvider(company, mutate=mutation))
                self.assertEqual(record["status"], "incomplete")
                self.assertTrue(record["issues"])
                self.assertFalse((directory / "model.xlsx").exists())

    def test_reviewer_can_request_original_context_independently(self):
        class Reader(ScriptedProvider):
            searched_context = False
            requested_context = False

            def complete(self, **kwargs):
                payload = kwargs["payload"]
                if payload["stage"] == "model_review" and not self.searched_context:
                    self.searched_context = True
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("model_review")
                    return {"checks": [], "matches": [], "issues": [], "search_query": "FY2024", "read_requests": []}
                if payload["stage"] == "model_review" and not self.requested_context:
                    self.requested_context = True
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("model_review")
                    if not any(item["unit_id"] == f"{self.company.sheet}!R2"
                               for item in payload["search_results"]["matches"]):
                        raise AssertionError("Reviewer did not receive requested archive search results")
                    return {"checks": [], "matches": [], "issues": [], "search_query": "", "read_requests": [
                        {"source_id": "annual", "unit_id": f"{self.company.sheet}!R2"}]}
                if payload["stage"] == "model_review":
                    if not any(item["unit_id"] == f"{self.company.sheet}!R2" for item in payload["evidence"]):
                        raise AssertionError("Reviewer did not receive requested source context")
                return super().complete(**kwargs)

        company = SyntheticCompany()
        provider = Reader(company)
        _, record, _ = self.run_case(company, provider)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertTrue(provider.searched_context)
        self.assertTrue(provider.requested_context)
        self.assertGreaterEqual(sum(call["stage"] == "model_review" for call in provider.calls), 3)

    def test_whole_documents_reach_extraction_and_independent_reviews_without_character_caps(self):
        company = SyntheticCompany()
        book = load_workbook(io.BytesIO(company.payload()))
        sheet = book.active
        # Make source context exceed both former read and AI payload caps, and
        # place a meaningful note outside the first 60 catalog previews.
        for _ in range(45):
            sheet.append(["Synthetic documentary context. " * 250])
        late_note = "MATERIAL LATE NOTE: reported debt excludes supplier financing."
        sheet.append([late_note])
        final_unit = f"{company.sheet}!R{sheet.max_row}"
        output = io.BytesIO()
        book.save(output)
        original = output.getvalue()
        supplement = dict(company.source, id="supplement", url="https://example.invalid/supplement.html",
                          title="SYNTHETIC SUPPLEMENT", kind="results")
        company.discovery["sources"].append(supplement)

        def fetch(url):
            if url == supplement["url"]:
                return b"<p>Synthetic subsequent documentary context.</p>", url, "text/html"
            return original, url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        company.fetch = fetch

        class DocumentReader(ScriptedProvider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                payload = kwargs["payload"]
                if payload["stage"] == "extract" and result["action"] == "propose":
                    if not any(unit["source_id"] == "supplement" for unit in payload["evidence"]):
                        result.update(action="read", candidate=None, read_requests=[
                            {"source_id": "supplement", "unit_id": "chunk:1"}])
                        return result
                if result.get("read_requests"):
                    # Selecting one unit must deliver every unit of that source.
                    result["read_requests"] = result["read_requests"][:1]
                return result

        directory, record, provider = self.run_case(company, DocumentReader(company))
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertTrue((directory / "model.xlsx").exists())
        extraction = [call for call in provider.calls if call["stage"] == "extract"]
        self.assertEqual({unit["source_id"] for unit in extraction[-1]["evidence"]},
                         {"annual", "supplement"})
        self.assertGreater(len(json.dumps(extraction[-1])), 260_000)
        for stage in ("extract", "source_review", "source_alignment", "coverage_inventory", "model_review"):
            calls = [call for call in provider.calls if call["stage"] == stage]
            if stage in ("source_review", "coverage_inventory", "model_review"):
                self.assertEqual(calls[0]["evidence"], [], stage)
            received = [call for call in calls if call["evidence"]]
            self.assertTrue(received, stage)
            for call in received:
                with self.subTest(stage=stage):
                    source = next(item for item in call["catalog"] if item["id"] == "annual")
                    self.assertTrue(source["catalog_truncated"])
                    self.assertNotIn(final_unit, [unit["id"] for unit in source["units"]])
                    annual = [unit for unit in call["evidence"] if unit["source_id"] == "annual"]
                    self.assertEqual(len(annual), source["unit_count"])
                    self.assertEqual(len({unit["unit_id"] for unit in annual}), len(annual))
                    self.assertIn(late_note, next(unit["text"] for unit in annual
                                                 if unit["unit_id"] == final_unit))
                    self.assertGreater(sum(len(unit["text"]) for unit in annual), 260_000)
                    if stage in ("source_review", "coverage_inventory", "model_review"):
                        self.assertNotIn("supplement", {unit["source_id"] for unit in call["evidence"]})

    def test_every_required_review_must_be_present_and_pass(self):
        for fault in ("missing_source", "duplicate_source", "missing_source_context", "missing_alignment",
                      "duplicate_alignment", "missing_alignment_context",
                      "unread_alignment_context",
                      "missing_model", "duplicate_model", "model_reject"):
            with self.subTest(fault=fault):
                company = SyntheticCompany()
                directory, record, _ = self.run_case(company, ScriptedProvider(company, review_fault=fault))
                self.assertEqual(record["status"], "incomplete")
                self.assertTrue(record["issues"])
                self.assertFalse((directory / "model.xlsx").exists())
                if fault == "model_reject":
                    self.assertIn("identity", " ".join(record["issues"]))

    def test_source_review_can_repair_a_semantic_mapping_without_a_human(self):
        def wrong_share_definition(candidate):
            rows = {row["id"]: row for row in candidate["rows"]}
            rows["shares"]["observations"] = copy.deepcopy(rows["share_capital"]["observations"])

        company = SyntheticCompany()
        provider = ScriptedProvider(company, mutate=wrong_share_definition, repair=True)
        directory, record, _ = self.run_case(company, provider)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(len(record["attempts"]), 2)
        self.assertIn("shares", " ".join(record["attempts"][0]["issues"]))
        shares = next(row for row in record["rows"] if row["id"] == "shares")
        self.assertEqual(shares["values"], {"2024": 100.0, "2025": 100.0})
        repairs = [call for call in provider.calls if call["stage"] == "extract" and call["feedback"]]
        self.assertTrue(repairs)
        self.assertTrue(record["attempts"][0]["reviews"])
        for attempt in record["attempts"]:
            self.assertTrue(any(review["stage"] == "source_review" for review in attempt["reviews"]))
        blind_calls = [call for call in provider.calls if call["stage"] == "source_review"]
        self.assertGreaterEqual(sum(not call["evidence"] for call in blind_calls), 4)
        log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len(log), record["call_count"])
        self.assertEqual(record["usage"]["calls"], record["call_count"])
        self.assertTrue(any("shares" in json.dumps(entry["result"].get("facts", []))
                            for entry in log if entry["stage"] == "source_review"))

    def test_independent_passes_hide_candidate_and_previous_verdicts(self):
        company = SyntheticCompany()
        directory, record, provider = self.run_case(company)
        self.assertEqual(record["status"], "accepted", record["issues"])
        blind = [call for call in provider.calls if call["stage"] == "source_review"]
        inventory = [call for call in provider.calls if call["stage"] == "coverage_inventory"]
        models = [call for call in provider.calls if call["stage"] == "model_review"]
        self.assertTrue(blind and inventory and models)
        for call in blind:
            self.assertEqual(set(call["discovery"]), {"company", "periods"})
            for concept in call["concepts"]:
                self.assertEqual(set(concept), {"id", "label", "section", "entity", "periods"})
            for hidden in ("rows", "candidate", "reporting_basis", "checks", "source_reviews", "feedback", "facts"):
                self.assertNotIn(hidden, call)
        for call in inventory:
            for hidden in ("rows", "concepts", "candidate", "coverage", "checks", "reviews", "feedback", "facts"):
                self.assertNotIn(hidden, call)
        self.assertEqual({call["section"] for call in inventory}, set(SECTIONS))
        first_model = next(index for index, call in enumerate(provider.calls) if call["stage"] == "model_review")
        self.assertEqual({call["section"] for call in provider.calls[:first_model]
                          if call["stage"] == "coverage_inventory"}, set(SECTIONS))
        for call in models:
            for hidden in ("checks", "source_reviews", "facts", "reviews", "feedback", "independent_checks"):
                self.assertNotIn(hidden, call)
            self.assertEqual({item["section"] for item in call["inventory"]}, set(SECTIONS))
        log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
        aligned = [entry for entry in log if entry["stage"] == "source_alignment"]
        self.assertTrue(aligned)
        for entry in aligned:
            for hidden in ("checks", "source_reviews", "reviews", "feedback", "independent_checks"):
                self.assertNotIn(hidden, entry["payload"])
            reconstructed = [fact for previous in log[:entry["call"] - 1]
                             if previous["stage"] == "source_review" for fact in previous["result"]["facts"]]
            self.assertTrue(all(fact in reconstructed for fact in entry["payload"]["facts"]))
        self.assertTrue(any(entry["stage"] == "coverage_inventory" and entry["result"]["items"] for entry in log))
        self.assertTrue(any(entry["stage"] == "source_review" and entry["result"]["facts"] for entry in log))

    def test_blind_reconstruction_rejects_valid_tokens_with_wrong_meaning(self):
        def wrong_number(candidate):
            rows = {row["id"]: row for row in candidate["rows"]}
            rows["shares"]["observations"] = copy.deepcopy(rows["share_capital"]["observations"])

        def wrong_year(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "share_capital")
            for observation in row["observations"]:
                observation["reference"]["number_index"] = 1 - observation["reference"]["number_index"]

        def wrong_scale(candidate):
            next(row for row in candidate["rows"] if row["id"] == "shares")["scale"] = 1000

        def wrong_entity(candidate):
            next(row for row in candidate["rows"] if row["id"] == "shares")["entity"] = "Unrelated legal entity"

        for mutation in (wrong_number, wrong_year, wrong_scale, wrong_entity):
            with self.subTest(mutation=mutation.__name__):
                company = SyntheticCompany()
                directory, record, provider = self.run_case(company, ScriptedProvider(company, mutate=mutation))
                self.assertEqual(record["status"], "incomplete")
                self.assertFalse((directory / "model.xlsx").exists())
                self.assertIn("source_review", [call["stage"] for call in provider.calls])
                self.assertTrue(record["attempts"][0]["reviews"])
                self.assertTrue(any(not check["passed"] and check["name"].startswith("blind_source:")
                                    for check in record["checks"]), record["issues"])

    def test_independent_inventory_blocks_omission_even_when_all_verdicts_approve(self):
        def omit_supplement(candidate):
            candidate["rows"] = [row for row in candidate["rows"] if row["id"] != "minority_interest"]

        company = SyntheticCompany()
        directory, record, provider = self.run_case(company, ScriptedProvider(company, mutate=omit_supplement))
        self.assertEqual(record["status"], "incomplete")
        self.assertFalse((directory / "model.xlsx").exists())
        self.assertTrue(all(check["passed"] for check in record["checks"]
                            if not check["name"].startswith(("inventory", "independent"))))
        model = next(review["result"] for review in record["reviews"] if review["stage"] == "model_review")
        self.assertTrue(all(check["passed"] for check in model["checks"]))
        missing = next(match for match in model["matches"] if match["inventory_id"] == "equity_minority_interest")
        self.assertEqual(missing["row_id"], "")
        self.assertIn("minority_interest", " ".join(record["issues"]))

    def test_equal_values_do_not_override_a_failed_definition_or_vintage_alignment(self):
        def wrong_definition(candidate):
            next(row for row in candidate["rows"] if row["id"] == "shares")["definition"] = "A different economic share measure."

        def wrong_vintage(candidate):
            candidate["reporting_basis"] = "Original FY2024 vintage, inconsistent with the FY2025 comparative source."

        for mutation, fault in ((wrong_definition, "definition_disagreement"), (wrong_vintage, "basis_disagreement")):
            with self.subTest(fault=fault):
                company = SyntheticCompany()
                directory, record, _ = self.run_case(company, ScriptedProvider(company, mutate=mutation, review_fault=fault))
                self.assertEqual(record["status"], "incomplete")
                blind = [check for check in record["checks"] if check["name"].startswith("blind_source:")]
                self.assertTrue(blind and all(check["passed"] for check in blind))
                alignment = [review["result"] for review in record["reviews"] if review["stage"] == "source_alignment"]
                self.assertTrue(any(not check["passed"] for result in alignment for check in result["checks"]))
                self.assertFalse((directory / "model.xlsx").exists())

    def test_one_source_unit_disagreement_at_equal_precision_is_not_rounding(self):
        def swap_years(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "shares")
            for observation in row["observations"]:
                observation["reference"]["number_index"] = 1 - observation["reference"]["number_index"]

        company = SyntheticCompany()
        company.display_values["shares"] = [100, 101]
        directory, record, _ = self.run_case(company, ScriptedProvider(company, mutate=swap_years))
        self.assertEqual(record["status"], "incomplete")
        blind = [check for check in record["checks"] if check["name"].startswith("blind_source:shares:")]
        self.assertEqual(len(blind), 2)
        self.assertTrue(all(abs(check["difference"]) == 1 and not check["passed"] for check in blind))
        self.assertFalse((directory / "model.xlsx").exists())

    def test_different_source_precision_accepts_equivalent_rounded_disclosure(self):
        company = SyntheticCompany()
        company.display_values["shares"] = [100.4, 101.4]
        rounded = dict(company.source, id="rounded", url="https://example.invalid/rounded-shares.xlsx",
                       title="SYNTHETIC ROUNDED SHARE SUMMARY", kind="results")
        company.discovery["sources"].append(rounded)
        original_fetch = company.fetch

        def fetch(url):
            if url != rounded["url"]:
                return original_fetch(url)
            book = Workbook()
            sheet = book.active
            sheet.title = "Rounded"
            sheet.append(["SYNTHETIC " + company.name + "; same share definition and FY2025 comparative basis; whole shares."])
            sheet.append(["Reported line", "FY2024", "FY2025"])
            sheet.append(["shares", 100, 101])
            output = io.BytesIO()
            book.save(output)
            return output.getvalue(), url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        def use_rounded_source(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "shares")
            row["definition_refs"] = [{"source_id": "rounded", "unit_id": f"Rounded!R{number}"}
                                      for number in (1, 2, 3)]
            for observation in row["observations"]:
                observation["reference"].update(source_id="rounded", unit_id="Rounded!R3")

        company.fetch = fetch
        _, record, _ = self.run_case(company, ScriptedProvider(company, mutate=use_rounded_source))
        self.assertEqual(record["status"], "accepted", record["issues"])
        blind = [check for check in record["checks"] if check["name"].startswith("blind_source:shares:")]
        self.assertEqual(len(blind), 2)
        self.assertTrue(all(check["passed"] and 0 < abs(check["difference"]) < check["tolerance"] for check in blind))

    def test_independent_sections_can_reference_one_shared_equity_total(self):
        class SharedTotalProvider(ScriptedProvider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                payload = kwargs["payload"]
                equity = next(row for row in self.company.rows if row["role"] == "total_equity")
                if payload["stage"] == "coverage_inventory" and payload["section"] == "equity" and result["items"]:
                    result["items"].append({"id": "equity_total_equity", "label": equity["label"],
                                            "entity": equity["entity"], "references": copy.deepcopy(equity["definition_refs"]),
                                            "periods": [observation["period"] for observation in equity["observations"]]})
                if payload["stage"] == "model_review" and result["checks"]:
                    result["matches"].append({"inventory_id": "equity_total_equity", "row_id": equity["id"],
                                              "reason": "The equity schedule repeats the same consolidated equity total.",
                                              "references": copy.deepcopy(equity["definition_refs"])})
                return result

        company = SyntheticCompany()
        _, record, _ = self.run_case(company, SharedTotalProvider(company))
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(sum(row["role"] == "total_equity" for row in record["rows"]), 1)
        names = {check["name"] for check in record["checks"] if check["passed"]}
        self.assertIn("inventory_match:balance_sheet_equity", names)
        self.assertIn("inventory_match:equity_total_equity", names)

    def test_incomplete_or_duplicate_independent_inventory_and_matches_fail(self):
        for fault in ("empty_inventory", "duplicate_inventory", "wrong_inventory_section",
                      "missing_match", "duplicate_match", "reused_match"):
            with self.subTest(fault=fault):
                company = SyntheticCompany()
                directory, record, _ = self.run_case(company, ScriptedProvider(company, review_fault=fault))
                self.assertEqual(record["status"], "incomplete")
                self.assertTrue(record["issues"])
                self.assertFalse((directory / "model.xlsx").exists())

    def test_each_independent_stage_must_read_its_own_evidence(self):
        class SkipsReading(ScriptedProvider):
            def __init__(self, company, stage):
                super().__init__(company)
                self.skipped_stage = stage

            def complete(self, **kwargs):
                payload = kwargs["payload"]
                if payload["stage"] == self.skipped_stage and not payload["evidence"]:
                    # Return final scripted selections immediately, without asking
                    # the workflow to read any source in this independent stage.
                    evidence = [{"source_id": source["id"], "unit_id": unit["id"]}
                                for source in payload["catalog"] for unit in source["units"]]
                    kwargs = dict(kwargs, payload=dict(payload, evidence=evidence))
                return super().complete(**kwargs)

        for stage in ("source_review", "coverage_inventory", "model_review"):
            with self.subTest(stage=stage):
                company = SyntheticCompany()
                directory, record, _ = self.run_case(company, SkipsReading(company, stage))
                self.assertEqual(record["status"], "incomplete")
                self.assertIn("not independently read", " ".join(record["issues"]))
                self.assertFalse((directory / "model.xlsx").exists())

    def test_new_independent_calls_respect_dollar_budget_and_preserve_logs(self):
        directory, record, provider = self.run_case(max_cost_usd=0.015)
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(record["call_count"], 5)
        self.assertEqual(len(provider.calls), 5)
        self.assertIn("budget", " ".join(record["issues"]))
        self.assertFalse((directory / "model.xlsx").exists())
        log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len(log), 5)
        self.assertTrue(any(entry["stage"] == "source_review" for entry in log))
        self.assertEqual(record["usage"]["calls"], 5)
        self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], 0.015)
        self.assertEqual(record["budget"]["status"], "exhausted")
        self.assertAlmostEqual(record["usage"]["by_stage"]["source_review"]["estimated_cost_usd"], 0.006)

    def deterministic_case(self, company, mutate=None):
        candidate = copy.deepcopy(company.candidate)
        if mutate:
            mutate(candidate)
        with tempfile.TemporaryDirectory(dir=self.root) as temporary:
            store = SourceStore(Path(temporary), fetcher=company.fetch)
            store.add(company.source)
            refs = [{"source_id": source["id"], "unit_id": unit["id"]}
                    for source in store.catalog() for unit in source["units"]]
            evidence = store.read(refs)
            return materialize_and_check(candidate, company.discovery, store,
                                         {(item["source_id"], item["unit_id"]) for item in evidence})

    def test_cash_and_segment_bridges_require_reported_distinct_anchors(self):
        def remove_cash_bridge(candidate):
            candidate["equations"] = [equation for equation in candidate["equations"]
                                      if equation["target"] != "balance_sheet_cash"]

        def remove_segment_bridge(candidate):
            candidate["equations"] = [equation for equation in candidate["equations"]
                                      if equation["target"] != "segment_group_profit"]

        def copy_cash_tokens(candidate):
            rows = {row["id"]: row for row in candidate["rows"]}
            rows["balance_sheet_cash"]["observations"] = copy.deepcopy(rows["cash_close"]["observations"])

        def derive_liabilities(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "liabilities")
            row.update(observations=[], terms=[{"row_id": "assets", "coefficient": 1},
                                               {"row_id": "equity", "coefficient": -1}])

        def duplicate_segment_entity(candidate):
            next(row for row in candidate["rows"] if row["id"] == "profit_b")["entity"] = "Alpha"

        for mutation, expected in (
            (remove_cash_bridge, "cash_balance_sheet_reconciliation"),
            (remove_segment_bridge, "segment_profit_reconciliation"),
            (copy_cash_tokens, "cash_balance_sheet_reconciliation"),
            (derive_liabilities, "balance_sheet_independent:2025"),
            (duplicate_segment_entity, "segment_entity_coverage"),
        ):
            with self.subTest(mutation=mutation.__name__):
                _, checks = self.deterministic_case(SyntheticCompany(), mutation)
                self.assertFalse(next(check for check in checks if check["name"] == expected)["passed"])

    def test_liability_total_can_sum_independently_reported_components(self):
        company = SyntheticCompany()
        liabilities = next(row for row in company.rows if row["id"] == "liabilities")
        for row_id, values in (("current_liabilities", [350, 400]), ("noncurrent_liabilities", [250, 320])):
            row = copy.deepcopy(liabilities)
            reference = {"source_id": "annual", "unit_id": f"{company.sheet}!R{len(company.rows) + 3}"}
            row.update(id=row_id, label=row_id.replace("_", " "), role="other",
                       definition="Separately reported " + row_id,
                       observations=[{"period": str(year), "reference": dict(reference, number_index=index,
                                                                        number_format="dot_decimal")}
                                     for index, year in enumerate((2024, 2025))],
                       definition_refs=[*liabilities["definition_refs"][:2], reference])
            company.rows.append(row)
            company.display_values[row_id] = values
        liabilities.update(observations=[],
                           definition="Calculated total from independently reported current and noncurrent liabilities.",
                           terms=[{"row_id": row_id, "coefficient": 1}
                                  for row_id in ("current_liabilities", "noncurrent_liabilities")])
        company.display_values["liabilities"] = [None, None]
        _, checks = self.deterministic_case(company)
        self.assertTrue(all(check["passed"] for check in checks), checks)
        balance = [check for check in checks if check["name"].startswith("balance_sheet:")]
        self.assertTrue(balance and all(check["basis"] == "corroboration" for check in balance))

    def test_balanced_cash_statements_still_fail_adjacent_year_cash_discontinuity(self):
        company = SyntheticCompany()
        for row_id, value in (("cash_open", 105), ("cash_close", 135),
                              ("balance_sheet_cash", 135), ("net_debt", 215)):
            company.display_values[row_id][1] = value
        _, checks = self.deterministic_case(company)
        failures = [check for check in checks if not check["passed"]]
        self.assertEqual([check["name"] for check in failures], ["cash_continuity:2025"])
        self.assertEqual(failures[0]["difference"], 5)

    def company_with_cash_adjustment(self, *, sparse=False):
        company = SyntheticCompany()
        adjustment = copy.deepcopy(next(row for row in company.rows if row["id"] == "minority_interest"))
        reference = {"source_id": "annual", "unit_id": f"{company.sheet}!R{len(company.rows) + 3}"}
        adjustment.update(id="cash_scope_adjustment", label="Reported opening cash scope adjustment",
                          section="cash_flow", definition="Reported cash reclassification into the comparative scope.",
                          time_basis="flow",
                          observations=[{"period": str(year), "reference": dict(reference, number_index=index,
                                                                           number_format="dot_decimal")}
                                        for index, year in enumerate((2024, 2025))],
                          definition_refs=[*adjustment["definition_refs"][:2], reference])
        if sparse:
            adjustment["observations"] = [{"period": "2025", "reference": dict(reference, number_index=0,
                                                                             number_format="dot_decimal")}]
            adjustment["unavailable"] = [{"period": "2024", "explanation": "Scope change applies only to FY2025 opening cash.",
                                           "references": [reference]}]
        company.rows.append(adjustment)
        company.display_values["cash_scope_adjustment"] = [None if sparse else 0, 5]
        company.candidate["cash_continuity"][0].update(
            terms=[{"row_id": "cash_scope_adjustment", "coefficient": 1}], references=[reference],
            explanation="The cited cash scope reclassification reconciles closing to next opening cash.")
        for row_id, value in (("cash_open", 105), ("cash_close", 135),
                              ("balance_sheet_cash", 135), ("net_debt", 215)):
            company.display_values[row_id][1] = value
        return company

    def test_adjacent_cash_bridge_accepts_cited_reported_adjustments(self):
        company = self.company_with_cash_adjustment()
        _, checks = self.deterministic_case(company)
        self.assertTrue(all(check["passed"] for check in checks), checks)
        continuity = next(check for check in checks if check["name"] == "cash_continuity:2025")
        self.assertEqual(continuity["basis"], "corroboration")

    def test_sparse_cash_adjustment_preserves_undisclosed_year_as_blank(self):
        company = self.company_with_cash_adjustment(sparse=True)
        directory, record, provider = self.run_case(company)
        self.assertEqual(record["status"], "accepted", record["issues"])
        adjustment = next(row for row in record["rows"] if row["id"] == "cash_scope_adjustment")
        self.assertEqual(adjustment["values"], {"2025": 5.0})
        concepts = [concept for call in provider.calls if call["stage"] == "source_review"
                    for concept in call["concepts"] if concept["id"] == "cash_scope_adjustment"]
        self.assertTrue(concepts and all(concept["periods"] == ["2025"] for concept in concepts))
        book = load_workbook(directory / "model.xlsx", data_only=False)
        try:
            row = next(row for row in book["cash_flow"] if row[0].value == "Reported opening cash scope adjustment")
            self.assertIsNone(row[3].value)
            self.assertEqual(row[4].value, 5)
            continuity = next(item for item in book["Checks"] if item[0].value == "cash_continuity:2025")
            self.assertIn(f"-'cash_flow'!{row[4].coordinate}*1", continuity[6].value)
            self.assertAlmostEqual(WorkbookTests.arithmetic_value(book, continuity[6]), 0)
            row[4].value = 4
            self.assertAlmostEqual(WorkbookTests.arithmetic_value(book, continuity[6]), 1)
        finally:
            book.close()

    def test_sparse_rows_cannot_hide_missing_evidence_or_supply_missing_equation_inputs(self):
        def reused_in_statement(candidate):
            candidate["equations"].append({"target": "minority_interest",
                                            "terms": [{"row_id": "cash_scope_adjustment", "coefficient": 1}],
                                            "explanation": "Invalid reuse of a sparse adjustment in annual statements."})

        def missing_required_period(candidate):
            next(row for row in candidate["rows"] if row["id"] == "cash_scope_adjustment")["observations"] = []

        def uncited_absence(candidate):
            next(row for row in candidate["rows"] if row["id"] == "cash_scope_adjustment")["unavailable"][0]["references"] = []

        company = self.company_with_cash_adjustment(sparse=True)
        with self.assertRaisesRegex(ValueError, "missing or repeated historical observation"):
            self.deterministic_case(company, missing_required_period)
        with self.assertRaisesRegex(ValueError, "cited"):
            self.deterministic_case(company, uncited_absence)
        _, checks = self.deterministic_case(company, reused_in_statement)
        self.assertFalse(next(check for check in checks if check["name"] == "equation:minority_interest:2024")["passed"])

    def test_constructed_cash_identity_is_labelled_and_period_gaps_are_not_compared(self):
        def derived_opening(candidate):
            row = next(row for row in candidate["rows"] if row["id"] == "cash_open")
            row.update(observations=[], terms=[{"row_id": "cash_close", "coefficient": 1},
                                               {"row_id": "cash_change", "coefficient": -1}])

        company = SyntheticCompany()
        _, checks = self.deterministic_case(company, derived_opening)
        cash = [check for check in checks if check["name"].startswith("cash_rollforward:")]
        self.assertTrue(all(check["passed"] and check["basis"] == "arithmetic_construction" for check in cash))
        company.discovery["periods"][1]["start"] = "2025-01-02"
        company.candidate["cash_continuity"] = []
        _, checks = self.deterministic_case(company)
        continuity = next(check for check in checks if check["name"] == "cash_continuity:2025")
        self.assertTrue(continuity["passed"])
        self.assertEqual(continuity["basis"], "not_applicable")

    def test_missing_api_configuration_creates_no_outputs_or_network_calls(self):
        with patch.dict(os.environ, {}, clear=True), patch("research_platform.ai.urlopen") as api:
            with self.assertRaisesRegex(AIError, "OPENAI_API_KEY"):
                initiate_company(self.root, "Arbitrary Unconfigured Company")
            api.assert_not_called()
        self.assertFalse((self.root / "outputs").exists())

    def test_dollar_budget_exhaustion_stops_without_a_workbook(self):
        directory, record, provider = self.run_case(max_cost_usd=0.003)
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(len(provider.calls), 1)
        self.assertEqual(record["call_count"], 1)
        self.assertIn("budget", " ".join(record["issues"]))
        self.assertFalse((directory / "model.xlsx").exists())

    def test_twenty_dollar_stop_includes_crossing_call_and_sends_no_next_request(self):
        class Expensive(ScriptedProvider):
            input_tokens = 1_000_000

            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                self.last_response["usage"] = {"input_tokens": self.input_tokens, "output_tokens": 0}
                return result

        for count, expected in ((1_000_000, 20.0), (1_050_000, 21.0)):
            with self.subTest(estimated_usd=expected):
                company = SyntheticCompany()
                provider = Expensive(company)
                provider.input_tokens = count
                directory, record, _ = self.run_case(company, provider)
                self.assertEqual(record["status"], "incomplete")
                self.assertEqual(record["call_count"], 5)
                self.assertEqual(len(provider.calls), 5)
                self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], expected)
                self.assertEqual(record["budget"]["status"], "exhausted")
                self.assertEqual(record["budget"]["estimated_remaining_usd"], 0)
                self.assertFalse((directory / "model.xlsx").exists())
                log = [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
                self.assertEqual(len(log), 5)
                self.assertEqual(log[-1]["stage"], "source_review")
                self.assertAlmostEqual(sum(group["estimated_cost_usd"] for group in record["usage"]["by_stage"].values()), expected)

    def test_budget_crossed_by_final_review_cannot_publish_accepted_workbook(self):
        class ExpensiveFinal(ScriptedProvider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                if kwargs["payload"]["stage"] == "model_review" and not result["read_requests"]:
                    self.last_response["usage"] = {"input_tokens": 0, "output_tokens": 2_000_000}
                return result

        company = SyntheticCompany()
        directory, record, provider = self.run_case(company, ExpensiveFinal(company))
        self.assertEqual(provider.calls[-1]["stage"], "model_review")
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(record["budget"]["status"], "exhausted")
        self.assertGreaterEqual(record["usage"]["by_stage"]["model_review"]["estimated_cost_usd"], 20)
        self.assertFalse((directory / "model.xlsx").exists())
        self.assertFalse((directory / "model.pending.xlsx").exists())

    def test_unknown_cost_or_stale_provider_response_stops_before_next_request(self):
        class UnknownCost(ScriptedProvider):
            fault = "model"

            def complete(self, **kwargs):
                if self.fault == "stale" and self.calls:
                    self.calls.append(copy.deepcopy(kwargs["payload"]))
                    raise AIError("Synthetic transport failure without a new response")
                result = super().complete(**kwargs)
                if self.fault == "model":
                    self.last_response["model"] = "unpriced-synthetic-model"
                elif self.fault == "tier":
                    self.last_response["service_tier"] = "unpriced-tier"
                elif self.fault == "counts":
                    self.last_response["usage"]["total_tokens"] = -1
                return result

        for fault in ("model", "tier", "counts", "stale"):
            with self.subTest(fault=fault):
                company = SyntheticCompany()
                provider = UnknownCost(company)
                provider.fault = fault
                directory, record, _ = self.run_case(company, provider)
                self.assertEqual(record["status"], "incomplete")
                self.assertEqual(record["call_count"], 2 if fault == "stale" else 1)
                self.assertEqual(record["budget"]["status"], "unknown_cost")
                self.assertFalse(record["usage"]["cost_is_complete"])
                self.assertFalse((directory / "model.xlsx").exists())
                if fault == "stale":
                    self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], 0.003)
                    self.assertEqual(record["usage"]["reported_calls"], 1)
                    self.assertEqual(record["usage"]["by_stage"]["extract"]["missing_usage"], 1)

    def test_transport_failure_retains_known_spend_and_stops_without_retry(self):
        company = SyntheticCompany()
        provider = OpenAIProvider(api_key="SYNTHETIC-SECRET", models={
            role: "gpt-6.1-sol" for role in ("researcher", "extractor", "reviewer")})
        discovery = ScriptedProvider.synthetic_response("discover")
        discovery["output"] = [
            {"type": "web_search_call", "status": "completed", "action": {"type": "search"}},
            {"type": "message", "content": [
                {"type": "output_text", "text": json.dumps(company.discovery)}]},
        ]
        responses = [io.BytesIO(json.dumps(discovery).encode("utf-8")),
                     URLError(TimeoutError("SYNTHETIC-SECRET"))]
        with patch("research_platform.ai.urlopen", side_effect=responses) as network:
            directory, record, _ = self.run_case(company, provider)
        self.assertEqual(network.call_count, 2)
        self.assertEqual(record["call_count"], 2)
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(record["budget"]["status"], "unknown_cost")
        self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], 0.013)
        self.assertEqual(record["usage"]["reported_calls"], 1)
        self.assertEqual(record["usage"]["missing_usage"], 1)
        self.assertFalse(record["usage"]["cost_is_complete"])
        self.assertIsNone(record["usage"]["by_stage"]["extract"]["estimated_cost_usd"])
        self.assertFalse((directory / "model.xlsx").exists())
        saved_calls = (directory / "calls.jsonl").read_text(encoding="utf-8")
        self.assertNotIn("SYNTHETIC-SECRET", saved_calls + json.dumps(record))
        calls = [json.loads(line) for line in saved_calls.splitlines()]
        self.assertEqual([call["stage"] for call in calls], ["discover", "extract"])
        failure = calls[-1]["provider_response"]
        self.assertEqual(failure["_transport"]["failure_kind"], "timeout")
        self.assertNotIn("usage", failure)
        self.assertIsNone(calls[-1]["usage"]["estimated_cost_usd"])
        self.assertFalse(calls[-1]["usage"]["reported"])

    def test_invalid_budget_or_unpriced_configuration_makes_no_request_or_run(self):
        company = SyntheticCompany()
        provider = ScriptedProvider(company)
        for amount in (True, 0, -1, 20.01, float("nan"), float("inf"), "20"):
            with self.subTest(amount=amount), self.assertRaises(ValueError):
                initiate_company(self.root, company.name, provider=provider, max_cost_usd=amount)
        with patch.object(provider, "metadata", return_value={"models": {"researcher": "unknown-model"}}):
            with self.assertRaisesRegex(AIError, "known pricing"):
                initiate_company(self.root, company.name, provider=provider)
        self.assertEqual(provider.calls, [])
        self.assertFalse((self.root / "outputs").exists())

    def test_useful_progress_can_exceed_former_request_and_navigation_ceilings(self):
        class ManySteps(ScriptedProvider):
            saves = 0
            reviewer_searches = 0

            def complete(self, **kwargs):
                payload = kwargs["payload"]
                if payload["stage"] == "extract" and self.saves < 65:
                    self.saves += 1
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("extract")
                    row = copy.deepcopy(self.company.rows[0])
                    row["definition"] += f" Synthetic draft revision {self.saves}."
                    return {"action": "save_rows", "read_requests": [], "sources": [], "rows": [row],
                            "delete_row_ids": [], "search_query": "", "candidate": None}
                if payload["stage"] == "model_review" and self.reviewer_searches < 7:
                    query = self.company.rows[self.reviewer_searches]["label"]
                    self.reviewer_searches += 1
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("model_review")
                    return {"checks": [], "matches": [], "issues": [], "search_query": query, "read_requests": []}
                return super().complete(**kwargs)

        company = SyntheticCompany()
        messages = []
        _, record, provider = self.run_case(company, ManySteps(company), progress=messages.append)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertGreater(record["call_count"], 80)
        self.assertGreater(sum(call["stage"] == "extract" for call in provider.calls), 30)
        self.assertGreater(sum(call["stage"] == "model_review" for call in provider.calls), 5)
        self.assertLess(record["usage"]["estimated_cost_usd"], 20)
        self.assertNotIn("max_calls", record["request"])
        self.assertNotIn("max_rounds", record["request"])
        self.assertTrue(any("Stage extract:" in message and "total" in message for message in messages))
        self.assertTrue(any("Run usage:" in message and "estimated" in message.lower() for message in messages))
        self.assertEqual(sum(group["total_tokens"] for group in record["usage"]["by_stage"].values()), record["usage"]["total_tokens"])

    def test_unchanged_failed_candidate_stops_with_original_findings(self):
        company = SyntheticCompany()
        _, record, provider = self.run_case(company, ScriptedProvider(company, review_fault="model_reject"))
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(provider.proposals, 2)
        self.assertIn("No progress", " ".join(record["issues"]))
        self.assertIn("identity", " ".join(record["issues"]))
        self.assertTrue(record["attempts"][0]["issues"])

    def test_repeated_navigation_without_new_context_stops(self):
        class RepeatingRead(ScriptedProvider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                if kwargs["payload"]["stage"] == "source_review":
                    result.update(facts=[], read_requests=[{"source_id": "annual", "unit_id": f"{self.company.sheet}!R1"}])
                return result

        company = SyntheticCompany()
        _, record, provider = self.run_case(company, RepeatingRead(company))
        self.assertEqual(record["status"], "incomplete")
        self.assertIn("no progress", " ".join(record["issues"]))
        self.assertEqual(sum(call["stage"] == "source_review" for call in provider.calls), 2)

    def test_reading_missing_evidence_allows_resubmitting_an_unchanged_candidate(self):
        class EarlyProposal(ScriptedProvider):
            repaired_context = False

            def complete(self, **kwargs):
                payload = kwargs["payload"]
                if payload["stage"] == "extract" and self.proposals == 1 and not self.repaired_context:
                    self.repaired_context = True
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("extract")
                    refs = [{"source_id": source["id"], "unit_id": unit["id"]}
                            for source in payload["catalog"] for unit in source["units"]]
                    return {"action": "read", "read_requests": refs, "sources": [], "rows": [],
                            "delete_row_ids": [], "search_query": "", "candidate": None}
                return super().complete(**kwargs)

        company = SyntheticCompany()
        provider = EarlyProposal(company)
        provider.read_requested = True  # Deliberately propose before reading original context.
        _, record, _ = self.run_case(company, provider)
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(len(record["attempts"]), 2)
        self.assertEqual(record["attempts"][0]["candidate"], record["attempts"][1]["candidate"])
        self.assertTrue(record["attempts"][0]["issues"])

    def test_corrected_drafts_can_exceed_former_three_attempt_limit(self):
        class Correcting(ScriptedProvider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                if kwargs["payload"]["stage"] == "extract" and result["candidate"] is not None and self.proposals < 4:
                    result["candidate"]["reporting_basis"] = f"Synthetic incorrect vintage {self.proposals}"
                return result

        company = SyntheticCompany()
        _, record, provider = self.run_case(company, Correcting(company, review_fault="basis_disagreement"))
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(provider.proposals, 4)
        self.assertEqual(len(record["attempts"]), 4)
        self.assertLess(record["usage"]["estimated_cost_usd"], 20)
        self.assertTrue(all(attempt["issues"] for attempt in record["attempts"][:3]))
        self.assertTrue(all(any(review["stage"] == "source_review" for review in attempt["reviews"])
                            for attempt in record["attempts"]))

    def test_charged_failure_preserves_error_and_stage_charge_when_budget_exhausted(self):
        class ChargedFailure(ScriptedProvider):
            def complete(self, **kwargs):
                super().complete(**kwargs)
                self.last_response["status"] = "incomplete"
                raise AIError("Synthetic charged provider failure")

        company = SyntheticCompany()
        directory, record, provider = self.run_case(company, ChargedFailure(company), max_cost_usd=0.003)
        self.assertEqual(len(provider.calls), 1)
        self.assertEqual(record["budget"]["status"], "exhausted")
        self.assertIn("Synthetic charged provider failure", " ".join(record["issues"]))
        self.assertAlmostEqual(record["usage"]["by_stage"]["discover"]["estimated_cost_usd"], 0.003)
        logged = json.loads((directory / "calls.jsonl").read_text(encoding="utf-8").strip())
        self.assertEqual(logged["error"], "Synthetic charged provider failure")
        self.assertIn("budget", logged["budget_stop"])

    def test_run_records_missing_usage_and_charged_incomplete_response(self):
        class MissingUsage(ScriptedProvider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                self.last_response.pop("usage")
                return result

        company = SyntheticCompany()
        directory, record, _ = self.run_case(company, MissingUsage(company))
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(record["call_count"], 1)
        self.assertEqual(record["usage"]["calls"], record["call_count"])
        self.assertEqual(record["usage"]["missing_usage"], record["call_count"])
        self.assertFalse(record["usage"]["cost_is_complete"])
        self.assertEqual(record["budget"]["status"], "unknown_cost")
        self.assertIsNone(record["usage"]["estimated_cost_usd"])
        self.assertFalse((directory / "model.xlsx").exists())

        class IncompleteProvider(ScriptedProvider):
            def metadata(self):
                return {"provider": "SYNTHETIC SCRIPTED TEST", "models": {
                    role: "gpt-6-luna" for role in ("researcher", "extractor", "reviewer")}}

            def complete(self, **kwargs):
                self.calls.append(copy.deepcopy(kwargs["payload"]))
                self.last_response = {"status": "incomplete", "model": "gpt-6-luna", "output": [],
                                      "usage": {"input_tokens": 1000, "output_tokens": 100,
                                                "total_tokens": 1100}}
                raise AIError("Synthetic output limit reached")

        company = SyntheticCompany()
        messages = []
        directory, record, _ = self.run_case(company, IncompleteProvider(company), progress=messages.append)
        self.assertEqual(record["status"], "incomplete")
        self.assertFalse((directory / "model.xlsx").exists())
        self.assertEqual(record["call_count"], 1)
        self.assertEqual(record["usage"]["reported_calls"], 1)
        self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], 0.00015)
        logged = json.loads((directory / "calls.jsonl").read_text(encoding="utf-8").strip())
        self.assertEqual(logged["provider_response"]["status"], "incomplete")
        self.assertEqual(logged["usage"]["output_tokens"], 100)
        self.assertAlmostEqual(logged["usage"]["estimated_cost_usd"], 0.00015)
        self.assertTrue(any("discover" in message for message in messages))

    def test_circular_definitions_and_incompatible_units_are_rejected(self):
        def cycle(candidate):
            candidate["equations"].append({"target": "expenses",
                                          "terms": [{"row_id": "net_income", "coefficient": 1}],
                                          "explanation": "A deliberately circular accounting relationship."})

        def units(candidate):
            next(row for row in candidate["rows"] if row["id"] == "debt")["unit"] = "shares"

        for mutation, message in ((cycle, "Circular"), (units, "units")):
            with self.subTest(mutation=mutation.__name__):
                company = SyntheticCompany()
                directory, record, _ = self.run_case(company, ScriptedProvider(company, mutate=mutation))
                self.assertEqual(record["status"], "incomplete")
                self.assertIn(message, " ".join(record["issues"]))
                self.assertFalse((directory / "model.xlsx").exists())


class ResumeTests(unittest.TestCase):
    class Provider(ScriptedProvider):
        """Use only supplied context, as a new remote request would after restart."""

        def complete(self, **kwargs):
            payload = kwargs["payload"]
            if payload["stage"] == "extract":
                self.read_requested = bool(payload["evidence"])
                self.saved_batches = (len(payload["draft_row_index"]) + 11) // 12
            return super().complete(**kwargs)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.company = SyntheticCompany()

    def tearDown(self):
        self.temporary.cleanup()

    def run_company(self, provider=None, **kwargs):
        provider = provider or self.Provider(self.company)
        kwargs.setdefault("historical_years", 2)
        directory = initiate_company(self.root, self.company.name, provider=provider,
                                     fetcher=self.company.fetch, **kwargs)
        return directory, read_json(directory / "record.json"), provider

    def interrupt_after(self, call_count, provider=None, **kwargs):
        provider = provider or self.Provider(self.company)
        completed = 0

        def progress(message):
            nonlocal completed
            if message.startswith("Run usage:"):
                completed += 1
                if completed == call_count:
                    raise KeyboardInterrupt("Synthetic process interruption after durable response")

        with self.assertRaises(KeyboardInterrupt):
            self.run_company(provider, progress=progress, **kwargs)
        directory = max((self.root / "outputs").glob("initiate_*"),
                        key=lambda path: read_json(path / "record.json")["started_at"])
        return directory, read_json(directory / "record.json"), provider

    @staticmethod
    def calls(directory):
        return [json.loads(line) for line in (directory / "calls.jsonl").read_text(encoding="utf-8").splitlines()]

    @staticmethod
    def write_record(directory, record):
        (directory / "record.json").write_text(json.dumps(record), encoding="utf-8")

    def assert_blocked_without_spending(self, error, pattern, **kwargs):
        provider = self.Provider(self.company)
        previous = set((self.root / "outputs").glob("initiate_*"))
        fetches = len(self.company.fetches)
        with self.assertRaisesRegex(error, pattern):
            self.run_company(provider, **kwargs)
        self.assertEqual(provider.calls, [])
        self.assertEqual(len(self.company.fetches), fetches)
        self.assertEqual(set((self.root / "outputs").glob("initiate_*")), previous)

    def test_interrupted_draft_resumes_without_research_refetch_or_duplicate_charges(self):
        parent, saved, first = self.interrupt_after(4, self.Provider(self.company, batch=True))
        parent_bytes = (parent / "record.json").read_bytes()
        original_calls = self.calls(parent)
        self.assertEqual([call["stage"] for call in original_calls], ["discover", "extract", "extract", "extract"])
        self.assertEqual(len(self.company.fetches), 1)
        directory, record, resumed = self.run_company(self.Provider(self.company, batch=True))
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertNotEqual(directory, parent)
        self.assertEqual(record["parent_run"], str(parent.resolve()))
        self.assertEqual((parent / "record.json").read_bytes(), parent_bytes)
        self.assertEqual(self.calls(parent), original_calls)
        self.assertEqual(self.calls(directory)[:len(original_calls)], original_calls)
        self.assertNotIn("discover", [call["stage"] for call in resumed.calls])
        self.assertEqual(len(self.company.fetches), 1)
        self.assertEqual(record["call_count"], len(first.calls) + len(resumed.calls))
        self.assertEqual(record["usage"]["total_tokens"], record["call_count"] * 1100)
        self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], record["call_count"] * 0.003)
        self.assertEqual(record["sources"], saved["sources"])
        self.assertTrue(Path(record["sources"][0]["path"]).is_relative_to(parent))
        self.assertTrue((directory / "model.xlsx").exists())

    def test_restart_inside_blind_review_restores_only_that_reviews_original_context(self):
        parent, saved, _ = self.interrupt_after(4)
        self.assertEqual(self.calls(parent)[-1]["stage"], "source_review")
        self.assertTrue(self.calls(parent)[-1]["result"]["read_requests"])
        directory, record, resumed = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        first = resumed.calls[0]
        self.assertEqual(first["stage"], "source_review")
        self.assertTrue(first["evidence"])
        self.assertNotIn("rows", first)
        self.assertNotIn("candidate", first)
        self.assertNotIn("reviews", first)
        self.assertTrue(all(set(concept) == {"id", "label", "section", "entity", "periods"}
                            for concept in first["concepts"]))
        self.assertEqual(record["usage"]["by_stage"]["discover"]["calls"], 1)
        self.assertEqual(record["usage"]["total_tokens"], 1100 * (saved["call_count"] + len(resumed.calls)))

    def test_multiple_continuations_reuse_finished_reviews_and_carry_one_budget(self):
        first, _, _ = self.interrupt_after(7)
        second, second_saved, _ = self.interrupt_after(3)
        self.assertEqual(second_saved["parent_run"], str(first.resolve()))
        self.assertEqual(second_saved["call_count"], 10)
        self.assertTrue(second_saved["reviews"])
        frozen = copy.deepcopy(second_saved["reviews"])
        directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(record["parent_run"], str(second.resolve()))
        self.assertEqual(record["call_count"], 10 + len(provider.calls))
        self.assertEqual(len(self.company.fetches), 1)
        self.assertEqual([item["stage"] for item in record["reviews"][:len(frozen)]],
                         [item["stage"] for item in frozen])
        self.assertEqual([item["result"] for item in record["reviews"][:len(frozen)]],
                         [item["result"] for item in frozen])
        self.assertEqual(self.calls(directory)[:10], self.calls(second))
        self.assertEqual(record["usage"]["total_tokens"], record["call_count"] * 1100)

    def test_interrupting_replay_preserves_future_inherited_calls_and_sources(self):
        first, first_record, _ = self.interrupt_after(3)

        def stop_during_replay(message):
            if message.startswith("Reused call 1:"):
                raise KeyboardInterrupt("Synthetic stop while replaying discovery")

        provider = self.Provider(self.company)
        with self.assertRaises(KeyboardInterrupt):
            self.run_company(provider, progress=stop_during_replay)
        self.assertEqual(provider.calls, [])
        second = max((self.root / "outputs").glob("initiate_*"),
                     key=lambda path: read_json(path / "record.json")["started_at"])
        second_record = read_json(second / "record.json")
        self.assertEqual(second_record["parent_run"], str(first.resolve()))
        self.assertEqual(second_record["sources"], first_record["sources"])
        self.assertEqual(second_record["source_attempts"], first_record["source_attempts"])
        self.assertEqual(self.calls(second), self.calls(first))
        directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(record["parent_run"], str(second.resolve()))
        self.assertEqual(provider.calls[0]["stage"], "source_review")
        self.assertEqual(len(self.company.fetches), 1)
        self.assertEqual(record["call_count"], 3 + len(provider.calls))

    def test_known_cost_failed_call_can_continue_but_its_charge_is_not_erased(self):
        class ChargedFailure(self.Provider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                if kwargs["payload"]["stage"] == "extract":
                    raise AIError("Synthetic response ended after reported usage")
                return result

        parent, before, _ = self.run_company(ChargedFailure(self.company))
        self.assertEqual(before["status"], "incomplete")
        self.assertEqual(before["call_count"], 2)
        self.assertTrue(before["usage"]["cost_is_complete"])
        directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(provider.calls[0]["stage"], "extract")
        self.assertEqual(record["usage"]["total_tokens"], (2 + len(provider.calls)) * 1100)
        self.assertEqual(self.calls(directory)[:2], self.calls(parent))
        self.assertIn("error", self.calls(directory)[1])

    def test_cost_before_interruption_counts_towards_same_twenty_dollar_stop(self):
        class Expensive(self.Provider):
            def complete(self, **kwargs):
                result = super().complete(**kwargs)
                self.last_response["usage"] = {"input_tokens": 1_000_000, "output_tokens": 0}
                return result

        parent, before, _ = self.interrupt_after(4, Expensive(self.company))
        self.assertAlmostEqual(before["usage"]["estimated_cost_usd"], 16.0)
        directory, record, provider = self.run_company(Expensive(self.company))
        self.assertEqual(len(provider.calls), 1)
        self.assertEqual(record["call_count"], 5)
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(record["budget"]["status"], "exhausted")
        self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], 20.0)
        self.assertFalse((directory / "model.xlsx").exists())
        self.assert_blocked_without_spending(AIError, "budget|reached|exhausted")

    def test_unknown_failed_request_blocks_restart_before_any_network_or_output(self):
        class UnknownFailure(self.Provider):
            def complete(self, **kwargs):
                if kwargs["payload"]["stage"] == "extract":
                    self.calls.append(copy.deepcopy(kwargs["payload"]))
                    raise AIError("Synthetic lost response without usage")
                return super().complete(**kwargs)

        parent, record, _ = self.run_company(UnknownFailure(self.company))
        original = (parent / "record.json").read_bytes()
        self.assertFalse(record["usage"]["cost_is_complete"])
        self.assert_blocked_without_spending(AIError, "unknown|usage|cost")
        self.assertEqual((parent / "record.json").read_bytes(), original)

    def test_journaled_response_recovers_stale_pending_checkpoint_without_recharging(self):
        parent, record, _ = self.interrupt_after(3)
        last = self.calls(parent)[-1]
        record["pending_call"] = {key: copy.deepcopy(last[key]) for key in
                                  ("call", "call_id", "stage", "role", "started_at")}
        record["call_count"] = 2
        record["usage"] = aggregate_usage(self.calls(parent)[:2])
        self.write_record(parent, record)
        directory, resumed, provider = self.run_company()
        self.assertEqual(resumed["status"], "accepted", resumed["issues"])
        self.assertEqual(provider.calls[0]["stage"], "source_review")
        self.assertEqual(resumed["call_count"], 3 + len(provider.calls))
        self.assertEqual(resumed["usage"]["total_tokens"], resumed["call_count"] * 1100)
        self.assertEqual(self.calls(directory)[:3], self.calls(parent))

    def test_unjournaled_inflight_request_blocks_restart_without_losing_saved_work(self):
        parent, record, _ = self.interrupt_after(2)
        record["pending_call"] = {"call": 3, "call_id": "synthetic-unanswered-request",
                                  "stage": "extract", "role": "extractor", "started_at": record["started_at"]}
        self.write_record(parent, record)
        original = (parent / "record.json").read_bytes()
        self.assert_blocked_without_spending(AIError, "in-flight|known charge|usage")
        self.assertEqual((parent / "record.json").read_bytes(), original)
        self.assertEqual(len(self.calls(parent)), 2)

    def test_code_environment_provider_and_source_drift_stop_before_continuation(self):
        parent, record, _ = self.interrupt_after(3)
        for key in ("code", "environment", "provider"):
            with self.subTest(drift=key):
                changed = copy.deepcopy(record)
                if key == "code":
                    changed[key]["initiation.py"] += "\n# Synthetic incompatible version\n"
                elif key == "environment":
                    changed[key]["openpyxl"] = "synthetic-other-version"
                else:
                    changed[key]["models"]["reviewer"] = "gpt-6.1-sol-2026-01-01"
                self.write_record(parent, changed)
                self.assert_blocked_without_spending(ValueError, "differs|differ|version|settings")
        self.write_record(parent, record)
        original_path = Path(record["sources"][0]["path"])
        original_path.write_bytes(original_path.read_bytes() + b"synthetic corruption")
        self.assert_blocked_without_spending(ValueError, "hash|content")

    def test_damaged_or_missing_journal_cannot_hide_previously_charged_calls(self):
        parent, record, _ = self.interrupt_after(3)
        journal = parent / "calls.jsonl"
        original = journal.read_bytes()
        journal.write_bytes(original + b'{"call":4,')
        self.assert_blocked_without_spending(AIError, "journal|uncertain")
        journal.write_text("", encoding="utf-8")
        self.assert_blocked_without_spending(AIError, "missing|journal")

    def test_accepted_run_is_reused_by_normalized_query_until_fresh_is_explicit(self):
        accepted, record, _ = self.run_company()
        original = (accepted / "record.json").read_bytes()
        provider = self.Provider(self.company)
        query = "  " + self.company.name.upper().replace(" ", "   ") + "  "
        reused = initiate_company(self.root, query, provider=provider, fetcher=self.company.fetch, historical_years=2)
        self.assertEqual(reused, accepted)
        self.assertEqual(provider.calls, [])
        self.assertEqual((accepted / "record.json").read_bytes(), original)
        self.assertEqual(len(self.company.fetches), 1)
        directory, fresh, provider = self.run_company(fresh=True)
        self.assertNotEqual(directory, accepted)
        self.assertEqual(fresh["status"], "accepted", fresh["issues"])
        self.assertIsNone(fresh["parent_run"])
        self.assertEqual(provider.calls[0]["stage"], "discover")
        self.assertEqual(len(self.company.fetches), 2)

    def test_explicit_resume_checks_request_settings_and_disallows_older_branch(self):
        first, _, _ = self.interrupt_after(3)
        second, _, _ = self.interrupt_after(1, resume=first / "record.json")
        self.assert_blocked_without_spending(ValueError, "newer|branch", resume=first)
        self.assert_blocked_without_spending(ValueError, "match", resume=second, years=[2024, 2025])
        self.assert_blocked_without_spending(ValueError, "spending limit", resume=second, max_cost_usd=5)
        self.assert_blocked_without_spending(ValueError, "either|both", resume=second, fresh=True)

    def test_resume_keeps_original_research_date_when_restarted_another_day(self):
        from datetime import date

        parent, before, _ = self.interrupt_after(3)

        class LaterDate(date):
            @classmethod
            def today(cls):
                return cls(2035, 1, 1)

        with patch("research_platform.initiation.date", LaterDate):
            directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(record["request"]["as_of"], before["request"]["as_of"])
        self.assertEqual(provider.calls[0]["stage"], "source_review")
        self.assertEqual(record["call_count"], 3 + len(provider.calls))

    def test_older_nonresumable_runs_need_explicit_fresh_instead_of_silent_paid_restart(self):
        parent, record, _ = self.interrupt_after(3)
        record.pop("resume_version")
        self.write_record(parent, record)
        original = (parent / "record.json").read_bytes()
        self.assert_blocked_without_spending(ValueError, "older|predates|fresh")
        directory, fresh, provider = self.run_company(fresh=True)
        self.assertEqual(fresh["status"], "accepted", fresh["issues"])
        self.assertIsNone(fresh["parent_run"])
        self.assertEqual(provider.calls[0]["stage"], "discover")
        self.assertEqual((parent / "record.json").read_bytes(), original)

    def test_interrupt_during_source_download_retries_only_unfinished_download(self):
        provider = self.Provider(self.company)
        with patch.object(self.company, "fetch", side_effect=KeyboardInterrupt("Synthetic interrupted download")):
            with self.assertRaises(KeyboardInterrupt):
                self.run_company(provider)
        parent = next((self.root / "outputs").glob("initiate_*"))
        self.assertEqual(len(self.calls(parent)), 1)
        directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(provider.calls[0]["stage"], "extract")
        self.assertEqual(len(self.company.fetches), 1)
        self.assertEqual(record["call_count"], len(provider.calls) + 1)
        self.assertEqual(self.calls(directory)[0], self.calls(parent)[0])

    def test_interrupt_during_source_parsing_does_not_create_false_completed_evidence(self):
        with patch.object(SourceStore, "_index", side_effect=KeyboardInterrupt("Synthetic interrupted parser")):
            with self.assertRaises(KeyboardInterrupt):
                self.run_company()
        parent = next((self.root / "outputs").glob("initiate_*"))
        self.assertEqual(len(self.calls(parent)), 1)
        self.assertEqual(len(self.company.fetches), 1)
        directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(provider.calls[0]["stage"], "extract")
        self.assertTrue(all(source["status"] == "ready" for source in record["sources"]))
        self.assertEqual(record["call_count"], len(provider.calls) + 1)
        self.assertEqual(self.calls(directory)[0], self.calls(parent)[0])

    def test_interrupt_during_excel_export_resumes_without_repeating_any_ai_call(self):
        with patch("research_platform.initiation.export_workbook",
                   side_effect=KeyboardInterrupt("Synthetic interrupted workbook export")):
            with self.assertRaises(KeyboardInterrupt):
                self.run_company()
        parent = next((self.root / "outputs").glob("initiate_*"))
        before = read_json(parent / "record.json")
        self.assertNotEqual(before["status"], "accepted")
        self.assertFalse((parent / "model.xlsx").exists())
        directory, record, provider = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(provider.calls, [])
        self.assertEqual(record["usage"], before["usage"])
        self.assertEqual(self.calls(directory), self.calls(parent))
        self.assertEqual(record["reviews"], before["reviews"])
        self.assertEqual(len(self.company.fetches), 1)
        self.assertTrue((directory / "model.xlsx").exists())

    def test_source_paths_and_extraction_metadata_must_belong_to_frozen_archive(self):
        parent, record, _ = self.interrupt_after(3)
        changed = copy.deepcopy(record)
        changed["sources"][0]["unit_count"] += 1
        changed["source_attempts"][0]["descriptor"]["unit_count"] += 1
        self.write_record(parent, changed)
        self.assert_blocked_without_spending(ValueError, "extraction metadata|changed")
        outside = self.root / "outside.xlsx"
        outside.write_bytes(Path(record["sources"][0]["path"]).read_bytes())
        changed = copy.deepcopy(record)
        changed["sources"][0]["path"] = str(outside)
        changed["source_attempts"][0]["descriptor"]["path"] = str(outside)
        self.write_record(parent, changed)
        self.assert_blocked_without_spending(ValueError, "outside|ancestry")

    def test_project_lock_prevents_concurrent_budget_branches_and_releases_after_error(self):
        from research_platform.initiation import _initiation_lock

        with _initiation_lock(self.root / "outputs"):
            self.assert_blocked_without_spending(AIError, "Another initiation|active")
        directory, record, _ = self.run_company()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertTrue((directory / "model.xlsx").exists())


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        company = SyntheticCompany()
        for row in company.rows:
            if row["unit"] != "shares":
                row["scale"] = 1000000
            if row["id"] == "debt":
                row["scale"] = 1000
                company.display_values["debt"] = [300000, 350000]
            if row["id"] == "payables":
                row["sign"] = -1
                company.display_values["payables"] = [80, 90]
        candidate = copy.deepcopy(company.candidate)
        rows = {row["id"]: row for row in candidate["rows"]}
        for target in ("net_debt", "working_capital"):
            equation = next(item for item in candidate["equations"] if item["target"] == target)
            rows[target]["terms"] = equation["terms"]
            rows[target]["observations"] = []
            candidate["equations"].remove(equation)
        linked_row = copy.deepcopy(rows["net_debt"])
        linked_row.update(id="linked_debt", label="Linked debt", role="other",
                          terms=[{"row_id": "net_debt", "coefficient": 1}])
        # Put the dependent row first to exercise forward cell references.
        candidate["rows"].insert(0, linked_row)
        store = SourceStore(self.root / "sources", fetcher=company.fetch)
        store.add(company.source)
        refs = [{"source_id": source["id"], "unit_id": unit["id"]}
                for source in store.catalog() for unit in source["units"]]
        evidence = store.read(refs)
        resolved, checks = materialize_and_check(candidate, company.discovery, store,
                                                 {(item["source_id"], item["unit_id"]) for item in evidence})
        self.assertTrue(all(check["passed"] for check in checks), checks)
        self.record = {"status": "accepted", "request": {"forecast_years": 5},
                       "discovery": company.discovery, "candidate": candidate,
                       "rows": resolved, "checks": checks, "sources": store.sources, "reviews": []}

    def export(self):
        before = copy.deepcopy(self.record)
        path = self.root / "model.xlsx"
        export_workbook(path, self.record)
        self.assertEqual(self.record, before, "The exported view must not modify the accepted Python record")
        book = load_workbook(path, data_only=False)
        self.addCleanup(book.close)
        return book

    def cells(self, book, section, row_id):
        sheet = book[section]
        id_column = next(cell.column for cell in sheet[4] if cell.value == "Row ID")
        return next(row for row in sheet.iter_rows(min_row=5) if row[id_column - 1].value == row_id)

    @staticmethod
    def arithmetic_value(book, cell):
        """Evaluate only the exported arithmetic branch for tests, not Excel.

        Resolve the actual saved cell links recursively and interpret a restricted
        arithmetic AST. No eval, arbitrary formula engine or production cache.
        """
        if cell.data_type != "f":
            if not isinstance(cell.value, (int, float)):
                raise ValueError("Missing numeric input")
            return cell.value
        branch = re.search(r",\((.+)\)/([0-9]+),NA\(\)\)$", cell.value)
        if branch is None:
            raise AssertionError("Expected an exported arithmetic formula")

        def reference(match):
            target = book[match[1] or cell.parent.title][match[2]]
            return repr(WorkbookTests.arithmetic_value(book, target))

        text = re.sub(r"(?:'([^']+)'!)?([A-Z]+[0-9]+)", reference, branch[1])

        def calculate(node):
            if isinstance(node, ast.Constant) and type(node.value) in (int, float):
                return node.value
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                return calculate(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
            if isinstance(node, ast.BinOp):
                left, right = calculate(node.left), calculate(node.right)
                if isinstance(node.op, ast.Add):
                    return left + right
                if isinstance(node.op, ast.Sub):
                    return left - right
                if isinstance(node.op, ast.Mult):
                    return left * right
                if isinstance(node.op, ast.Div):
                    return left / right
            raise AssertionError("Unexpected arithmetic syntax in exported formula")

        return calculate(ast.parse(f"({text})/{branch[2]}", mode="eval").body)

    def test_generated_arithmetic_evaluates_to_python_results_and_detects_changed_inputs(self):
        book = self.export()
        for row in self.record["rows"]:
            cells = self.cells(book, row["section"], row["id"])
            for column, year in ((3, "2024"), (4, "2025")):
                self.assertAlmostEqual(self.arithmetic_value(book, cells[column]),
                                       row["values"][year] / row["scale"])
        for check in book["Checks"].iter_rows(min_row=2):
            if check[6].data_type == "f":
                self.assertAlmostEqual(self.arithmetic_value(book, check[6]), check[3].value)
        # A workbook-only edit breaks a reconciliation, without rewriting its
        # recorded Python PASS or original numeric difference.
        revenue = self.cells(book, "income_statement", "revenue")[3]
        revenue.value += 1
        check = next(row for row in book["Checks"] if row[0].value == "equation:revenue:2024")
        self.assertEqual(self.arithmetic_value(book, check[6]), 1000000)
        self.assertEqual(check[1].value, "PASS")
        self.assertEqual(check[3].value, 0)
        revenue.value = None
        with self.assertRaisesRegex(ValueError, "Missing numeric input"):
            self.arithmetic_value(book, check[6])

    def test_excel_formula_length_limit_stops_export_instead_of_writing_an_invalid_workbook(self):
        source = next(row for row in self.record["rows"] if row["id"] == "debt")
        target = next(row for row in self.record["rows"] if row["id"] == "net_debt")
        target["terms"] = []
        for index in range(300):
            component = copy.deepcopy(source)
            component["id"] = f"debt_component_{index}"
            self.record["rows"].append(component)
            target["terms"].append({"row_id": component["id"], "coefficient": 1})
        path = self.root / "too_long.xlsx"
        with self.assertRaisesRegex(ValueError, "8,192-character formula limit"):
            export_workbook(path, self.record)
        self.assertFalse(path.exists())

    def test_derived_formulas_preserve_cross_sheet_signs_scales_and_forward_links(self):
        book = self.export()
        debt = self.cells(book, "balance_sheet", "debt")
        cash = self.cells(book, "cash_flow", "cash_close")
        net_debt = self.cells(book, "net_debt", "net_debt")
        working_capital = self.cells(book, "working_capital", "working_capital")
        linked = self.cells(book, "net_debt", "linked_debt")
        self.assertLess(linked[0].row, net_debt[0].row)
        by_id = {row["id"]: row for row in self.record["rows"]}
        for column, year in ((3, "2024"), (4, "2025")):
            debt_ref, cash_ref = f"'balance_sheet'!{debt[column].coordinate}", f"'cash_flow'!{cash[column].coordinate}"
            self.assertEqual(net_debt[column].value,
                             f"=IF(COUNT({debt_ref},{cash_ref})=2,(+{debt_ref}*1000-{cash_ref}*1000000)/1000000,NA())")
            self.assertEqual(net_debt[column].data_type, "f")
            # The linked inputs imply exactly the normalized Python amount.
            self.assertEqual(debt[column].value * 1000 - cash[column].value * 1000000,
                             by_id["net_debt"]["values"][year])
            self.assertEqual(linked[column].value,
                             f"=IF(COUNT({net_debt[column].coordinate})=1,(+{net_debt[column].coordinate}*1000000)/1000000,NA())")
            payables = self.cells(book, "balance_sheet", "payables")[column]
            self.assertLess(payables.value, 0)
            self.assertIn(f"+'balance_sheet'!{payables.coordinate}*1000000", working_capital[column].value)
        self.assertEqual(net_debt[3].font.color.rgb, "00008000")
        self.assertEqual(linked[3].font.color.rgb, "00000000")
        self.assertEqual(debt[3].font.color.rgb, "000000FF")

    def test_reported_anchors_stay_inputs_and_live_checks_are_separate_from_baseline(self):
        book = self.export()
        revenue = self.cells(book, "income_statement", "revenue")
        self.assertEqual(revenue[3].value, 1000)
        self.assertEqual(revenue[3].data_type, "n")
        self.assertIn('"source_id": "annual"', revenue[3].comment.text)
        check = next(row for row in book["Checks"] if row[0].value == "equation:revenue:2024")
        self.assertEqual(check[1].value, "PASS")
        self.assertEqual(check[3].value, 0)
        self.assertEqual(check[6].data_type, "f")
        self.assertIn(f"'income_statement'!{revenue[3].coordinate}*1000000", check[6].value)
        self.assertIn("'segments'!", check[6].value)
        self.assertIn("ABS(G", check[7].value)
        self.assertIn("INPUT ERROR", check[7].value)
        structural = next(row for row in book["Checks"] if row[0].value == "required_group_roles")
        self.assertIsNone(structural[6].value)
        self.assertIsNone(structural[7].value)
        instructions = " ".join(str(cell.value) for row in book["Start here"] for cell in row)
        self.assertIn("do not change saved Python", instructions)
        self.assertIn("cannot renew acceptance", instructions)
        self.assertEqual(book.calculation.calcMode, "auto")
        self.assertTrue(book.calculation.fullCalcOnLoad)
        self.assertTrue(book.calculation.forceFullCalc)
        cached = load_workbook(self.root / "model.xlsx", data_only=True)
        self.addCleanup(cached.close)
        self.assertIsNone(self.cells(cached, "net_debt", "net_debt")[3].value)

    def test_untrusted_text_cannot_become_excel_formulas(self):
        attack = '=HYPERLINK("https://example.invalid/leak","open")'
        self.record["discovery"]["company"]["legal_name"] = attack
        self.record["candidate"]["reporting_basis"] = attack
        for row in self.record["rows"]:
            row["label"] = row["definition"] = attack
        self.record["sources"][0]["title"] = attack
        self.record["sources"][0]["url"] = attack
        self.record["checks"][0]["detail"] = attack
        self.record["reviews"] = [{"stage": "model_review", "result": {
            "checks": [{"topic": "entity", "passed": True, "reason": attack}]}}]
        book = self.export()
        injected_cells = [cell for sheet in book for row in sheet for cell in row if cell.value == attack]
        self.assertGreater(len(injected_cells), 20)
        self.assertTrue(all(cell.data_type == "s" for cell in injected_cells))
        formulas = [cell.value for sheet in book for row in sheet for cell in row if cell.data_type == "f"]
        self.assertTrue(formulas)
        self.assertTrue(all(formula.startswith(("=IF(COUNT(", "=IFERROR(IF(ABS(")) for formula in formulas))

    def test_financial_layout_hides_audit_columns_and_keeps_forecasts_and_missing_periods_blank(self):
        sparse = copy.deepcopy(next(row for row in self.record["rows"] if row["id"] == "minority_interest"))
        sparse.update(id="adjustment", label="One-period adjustment", section="cash_flow",
                      values={"2025": 5000000}, observations=[sparse["observations"][1]],
                      unavailable=[{"period": "2024", "explanation": "The adjustment is not disclosed for this period.",
                                    "references": copy.deepcopy(sparse["definition_refs"])}])
        self.record["rows"].append(sparse)
        book = self.export()
        sparse_cells = self.cells(book, "cash_flow", "adjustment")
        self.assertIsNone(sparse_cells[3].value)
        self.assertIn("Unavailable: The adjustment is not disclosed for this period.", sparse_cells[3].comment.text)
        self.assertIn("Statements!R28", sparse_cells[3].comment.text)
        self.assertEqual(sparse_cells[4].value, 5)
        for section in SECTIONS:
            sheet = book[section]
            self.assertEqual(sheet.freeze_panes, "D5")
            self.assertFalse(sheet.sheet_view.showGridLines)
            forecast = [cell.column for cell in sheet[4] if "forecast (blank)" in str(cell.value)]
            self.assertEqual(len(forecast), 5)
            for row in sheet.iter_rows(min_row=5):
                for column in forecast:
                    self.assertIsNone(row[column - 1].value)
                    self.assertEqual(row[column - 1].fill.fgColor.rgb, "00FFF2CC")
            audit_column = next(cell.column_letter for cell in sheet[4] if cell.value == "Row ID")
            self.assertTrue(sheet.column_dimensions[audit_column].hidden)
            self.assertEqual(sheet.column_dimensions[audit_column].outlineLevel, 1)


class EvidenceTests(unittest.TestCase):
    def test_document_reads_preserve_all_units_order_citations_and_isolation(self):
        first, second = SyntheticCompany(), SyntheticCompany(1)
        sources = {first.source["url"]: first, second.source["url"]: second}

        def fetch(url):
            return sources[url].fetch(url)

        with tempfile.TemporaryDirectory() as temporary:
            store = SourceStore(Path(temporary), fetcher=fetch)
            store.add(first.source)
            store.add(dict(second.source, id="comparison"))
            selected = [
                {"source_id": "comparison", "unit_id": f"{second.sheet}!R3"},
                {"source_id": "annual", "unit_id": f"{first.sheet}!R3"},
                {"source_id": "comparison", "unit_id": f"{second.sheet}!R1"},
            ]
            result = store.read_documents(selected)
            catalog = {source["id"]: source for source in store.catalog()}
            expected = [(source_id, unit["id"]) for source_id in ("comparison", "annual")
                        for unit in catalog[source_id]["units"]]
            self.assertEqual([(unit["source_id"], unit["unit_id"]) for unit in result], expected)
            self.assertEqual(len(store.read([selected[0]])), 1)
            self.assertEqual(store.read_documents([]), [])
            reference = first.rows[0]["observations"][0]["reference"]
            token = next(unit for unit in result if unit["source_id"] == "annual"
                         and unit["unit_id"] == reference["unit_id"])["numbers"][0]
            self.assertEqual(token, store.read([reference])[0]["numbers"][0])
            token["raw"] = 999999
            self.assertEqual(store.resolve(reference), first.display_values[first.rows[0]["id"]][0])
            for invalid in ({"source_id": "annual", "unit_id": "missing"},
                            {"source_id": "missing", "unit_id": f"{first.sheet}!R1"}):
                with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                    store.read_documents([*selected, invalid])

    def test_large_html_line_is_indexed_and_read_without_character_truncation(self):
        context = "Synthetic qualitative context. " * 100_001
        tail = "FINAL DISCLOSURE: reported revenue 42"
        self.assertGreater(len(context), 3_000_000)
        payload = ("<p>" + context + tail + "</p>").encode("utf-8")
        company = SyntheticCompany()
        with tempfile.TemporaryDirectory() as temporary:
            store = SourceStore(Path(temporary), fetcher=lambda url: (payload, url, "text/html"))
            store.add(company.source)
            reference = {"source_id": "annual", "unit_id": "chunk:1"}
            exact = store.read([reference])
            whole = store.read_documents([reference])
            self.assertEqual(exact, whole)
            self.assertEqual(len(whole), 1)
            self.assertEqual(whole[0]["text"], context + tail)
            number = next(token for token in whole[0]["numbers"] if token["raw"] == "42")
            self.assertEqual(store.resolve(dict(reference, number_index=number["index"],
                                                number_format="dot_decimal")), 42)

    def test_native_spreadsheet_numbers_use_displayed_precision(self):
        book = Workbook()
        sheet = book.active
        formats = (
            (1200, "0.00", 0.005),
            (1200, "#,##0.00", 0.005),
            (-1200, "0.00;[Red](0.00)", 0.005),
            (0.12, "0.00%", 0.00005),
        )
        for raw, number_format, _ in formats:
            sheet.append(["Synthetic amount", raw])
            sheet.cell(sheet.max_row, 2).number_format = number_format
        output = io.BytesIO()
        book.save(output)
        company = SyntheticCompany()
        with tempfile.TemporaryDirectory() as temporary:
            store = SourceStore(Path(temporary), fetcher=lambda url: (
                output.getvalue(), url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
            store.add(company.source)
            for row, (raw, number_format, expected) in enumerate(formats, 1):
                with self.subTest(number_format=number_format):
                    reference = {"source_id": "annual", "unit_id": f"Sheet!R{row}",
                                 "number_index": 0, "number_format": "dot_decimal"}
                    token = store.read([reference])[0]["numbers"][0]
                    self.assertEqual(token["display_format"], number_format)
                    self.assertEqual(store.resolve(reference), raw)
                    self.assertAlmostEqual(_rounding_bound(reference, 1, store), expected)
                    self.assertAlmostEqual(_rounding_bound(reference, 1000000, store), expected * 1000000)

    def test_scaled_scientific_fraction_and_date_formats_cannot_be_accepted_as_plain_amounts(self):
        book = Workbook()
        sheet = book.active
        formats = ("#,##0,", "0,,", "0.00E+00", "# ?/?", "yyyy-mm-dd")
        for number_format in formats:
            sheet.append(["Synthetic amount", 1200])
            sheet.cell(sheet.max_row, 2).number_format = number_format
        output = io.BytesIO()
        book.save(output)
        company = SyntheticCompany()
        with tempfile.TemporaryDirectory() as temporary:
            store = SourceStore(Path(temporary), fetcher=lambda url: (
                output.getvalue(), url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
            store.add(company.source)
            for row, number_format in enumerate(formats, 1):
                reference = {"source_id": "annual", "unit_id": f"Sheet!R{row}",
                             "number_index": 0, "number_format": "dot_decimal"}
                with self.subTest(number_format=number_format), self.assertRaises(ValueError):
                    # Date-formatted Excel values become datetime objects and are
                    # not numeric evidence; the other formats fail precision checks.
                    store.resolve(reference)
                    _rounding_bound(reference, 1, store)

    def test_archived_cells_are_exact_and_read_results_cannot_mutate_sources(self):
        company = SyntheticCompany(1)
        with tempfile.TemporaryDirectory() as temporary:
            store = SourceStore(Path(temporary), fetcher=company.fetch)
            source = store.add(company.source)
            row = next(row for row in company.rows if row["id"] == "revenue")
            reference = row["observations"][0]["reference"]
            self.assertEqual(store.resolve(reference), 1000)
            self.assertEqual(source["sha256"],
                             hashlib.sha256(Path(source["path"]).read_bytes()).hexdigest())
            preview = store.catalog(max_units=1)[0]
            self.assertTrue(preview["catalog_truncated"])
            self.assertEqual(len(preview["units"]), 1)
            result = store.search("REVENUE")
            self.assertTrue(any(item["unit_id"] == reference["unit_id"] for item in result["matches"]))
            self.assertFalse(result["truncated"])
            read = store.read([reference])
            read[0]["numbers"][0]["raw"] = 999999
            self.assertEqual(store.resolve(row["observations"][1]["reference"]), 1100)
            for index in (-1, True, 10000):
                with self.subTest(index=index), self.assertRaises(ValueError):
                    store.resolve(dict(reference, number_index=index))
            with self.assertRaises(ValueError):
                store.add(dict(company.source, url="https://example.invalid/different.xlsx"))

    def test_html_retains_table_context_and_excludes_script_numbers(self):
        payload = (b"<html><script>99999</script><table><tr><th>Metric</th><th>FY2025</th></tr>"
                   b"<tr><td>Revenue EUR thousand</td><td>1.234,50</td></tr></table></html>")
        company = SyntheticCompany()
        with tempfile.TemporaryDirectory() as temporary:
            store = SourceStore(Path(temporary), fetcher=lambda url: (payload, url, "text/html"))
            store.add(company.source)
            unit = store.read([{"source_id": "annual", "unit_id": "chunk:1"}])[0]
            self.assertIn("Revenue EUR thousand", unit["text"])
            self.assertNotIn("99999", unit["text"])
            number = next(token for token in unit["numbers"] if token["raw"] == "1.234,50")
            self.assertEqual(store.resolve({"source_id": "annual", "unit_id": "chunk:1",
                                            "number_index": number["index"],
                                            "number_format": "comma_decimal"}), 1234.5)
        with patch("research_platform.evidence.build_opener") as opener:
            for url in ("file:///private", "http://example.invalid", "https://user:password@example.invalid"):
                with self.subTest(url=url), self.assertRaises(ValueError):
                    download(url)
            opener.assert_not_called()


class AIAdapterTests(unittest.TestCase):
    def provider(self):
        return OpenAIProvider(api_key="SYNTHETIC-SECRET", models={
            "researcher": "test-researcher", "extractor": "test-extractor", "reviewer": "test-reviewer"})

    def complete(self, provider, response, *, web=False):
        stream = io.BytesIO(json.dumps(response).encode("utf-8"))
        with patch("research_platform.ai.urlopen", return_value=stream) as network:
            result = provider.complete(role="researcher", instructions="Synthetic instructions",
                                       payload={"synthetic": True},
                                       schema={"type": "object", "properties": {}, "required": [],
                                               "additionalProperties": False}, web=web)
            request = network.call_args.args[0]
        return result, json.loads(request.data)

    def test_env_file_is_literal_allowlisted_and_process_environment_wins(self):
        with tempfile.TemporaryDirectory() as temporary:
            env_file = Path(temporary) / ".env"
            literal_key = "SYNTHETIC-SECRET-${HOME}-$(Write-Output ignored)#literal"
            env_file.write_text(
                "# Synthetic offline settings only\n"
                f'OPENAI_API_KEY="{literal_key}"\n'
                "RESEARCH_MODEL='file-researcher'\n"
                "EXTRACTION_MODEL=file-extractor\n"
                "REVIEW_MODEL=file-reviewer\n", encoding="utf-8")
            with patch.dict(os.environ, {"RESEARCH_MODEL": "environment-researcher"}, clear=True), \
                    patch("os.system") as shell, patch("research_platform.ai.urlopen") as network:
                before = dict(os.environ)
                provider = OpenAIProvider.from_env(env_file)
                provider.validate_configuration()
                self.assertEqual(provider._api_key, literal_key)
                self.assertEqual(provider.models, {
                    "researcher": "environment-researcher", "extractor": "file-extractor",
                    "reviewer": "file-reviewer"})
                self.assertEqual(dict(os.environ), before)
                self.assertNotIn(literal_key, json.dumps(provider.metadata()))
                shell.assert_not_called()
                network.assert_not_called()

    def test_env_configuration_errors_do_not_reveal_values(self):
        with tempfile.TemporaryDirectory() as temporary:
            env_file = Path(temporary) / ".env"
            with patch.dict(os.environ, {}, clear=True), patch("research_platform.ai.urlopen") as network:
                env_file.write_text("OPENAI_API_KEY=SYNTHETIC-SECRET\n", encoding="utf-8")
                provider = OpenAIProvider.from_env(env_file)
                provider.validate_configuration()
                self.assertEqual(set(provider.models.values()), {"gpt-6-luna"})
                self.assertNotIn("SYNTHETIC-SECRET", json.dumps(provider.metadata()))
                for invalid in ('OPENAI_API_KEY="SYNTHETIC-SECRET\n',
                                'UNRECOGNIZED_SETTING=SYNTHETIC-SECRET\n',
                                'OPENAI_API_KEY=SYNTHETIC-SECRET\nOPENAI_API_KEY=duplicate\n'):
                    with self.subTest(invalid=invalid):
                        env_file.write_text(invalid, encoding="utf-8")
                        with self.assertRaises(AIError) as raised:
                            OpenAIProvider.from_env(env_file)
                        self.assertNotIn("SYNTHETIC-SECRET", str(raised.exception))
                self.assertEqual(dict(os.environ), {})
                network.assert_not_called()

    def test_adapter_requires_real_completed_search_and_strict_output(self):
        provider = self.provider()
        result, sent = self.complete(provider, {"status": "completed", "output": [
            {"type": "web_search_call", "status": "completed"},
            {"type": "message", "content": [{"type": "output_text", "text": "{}"}]}]}, web=True)
        self.assertEqual(result, {})
        self.assertEqual(sent["tools"], [{"type": "web_search"}])
        self.assertEqual(sent["tool_choice"], "required")
        self.assertTrue(sent["text"]["format"]["strict"])
        self.assertFalse(sent["store"])
        self.assertNotIn("SYNTHETIC-SECRET", json.dumps(provider.metadata()))
        with self.assertRaisesRegex(AIError, "web search"):
            self.complete(provider, {"status": "completed", "output": [
                {"type": "message", "content": [{"type": "output_text", "text": "{}"}]}]}, web=True)

    def test_request_diagnostics_have_unique_ids_and_preserve_reported_usage(self):
        provider = self.provider()
        response = {"status": "completed", "usage": {"input_tokens": 9, "output_tokens": 2},
                    "output": [{"type": "message", "content": [
                        {"type": "output_text", "text": "{}"}]}]}
        request_ids = []
        for _ in range(2):
            stream = io.BytesIO(json.dumps(response).encode("utf-8"))
            with patch("research_platform.ai.urlopen", return_value=stream) as network:
                result = provider.complete(role="researcher", instructions="test", payload={},
                                           schema={"type": "object"})
            self.assertEqual(result, {})
            network.assert_called_once()
            self.assertEqual(network.call_args.kwargs["timeout"], 120)
            diagnostics = provider.last_response["_transport"]
            request_id = diagnostics["client_request_id"]
            self.assertEqual(str(UUID(request_id)), request_id)
            headers = {name.lower(): value for name, value in network.call_args.args[0].header_items()}
            self.assertEqual(headers["x-client-request-id"], request_id)
            self.assertEqual(diagnostics["timeout_seconds"], 120)
            self.assertGreaterEqual(diagnostics["elapsed_seconds"], 0)
            self.assertNotIn("failure_kind", diagnostics)
            self.assertEqual(provider.last_response["usage"], response["usage"])
            self.assertNotIn("SYNTHETIC-SECRET", json.dumps(provider.last_response))
            request_ids.append(request_id)
        self.assertNotEqual(*request_ids)

    def test_transport_diagnostics_classify_failures_without_retry_or_secret_leaks(self):
        untrusted_codes = OSError("SYNTHETIC-SECRET")
        untrusted_codes.errno = "SYNTHETIC-SECRET"
        untrusted_codes.winerror = "SYNTHETIC-SECRET"
        cases = [
            (URLError(TimeoutError("SYNTHETIC-SECRET")), "timeout", False),
            (TimeoutError("SYNTHETIC-SECRET"), "timeout", True),
            (URLError(socket.gaierror(-2, "SYNTHETIC-SECRET")), "dns_error", False),
            (URLError(ssl.SSLCertVerificationError(1, "SYNTHETIC-SECRET")), "tls_certificate_error", False),
            (URLError(ssl.SSLError(1, "SYNTHETIC-SECRET")), "tls_error", False),
            (URLError(ConnectionResetError(10054, "SYNTHETIC-SECRET")), "connection_error", False),
            (URLError("SYNTHETIC-SECRET"), "network_error", False),
            (URLError(untrusted_codes), "network_error", False),
            (IncompleteRead(b"SYNTHETIC-SECRET", 20), "interrupted_response", True),
            (BadStatusLine("SYNTHETIC-SECRET"), "protocol_error", True),
        ]
        for error, expected, during_read in cases:
            with self.subTest(kind=expected, during_read=during_read):
                provider = self.provider()
                stream = io.BytesIO(b"")
                with patch("research_platform.ai.urlopen", return_value=stream,
                           side_effect=None if during_read else error) as network, \
                        patch.object(stream, "read", side_effect=error), \
                        self.assertRaises(AIError) as raised:
                    provider.complete(role="researcher", instructions="test", payload={},
                                      schema={"type": "object"})
                network.assert_called_once()
                self.assertEqual(provider.last_response["status"], "transport_error")
                diagnostics = provider.last_response["_transport"]
                self.assertEqual(diagnostics["failure_kind"], expected)
                self.assertEqual(str(UUID(diagnostics["client_request_id"])),
                                 diagnostics["client_request_id"])
                self.assertEqual(diagnostics["timeout_seconds"], 120)
                self.assertGreaterEqual(diagnostics["elapsed_seconds"], 0)
                for key in ("errno", "winerror"):
                    if key in diagnostics:
                        self.assertIs(type(diagnostics[key]), int)
                self.assertNotIn("usage", provider.last_response)
                usage = response_usage(provider.last_response, "gpt-6.1-sol")
                self.assertFalse(usage["reported"])
                self.assertIsNone(usage["estimated_cost_usd"])
                self.assertNotIn("SYNTHETIC-SECRET",
                                 str(raised.exception) + json.dumps(provider.last_response))

    def test_adapter_rejects_refusal_incomplete_and_malformed_results(self):
        provider = self.provider()
        failures = [
            {"status": "incomplete", "output": []},
            {"status": "completed", "output": [{"type": "message", "content": [{"type": "refusal", "refusal": "No"}]}]},
            {"status": "completed", "output": []},
            ["not an API object"],
        ]
        for invalid_json in ("not JSON", '{"a":1,"a":2}', '{"a":NaN}', '{"a":1e999}', "[]"):
            failures.append({"status": "completed", "output": [
                {"type": "message", "content": [{"type": "output_text", "text": invalid_json}]}]})
        for failure in failures:
            with self.subTest(failure=failure), self.assertRaises(AIError):
                self.complete(provider, failure)
        for read_failure in (None, BadStatusLine("SYNTHETIC-SECRET")):
            error = HTTPError("https://example.invalid/SYNTHETIC-SECRET", 429, "SYNTHETIC-SECRET", {}, None)
            with self.subTest(error_body_interrupted=read_failure is not None), \
                    patch("research_platform.ai.urlopen", side_effect=error) as network, \
                    patch.object(error, "read", return_value=b"", side_effect=read_failure):
                with self.assertRaises(AIError) as raised:
                    provider.complete(role="researcher", instructions="test", payload={},
                                      schema={"type": "object"})
                self.assertEqual(network.call_count, 1)
                self.assertNotIn("SYNTHETIC-SECRET", str(raised.exception))
                self.assertEqual(provider.last_response["status"], "transport_error")
                self.assertEqual(provider.last_response["http_status"], 429)
                self.assertEqual(provider.last_response["_transport"]["failure_kind"], "http_error")
                self.assertNotIn("SYNTHETIC-SECRET", json.dumps(provider.last_response))
                usage = response_usage(provider.last_response, "gpt-6.1-sol")
                self.assertFalse(usage["reported"])
                self.assertIsNone(usage["estimated_cost_usd"])


class AIUsageTests(unittest.TestCase):
    @staticmethod
    def response(model="gpt-6.1-sol", *, input_tokens=1000, output_tokens=100,
                 cached=0, cache_write=0, reasoning=0):
        return {"status": "completed", "model": model, "output": [], "usage": {
            "input_tokens": input_tokens,
            "input_tokens_details": {"cached_tokens": cached, "cache_write_tokens": cache_write},
            "output_tokens": output_tokens,
            "output_tokens_details": {"reasoning_tokens": reasoning},
            "total_tokens": input_tokens + output_tokens}}

    def test_standard_model_rates_partition_cache_without_rebilling_reasoning(self):
        for model, expected in (("gpt-6.1-sol", 0.0267), ("gpt-6-luna", 0.001345),
                                ("gpt-6-astra", 0.1345)):
            with self.subTest(model=model):
                response = self.response(model, input_tokens=10000, output_tokens=1000,
                                         cached=2000, cache_write=1000, reasoning=400)
                usage = response_usage(response, model)
                self.assertTrue(usage["reported"])
                self.assertEqual(usage["input_tokens"], 10000)
                self.assertEqual(usage["cached_input_tokens"], 2000)
                self.assertEqual(usage["cache_write_tokens"], 1000)
                self.assertEqual(usage["output_tokens"], 1000)
                self.assertEqual(usage["reasoning_tokens"], 400)
                self.assertEqual(usage["total_tokens"], 11000)
                self.assertAlmostEqual(usage["estimated_cost_usd"], expected)
                self.assertTrue(usage["pricing_date"])
                self.assertIn("openai.com", usage["pricing_source"])
        dated = response_usage(self.response("gpt-6.1-sol-2026-10-01"), "gpt-6.1-sol")
        self.assertEqual(dated["model"], "gpt-6.1-sol-2026-10-01")
        self.assertAlmostEqual(dated["estimated_cost_usd"], 0.003)

    def test_known_pricing_requires_an_exact_model_or_dated_snapshot(self):
        from research_platform.ai import has_known_pricing

        for model in ("gpt-6.1-sol", "gpt-6-luna", "gpt-6-astra", "gpt-6.1-sol-2026-10-01"):
            self.assertTrue(has_known_pricing(model))
        for model in ("unpriced-model", "gpt-6.1-sol-extra", "prefix-gpt-6-luna", "", None):
            self.assertFalse(has_known_pricing(model))

    def test_long_context_threshold_changes_entire_request_above_272000(self):
        for model, rates in (("gpt-6.1-sol", (2, 0.1, 2.5, 10)),
                             ("gpt-6-luna", (0.1, 0.01, 0.125, 0.5)),
                             ("gpt-6-astra", (10, 1, 12.5, 50))):
            ordinary_rate, cached_rate, write_rate, output_rate = rates
            for count, input_multiplier, output_multiplier in ((272000, 1, 1), (272001, 2, 1.5)):
                with self.subTest(model=model, input_tokens=count):
                    usage = response_usage(self.response(model, input_tokens=count, output_tokens=1000,
                                                         cached=20000, cache_write=12000), model)
                    expected = (((count - 32000) * ordinary_rate + 20000 * cached_rate
                                 + 12000 * write_rate) * input_multiplier
                                + 1000 * output_rate * output_multiplier) / 1000000
                    self.assertAlmostEqual(usage["estimated_cost_usd"], expected)

    def test_search_actions_are_charged_and_unknown_pricing_stays_unknown(self):
        response = self.response()
        response["output"] = [{"type": "web_search_call", "status": "completed",
                               "action": {"type": action}}
                              for action in ("search", "open_page", "find_in_page")]
        usage = response_usage(response, "gpt-6.1-sol")
        self.assertEqual(usage["web_search_calls"], 1)
        self.assertEqual(usage["web_tool_calls"], 3)
        self.assertAlmostEqual(usage["estimated_cost_usd"], 0.013)
        unknown_cases = []
        for action in (None, {"type": "future_action"}):
            unknown = copy.deepcopy(response)
            unknown["output"].append({"type": "web_search_call", "action": action})
            unknown_cases.append(unknown)
        unknown_cases.extend((dict(response, service_tier="priority"),
                              dict(response, model="unpriced-model")))
        for unknown in unknown_cases:
            with self.subTest(unknown=unknown):
                usage = response_usage(unknown, "gpt-6.1-sol")
                self.assertTrue(usage["reported"])
                self.assertIsNone(usage["estimated_cost_usd"])

    def test_aggregate_retains_known_subtotal_and_does_not_fabricate_missing_usage(self):
        entries = [response_usage(self.response(model), model)
                   for model in ("gpt-6.1-sol", "gpt-6-luna", "unpriced-model")]
        missing = response_usage(None, "gpt-6.1-sol")
        self.assertFalse(missing["reported"])
        self.assertIsNone(missing["input_tokens"])
        self.assertIsNone(missing["output_tokens"])
        self.assertIsNone(missing["estimated_cost_usd"])
        total = aggregate_usage([{"usage": entries[0]}, entries[1], entries[2], missing])
        self.assertEqual(total["calls"], 4)
        self.assertEqual(total["reported_calls"], 3)
        self.assertEqual(total["missing_usage"], 1)
        self.assertEqual(total["missing_pricing"], 1)
        self.assertFalse(total["cost_is_complete"])
        self.assertEqual(total["input_tokens"], 3000)
        self.assertEqual(total["output_tokens"], 300)
        self.assertEqual(total["total_tokens"], 3300)
        self.assertAlmostEqual(total["estimated_cost_usd"], 0.00315)
        self.assertIn("gpt-6-luna", total["by_model"])
        self.assertIsNone(aggregate_usage([missing])["estimated_cost_usd"])
        for invalid in ({"input_tokens": -1, "output_tokens": 10},
                        {"input_tokens": True, "output_tokens": 10},
                        {"input_tokens": 10, "output_tokens": 10, "total_tokens": 19}):
            with self.subTest(invalid=invalid):
                usage = response_usage(dict(self.response(), usage=invalid), "gpt-6.1-sol")
                self.assertFalse(usage["reported"])
                self.assertIsNone(usage["estimated_cost_usd"])

    def test_stage_totals_include_revisions_and_charged_failed_calls(self):
        discovery = self.response(input_tokens=2000, output_tokens=300, reasoning=100)
        discovery["output"] = [{"type": "web_search_call", "status": "completed",
                                 "action": {"type": "search"}}]
        extraction = self.response(input_tokens=1000, output_tokens=200, cached=400,
                                   cache_write=100, reasoning=50)
        failed_review = dict(self.response(input_tokens=500, output_tokens=100,
                                         reasoning=90), status="incomplete")
        entries = [
            {"stage": stage, "attempt": attempt, "usage": response_usage(response, "gpt-6.1-sol")}
            for stage, attempt, response in (
                ("discover", 1, discovery), ("extract", 1, extraction),
                ("model_review", 1, failed_review), ("extract", 2, extraction))
        ]
        summary = aggregate_usage(entries)
        self.assertEqual(summary["calls"], 4)
        self.assertEqual(summary["total_tokens"], 5300)
        self.assertEqual(summary["reasoning_tokens"], 290)
        self.assertEqual(summary["total_tokens"], summary["input_tokens"] + summary["output_tokens"])
        self.assertEqual(summary["by_stage"]["extract"]["calls"], 2)
        self.assertEqual(summary["by_stage"]["extract"]["total_tokens"], 2400)
        self.assertEqual(summary["by_stage"]["extract"]["cached_input_tokens"], 800)
        self.assertEqual(summary["by_stage"]["model_review"]["total_tokens"], 600)
        self.assertEqual(summary["by_stage"]["discover"]["web_search_calls"], 1)
        self.assertAlmostEqual(summary["estimated_cost_usd"], 0.02558)
        for key in ("calls", "input_tokens", "output_tokens", "total_tokens", "web_search_calls"):
            self.assertEqual(summary[key], sum(stage[key] for stage in summary["by_stage"].values()))
        self.assertAlmostEqual(summary["estimated_cost_usd"], sum(
            stage["estimated_cost_usd"] for stage in summary["by_stage"].values()))
        self.assertTrue(summary["usage_is_complete"])
        self.assertTrue(all(stage["cost_is_complete"] for stage in summary["by_stage"].values()))

    def test_stage_gaps_distinguish_missing_tokens_from_missing_prices(self):
        known = response_usage(self.response(), "gpt-6.1-sol")
        unknown_price = response_usage(self.response("unpriced-model"), "unpriced-model")
        missing = response_usage(None, "gpt-6.1-sol")
        summary = aggregate_usage([
            {"stage": "extract", "usage": known}, {"stage": "extract", "usage": missing},
            {"stage": "model_review", "usage": unknown_price},
            {"stage": "discover", "usage": missing},
        ])
        extraction = summary["by_stage"]["extract"]
        self.assertEqual(extraction["total_tokens"], 1100)
        self.assertFalse(extraction["usage_is_complete"])
        self.assertFalse(extraction["cost_is_complete"])
        self.assertAlmostEqual(extraction["estimated_cost_usd"], 0.003)
        self.assertTrue(summary["by_stage"]["model_review"]["usage_is_complete"])
        self.assertFalse(summary["by_stage"]["model_review"]["cost_is_complete"])
        self.assertIsNone(summary["by_stage"]["model_review"]["estimated_cost_usd"])
        self.assertIsNone(summary["by_stage"]["discover"]["estimated_cost_usd"])
        self.assertFalse(summary["usage_is_complete"])
        self.assertEqual(summary["missing_usage"], 2)
        self.assertEqual(summary["missing_pricing"], 1)
        self.assertAlmostEqual(summary["estimated_cost_usd"], 0.003)
        self.assertIn("unspecified", aggregate_usage([known])["by_stage"])

    def test_terminal_usage_shows_total_and_stage_subtotals_without_zero_for_unknown(self):
        from research_platform.ai import format_stage_usage, format_usage

        known = response_usage(self.response(reasoning=80), "gpt-6.1-sol")
        missing = response_usage(None, "gpt-6.1-sol")
        summary = aggregate_usage([
            {"stage": "extract", "usage": known}, {"stage": "extract", "usage": missing},
            {"stage": "model_review", "usage": missing},
        ])
        total_line = format_usage(summary)
        stage_lines = format_stage_usage(summary).splitlines()
        self.assertIn("Reported token subtotal: 1,100 total", total_line)
        self.assertIn("2 call(s) missing usage", total_line)
        self.assertIn("80 reasoning included", total_line)
        self.assertIn("extract: 2 call(s); 1,100 total tokens", stage_lines[0])
        self.assertIn("subtotal, 1 call(s) missing usage", stage_lines[0])
        self.assertIn("$0.003000 USD (incomplete)", stage_lines[0])
        self.assertIn("model_review: 1 call(s); tokens unavailable", stage_lines[1])
        self.assertIn("estimated cost unknown (incomplete)", stage_lines[1])
        self.assertNotIn("$0.000000", stage_lines[1])

    def test_rejected_api_responses_still_retain_billable_usage(self):
        adapter = AIAdapterTests()
        provider = adapter.provider()
        for rejected in (
            dict(self.response("gpt-6-luna"), status="incomplete"),
            dict(self.response("gpt-6-luna"), output=[{"type": "message", "content": [
                {"type": "refusal", "refusal": "Synthetic refusal"}]}]),
        ):
            with self.subTest(status=rejected["status"]):
                with self.assertRaises(AIError):
                    adapter.complete(provider, rejected)
                usage = response_usage(provider.last_response, "gpt-6-luna")
                self.assertTrue(usage["reported"])
                self.assertEqual(usage["total_tokens"], 1100)
                self.assertAlmostEqual(usage["estimated_cost_usd"], 0.00015)


class ApiCostEffortTests(unittest.TestCase):
    def provider(self, **kwargs):
        return OpenAIProvider(api_key="SYNTHETIC-SECRET", models={
            role: "gpt-6.1-sol" for role in ("researcher", "extractor", "reviewer")}, **kwargs)

    def arguments(self, payload):
        return {"role": "extractor", "instructions": "Synthetic document mapping",
                "payload": payload, "schema": {"type": "object"}}

    def test_full_documents_have_stable_boundaries_without_duplicate_dynamic_context(self):
        provider = self.provider()
        units = [{"source_id": "a", "unit_id": "page-1", "text": "Original text " * 30000,
                  "numbers": [{"raw": "12", "index": 0}]},
                 {"source_id": "a", "unit_id": "page-2", "text": "Important late note", "numbers": []}]
        original = {"evidence": units, "feedback": ["first"]}
        untouched = copy.deepcopy(original)
        first = provider._request_body(**self.arguments(original))
        appended = {"source_id": "b", "unit_id": "sheet!R1", "text": "Other document", "numbers": []}
        second = provider._request_body(**self.arguments(
            {"evidence": [*units, appended], "feedback": ["changed"]}))
        self.assertEqual(original, untouched)
        self.assertEqual(first["input"][0], second["input"][0])
        self.assertEqual(len(first["input"]), 2)
        self.assertEqual(len(second["input"]), 3)
        supplied = json.loads(first["input"][0]["content"][0]["text"])
        self.assertEqual(supplied["evidence"], units)
        dynamic = json.loads(second["input"][-1]["content"][0]["text"])
        self.assertEqual(dynamic["evidence"]["source_ids"], ["a", "b"])
        self.assertNotIn("Important late note", second["input"][-1]["content"][0]["text"])
        self.assertEqual(json.loads(provider._request_body(**self.arguments(
            {"evidence": [], "feedback": []}))["input"][0]["content"][0]["text"])["evidence"], [])

    def test_identical_references_deduplicate_but_conflicting_evidence_never_does(self):
        provider = self.provider()
        unit = {"source_id": "a", "unit_id": "page-1", "text": "10", "numbers": []}
        body = provider._request_body(**self.arguments({"evidence": [unit, copy.deepcopy(unit)]}))
        self.assertEqual(json.loads(body["input"][0]["content"][0]["text"])["evidence"], [unit])
        with self.assertRaisesRegex(AIError, "Conflicting"):
            provider._request_body(**self.arguments({"evidence": [unit, dict(unit, text="20")]}))

    def test_catalog_is_an_intact_stable_prefix_before_selected_documents(self):
        provider = self.provider()
        catalog = [{"id": "a", "title": "Whole annual report", "category": "financial_reports",
                    "units": [{"id": "page-1", "preview": "Unique navigation preview", "number_count": 2}]},
                   {"id": "b", "title": "Official historical tables", "category": "financial_reports", "units": []}]
        first_unit = {"source_id": "a", "unit_id": "page-1", "text": "Complete original text", "numbers": []}
        second_unit = {"source_id": "b", "unit_id": "Table!R1", "text": "Full spreadsheet row", "numbers": []}
        original = {"catalog": catalog, "evidence": [], "feedback": []}
        untouched = copy.deepcopy(original)
        first = provider._request_body(**self.arguments(original))
        after_read = provider._request_body(**self.arguments(
            {"catalog": catalog, "evidence": [first_unit], "feedback": ["changed task state"]}))
        after_second = provider._request_body(**self.arguments(
            {"catalog": catalog, "evidence": [first_unit, second_unit], "feedback": ["changed again"]}))
        self.assertEqual(original, untouched)
        self.assertEqual(first["input"][0], after_read["input"][0])
        self.assertEqual(after_read["input"][:2], after_second["input"][:2])
        self.assertEqual(json.loads(first["input"][0]["content"][0]["text"])["catalog"], catalog)
        self.assertEqual(json.loads(after_read["input"][1]["content"][0]["text"])["evidence"], [first_unit])
        self.assertNotIn("Unique navigation preview", after_second["input"][-1]["content"][0]["text"])
        self.assertEqual(json.loads(after_second["input"][-1]["content"][0]["text"])["catalog"]["source_count"], 2)

    def test_filtered_review_catalog_never_inherits_other_roles_sources_or_reads(self):
        provider = self.provider()
        primary = {"id": "annual", "category": "financial_reports", "title": "Public original"}
        vendor = {"id": "vendor", "category": "visible_alpha", "title": "Vendor-only context"}
        provider._request_body(**self.arguments({"catalog": [primary, vendor], "evidence": [
            {"source_id": "vendor", "unit_id": "sheet!R1", "text": "Prior proposed value", "numbers": []}]}))
        arguments = dict(self.arguments({"catalog": [primary], "evidence": [], "concepts": []}), role="reviewer")
        blind = provider._request_body(**arguments)
        self.assertEqual(len(blind["input"]), 2)
        self.assertEqual(json.loads(blind["input"][0]["content"][0]["text"])["catalog"], [primary])
        self.assertEqual(json.loads(blind["input"][-1]["content"][0]["text"])["evidence"], [])
        self.assertNotIn("Vendor-only context", json.dumps(blind))
        self.assertNotIn("Prior proposed value", json.dumps(blind))
        self.assertIn("catalog-and-whole-document-prefix", provider.metadata()["prompt_layout"])

    def test_effort_defaults_and_settings_are_explicit_in_request_and_resume_metadata(self):
        baseline = self.provider()
        self.assertEqual(baseline._request_body(**self.arguments({}))["reasoning"],
                         {"effort": "medium", "mode": "standard"})
        self.assertEqual(baseline.metadata()["reasoning_efforts"],
                         dict.fromkeys(("researcher", "extractor", "reviewer"), "medium"))
        changed = self.provider(reasoning_efforts={"extractor": "high"})
        self.assertNotEqual(changed.metadata(), baseline.metadata())
        with tempfile.TemporaryDirectory() as temporary:
            env_file = Path(temporary) / ".env"
            env_file.write_text("OPENAI_API_KEY=SYNTHETIC-SECRET\n"
                "RESEARCH_MODEL=gpt-6.1-sol\nEXTRACTION_MODEL=gpt-6.1-sol\nREVIEW_MODEL=gpt-6.1-sol\n"
                "EXTRACTION_REASONING_EFFORT=high\nREVIEW_REASONING_EFFORT=high\n", encoding="utf-8")
            with patch.dict(os.environ, {"REVIEW_REASONING_EFFORT": "medium"}, clear=True):
                configured = OpenAIProvider.from_env(env_file)
                configured.validate_configuration()
            self.assertEqual(configured.reasoning_efforts,
                             {"researcher": "medium", "extractor": "high", "reviewer": "medium"})
            self.assertNotIn("SYNTHETIC-SECRET", json.dumps(configured.metadata()))

    def test_invalid_effort_is_rejected_before_a_request(self):
        for effort in ("", "extreme", None, [], "SYNTHETIC-SECRET"):
            with self.subTest(effort=effort), patch("research_platform.ai.urlopen") as network:
                provider = self.provider(reasoning_efforts={"extractor": effort})
                with self.assertRaises(AIError) as raised:
                    provider.complete(**self.arguments({}))
                self.assertNotIn("SYNTHETIC-SECRET", str(raised.exception))
                network.assert_not_called()

    def test_local_allowance_reserves_input_and_full_output_without_spending_or_usage(self):
        provider = self.provider()
        with patch("research_platform.ai.urlopen") as network:
            small = provider.estimate_request_cost(**self.arguments({"text": "tiny"}))
            large = provider.estimate_request_cost(**self.arguments({"text": "x" * 280000}))
            unknown_web = provider.estimate_request_cost(**self.arguments({}), web=True)
            network.assert_not_called()
        self.assertIsNone(provider.last_response)
        self.assertFalse(small["is_exact"])
        self.assertGreaterEqual(small["estimated_input_tokens"], 2048)
        self.assertAlmostEqual(small["estimated_cost_usd"],
                               (small["estimated_input_tokens"] * 2.5 + 16000 * 10) / 1_000_000)
        self.assertAlmostEqual(large["estimated_cost_usd"],
                               (large["estimated_input_tokens"] * 5 + 16000 * 15) / 1_000_000)
        self.assertGreater(large["estimated_cost_usd"], small["estimated_cost_usd"])
        self.assertIsNone(unknown_web["estimated_cost_usd"])


class LocalSourceTests(unittest.TestCase):
    """Local originals remain reproducible evidence, never implicit API input."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.archive = self.root / "company"
        (self.archive / "financial_reports").mkdir(parents=True)

    def inventory(self):
        from research_platform.evidence import inventory_local_sources
        return inventory_local_sources(self.archive)

    @staticmethod
    def legacy_workbook(*, cached_formula=False, empty_formula=False):
        """A tiny independent BIFF8 fixture: a label and one numeric/formula cell."""
        import struct

        def record(code, data):
            return struct.pack("<HH", code, len(data)) + data

        def beginning(kind):
            return record(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 1997, 0, 0x0600))

        end = record(0x000A, b"")
        # Built-in number format 2 is 0.00; no workbook writer is required.
        style = record(0x00E0, struct.pack("<HH", 0, 2) + bytes(16))
        label = record(0x0204, struct.pack("<HHHHB", 0, 0, 0, 7, 0) + b"Revenue")
        if cached_formula or empty_formula:
            cache = b"\x03\x00\x00\x00\x00\x00\xff\xff" if empty_formula else struct.pack("<d", 123.5)
            formula = b"\x1f" + struct.pack("<d", 123.5)
            number = record(0x0006, struct.pack("<HHH", 0, 1, 0) + cache
                            + struct.pack("<HIH", 0, 0, len(formula)) + formula)
        else:
            number = record(0x0203, struct.pack("<HHHd", 0, 1, 0, 123.5))
        sheet = (beginning(0x0010) + record(0x0200, struct.pack("<IIHHH", 0, 1, 0, 2, 0))
                 + label + number + end)

        def sheet_reference(offset):
            return record(0x0085, struct.pack("<IBBBB", offset, 0, 0, 5, 0) + b"Sheet")

        globals_ = beginning(0x0005) + record(0x0042, struct.pack("<H", 1200)) + style
        offset = len(globals_ + sheet_reference(0) + end)
        return globals_ + sheet_reference(offset) + end + sheet

    def test_legacy_xls_preserves_original_and_numeric_cell_precision_on_restore(self):
        path = self.archive / "financial_reports" / "original.xls"
        path.write_bytes(self.legacy_workbook())
        spec = self.inventory()["sources"][0]
        self.assertEqual(spec["format"], "xls")
        store = SourceStore(self.root / "run" / "sources")
        descriptor = store.add_local(spec, root=self.archive)
        reference = {"source_id": spec["id"], "unit_id": "Sheet!R1", "number_index": 0,
                     "number_format": "dot_decimal"}
        self.assertEqual(store.resolve(reference), 123.5)
        self.assertEqual(store.read([reference])[0]["numbers"][0]["display_format"], "0.00")
        self.assertEqual(_rounding_bound(reference, 1, store), 0.005)
        self.assertEqual(Path(descriptor["path"]).read_bytes(), path.read_bytes())
        restored = SourceStore(self.root / "continuation")
        restored.restore(descriptor, archive_root=self.root / "run")
        self.assertEqual(restored.resolve(reference), 123.5)

    def test_legacy_xls_formula_cache_is_read_and_never_recalculated(self):
        path = self.archive / "financial_reports" / "original.xls"
        path.write_bytes(self.legacy_workbook(cached_formula=True))
        spec = self.inventory()["sources"][0]
        store = SourceStore(self.root / "run")
        descriptor = store.add_local(spec, root=self.archive)
        reference = {"source_id": spec["id"], "unit_id": "Sheet!R1", "number_index": 0,
                     "number_format": "dot_decimal"}
        self.assertEqual(store.resolve(reference), 123.5)
        self.assertIn("not recalculated", descriptor["extraction_warnings"][0]["messages"][0])

    def test_legacy_xls_missing_formula_cache_fails_explicitly(self):
        path = self.archive / "financial_reports" / "original.xls"
        path.write_bytes(self.legacy_workbook(empty_formula=True))
        spec = self.inventory()["sources"][0]
        store = SourceStore(self.root / "run")
        with self.assertRaisesRegex(ValueError, "no usable cached result.*B1"):
            store.add_local(spec, root=self.archive)
        self.assertEqual(store.sources[0]["status"], "unreadable")
        self.assertEqual(Path(store.sources[0]["path"]).read_bytes(), path.read_bytes())

    def test_primary_search_filters_sources_before_limiting_matches(self):
        self.report(value="100")
        vendor = self.archive / "visible_alpha"
        vendor.mkdir()
        (vendor / "actuals.csv").write_text(
            "Concept,Actual\n" + "".join(f"Revenue {index},{index}\n" for index in range(40)),
            encoding="utf-8")
        specs = self.inventory()["sources"]
        vendor_spec = next(spec for spec in specs if spec["category"] == "visible_alpha")
        primary_spec = next(spec for spec in specs if spec["category"] == "financial_reports")
        store = SourceStore(self.root / "run")
        # The crowded vendor is deliberately indexed before primary evidence.
        store.add_local(vendor_spec, root=self.archive)
        store.add_local(primary_spec, root=self.archive)
        unrestricted = store.search("Revenue")
        self.assertEqual(len(unrestricted["matches"]), 30)
        self.assertTrue(unrestricted["truncated"])
        self.assertEqual({match["source_id"] for match in unrestricted["matches"]}, {vendor_spec["id"]})
        primary = store.search("Revenue", source_ids={primary_spec["id"]})
        self.assertEqual(len(primary["matches"]), 1)
        self.assertEqual(primary["matches"][0]["source_id"], primary_spec["id"])
        self.assertFalse(primary["truncated"])
        self.assertEqual(store.search("Revenue", source_ids=[])["matches"], [])
        for invalid in ({"unknown"}, primary_spec["id"], [None]):
            with self.subTest(source_ids=invalid):
                with self.assertRaises(ValueError):
                    store.search("Revenue", source_ids=invalid)

    def report(self, name="company_2025-fy_results_pub-2026-02-19.html", value="8\u2019639"):
        path = self.archive / "financial_reports" / name
        path.write_text(f"<html><body><h1>Consolidated company accounts</h1><p>Revenue  {value}</p></body></html>",
                        encoding="utf-8")
        return path

    def test_local_inventory_archives_originals_without_network_and_restores(self):
        path = self.report()
        manifest = self.inventory()
        spec = manifest["sources"][0]
        self.assertEqual(spec["category"], "financial_reports")
        self.assertIsNone(spec["published_at"])
        self.assertEqual(spec["filename_publication_hint"], "2026-02-19")
        self.assertEqual(spec["metadata_verification"], "user_supplied_not_verified")
        self.assertEqual(manifest, self.inventory())
        store = SourceStore(self.root / "run" / "sources", fetcher=lambda url: self.fail("Unexpected network access"))
        descriptor = store.add_local(spec, root=self.archive)
        self.assertEqual(Path(descriptor["path"]).read_bytes(), path.read_bytes())
        self.assertEqual(store.add_local(spec, root=self.archive), descriptor)
        self.assertEqual(store.catalog()[0]["category"], "financial_reports")
        unit = store.read_documents([{"source_id": spec["id"], "unit_id": "chunk:1"}])[0]
        reference = {"source_id": spec["id"], "unit_id": "chunk:1", "number_index": 0,
                     "number_format": "dot_decimal"}
        self.assertEqual(unit["numbers"][0]["raw"], "8\u2019639")
        self.assertEqual(store.resolve(reference), 8639)
        restored = SourceStore(self.root / "continuation")
        restored.restore(descriptor, archive_root=self.root / "run")
        self.assertEqual(restored.resolve(reference), 8639)

    def test_manifest_detects_modified_added_removed_files_and_metadata(self):
        path = self.report()
        before = self.inventory()
        path.write_text("<html>Revenue  8640</html>", encoding="utf-8")
        changed = self.inventory()
        self.assertNotEqual(before["manifest_sha256"], changed["manifest_sha256"])
        with self.assertRaisesRegex(ValueError, "changed after inventory"):
            SourceStore(self.root / "run").add_local(before["sources"][0], root=self.archive)
        extra = self.report("another.html", "123")
        added = self.inventory()
        self.assertNotEqual(changed["manifest_sha256"], added["manifest_sha256"])
        extra.unlink()
        self.assertEqual(changed["manifest_sha256"], self.inventory()["manifest_sha256"])
        (self.archive / "sources.csv").write_text(
            "relative_filename,original_download_url,verified_publication_date,reporting_basis\n"
            f"financial_reports/{path.name},https://example.com/results,2026-02-19,restated\n", encoding="utf-8")
        indexed = self.inventory()
        self.assertNotEqual(changed["manifest_sha256"], indexed["manifest_sha256"])
        self.assertEqual(indexed["sources"][0]["published_at"], "2026-02-19")
        self.assertEqual(indexed["sources"][0]["supplied_metadata"]["reporting_basis"], "restated")
        self.assertEqual(indexed["sources"][0]["url"], "https://example.com/results")

    def test_local_archive_rejects_unsupported_files_and_modified_claimed_hash(self):
        path = self.report()
        unsupported = self.archive / "financial_reports" / "report.docx"
        unsupported.write_bytes(b"not a supported original")
        with self.assertRaisesRegex(ValueError, "Unsupported local source"):
            self.inventory()
        unsupported.unlink()
        (self.archive / "sources.csv").write_text(
            "relative_filename,sha256\n" + f"financial_reports/{path.name},{'0' * 64}\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "hash disagrees"):
            self.inventory()

    def test_local_paths_cannot_escape_archive_or_read_hidden_configuration(self):
        self.report()
        hidden = self.archive / ".private"
        hidden.write_text("synthetic secret must not be indexed", encoding="utf-8")
        manifest = self.inventory()
        self.assertEqual(len(manifest["sources"]), 1)
        spec = dict(manifest["sources"][0], local_path="../outside.html")
        with self.assertRaisesRegex(ValueError, "relative"):
            SourceStore(self.root / "run").add_local(spec, root=self.archive)
        (self.archive / "sources.csv").write_text("relative_filename\n../outside.html\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "relative filenames"):
            self.inventory()

    def test_local_history_can_archive_and_restore_more_than_twelve_documents(self):
        for year in range(2010, 2024):
            self.report(f"company_{year}-fy.html", str(year))
        manifest = self.inventory()
        store = SourceStore(self.root / "run" / "sources")
        restored = SourceStore(self.root / "resumed")
        for spec in manifest["sources"]:
            descriptor = store.add_local(spec, root=self.archive)
            restored.restore(descriptor, archive_root=self.root / "run")
        self.assertEqual(len(store.sources), 14)
        self.assertEqual(restored.catalog(), store.catalog())

    def test_vendor_formula_without_cached_result_is_an_explicit_failure(self):
        folder = self.archive / "visible_alpha"
        folder.mkdir()
        book = Workbook()
        sheet = book.active
        sheet.append(["Historical sales", "2025 Actual"])
        sheet.append(["Revenue", "=100+2"])
        book.save(folder / "original.xlsx")
        book.close()
        manifest = self.inventory()
        store = SourceStore(self.root / "run")
        with self.assertRaisesRegex(ValueError, "no cached result.*B2"):
            store.add_local(manifest["sources"][0], root=self.archive)
        self.assertEqual(store.sources[0]["status"], "unreadable")
        self.assertTrue(Path(store.sources[0]["path"]).is_file())

    def test_cached_vendor_formula_is_read_without_recalculation_and_is_flagged(self):
        import zipfile

        folder = self.archive / "visible_alpha"
        folder.mkdir()
        path = folder / "original.xlsx"
        book = Workbook()
        book.active.append(["Revenue", "=100+2"])
        book.save(path)
        book.close()
        # openpyxl deliberately does not calculate formulas. This fixture emulates
        # the cache written by Excel, independently of the reader under test.
        rewritten = io.BytesIO()
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(rewritten, "w") as output:
            for entry in original.infolist():
                payload = original.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    payload = re.sub(rb"<f>100\+2</f><v(?:\s*/>|></v>)", b"<f>100+2</f><v>102</v>", payload)
                output.writestr(entry, payload)
        path.write_bytes(rewritten.getvalue())
        spec = self.inventory()["sources"][0]
        store = SourceStore(self.root / "run")
        descriptor = store.add_local(spec, root=self.archive)
        reference = {"source_id": spec["id"], "unit_id": "Sheet!R1", "number_index": 0,
                     "number_format": "dot_decimal"}
        self.assertEqual(store.resolve(reference), 102)
        self.assertIn("not recalculated", descriptor["extraction_warnings"][0]["messages"][0])
        self.assertEqual(Path(descriptor["path"]).read_bytes(), path.read_bytes())

    def test_local_csv_and_transcript_have_exact_numeric_citations(self):
        vendor = self.archive / "visible_alpha"
        vendor.mkdir()
        (vendor / "actuals.csv").write_text("Concept,2025 Actual\nRevenue,100\n", encoding="utf-8")
        transcripts = self.archive / "transcripts"
        transcripts.mkdir()
        (transcripts / "remarks.txt").write_text("Management: Sales were 100.\nQ&A: definition follows.", encoding="utf-8")
        store = SourceStore(self.root / "run")
        manifest = self.inventory()
        for spec in manifest["sources"]:
            store.add_local(spec, root=self.archive)
        vendor_spec = next(spec for spec in manifest["sources"] if spec["category"] == "visible_alpha")
        reference = {"source_id": vendor_spec["id"], "unit_id": "row:2", "number_index": 0,
                     "number_format": "dot_decimal"}
        self.assertEqual(store.resolve(reference), 100)
        self.assertEqual(store.read([reference])[0]["numbers"][0]["locator"], "B2")
        self.assertEqual(len(store.sources), 2)

    def test_apostrophe_grouping_is_one_token_and_bad_grouping_is_not_split(self):
        from research_platform.evidence import _tokens
        for raw, expected, number_format in (("8'639", 8639, "dot_decimal"),
                                              ("(1\u2019234.50)", -1234.5, "dot_decimal"),
                                              ("1\u2018234,50", 1234.5, "comma_decimal")):
            with self.subTest(raw=raw):
                self.assertEqual(parse_number(raw, number_format), expected)
                self.assertEqual([token["raw"] for token in _tokens("Revenue  " + raw)], [raw])
        for invalid in ("1'23", "12\u201934", "1'234\u201956"):
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    parse_number(invalid)
                self.assertEqual(_tokens(invalid), [])


class LocalInitiationTests(unittest.TestCase):
    """Synthetic local originals exercise the full no-browser initiation route."""

    class LocalProvider(ScriptedProvider):
        def __init__(self, company, **kwargs):
            super().__init__(company, **kwargs)
            self.full_company = copy.deepcopy(company)
            self.web_flags = []

        @staticmethod
        def scope(company, periods):
            scoped = copy.deepcopy(company)
            scoped.discovery["periods"] = [item for item in scoped.discovery["periods"] if item["id"] in periods]
            fiscal_years = {item["year"] for item in scoped.discovery["periods"]}
            scoped.discovery["period_coverage"] = [item for item in scoped.discovery["period_coverage"]
                                                    if item["year"] in fiscal_years]
            for row in scoped.rows:
                for field in ("observations", "unavailable", "period_bases", "period_terms"):
                    row[field] = [item for item in row[field] if item["period"] in periods]
            scoped.candidate["rows"] = scoped.rows
            scoped.candidate["coverage"] = [item for item in scoped.candidate["coverage"] if item["period"] in periods]
            scoped.candidate["cash_continuity"] = [item for item in scoped.candidate["cash_continuity"]
                                                    if item["period"] in periods and item["previous_period"] in periods]
            for equation in scoped.candidate["equations"]:
                equation["periods"] = [period for period in equation["periods"] if period in periods]
            return scoped

        def complete(self, *, role, instructions, payload, schema, web=False):
            self.web_flags.append(web)
            if web:
                raise AssertionError("Local initiation must never request browser search")
            stage = payload["stage"]
            if stage == "discover_local":
                self.calls.append(copy.deepcopy(payload))
                self.last_response = self.synthetic_response(stage)
                if not payload.get("evidence"):
                    return {"search_query": "", "read_requests": [
                        {"source_id": source["id"], "unit_id": source["units"][0]["id"]}
                        for source in payload["catalog"] if source.get("category") == "financial_reports"],
                        "discovery": None}
                return {"search_query": "", "read_requests": [],
                        "discovery": copy.deepcopy(self.full_company.discovery)}
            if stage == "history_review":
                self.calls.append(copy.deepcopy(payload))
                self.last_response = self.synthetic_response(stage)
                result = {"search_query": "", "read_requests": [], "checks": [], "references": [], "issues": []}
                if not payload["evidence"]:
                    result["read_requests"] = [{"source_id": source["id"], "unit_id": source["units"][0]["id"]}
                                               for source in payload["catalog"]]
                else:
                    result["checks"] = [{"topic": topic, "passed": True, "reason": "Synthetic documented history checked."}
                                        for topic in payload["required_topics"]]
                    result["references"] = [{"source_id": unit["source_id"], "unit_id": unit["unit_id"]}
                                            for unit in payload["evidence"]]
                return result
            discovery = payload.get("discovery", self.full_company.discovery)
            periods = {item["id"] for item in discovery["periods"]}
            self.company = self.scope(self.full_company, periods)
            if stage == "extract":
                self.read_requested = bool(payload["evidence"])
            result = super().complete(role=role, instructions=instructions, payload=payload, schema=schema, web=False)
            if stage == "source_review":
                for fact in result["facts"]:
                    for observation in fact["observations"]:
                        observation["reporting_basis"] = self.company.candidate["reporting_basis"]
            if stage == "model_review" and result["checks"]:
                result["checks"] = [{"topic": topic, "passed": True,
                                     "reason": "Synthetic documentary structure and basis checked."}
                                    for topic in payload.get("required_topics", REVIEW_TOPICS)]
            return result

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.archive = self.root / "local_store" / "synthetic"
        reports = self.archive / "financial_reports"
        reports.mkdir(parents=True)
        self.original = reports / "synthetic_2025-fy_financial-statements_pub-2026-02-10.xlsx"
        self.prepare_company(SyntheticCompany())

    def prepare_company(self, company):
        from research_platform.evidence import inventory_local_sources

        self.company = company
        self.original.write_bytes(self.company.payload())
        source = inventory_local_sources(self.archive)["sources"][0]

        def remap(value):
            if isinstance(value, dict):
                for key, item in value.items():
                    if key == "source_id" and item == "annual":
                        value[key] = source["id"]
                    elif key == "source_ids":
                        value[key] = [source["id"] if source_id == "annual" else source_id for source_id in item]
                    else:
                        remap(item)
            elif isinstance(value, list):
                for item in value:
                    remap(item)

        remap(self.company.discovery)
        remap(self.company.candidate)
        self.company.source = {key: source[key] for key in ("id", "url", "title", "published_at", "reason")}
        self.company.source["kind"] = source["kind"]
        self.company.discovery["sources"] = [self.company.source]
        context = {"source_id": source["id"], "unit_id": f"{self.company.sheet}!R1"}
        candidate = self.company.candidate
        candidate["reporting_bases"] = [{"id": "original", "label": "Synthetic unchanged reporting basis",
            "accounting_standard": "Synthetic accounting", "consolidation_scope": "Consolidated group",
            "segment_structure": "Alpha and Beta operating segments", "measurement_basis": "reported",
            "vintage_date": "2026-02-10", "references": [context]}]
        candidate["basis_changes"] = []
        for row in self.company.rows:
            row["period_terms"] = []
            row["period_bases"] = [{"period": observation["period"], "basis_id": "original"}
                                   for observation in row["observations"]]
            for observation in row["observations"]:
                observation.update(actual_or_estimate="actual", scale=row["scale"], sign=row["sign"])
        for equation in candidate["equations"]:
            relevant = {equation["target"], *(term["row_id"] for term in equation["terms"])}
            common = set.intersection(*({observation["period"] for observation in row["observations"]}
                                        for row in self.company.rows if row["id"] in relevant))
            equation.update(periods=sorted(common), basis_bridge_refs=[])
        for continuity in candidate["cash_continuity"]:
            continuity["basis_treatment"] = "same_basis"

    def tearDown(self):
        self.temporary.cleanup()

    def run_local(self, *, company=None, provider=None, **kwargs):
        company = company or self.company
        provider = provider or self.LocalProvider(company)
        with patch("research_platform.ai.urlopen") as ai_network, \
                patch("research_platform.evidence.download") as browser:
            directory = initiate_company(self.root, company.name, provider=provider,
                                         source_directory=self.archive, **kwargs)
            ai_network.assert_not_called()
            browser.assert_not_called()
        return directory, read_json(directory / "record.json"), provider

    def test_single_year_uses_local_originals_full_documents_and_independent_checks(self):
        company = self.LocalProvider.scope(self.company, {"2025"})
        directory, record, provider = self.run_local(company=company, years=[2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertTrue((directory / "model.xlsx").is_file())
        self.assertEqual(record["request"]["input_mode"], "local")
        self.assertTrue(record["request"]["source_manifest_sha256"])
        self.assertTrue(all(check["passed"] for check in record["checks"]))
        self.assertFalse(any(provider.web_flags))
        self.assertNotIn("discover", [call["stage"] for call in provider.calls])
        self.assertIn("discover_local", [call["stage"] for call in provider.calls])
        self.assertEqual(next(row for row in record["rows"] if row["id"] == "revenue")["values"], {"2025": 1100.0})
        for stage in ("extract", "source_review", "coverage_inventory", "model_review"):
            evidence_calls = [call for call in provider.calls if call["stage"] == stage and call.get("evidence")]
            self.assertTrue(evidence_calls, stage)
            self.assertTrue(all(len(call["evidence"]) == len(self.company.rows) + 2 for call in evidence_calls), stage)
        self.assertEqual(record["usage"]["web_search_calls"], 0)

    def test_two_year_scopes_merge_and_preserve_cited_actuals(self):
        directory, record, provider = self.run_local(years=[2024, 2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(next(row for row in record["rows"] if row["id"] == "revenue")["values"],
                         {"2024": 1000.0, "2025": 1100.0})
        self.assertEqual({tuple(call["request"]["years"]) for call in provider.calls if call["stage"] == "extract"},
                         {(2024,), (2025,)})
        self.assertTrue(any(check["name"] == "cash_continuity:2025" and check["passed"] for check in record["checks"]))
        self.assertTrue((directory / "model.xlsx").exists())

    def test_default_ten_year_request_does_not_silently_accept_two_year_archive(self):
        _, record, provider = self.run_local()
        discovery = next(call for call in provider.calls if call["stage"] == "discover_local")
        self.assertEqual(discovery.get("request", discovery)["historical_years"], 10)
        self.assertEqual(record["request"]["historical_years"], 10)
        self.assertEqual(record["status"], "incomplete")
        self.assertIn("historical years", "; ".join(record["issues"]))
        self.assertFalse(any(call["stage"] == "extract" for call in provider.calls))

    def test_ten_year_local_history_accepts_with_reported_interims_and_separate_history_review(self):
        company = SyntheticCompany().include_interims()
        earlier = list(range(2016, 2024))
        company.column_periods = [str(year) for year in earlier] + company.column_periods
        earlier_cash_values = {"cash_close": 70, "balance_sheet_cash": 70, "cash_open": 70,
                               "cash_change": 0, "financing": -40, "net_debt": 230}
        for row in company.rows:
            current = company.display_values[row["id"]]
            company.display_values[row["id"]] = [earlier_cash_values.get(row["id"], current[0])] * len(earlier) + current
            reference = copy.deepcopy(row["observations"][0]["reference"])
            for observation in row["observations"]:
                observation["reference"]["number_index"] += len(earlier)
            row["observations"] = [{"period": str(year), "reference": dict(reference, number_index=index)}
                                   for index, year in enumerate(earlier)] + row["observations"]
        company.discovery["periods"] = [
            {"id": str(year), "year": year, "kind": "annual", "start": f"{year}-01-01", "end": f"{year}-12-31"}
            for year in earlier] + company.discovery["periods"]
        old_coverage = [item for item in company.discovery["period_coverage"] if item["year"] == 2024]
        company.discovery["period_coverage"] = [dict(copy.deepcopy(item), year=year)
            for year in earlier for item in old_coverage] + company.discovery["period_coverage"]
        old_sections = [item for item in company.candidate["coverage"] if item["period"] == "2024"]
        company.candidate["coverage"] = [dict(copy.deepcopy(item), period=str(year))
            for year in earlier for item in old_sections] + company.candidate["coverage"]
        self.prepare_company(company)
        directory, record, provider = self.run_local()
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(record["request"]["historical_years"], 10)
        annuals = [period["year"] for period in record["discovery"]["periods"] if period["kind"] == "annual"]
        self.assertEqual(annuals, list(range(2016, 2026)))
        self.assertEqual(len(record["historical_scopes"]), 10)
        revenue = next(row for row in record["rows"] if row["id"] == "revenue")
        self.assertEqual(len(revenue["values"]), 13)
        self.assertEqual(revenue["values"]["2025-9M"], 825.0)
        self.assertEqual(revenue["values"]["2025-Q1"], 220.0)
        self.assertNotIn("2025-Q3", revenue["values"])
        net_debt = next(row for row in record["rows"] if row["id"] == "net_debt")
        self.assertNotIn("2025-9M", net_debt["values"])
        self.assertIn("2025-H1", net_debt["values"])
        checks = {check["name"]: check for check in record["checks"]}
        self.assertTrue(all(checks[f"cash_continuity:{year}"]["passed"] for year in range(2017, 2026)))
        historical_review = [call for call in provider.calls if call["stage"] == "history_review"]
        self.assertEqual(len(historical_review), 2)
        self.assertEqual(historical_review[0]["evidence"], [])
        self.assertTrue(historical_review[1]["evidence"])
        self.assertTrue(all("reviews" not in call and "checks" not in call for call in historical_review))
        structural_calls = [call for call in provider.calls if call["stage"] == "model_review"]
        self.assertTrue(structural_calls)
        self.assertTrue(all(call["catalog"] and "sources" not in call for call in structural_calls))
        self.assertTrue(all(len(call["request"]["years"]) == 1 for call in provider.calls if call["stage"] == "extract"))
        self.assertFalse(any(provider.web_flags))
        self.assertTrue((directory / "model.xlsx").is_file())

    def test_estimated_observation_cannot_be_accepted_as_history(self):
        company = self.LocalProvider.scope(self.company, {"2025"})
        def insert_estimate(candidate):
            candidate["rows"][0]["observations"][0]["actual_or_estimate"] = "estimate"
        provider = self.LocalProvider(company, mutate=insert_estimate)
        directory, record, _ = self.run_local(company=company, provider=provider, years=[2025])
        self.assertEqual(record["status"], "incomplete")
        self.assertFalse((directory / "model.xlsx").exists())
        self.assertTrue(any("estimates cannot enter historical actuals" in "; ".join(call.get("feedback", []))
                            for call in provider.calls if call["stage"] == "extract"), record["issues"])

    def test_accepted_archive_reuses_without_calls_and_changed_bytes_block_resume(self):
        company = self.LocalProvider.scope(self.company, {"2025"})
        first, record, _ = self.run_local(company=company, years=[2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        reused = self.LocalProvider(company)
        second, _, _ = self.run_local(company=company, provider=reused, years=[2025])
        self.assertEqual(first, second)
        self.assertEqual(reused.calls, [])
        book = load_workbook(self.original)
        book.active["A1"] = "Changed source context"
        book.save(self.original)
        with self.assertRaisesRegex(ValueError, "changed|fresh"), \
                patch.object(reused, "complete", side_effect=AssertionError("No call may be dispatched")):
            self.run_local(company=company, provider=reused, years=[2025])

    def test_preflight_allowance_stops_before_charged_or_pending_request(self):
        class ExpensiveProvider(self.LocalProvider):
            def estimate_request_cost(self, **kwargs):
                return {"estimated_cost_usd": 21.0, "estimated_input_tokens": 10000000,
                        "max_output_tokens": 16000, "is_exact": False, "method": "synthetic allowance"}
        provider = ExpensiveProvider(self.company)
        directory, record, _ = self.run_local(provider=provider, years=[2024, 2025])
        self.assertEqual(record["status"], "incomplete")
        self.assertEqual(provider.calls, [])
        self.assertIsNone(record["pending_call"])
        self.assertEqual(record["usage"]["calls"], 0)
        self.assertEqual(record["usage"]["estimated_cost_usd"], 0)
        self.assertFalse((directory / "model.xlsx").exists())

    def test_identical_archive_at_another_location_cannot_reset_existing_budget(self):
        import shutil

        company = self.LocalProvider.scope(self.company, {"2025"})
        first, record, _ = self.run_local(company=company, years=[2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        moved = self.root / "local_store" / "same_documents_elsewhere"
        shutil.copytree(self.archive, moved)
        provider = self.LocalProvider(company)
        with self.assertRaisesRegex(ValueError, "Source archive or intake mode differs"), \
                patch.object(provider, "complete", side_effect=AssertionError("No AI request may be sent")):
            initiate_company(self.root, company.name, provider=provider,
                             source_directory=moved, years=[2025])
        self.assertEqual(provider.calls, [])
        self.assertEqual(list((self.root / "outputs").glob("initiate_*")), [first])
        self.assertEqual(read_json(first / "record.json"), record)

    def test_switching_saved_injected_research_to_local_requires_explicit_fresh(self):
        company = SyntheticCompany()
        legacy_provider = ScriptedProvider(company)
        first = initiate_company(self.root, company.name, provider=legacy_provider,
                                 fetcher=company.fetch, years=[2024, 2025])
        record = read_json(first / "record.json")
        self.assertEqual(record["status"], "accepted", record["issues"])
        provider = self.LocalProvider(self.company)
        with self.assertRaisesRegex(ValueError, "Source archive or intake mode differs"), \
                patch.object(provider, "complete", side_effect=AssertionError("No AI request may be sent")):
            self.run_local(provider=provider, years=[2024, 2025])
        self.assertEqual(provider.calls, [])
        self.assertEqual(list((self.root / "outputs").glob("initiate_*")), [first])
        self.assertEqual(read_json(first / "record.json"), record)

    def test_changing_saved_numeric_reference_is_real_progress_with_same_row_identity(self):
        class EditingProvider(self.LocalProvider):
            saved_revisions = 0
            def complete(self, **kwargs):
                payload = kwargs["payload"]
                if payload["stage"] == "extract" and payload["evidence"] and self.saved_revisions < 2:
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("extract")
                    scoped = self.scope(self.full_company, {item["id"] for item in payload["discovery"]["periods"]})
                    row = copy.deepcopy(next(row for row in scoped.rows if row["id"] == "revenue"))
                    if self.saved_revisions == 0:
                        row["observations"][0]["reference"]["number_index"] = 0
                    self.saved_revisions += 1
                    return {"action": "save_rows", "read_requests": [], "sources": [], "rows": [row],
                            "delete_row_ids": [], "search_query": "", "candidate": None}
                return super().complete(**kwargs)
        company = self.LocalProvider.scope(self.company, {"2025"})
        provider = EditingProvider(company)
        _, record, _ = self.run_local(company=company, provider=provider, years=[2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertEqual(provider.saved_revisions, 2)
        revenue = next(row for row in record["rows"] if row["id"] == "revenue")
        self.assertEqual(revenue["values"], {"2025": 1100.0})

    def test_unreadable_primary_archive_stops_before_ai_and_preserves_failed_original(self):
        invalid_xlsx = b"PK\x03\x04SYNTHETIC INVALID XLSX ORIGINAL"
        self.original.write_bytes(invalid_xlsx)
        directory, record, provider = self.run_local(years=[2024, 2025])
        self.assertEqual(record["status"], "incomplete")
        self.assertIn("No readable original financial_reports", "; ".join(record["issues"]))
        self.assertEqual(provider.calls, [])
        self.assertIsNone(record["pending_call"])
        self.assertEqual(record["usage"]["calls"], 0)
        self.assertEqual(record["usage"]["estimated_cost_usd"], 0)
        self.assertTrue(record["archive_errors"])
        self.assertEqual(record["sources"][0]["status"], "unreadable")
        self.assertEqual(Path(record["sources"][0]["path"]).read_bytes(), invalid_xlsx)
        self.assertFalse((directory / "model.xlsx").exists())


    def test_interrupted_local_extraction_reuses_saved_discovery_and_charges_once(self):
        company = self.LocalProvider.scope(self.company, {"2025"})
        original_provider = self.LocalProvider(company)
        def stop_after_saved_read(message):
            if message.startswith("Stage extract:"):
                raise KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):
            self.run_local(company=company, provider=original_provider, years=[2025], progress=stop_after_saved_read)
        original_directory = next((self.root / "outputs").glob("initiate_*"))
        saved = read_json(original_directory / "record.json")
        self.assertEqual(saved["status"], "incomplete")
        self.assertIsNone(saved["pending_call"])
        continuation = self.LocalProvider(company)
        directory, record, _ = self.run_local(company=company, provider=continuation, years=[2025])
        self.assertEqual(record["status"], "accepted", record["issues"])
        self.assertNotEqual(directory, original_directory)
        self.assertEqual(record["parent_run"], str(original_directory))
        self.assertFalse(any(call["stage"] == "discover_local" for call in continuation.calls))
        self.assertEqual(continuation.calls[0]["stage"], "extract")
        self.assertTrue(continuation.calls[0]["evidence"])
        self.assertEqual(record["usage"]["calls"], len(original_provider.calls) + len(continuation.calls))
        self.assertAlmostEqual(record["usage"]["estimated_cost_usd"], record["usage"]["calls"] * 0.003)
        self.assertEqual(read_json(original_directory / "record.json"), saved)


class HistoricalBasisTests(unittest.TestCase):
    """Versioned financial history must not quietly splice unlike observations."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.company = SyntheticCompany()
        self.store = SourceStore(Path(self.temporary.name), fetcher=self.company.fetch)
        self.store.add(self.company.source)
        self.store.sources[0]["category"] = "financial_reports"
        self.ref = {"source_id": "annual", "unit_id": "Statements!R1"}
        self.read_units = {(item["source_id"], item["unit_id"])
                           for item in self.store.read_documents([self.ref])}
        self.candidate = copy.deepcopy(self.company.candidate)
        self.candidate["reporting_bases"] = [{
            "id": "reported", "label": "Synthetic consolidated reporting framework",
            "accounting_standard": "Synthetic accounting policies", "consolidation_scope": "Synthetic group",
            "segment_structure": "Synthetic Alpha and Beta business segmentation", "measurement_basis": "reported",
            "vintage_date": "2026-02-10", "references": [copy.deepcopy(self.ref)]}]
        self.candidate["basis_changes"] = []
        for row in self.candidate["rows"]:
            row["period_bases"] = [{"period": str(year), "basis_id": "reported"} for year in (2024, 2025)]
            row["period_terms"] = []
            for obs in row["observations"]:
                obs.update(actual_or_estimate="actual", scale=row["scale"], sign=row["sign"])
        for equation in self.candidate["equations"]:
            equation.update(periods=[], basis_bridge_refs=[])
        for bridge in self.candidate["cash_continuity"]:
            bridge["basis_treatment"] = "same_basis"

    def materialize(self, candidate=None):
        return materialize_and_check(candidate or self.candidate, self.company.discovery,
                                     self.store, self.read_units)

    def second_basis(self):
        basis = copy.deepcopy(self.candidate["reporting_bases"][0])
        basis.update(id="revised", label="Synthetic revised framework", vintage_date="2026-03-01")
        self.candidate["reporting_bases"].append(basis)

    def test_local_history_requires_actuals_complete_basis_assignments_and_primary_context(self):
        _, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks))
        for mutation in ("estimate", "missing_basis", "unknown_basis", "vendor_only_basis"):
            candidate = copy.deepcopy(self.candidate)
            if mutation == "estimate":
                candidate["rows"][0]["observations"][0]["actual_or_estimate"] = "estimate"
            elif mutation == "missing_basis":
                candidate["rows"][0]["period_bases"].pop()
            elif mutation == "unknown_basis":
                candidate["rows"][0]["period_bases"][0]["basis_id"] = "invented"
            else:
                self.store.sources[0]["category"] = "visible_alpha"
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                self.materialize(candidate)
            self.store.sources[0]["category"] = "financial_reports"

    def test_basis_transition_requires_ledger_even_when_amounts_are_unchanged(self):
        self.second_basis()
        shares = next(row for row in self.candidate["rows"] if row["id"] == "shares")
        shares["period_bases"][1]["basis_id"] = "revised"
        with self.assertRaisesRegex(ValueError, "change ledger"):
            self.materialize()
        self.candidate["basis_changes"] = [{
            "from_basis": "reported", "to_basis": "revised", "effective_date": "2025-01-01",
            "kind": "definition", "treatment": "explicit_break", "affected_rows": ["shares"],
            "explanation": "Synthetic share definition changed; equal amounts do not prove comparability.",
            "references": [copy.deepcopy(self.ref)]}]
        rows, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks))
        self.assertEqual(next(row for row in rows if row["id"] == "shares")["values"], {"2024": 100, "2025": 100})

    def test_mixed_basis_equation_needs_explicit_cited_bridge_and_core_identity_cannot_be_waived(self):
        self.second_basis()
        expenses = next(row for row in self.candidate["rows"] if row["id"] == "expenses")
        for assignment in expenses["period_bases"]:
            assignment["basis_id"] = "revised"
        with self.assertRaisesRegex(ValueError, "incompatible reporting bases"):
            self.materialize()
        equation = next(item for item in self.candidate["equations"] if item["target"] == "net_income")
        equation["basis_bridge_refs"] = [copy.deepcopy(self.ref)]
        _, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks))
        assets = next(row for row in self.candidate["rows"] if row["id"] == "assets")
        for assignment in assets["period_bases"]:
            assignment["basis_id"] = "revised"
        with self.assertRaisesRegex(ValueError, "incompatible reporting bases"):
            self.materialize()

    def test_segment_reconciliation_uses_each_periods_actual_segment_rows(self):
        for old_id, new_id in (("segment_a", "segment_a_v2"), ("profit_a", "profit_a_v2")):
            old = next(row for row in self.candidate["rows"] if row["id"] == old_id)
            new = copy.deepcopy(old)
            # The disclosed segment label may stay identical after a reorganization;
            # the row ID and dated framework, not an invented entity name, distinguish it.
            new.update(id=new_id)
            old["observations"] = [old["observations"][0]]
            old["period_bases"] = [old["period_bases"][0]]
            old["unavailable"] = [{"period": "2025", "explanation": "Synthetic later version separately represented.", "references": [copy.deepcopy(self.ref)]}]
            new["observations"] = [new["observations"][1]]
            new["period_bases"] = [new["period_bases"][1]]
            new["unavailable"] = [{"period": "2024", "explanation": "Synthetic later version not reported in earlier period.", "references": [copy.deepcopy(self.ref)]}]
            self.candidate["rows"].append(new)
        originals = [item for item in self.candidate["equations"] if item["target"] in ("revenue", "segment_group_profit")]
        for original in originals:
            original["periods"] = ["2024"]
            later = copy.deepcopy(original)
            later["periods"] = ["2025"]
            for term in later["terms"]:
                if term["row_id"] in ("segment_a", "profit_a"):
                    term["row_id"] += "_v2"
            self.candidate["equations"].append(later)
        _, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks), [check for check in checks if not check["passed"]])
        originals[0]["periods"] = []
        with self.assertRaises(ValueError):
            self.materialize()

    def test_period_specific_calculations_export_formulas_and_ledger(self):
        row = next(row for row in self.candidate["rows"] if row["id"] == "net_debt")
        self.candidate["equations"] = [item for item in self.candidate["equations"] if item["target"] != "net_debt"]
        row["observations"] = []
        ordinary = [{"row_id": "debt", "coefficient": 1}, {"row_id": "cash_close", "coefficient": -1}]
        row["period_terms"] = [{"period": "2024", "terms": ordinary},
                               {"period": "2025", "terms": ordinary + [{"row_id": "minority_interest", "coefficient": 1}]}]
        rows, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks))
        calculated = next(item for item in rows if item["id"] == "net_debt")
        self.assertEqual(calculated["kind"], "calculated")
        self.assertEqual(calculated["values"], {"2024": 200, "2025": 245})
        path = Path(self.temporary.name) / "model.xlsx"
        export_workbook(path, {"request": {"forecast_years": 5}, "discovery": self.company.discovery,
                               "candidate": self.candidate, "rows": rows, "checks": checks,
                               "status": "accepted", "reviews": [], "sources": self.store.sources})
        book = load_workbook(path, data_only=False)
        self.addCleanup(book.close)
        self.assertIn("Reporting basis", book.sheetnames)
        self.assertEqual(book["net_debt"]["D5"].data_type, "f")
        self.assertEqual(book["net_debt"]["E5"].data_type, "f")
        self.assertNotIn("'equity'!", book["net_debt"]["D5"].value)
        self.assertIn("'equity'!", book["net_debt"]["E5"].value)
        self.assertIn("reporting basis=reported", book["net_debt"]["E5"].comment.text)

    def test_source_scale_is_observation_specific_and_display_scale_is_separate(self):
        row = next(row for row in self.candidate["rows"] if row["id"] == "shares")
        row["observations"][1]["scale"] = 1000
        rows, _ = self.materialize()
        result = next(item for item in rows if item["id"] == "shares")
        self.assertEqual(result["values"], {"2024": 100, "2025": 100000})
        self.assertEqual(result["scale"], 1)

    def test_review_scope_omits_unavailable_subtotal_equations_without_broadening_periods(self):
        from research_platform.initiation import _scope_model

        subtotal = copy.deepcopy(next(row for row in self.candidate["rows"] if row["id"] == "expenses"))
        subtotal.update(id="old_subtotal", label="Earlier disclosed expense subtotal",
                        definition="Synthetic reported supporting subtotal available only in 2024.")
        subtotal["observations"] = [subtotal["observations"][0]]
        subtotal["period_bases"] = [subtotal["period_bases"][0]]
        subtotal["unavailable"] = [{"period": "2025", "explanation": "The later report does not disclose this supporting subtotal.",
                                    "references": [copy.deepcopy(self.ref)]}]
        self.candidate["rows"].append(subtotal)
        equation = {"target": "old_subtotal", "terms": [{"row_id": "expenses", "coefficient": 1}],
                    "explanation": "Synthetic supporting composition", "periods": [], "basis_bridge_refs": []}
        self.candidate["equations"].append(equation)
        rows, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks))
        check = next(check for check in checks if check["name"] == "equation:old_subtotal:2025")
        self.assertEqual(check["basis"], "not_applicable")

        later, later_rows = _scope_model(self.candidate, rows, {"2025"})
        self.assertNotIn("old_subtotal", {row["id"] for row in later_rows})
        self.assertNotIn("old_subtotal", {item["target"] for item in later["equations"]})
        self.assertTrue(all(item["periods"] == ["2025"] for item in later["equations"]))
        both, _ = _scope_model(self.candidate, rows, {"2024", "2025"})
        scoped = next(item for item in both["equations"] if item["target"] == "old_subtotal")
        self.assertEqual(scoped["periods"], ["2024"])
        self.assertEqual(scoped["terms"], equation["terms"])
        self.assertEqual(equation["periods"], [])

    def test_cross_basis_cash_needs_documented_no_effect_not_equality_alone(self):
        self.second_basis()
        # The immutable synthetic source explicitly documents the lack of cash
        # impact; a reviewer must still check this claim in a real initiation.
        content = (b"<h1>SYNTHETIC accounting policy transition</h1>"
                   b"<p>Effective 1 January 2025, the synthetic reporting policy changes. "
                   b"The change has no effect on opening cash or the preceding closing cash balance. "
                   b"No cash adjustment is required.</p><p>A separate synthetic control balance is 101</p>")
        self.store.fetcher = lambda url: (content, url, "text/html")
        self.store.add({"id": "policy", "url": "https://example.invalid/policy.html", "title": "Synthetic policy note",
                        "published_at": "2026-02-10", "reason": "Explicit cash treatment for a policy transition", "kind": "filing"})
        self.store.sources[-1]["category"] = "financial_reports"
        catalog = next(item for item in self.store.catalog() if item["id"] == "policy")
        ref = {"source_id": "policy", "unit_id": catalog["units"][0]["id"]}
        self.read_units.update((item["source_id"], item["unit_id"]) for item in self.store.read_documents([ref]))
        for row in self.candidate["rows"]:
            row["period_bases"][1]["basis_id"] = "revised"
        self.candidate["basis_changes"] = [{
            "from_basis": "reported", "to_basis": "revised", "effective_date": "2025-01-01",
            "kind": "accounting_policy", "treatment": "explicit_break",
            "affected_rows": [row["id"] for row in self.candidate["rows"]],
            "explanation": "The synthetic policy changes, with explicitly disclosed no effect on cash.", "references": [ref]}]
        bridge = self.candidate["cash_continuity"][0]
        bridge.update(explanation="The original policy note explicitly states that opening and preceding closing cash are unchanged.",
                      references=[ref])
        with self.assertRaisesRegex(ValueError, "explicit treatment"):
            self.materialize()
        bridge["basis_treatment"] = "reported_adjustment"
        with self.assertRaisesRegex(ValueError, "reported adjustment terms"):
            self.materialize()
        bridge["basis_treatment"] = "documented_no_effect"
        _, checks = self.materialize()
        self.assertTrue(all(check["passed"] for check in checks))
        wrong = copy.deepcopy(self.candidate)
        control_unit = next(unit for unit in self.store.read_documents([ref])
                            if any(str(token["raw"]).rstrip(".") == "101" for token in unit["numbers"]))
        number_index = next(index for index, token in enumerate(control_unit["numbers"])
                            if str(token["raw"]).rstrip(".") == "101")
        opening = next(row for row in wrong["rows"] if row["id"] == "cash_open")
        opening["observations"][1]["reference"] = {"source_id": "policy", "unit_id": control_unit["unit_id"],
                                                    "number_index": number_index, "number_format": "dot_decimal"}
        _, wrong_checks = self.materialize(wrong)
        no_effect = next(check for check in wrong_checks if check["name"] == "cash_continuity:2025")
        self.assertFalse(no_effect["passed"])
        self.assertEqual(no_effect["tolerance"], 0)
        self.store.sources[-1]["category"] = "transcripts"
        with self.assertRaisesRegex(ValueError, "financial-report context"):
            self.materialize()


class VisibleAlphaTests(unittest.TestCase):
    """Vendor layout and column evidence constrain AI claims about historical cells."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.archive = self.root / "company"
        (self.archive / "visible_alpha").mkdir(parents=True)
        self.path = self.archive / "visible_alpha" / "synthetic_vendor.xlsx"
        self.addresses = self.make_workbook(self.path)
        from research_platform.evidence import inventory_local_sources
        self.spec = inventory_local_sources(self.archive)["sources"][0]
        self.store = SourceStore(self.root / "originals")
        self.descriptor = self.store.add_local(self.spec, root=self.archive)

    @staticmethod
    def make_workbook(path, *, recognized=True):
        """A vendor-format fixture with deliberately interleaved actuals/estimates."""
        book = Workbook()
        sheet = book.active
        sheet.title = "Vendor data"
        sheet["C2"] = "SYNTHETIC Manufacturing Group"
        periods = ["FY-2024", "FY-2026E", "FY-2024", "FY-2025", "FY-2025E",
                   "1HFY-2024", "1QFY-2024"]
        series = ["VA Actuals", "Consensus", "Filings", "Filings", "VA Actuals",
                  "Filings", "Filings"]

        def headers(first):
            for column, (period, source) in enumerate(zip(periods, series), 7):
                sheet.cell(first, column, period)
                sheet.cell(first + 1, column, source)
            sheet.cell(first + 1, 5, "Units")

        def metric(row, label, values):
            sheet.cell(row, 4, label)
            sheet.cell(row, 5, "USDM")
            sheet.cell(row, 6, "Analytics")
            for column, value in enumerate(values, 7):
                if value is not None:
                    sheet.cell(row, column, value)

        headers(4)
        sheet["D3"] = "Standardized" if recognized else "Unknown tables"
        sheet["D6"] = "Income statement"
        metric(7, "Standard revenue", [777777, 900001, 777777, 888888, 900002, 444444, 222222])
        headers(10)
        sheet["D12"] = "Ratios" if recognized else "Other analytics"
        metric(13, "Standard margin", [123456, 900003, 123456, 234567, 900004, 345678, 456789])
        sheet["D15"] = "Company Reported" if recognized else "Unclassified section"
        headers(16)
        sheet["D18"] = "Company business divisions"
        metric(19, "Manufacturing sales", [100, 900005, 100, 110, 900006, 45, 20])
        sheet["F20"] = "YoY Growth"
        for column in range(7, 14):
            sheet.cell(20, column, 555555)
        metric(21, "Discontinued division: no reported observation", [None] * 7)
        sheet["D23"] = "Reporting note: legacy division definitions remain distinct."
        headers(25)
        sheet["D27"] = "Income statement"
        metric(28, "Group sales", [100, 900007, 100, 110, 900008, 45, 20])
        key = book.create_sheet("Key")
        key["C3"] = "Legend"
        key["D4"] = "Direct Inferred means additional non-analyst source information was used."
        key["C6"] = "Source Information"
        key["D6"] = "Historical data is provided by VA Actuals"
        book.save(path)
        book.close()
        return {"standard": "I7", "ratio": "I13", "actual": "G19", "filing": "I19",
                "future": "H19", "conflicting_estimate": "K19", "later_actual": "J19",
                "half_year": "L19", "quarter": "M19", "helper": "I20", "group": "I28"}

    def reference(self, name, *, store=None, descriptor=None):
        store = store or self.store
        descriptor = descriptor or self.descriptor
        address = self.addresses.get(name, name)
        row = int(re.search(r"\d+$", address).group())
        ref = {"source_id": descriptor["id"], "unit_id": f"Vendor data!R{row}"}
        unit = store.read([ref])[0]
        token = next(token for token in unit["numbers"] if token["locator"] == address)
        return dict(ref, number_index=token["index"], number_format="dot_decimal")

    @staticmethod
    def period(year=2024, slot=None):
        value = {"id": str(year) if slot is None else f"{year}-{slot}", "year": year,
                 "kind": "annual" if slot is None else "half_year" if slot == "H1" else "quarter",
                 "start": f"{year}-01-01", "end": f"{year}-12-31" if slot is None else
                 f"{year}-06-30" if slot == "H1" else f"{year}-03-31"}
        if slot is not None:
            value.update(fiscal_year_start=f"{year}-01-01", fiscal_year_end=f"{year}-12-31")
        return value

    def test_recognized_layout_preserves_original_bytes_and_reconstructs_on_restore(self):
        original = self.path.read_bytes()
        self.assertEqual(self.descriptor["vendor_layout"]["status"], "recognized")
        self.assertEqual(Path(self.descriptor["path"]).read_bytes(), original)
        self.assertEqual(self.descriptor["sha256"], hashlib.sha256(original).hexdigest())
        reference = self.reference("filing")
        before = self.store.read_documents([reference], historical_years=[2024])
        restored = SourceStore(self.root / "unused")
        restored.restore(self.descriptor, archive_root=self.root)
        self.assertEqual(restored.read_documents([reference], historical_years=[2024]), before)
        self.assertEqual(restored.resolve(reference), 100)
        self.assertEqual(self.path.read_bytes(), original)

    def test_historical_view_retains_original_numeric_indexes_and_full_read_is_unchanged(self):
        reference = self.reference("filing")
        full = self.store.read_documents([reference])
        unit = next(item for item in full if item["unit_id"] == reference["unit_id"])
        self.assertEqual(reference["number_index"], 2)
        self.assertEqual(unit["numbers"][1]["raw"], 900005)
        view = self.store.read_documents([reference], historical_years=[2024])
        historical = next(item for item in view if item["unit_id"] == reference["unit_id"])
        token = next(token for token in historical["numbers"] if token["locator"] == "I19")
        self.assertEqual(token["index"], reference["number_index"])
        self.assertEqual(token["raw"], 100)
        self.assertNotIn("H19", {token["locator"] for token in historical["numbers"]})
        self.assertEqual(self.store.resolve(reference), token["raw"])
        historical["text"] = "MUTATED PROMPT COPY"
        token["raw"] = -123
        self.assertEqual(self.store.read_documents([reference]), full)

    def test_historical_view_keeps_notes_legends_and_empty_company_rows_without_estimate_values(self):
        view = self.store.read_documents([self.reference("filing")], historical_years=[2024])
        text = "\n".join(item["text"] for item in view)
        self.assertIn("legacy division definitions remain distinct", text)
        self.assertIn("Discontinued division: no reported observation", text)
        self.assertIn("Direct Inferred", text)
        self.assertIn("Historical data is provided by VA Actuals", text)
        actuals = [token["raw"] for unit in view for token in unit["numbers"]]
        for excluded in (777777, 123456, 555555, 900005, 900006, 900007, 900008, 110):
            self.assertNotIn(excluded, actuals)
            self.assertNotIn(str(excluded), text)
        self.assertIn(100, actuals)
        self.assertIn(45, actuals)
        self.assertIn(20, actuals)

    def test_eligible_company_actuals_and_filings_have_exact_numeric_support(self):
        for name in ("actual", "filing", "group"):
            with self.subTest(name=name):
                reference = self.reference(name)
                self.assertIsInstance(self.store.validate_historical_reference(reference, self.period()), dict)
                self.assertEqual(self.store.resolve(reference), 100)
        self.store.validate_historical_reference(self.reference("half_year"), self.period(slot="H1"))
        self.store.validate_historical_reference(self.reference("quarter"), self.period(slot="Q1"))

    def test_standardized_ratios_consensus_conflicting_estimates_and_helpers_are_ineligible(self):
        for name in ("standard", "ratio", "future", "conflicting_estimate", "helper"):
            with self.subTest(name=name):
                # A manually supplied hidden citation must not bypass the view.
                with self.assertRaises(ValueError):
                    self.store.validate_historical_reference(self.reference(name), self.period())

    def test_year_and_quarter_half_year_confusions_are_rejected(self):
        cases = (("later_actual", self.period()), ("filing", self.period(2025)),
                 ("half_year", self.period(slot="Q1")), ("quarter", self.period(slot="H1")),
                 ("half_year", self.period()), ("filing", self.period(slot="H1")))
        for name, period in cases:
            with self.subTest(name=name, period=period["id"]):
                with self.assertRaises(ValueError):
                    self.store.validate_historical_reference(self.reference(name), period)

    def test_unrecognized_vendor_layout_cannot_supply_historical_cells(self):
        path = self.archive / "visible_alpha" / "unrecognized.xlsx"
        self.make_workbook(path, recognized=False)
        from research_platform.evidence import inventory_local_sources
        spec = next(item for item in inventory_local_sources(self.archive)["sources"]
                    if item["local_path"].endswith("unrecognized.xlsx"))
        descriptor = self.store.add_local(spec, root=self.archive)
        self.assertEqual(descriptor["vendor_layout"]["status"], "unsupported")
        with self.assertRaises(ValueError):
            self.store.validate_historical_reference(self.reference("filing", descriptor=descriptor), self.period())
        self.assertEqual(Path(descriptor["path"]).read_bytes(), path.read_bytes())

    def test_unclassified_forecast_sheet_is_preserved_but_excluded_from_historical_prompt(self):
        book = load_workbook(self.path)
        book.create_sheet("Forecasts").append(["Consensus FY2027", 900009])
        book.save(self.path)
        book.close()
        from research_platform.evidence import inventory_local_sources
        spec = inventory_local_sources(self.archive)["sources"][0]
        descriptor = self.store.add_local(spec, root=self.archive)
        reference = self.reference("filing", descriptor=descriptor)
        full = self.store.read_documents([reference])
        self.assertIn("Forecasts!R1", {unit["unit_id"] for unit in full})
        view = self.store.read_documents([reference], historical_years=[2024])
        self.assertNotIn("Forecasts!R1", {unit["unit_id"] for unit in view})
        self.assertNotIn("900009", json.dumps(view))
        forecast = {"source_id": descriptor["id"], "unit_id": "Forecasts!R1",
                    "number_index": 0, "number_format": "dot_decimal"}
        self.assertEqual(self.store.resolve(forecast), 900009)
        with self.assertRaises(ValueError):
            self.store.validate_historical_reference(forecast, self.period())
        self.assertEqual(Path(descriptor["path"]).read_bytes(), self.path.read_bytes())

    def test_vendor_dash_is_not_zero_and_nonnumeric_cell_explanations_remain_visible(self):
        book = load_workbook(self.path)
        book["Vendor data"]["I19"] = "-"
        book["Vendor data"]["G19"] = "Not disclosed: refer to reporting note"
        book.save(self.path)
        book.close()
        from research_platform.evidence import inventory_local_sources
        spec = inventory_local_sources(self.archive)["sources"][0]
        descriptor = self.store.add_local(spec, root=self.archive)
        reference = self.reference("filing", descriptor=descriptor)
        # The generic reader retains its original semantics. Vendor initiation
        # must separately reject ambiguous missing markers before calculation.
        self.assertEqual(self.store.resolve(reference), 0)
        with self.assertRaises(ValueError):
            self.store.validate_historical_reference(reference, self.period())
        view = self.store.read_documents([reference], historical_years=[2024])
        row = next(unit for unit in view if unit["unit_id"] == reference["unit_id"])
        self.assertNotIn("I19", {token["locator"] for token in row["numbers"]})
        self.assertNotIn("G19", {token["locator"] for token in row["numbers"]})
        self.assertIn("Not disclosed: refer to reporting note", row["text"])
        self.assertIn("-", row["text"])
        self.assertEqual(Path(descriptor["path"]).read_bytes(), self.path.read_bytes())

    def test_interim_dates_use_discovered_fiscal_year_not_an_assumed_december_year_end(self):
        period = {"id": "2024-H1", "year": 2024, "kind": "half_year",
                  "start": "2023-07-01", "end": "2023-12-31",
                  "fiscal_year_start": "2023-07-01", "fiscal_year_end": "2024-06-30"}
        reference = self.reference("half_year")
        self.assertIsInstance(self.store.validate_historical_reference(reference, period), dict)
        wrong_half = dict(period, start="2024-01-01", end="2024-06-30")
        with self.assertRaises(ValueError):
            self.store.validate_historical_reference(reference, wrong_half)

    def test_materialization_cannot_accept_a_vendor_forecast_by_calling_it_an_actual(self):
        fixture = HistoricalBasisTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        vendor = fixture.store.add_local(self.spec, root=self.archive)
        reference = self.reference("future", store=fixture.store, descriptor=vendor)
        fixture.read_units.update((unit["source_id"], unit["unit_id"])
                                  for unit in fixture.store.read_documents([reference]))
        shares = next(row for row in fixture.candidate["rows"] if row["id"] == "shares")
        observation = next(item for item in shares["observations"] if item["period"] == "2024")
        observation["reference"] = reference
        observation["actual_or_estimate"] = "actual"
        with self.assertRaises(ValueError):
            fixture.materialize()

    def test_workflow_preloads_vendor_history_but_keeps_blind_reviews_on_primary_sources(self):
        fixture = LocalInitiationTests()
        fixture.setUp()
        self.addCleanup(fixture.tearDown)
        vendor_folder = fixture.archive / "visible_alpha"
        vendor_folder.mkdir()
        self.make_workbook(vendor_folder / "vendor.xlsx")
        company = fixture.LocalProvider.scope(fixture.company, {"2024"})
        primary_id = company.source["id"]

        class PreloadedProvider(fixture.LocalProvider):
            def complete(self, **kwargs):
                payload = kwargs["payload"]
                if payload["stage"] == "extract" and not any(
                        unit["source_id"] == primary_id for unit in payload["evidence"]):
                    self.calls.append(copy.deepcopy(payload))
                    self.last_response = self.synthetic_response("extract")
                    return {"action": "read", "read_requests": [{"source_id": primary_id,
                            "unit_id": f"{self.company.sheet}!R1"}], "sources": [],
                            "rows": [], "delete_row_ids": [], "search_query": "", "candidate": None}
                return super().complete(**kwargs)

        provider = PreloadedProvider(company)
        _, record, _ = fixture.run_local(company=company, provider=provider, years=[2024])
        self.assertEqual(record["status"], "accepted", record["issues"])
        first = next(call for call in provider.calls if call["stage"] == "extract")
        self.assertTrue(first["evidence"])
        self.assertTrue(any("vendor_view" in unit for unit in first["evidence"]))
        self.assertFalse(any(unit["source_id"] == primary_id for unit in first["evidence"]))
        vendor_id = next(source["id"] for source in record["sources"] if source["category"] == "visible_alpha")
        for call in provider.calls:
            if call["stage"] in ("source_review", "coverage_inventory"):
                self.assertNotIn(vendor_id, {item["id"] for item in call["catalog"]})
                self.assertNotIn(vendor_id, {item["source_id"] for item in call["evidence"]})

    def test_primary_company_spreadsheet_does_not_need_vendor_classification(self):
        reports = self.archive / "financial_reports"
        reports.mkdir()
        path = reports / "original.xlsx"
        book = Workbook()
        book.active.append(["SYNTHETIC original company sales", 2024, 100])
        book.save(path)
        book.close()
        from research_platform.evidence import inventory_local_sources
        spec = next(item for item in inventory_local_sources(self.archive)["sources"]
                    if item["category"] == "financial_reports")
        descriptor = self.store.add_local(spec, root=self.archive)
        reference = {"source_id": descriptor["id"], "unit_id": "Sheet!R1",
                     "number_index": 1, "number_format": "dot_decimal"}
        self.assertIsNone(self.store.validate_historical_reference(reference, self.period()))
        self.assertEqual(self.store.read_documents([reference], historical_years=[2024]),
                         self.store.read_documents([reference]))


if __name__ == "__main__":
    unittest.main()
