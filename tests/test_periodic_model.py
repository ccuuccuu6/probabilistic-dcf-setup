"""A different company's fiscal calendar, reporting change and forecast economics."""
import copy
from contextlib import closing
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tests.test_integrated_model import second_company
from tests.test_fx import source_workbook
from research_platform.fx import load_fx_source
from research_platform.integrated import build_periodic_spec,segment_research_context
from research_platform.periods import build_periods,calculate_periodic,validate_periodic


def different_period_company():
    data=second_company();data['company']='SYNTHETIC Harbour Components plc';data['currency']='GBP'
    data['forecast_basis']='SYNTHETIC current_components_basis';data['fiscal_calendar']={'fiscal_year_end':'06-30'}
    data['model_configuration']=dict(forecast_frequency='half_year',full_statement_frequency='quarterly',
        actual_through='2025-06-30',current_reporting_basis='current_components_basis',fx_currencies=['USD'],
        row_overrides={'cash_interest':dict(expression='lag("cash",1) * cash_yield',
            note='Company policy: short-term investments earn no interest; only opening bank cash earns the annual yield.')})
    data['schedule_anchor_policy']={'assumptions':[]}
    data['model_configuration']['as_of']='2026-10-23'
    data['source_freshness']=dict(checked_as_of='2026-10-23',official_indexes=[{'url':'https://synthetic.invalid/releases','evidence':'Explicit synthetic fixture; no real-company discovery claimed'}],releases=[])
    for role in data['historical'].values():
        role['basis']='synthetic_group_basis'
        for refs in role['sources'].values():
            for ref in refs:ref['unit']='GBPm'
    periods=build_periods([2024,2025,2026,2027],data['fiscal_calendar']);lookup={(p['fiscal_year'],p['kind']):p for p in periods}
    data['active_basis_by_period']={p['id']:dict(group='synthetic_group_basis',
        segments='legacy_combined_basis' if p['fiscal_year']==2024 else 'current_components_basis') for p in periods}
    data['entity_versions']=[dict(id='old_north',model_entity_id='north',label='Legacy combined components',
        period_ids=['2024-FY','2025-FY'],bases=['legacy_combined_basis'],scope='both_current_businesses',current_basis_compatible=False),
        dict(id='new_north',model_entity_id='north',label='Northern components',period_ids=['2025-FY'],
             bases=['current_components_basis'],scope='north_only',current_basis_compatible=True),
        dict(id='new_specialty',model_entity_id='specialty',label='Specialty components',period_ids=['2025-FY'],
             bases=['current_components_basis'],scope='specialty_only',current_basis_compatible=True)]
    observations=[]
    def put(row,year,kind,value,metric_type='flow',entity='group',basis='synthetic_group_basis',scope='whole_company',unit='GBPm',**extras):
        period=lookup[year,kind];identity=f'synthetic_{row}_{year}_{kind}_{entity}'
        item=dict(id=identity,row_id=row,metric=row,label=row.replace('_',' ').title(),entity_id=entity,model_entity_id=entity,
            period_id=period['id'],fiscal_year=year,kind=kind,start=period['start'],end=period['end'],
            metric_type=metric_type,value=value,currency=None if unit=='shares_m' else 'GBP',unit=unit,
            definition='Synthetic '+row,definition_id=row,basis=basis,compatibility_key=basis,scope=scope,
            restatement='original',status='reported',source_refs=[dict(source_id='synthetic_primary_statements',
                sheet='SYNTHETIC fiscal disclosures',cell=identity,raw=value,multiplier=1)],**extras)
        observations.append(item);return identity
    stock={'cash','short_investments','inventories','receivables','other_current_assets','current_assets','ppe','goodwill',
        'intangibles','associates','other_noncurrent_assets','noncurrent_assets','total_assets','short_debt','long_debt',
        'trade_payables','current_tax_liabilities','other_current_liabilities','current_liabilities','other_noncurrent_liabilities',
        'noncurrent_liabilities','total_liabilities','share_capital','treasury_stock','reserves','retained_earnings','parent_equity',
        'nci_equity','total_equity','liabilities_equity','owned_ppe','right_of_use_assets','lease_liabilities','held_for_sale_cash','closing_cash'}
    cumulative={'Q1':.15,'H1':.4,'9M':.6,'FY':1}
    closing={'Q1':.15,'Q2':.4,'H1':.4,'Q3':.6,'9M':.6,'Q4':1,'H2':1,'FY':1}
    for year in (2024,2025):
        annual={rid:role['values'][str(year)] for rid,role in data['historical'].items()}
        cash_change=annual['cash_change']
        for rid,value in annual.items():
            if rid=='opening_cash':
                for kind,weight in cumulative.items():put(rid,year,kind,value,'stock',measurement_date=f'{year-1}-06-30')
            elif rid in stock:
                for kind,weight in closing.items():
                    adjusted=value
                    if rid in {'closing_cash','cash','current_assets','total_assets','retained_earnings','parent_equity','total_equity','liabilities_equity'}:
                        adjusted=value-cash_change*(1-weight)
                    put(rid,year,kind,adjusted,'stock')
            elif rid=='shares_basic':
                for kind in closing:put(rid,year,kind,value,'weighted_average',unit='shares_m')
            elif rid=='eps_basic':
                for kind,weight in cumulative.items():put(rid,year,kind,value*weight,'ratio',unit='GBP_per_share')
            else:
                for kind,weight in cumulative.items():put(rid,year,kind,value*weight)
    # The predecessor shares a familiar name but covers BOTH successor businesses.
    put('north_sales',2024,'FY',180,entity='old_north',basis='legacy_combined_basis',scope='both_current_businesses')
    for segment in data['segments']:
        key=segment['id'];segment['label']='Northern components' if key=='north' else 'Specialty components'
        for metric in ('sales','utop','da'):
            for kind,weight in cumulative.items():
                put(key+'_'+metric,2025,kind,segment[metric+'2025']*weight,entity='new_'+key,
                    basis='current_components_basis',scope=key+'_only',current_basis_compatible=True)
        for metric in ('pricing','rig'):
            for kind in ('Q1','H1','FY'):
                value=segment[metric+'2025'] if kind!='H1' else segment[metric+'2025']+.01
                put(key+'_'+metric,2025,kind,value,'ratio',entity='new_'+key,basis='current_components_basis',
                    scope=key+'_only',unit='ratio',current_basis_compatible=True)
    for metric in ('utop','da'):
        for kind,weight in cumulative.items():put('corporate_'+metric,2025,kind,data['corporate'][metric+'2025']*weight,
            entity='corporate',basis='current_components_basis',current_basis_compatible=True)
    old_id=put('legacy_north_sales',2025,'FY',180,entity='old_north',basis='legacy_combined_basis',scope='both_current_businesses')
    data['reporting_bridges']=[dict(id='synthetic_split',effective_date='2024-07-01',published_at='2025-08-01',
        period_id='2025-FY',metric='sales',unit='GBPm',predecessor_basis='legacy_combined_basis',successor_basis='current_components_basis',
        predecessors=[dict(entity_id='old_north',value=180,observation_id=old_id)],
        successors=[dict(entity_id='new_north',value=100),dict(entity_id='new_specialty',value=80)],
        relationships=[dict(predecessor='old_north',successor='new_north',value=100,weight=100/180),
                       dict(predecessor='old_north',successor='new_specialty',value=80,weight=80/180)],
        residual=0,status='exact_published_split',limitations=['No supported allocation for fiscal 2024; retain the predecessor history separately.'],
        source_refs=[dict(source_id='synthetic_primary_segment_note',sheet='SYNTHETIC split bridge',cell='A1')])]
    data['observations']=observations
    return data


def _write_synthetic_originals(data, archive):
    """Materialize this explicitly synthetic fixture as real cached source cells.

    The entrypoint test verifies the actual source reader and exporter. These
    originals establish reproducible fixture provenance, not company research.
    """
    folder=archive/'financial_reports';folder.mkdir(parents=True)
    books={}
    for source_id in ('synthetic_primary_statements','synthetic_primary_segment_note'):
        book=Workbook();sheet=book.active;sheet.title='SYNTHETIC disclosures'
        sheet.append(['SYNTHETIC Harbour Components plc - GBP millions unless stated'])
        sheet.append(['Concept and period / fixture path','Reported numeric value'])
        books[source_id]=book

    def visit(item, path=''):
        if isinstance(item,dict):
            if item.get('source_id') in books:
                raw=item.get('raw',item.get('original_value'))
                if raw is None:
                    # The only unnumbered reference is the published split bridge.
                    assert path.startswith('/reporting_bridges/')
                    raw=180.
                sheet=books[item['source_id']].active
                sheet.append([path,raw]);cell=f'B{sheet.max_row}'
                item.update(sheet=sheet.title,cell=cell,unit_id='sheet:'+sheet.title,
                            raw=raw,multiplier=1)
            for key,value in item.items():visit(value,path+'/'+key)
        elif isinstance(item,list):
            for i,value in enumerate(item):visit(value,path+'/'+str(i))

    data['corporate']['source_refs']={key:dict(source_id='synthetic_primary_segment_note',raw=value)
        for key,value in data['corporate'].items() if key.endswith('2025')}
    visit(data)
    sources={}
    for source_id,book in books.items():
        path=folder/(source_id+'.xlsx');book.save(path);book.close()
        sources[source_id]=dict(id=source_id,local_path=path.relative_to(archive).as_posix(),
            sha256=hashlib.sha256(path.read_bytes()).hexdigest(),format='xlsx',category='financial_reports',
            title='SYNTHETIC Harbour Components fixture original',published_at='2025-08-01')
    data['sources']=sources
    return sources


class DifferentCompanyPeriodModelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data=different_period_company();cls.spec=build_periodic_spec(cls.data,2);cls.values=calculate_periodic(cls.spec)

    def test_detail_horizon_follows_june_fiscal_calendar_and_transitions_to_annual(self):
        data=different_period_company()
        data['model_configuration']['detailed_forecast_years']=2
        spec=build_periodic_spec(data,5);values=calculate_periodic(spec)
        self.assertEqual(spec['forecast_horizon']['current_fiscal_year'],2027)
        self.assertEqual(spec['forecast_horizon']['detailed_through'],2028)
        self.assertEqual([p['id'] for p in spec['periods'] if p['fiscal_year']>=2029],['2029-FY','2030-FY'])
        self.assertFalse(any('2029-H' in str(c.get('bindings')) for r in spec['rows'] for c in r['cells'].values()))
        self.assertAlmostEqual(values['debt_interest']['2029-FY'],values['nonlease_debt']['2028-FY']*values['borrowing_rate']['2029-FY'])
        self.assertAlmostEqual(values['owned_ppe_depreciation']['2029-FY'],values['owned_ppe']['2028-FY']*values['owned_depreciation_rate']['2029-FY'])
        self.assertAlmostEqual(values['opening_cash']['2029-FY'],values['closing_cash']['2028-FY'])
        self.assertFalse([c for c in validate_periodic(spec,values) if c['status'] in ('FAIL','MISSING')])

    def test_quarterly_sales_and_half_year_statements_share_the_causal_cost_cadence(self):
        data=different_period_company()
        data['model_configuration'].update(forecast_frequency='quarterly',primary_sales_frequency='quarterly',full_statement_frequency='half_year')
        spec=build_periodic_spec(data,4);values=calculate_periodic(spec)
        self.assertIsNotNone(values['north_sales']['2026-Q1'])
        self.assertIsNone(values['north_pool_cogs']['2026-Q1'])
        self.assertIsNotNone(values['north_pool_cogs']['2026-H1'])
        self.assertIsNone(values['utop']['2026-Q1'])
        self.assertIsNone(values['cash']['2026-H2'])
        self.assertFalse([c for c in validate_periodic(spec,values) if c['status'] in ('FAIL','MISSING')])

    def test_annual_reporter_uses_direct_flows_and_prior_closing_balances(self):
        data=different_period_company()
        data['model_configuration'].update(forecast_frequency='annual',primary_sales_frequency='annual',
            full_statement_frequency='annual',statement_forecast_frequency='annual')
        spec=build_periodic_spec(data,3);values=calculate_periodic(spec)
        future=[p['id'] for p in spec['periods'] if p['fiscal_year'] in spec['forecast_years']]
        self.assertEqual(future,['2026-FY','2027-FY','2028-FY'])
        self.assertAlmostEqual(values['opening_cash']['2027-FY'],values['closing_cash']['2026-FY'])
        self.assertAlmostEqual(values['debt_interest']['2027-FY'],values['nonlease_debt']['2026-FY']*values['borrowing_rate']['2027-FY'])
        self.assertFalse([c for c in validate_periodic(spec,values) if c['status'] in ('FAIL','MISSING')])


    def test_june_year_end_half_year_forecast_and_company_specific_interest(self):
        spec,values=self.spec,self.values;p={p['id']:p for p in spec['periods']}
        self.assertEqual(p['2026-H1']['start'],'2025-07-01');self.assertEqual(p['2026-FY']['end'],'2026-06-30')
        self.assertEqual(spec['forecast_periods'],['2026-H1','2026-H2','2027-H1','2027-H2'])
        self.assertAlmostEqual(values['north_sales']['2026-H1'],40*1.03)
        self.assertAlmostEqual(values['north_sales']['2026-H2'],60*1.03)
        self.assertAlmostEqual(values['sales']['2026-FY'],187)
        # Legacy pooled growth is inactive; corporate functional budgets now
        # have their own economic controls instead of a dormant editable box.
        self.assertIsNone(values['corporate_cost_growth']['2026-FY'])
        self.assertIsNone(values['corporate_cost_growth']['2026-H1'])
        self.assertIsNotNone(values['corporate_pool_marketing_admin']['2026-H1'])
        # Keep the disclosed blended 2025 yield, but apply the company policy to
        # bank cash only; short-term investments have no forecast interest income.
        self.assertAlmostEqual(values['cash_interest']['2026-H1'],29.5*(.8/34.5)*(184/365))
        for rid in ('sales','parent_profit','cfo','cfi','cff'):
            self.assertAlmostEqual(values[rid]['2026-FY'],values[rid]['2026-H1']+values[rid]['2026-H2'])
        self.assertIsNone(values['cash']['2026-H2'])
        self.assertIsNotNone(values['cash']['2026-FY'])
        self.assertIsNone(values['parent_profit']['2026-Q1'])
        failures=[c for c in validate_periodic(spec,values) if c['status'] in ('FAIL','MISSING')]
        self.assertFalse(failures,failures[:5])

    def test_reorganisation_preserves_predecessor_without_false_comparable_splice(self):
        rows={r['id']:r for r in self.spec['rows']}
        self.assertIsNone(self.values['north_sales']['2024-FY'])
        self.assertEqual(rows['north_sales']['cells']['2024-FY']['status'],'not_comparable')
        self.assertTrue(any(o['entity_id']=='old_north' and o['value']==180 for o in self.spec['observations']))
        bridge=self.spec['reporting_bridges'][0]
        self.assertEqual(sum(s['value'] for s in bridge['successors']),bridge['predecessors'][0]['value'])
        self.assertEqual(bridge['residual'],0)
        self.assertEqual(self.values['sales']['2025-FY'],180)
        self.assertEqual(len(self.spec['entity_versions']),3)

    def test_research_context_exposes_numeric_predecessor_without_consolidating_it(self):
        before=copy.deepcopy(self.data)
        context=segment_research_context(self.data,'new_north',periods=['2024-FY','2025-FY'])
        self.assertEqual({v['id'] for v in context['entity_versions']},{'old_north','new_north'})
        predecessor=[o for o in context['numerical_history'] if o['entity_id']=='old_north']
        self.assertTrue(predecessor)
        self.assertTrue(all(o['value']==180 and o['research_use']=='context_only' for o in predecessor))
        successor=[o for o in context['numerical_history'] if o['entity_id']=='new_north' and o['row_id']=='north_sales']
        self.assertEqual([(o['period_id'],o['value'],o['research_use']) for o in successor],[('2025-FY',100,'direct_consolidation')])
        self.assertEqual(context['reporting_bridges'][0]['id'],'synthetic_split')
        self.assertEqual(sum(x['value'] for x in context['reporting_bridges'][0]['successors']),180)
        self.assertEqual(self.data,before)
        self.assertEqual(self.values['north_sales']['2025-FY'],100)
        self.assertEqual(self.values['sales']['2025-FY'],180)

    def test_supported_flow_derivations_are_live_and_growth_is_not_subtracted(self):
        rows={r['id']:r for r in self.spec['rows']}
        self.assertAlmostEqual(self.values['sales']['2025-Q2'],45)
        self.assertAlmostEqual(self.values['sales']['2025-H2'],108)
        self.assertIn('expression',rows['sales']['cells']['2025-Q2'])
        self.assertEqual(rows['sales']['cells']['2025-Q2']['status'],'derived_actual')
        self.assertIsNone(self.values['north_pricing']['2025-H2'])
        self.assertFalse(any(o.get('is_derived') and o['row_id'].endswith(('_pricing','_rig','eps_basic')) for o in self.spec['observations']))
        self.assertAlmostEqual(self.values['inventories']['2025-H1'],30)

    def test_authoritative_assumptions_change_relevant_half_year_and_annual_statements(self):
        rows={r['id']:r for r in self.spec['rows']}
        self.assertTrue(rows['north_pricing']['cells']['2026-H1']['editable'])
        changed=calculate_periodic(self.spec,{'north_pricing':{'2026-H1':.12}})
        self.assertAlmostEqual(changed['north_sales']['2026-H1']-self.values['north_sales']['2026-H1'],4)
        self.assertAlmostEqual(changed['sales']['2026-FY']-self.values['sales']['2026-FY'],4)
        self.assertEqual(changed['north_sales']['2026-H2'],self.values['north_sales']['2026-H2'])
        self.assertNotEqual(changed['cash']['2026-FY'],self.values['cash']['2026-FY'])
        self.assertFalse([c for c in validate_periodic(self.spec,changed) if c['status']=='FAIL'])
        changed_ratio=calculate_periodic(self.spec,{'inventory_ratio':{'2026-H1':self.values['inventory_ratio']['2026-H1']+.02}})
        self.assertGreater(changed_ratio['inventories']['2026-H1'],self.values['inventories']['2026-H1'])
        self.assertLess(changed_ratio['cfo']['2026-H1'],self.values['cfo']['2026-H1'])
        self.assertFalse([c for c in validate_periodic(self.spec,changed_ratio) if c['status']=='FAIL'])

    def test_gbp_fx_selection_and_source_coverage_are_independent_of_company_identity(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'synthetic-fx.xlsx';source_workbook(path)
            fx=load_fx_source(path);spec=build_periodic_spec(self.data,2,fx_source=fx);values=calculate_periodic(spec)
        self.assertEqual(values['fx_average_usd']['2024-H2'],2)
        self.assertIsNone(values['fx_average_usd']['2026-H1'])
        rows={r['id']:r for r in spec['rows']}
        self.assertTrue(rows['fx_average_usd']['label'].startswith('GBP per USD'))
        self.assertFalse(any('CHF' in r['unit'] or 'Nestl' in r['label'] for r in spec['rows']))
        self.assertEqual(values['sales'],self.values['sales'])
        self.assertEqual(values['cash'],self.values['cash'])

    def test_company_calibration_override_controls_forecast_assumption(self):
        data=copy.deepcopy(self.data)
        data['model_configuration']['row_overrides']['cash_yield']={'initialization_expression':'cash_interest / cash'}
        spec=build_periodic_spec(data,1);values=calculate_periodic(spec)
        expected=.8/29.5
        self.assertAlmostEqual(values['cash_yield']['2025-FY'],expected)
        self.assertAlmostEqual(values['cash_yield']['2026-H1'],expected)
        self.assertAlmostEqual(values['cash_interest']['2026-H1'],.8*(184/365))

    def test_estimate_observation_keeps_estimate_classification(self):
        data=copy.deepcopy(self.data)
        source=next(o for o in data['observations'] if o['row_id']=='north_pricing' and o['period_id']=='2025-Q1')
        source['status']='estimate'
        spec=build_periodic_spec(data,1)
        cell=next(r for r in spec['rows'] if r['id']=='north_pricing')['cells']['2025-Q1']
        self.assertEqual(cell['status'],'estimate')
        self.assertEqual(cell['value'],source['value'])

    def test_preferred_reference_cannot_select_another_metric_or_period(self):
        data=copy.deepcopy(self.data)
        wrong=next(o for o in data['observations'] if o['row_id']=='cash' and o['period_id']=='2025-FY')
        data['model_configuration']['preferred_observation_ids']={'sales:2025-Q1':wrong['id']}
        with self.assertRaises(ValueError):build_periodic_spec(data,1)

    def test_evidence_only_version_preserves_selected_history_and_derivations(self):
        data=copy.deepcopy(self.data)
        original=next(o for o in data['observations'] if o['row_id']=='sales' and o['period_id']=='2025-Q1')
        alternative=copy.deepcopy(original)
        alternative.update(id='independently_reconstructed_alternative',value=30.,restatement='restated',
            eligible=True,model_selection_eligible=False)
        alternative['source_refs'][0]['raw']=30.
        data['observations'].append(alternative)
        before=copy.deepcopy(data)
        spec=build_periodic_spec(data,2);values=calculate_periodic(spec)
        self.assertEqual(data,before)
        retained=next(o for o in spec['observations'] if o['id']==alternative['id'])
        self.assertTrue(retained['eligible'])
        self.assertFalse(retained['model_selection_eligible'])
        self.assertEqual(retained['value'],30.)
        self.assertEqual(retained['source_refs'],alternative['source_refs'])
        self.assertEqual(spec['observation_selection'],self.spec['observation_selection'])
        self.assertEqual(values,self.values)
        self.assertEqual(values['sales']['2025-Q1'],27.)
        self.assertEqual(values['sales']['2025-Q2'],45.)
        self.assertFalse([c for c in validate_periodic(spec,values) if c['status'] in ('FAIL','MISSING')])
        data['model_configuration']['preferred_observation_ids']={'sales:2025-Q1':alternative['id']}
        with self.assertRaisesRegex(ValueError,'Preferred observation is ineligible'):
            build_periodic_spec(data,1)

    def test_evidence_only_version_does_not_suppress_sourced_annual_fallback(self):
        data=copy.deepcopy(self.data)
        alternative=copy.deepcopy(next(o for o in data['observations'] if o['row_id']=='sales' and o['period_id']=='2025-FY'))
        data['observations']=[o for o in data['observations'] if not (o['row_id']=='sales' and o['period_id']=='2025-FY')]
        baseline=build_periodic_spec(data,1)
        alternative.update(id='unselected_annual_restatement',value=190.,model_selection_eligible=False)
        alternative['source_refs'][0]['raw']=190.
        data['observations'].append(alternative)
        spec=build_periodic_spec(data,1)
        self.assertEqual(spec['observation_selection'],baseline['observation_selection'])
        self.assertEqual(calculate_periodic(spec),calculate_periodic(baseline))
        cell=next(r for r in spec['rows'] if r['id']=='sales')['cells']['2025-FY']
        self.assertEqual(cell['observation_id'],'annual_sales_2025')
        self.assertEqual(cell['value'],180.)
        self.assertTrue(any(o['id']==alternative['id'] for o in spec['observations']))

    def test_group_basis_selection_does_not_fall_back_to_incompatible_perimeter(self):
        data=copy.deepcopy(self.data)
        source=next(o for o in data['observations'] if o['row_id']=='sales' and o['period_id']=='2025-Q1')
        source.update(basis='incompatible_group',compatibility_key='incompatible_group',scope='unrelated_perimeter')
        spec=build_periodic_spec(data,1)
        cell=next(r for r in spec['rows'] if r['id']=='sales')['cells']['2025-Q1']
        self.assertIsNone(cell.get('value'))
        self.assertNotEqual(cell.get('observation_id'),source['id'])

    def test_ordinary_source_selection_rejects_incompatible_units(self):
        data=copy.deepcopy(self.data)
        source=next(o for o in data['observations'] if o['row_id']=='north_pricing' and o['period_id']=='2025-Q1')
        source['unit']='GBPm'
        with self.assertRaises(ValueError):build_periodic_spec(data,1)

    def test_inactive_lease_and_nci_calibrations_remain_explicit_policies(self):
        cases=[(('lease_liabilities','lease_interest_expense'),'lease_rate'),
               (('nci_profit','dividends_nci'),'nci_dividend_payout')]
        for zero_roles,driver in cases:
            with self.subTest(driver=driver):
                data=copy.deepcopy(self.data)
                for rid in zero_roles:
                    data['historical'][rid]['values']={y:0 for y in data['historical'][rid]['values']}
                    for o in data['observations']:
                        if o['row_id']==rid:o['value']=0
                spec=build_periodic_spec(data,1);values=calculate_periodic(spec)
                self.assertEqual(values[driver]['2026-H1'],0)
                self.assertEqual(values[driver]['2026-H2'],0)
                cell=next(r for r in spec['rows'] if r['id']==driver)['cells']['2026-H1']
                self.assertEqual(cell['editable'],driver=='lease_rate')
                self.assertEqual(cell['status'],'forecast')
                self.assertIn('inactive',cell['reason'].lower())
                if driver=='nci_dividend_payout':
                    pool=next(r for r in spec['rows'] if r['id']=='dividends_nci_annual_pool')['cells']['2026-FY']
                    self.assertTrue(pool['editable'])
                    self.assertEqual(values['dividends_nci_annual_pool']['2026-FY'],0)

    def test_full_construction_entrypoint_uses_real_synthetic_sources_and_exports_different_company(self):
        """Offline boundary: provider verdicts and Excel COM are mocked only.

        Inventory, file hashes, cached source cells, budget journal, both model
        compilers, reconciliation, sensitivities and XLSX export run normally.
        This proves workflow reuse; it does not validate AI research or Excel COM.
        """
        from research_platform.ai import OpenAIProvider
        from research_platform.integrated import construct_integrated
        data=different_period_company()
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);archive=root/'originals'
            sources=_write_synthetic_originals(data,archive)
            inputs=root/'synthetic_inputs.json';inputs.write_text(json.dumps(data),encoding='utf-8')
            provider=OpenAIProvider(api_key='SYNTHETIC-OFFLINE-TEST-KEY')
            def review(**kwargs):
                provider.last_response=dict(model='gpt-6-luna',usage=dict(input_tokens=20,output_tokens=10,total_tokens=30),output=[])
                return dict(findings=[],summary='SYNTHETIC mock verdict; no external research performed.')
            with patch('research_platform.ai.urlopen') as network, \
                    patch.object(provider,'complete',side_effect=review) as calls, \
                    patch('research_platform.integrated.recalculate_workbook',return_value={
                        'status':'passed','boundary':'MOCK: desktop Excel is not exercised'}) as recalc, \
                    patch('research_platform.model_excel.verify_excel_sensitivities',return_value={
                        'status':'passed','boundary':'MOCK: desktop Excel sensitivities are not exercised'}) as excel_sensitivity:
                directory=construct_integrated(root,data['company'],archive,inputs,provider=provider,forecast_years=2)
                self.assertEqual(calls.call_count,5)
                self.assertEqual({call.kwargs['payload']['stage'] for call in calls.call_args_list},
                    {'integrated_segments','integrated_income_review','integrated_cashflow_review',
                     'integrated_balance_review','integrated_schedules'})
                recalc.assert_called_once();excel_sensitivity.assert_called_once();network.assert_not_called()
                record=json.loads((directory/'record.json').read_text(encoding='utf-8'))
                self.assertEqual(record['status'],'complete_provisional')
                self.assertEqual(record['configuration']['mode'],'integrated_periodic_v2')
                source_checks=json.loads((directory/'source_validation.json').read_text(encoding='utf-8'))
                self.assertGreater(len(source_checks),500)
                self.assertTrue(all(c['status']=='PASS' for c in source_checks))
                for source_id,source in sources.items():
                    frozen=directory/'sources'/(source_id+'.xlsx')
                    self.assertEqual(hashlib.sha256(frozen.read_bytes()).hexdigest(),source['sha256'])
                spec=json.loads((directory/'model_specification.json').read_text(encoding='utf-8'))
                values=json.loads((directory/'forecast.json').read_text(encoding='utf-8'))
                inspection=json.loads((directory/'workbook_validation.json').read_text(encoding='utf-8'))
                self.assertEqual(spec['currency'],'GBP')
                self.assertEqual(spec['forecast_periods'],['2026-H1','2026-H2','2027-H1','2027-H2'])
                self.assertEqual(next(p for p in spec['periods'] if p['id']=='2026-FY')['end'],'2026-06-30')
                self.assertEqual({r['section'] for r in spec['rows'] if r['section'].startswith('segment_')},
                                 {'segment_north','segment_specialty'})
                self.assertAlmostEqual(values['sales']['2026-FY'],187)
                self.assertIsNone(values['north_sales']['2024-FY'])
                self.assertEqual(spec['reporting_bridges'][0]['predecessors'][0]['value'],180)
                self.assertEqual(spec['reporting_bridges'][0]['residual'],0)
                with closing(load_workbook(directory/'model.xlsx',data_only=False)) as book:
                    self.assertTrue({'Model','Assumptions','Historics','FX','Historical Segments'}<=set(book.sheetnames))
                    cols=inspection['period_columns'];positions=inspection['model_rows']
                    self.assertIsNone(book['Model'][cols['2024-FY']+str(positions['north_sales'])].value)
                    cell=book['Model'][cols['2026-H1']+str(positions['north_pricing'])]
                    self.assertIn("'Assumptions'!",cell.value)
                    self.assertTrue(book['Model'][cols['2026-FY']+str(positions['sales'])].value.startswith('='))
                    self.assertFalse(any('CHF' in str(c.value) or 'Nestl' in str(c.value)
                        for row in book['Model'] for c in row if c.value is not None))
                # A changed candidate value cannot bypass real source validation.
                bad=copy.deepcopy(data)
                next(o for o in bad['observations'] if o['row_id']=='north_sales' and o['period_id']=='2025-FY')['value']=101
                inputs.write_text(json.dumps(bad),encoding='utf-8')
                with self.assertRaisesRegex(ValueError,'Source validation has discrepancies'):
                    construct_integrated(root,data['company'],archive,inputs,provider=provider,forecast_years=2)
                self.assertEqual(calls.call_count,5)
                self.assertEqual(json.loads((directory/'record.json').read_text(encoding='utf-8'))['status'],'complete_provisional')
                network.assert_not_called()


if __name__=='__main__':unittest.main()
