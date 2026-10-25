"""Financial-renderer checks use small explicit synthetic specifications."""
import copy
import ast
import re
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from research_platform.model_excel import export_integrated,validate_workbook_layout
from research_platform.periods import build_periods, calculate_periodic


def example():
    def row(key, label, section, historical=None, expression=None, **kwargs):
        result = dict(id=key, label=label, section=section, unit="CHFm", kind="calculated" if expression else "reported",
                      historical=historical or {}, source_refs={}, basis="SYNTHETIC reference", note="Synthetic only", style="normal")
        if expression:
            result["expression"] = expression
        result.update(kwargs)
        return result
    rows = [
        row("north_sales", "Sales", "north", {"2025": 100}, 'lag("north_sales",1)*(1+pricing)', style="subtotal", source_refs={"2025": [{"source_id": "demo", "cell": "A1", "value": 100}]}),
        row("pricing", "Pricing — temporary assumption", "north", {"2025": 0}, unit="ratio", kind="assumption", inputs={"2026": 0.1, "2027": 0}),
        row("north_costs", "Operating costs", "north", {"2025": -60}, 'lag("north_costs",1)'),
        row("north_profit", "Operating profit", "north", {}, "north_sales+north_costs", historical_expression="north_sales+north_costs", style="subtotal"),
        row("south_sales", "Sales", "south", {"2025": 50}, 'lag("south_sales",1)'),
        row("group_sales", "Group sales", "income", {"2025": 150}, "north_sales+south_sales", historical_expression="north_sales+south_sales", style="subtotal"),
        row("inventory", "Inventories", "balance", {"2025": 30}, "group_sales*inventory_ratio"),
        row("inventory_ratio", "Inventory / sales", "balance", {"2025": 0.2}, unit="ratio", kind="assumption", inputs={"2026": 0.2, "2027": 0.2}),
        row("inventory_cf", "Inventory cash movement", "cash", {"2025": -5}, 'lag("inventory",1)-inventory'),
        row("missing_control", "Missing, not zero", "cash", {}, kind="assumption", inputs={"2026": None, "2027": None}),
        row("missing_result", "Missing flow", "cash", {}, "missing_control+inventory_cf"),
    ]
    spec = dict(company="SYNTHETIC Manufacturing", currency="CHF", historical_years=[2024, 2025], forecast_years=[2026, 2027],
                sections=[dict(id=key, label=label) for key, label in [("north", "NORTH"), ("south", "SOUTH"), ("income", "GROUP INCOME STATEMENT"), ("cash", "CASH FLOW"), ("balance", "BALANCE SHEET")]], rows=rows,
                identities=[dict(id="sales", label="Segments to group sales", expression="group_sales-north_sales-south_sales", periods=["2025", "2026", "2027"], tolerance=0.5, independent=False)],
                sources=[dict(id="demo", title="Synthetic fixture")], gaps=["Synthetic test only"],
                reporting_changes=[{"Effective date": "2025-01-01", "Change": "Synthetic segment basis change", "Evidence": "demo"}],
                interims={"columns": ["Metric", "Period", "Value"], "rows": [["Sales", "H1 2025", 75]]},
                historical_segments=[{"Basis": "Old structure", "Period": 2024, "Sales": 140}])
    values = {r["id"]: copy.deepcopy(r["historical"]) for r in rows}
    values.update(north_sales={"2025": 100, "2026": 110, "2027": 110}, north_costs={"2025": -60, "2026": -60, "2027": -60},
                  north_profit={"2025": 40, "2026": 50, "2027": 50}, south_sales={"2025": 50, "2026": 50, "2027": 50},
                  group_sales={"2025": 150, "2026": 160, "2027": 160}, inventory={"2025": 30, "2026": 32, "2027": 32},
                  inventory_cf={"2025": -5, "2026": -2, "2027": 0})
    return spec, values


class IntegratedExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "model.xlsx"
        self.spec, self.values = example()

    def render(self, checks=None):
        inspection = export_integrated(self.path, self.spec, self.values, checks or [])
        wb = load_workbook(self.path, data_only=False)
        self.addCleanup(wb.close)
        return inspection, wb

    def test_assumption_is_editable_once_and_drives_segment_group_and_cash(self):
        result, wb = self.render()
        positions = result["model_rows"]
        self.assertEqual(result["assumption_count"], 3)
        pricing = wb["Model"][f"F{positions['pricing']}"]
        self.assertEqual(pricing.value, '=IF(ISNUMBER(Assumptions!D6),Assumptions!D6,NA())')
        self.assertEqual(pricing.fill.fgColor.rgb[-6:], "FFFFCC")
        self.assertEqual(pricing.font.color.rgb[-6:], "008000")
        self.assertEqual(pricing.hyperlink.target, "#'Assumptions'!D6")
        self.assertIsNotNone(pricing.border.left.style)
        self.assertIn(f"F{positions['pricing']}", wb["Model"][f"F{positions['north_sales']}"].value)
        self.assertIn(f"F{positions['north_sales']}", wb["Model"][f"F{positions['group_sales']}"].value)
        self.assertIn(f"F{positions['inventory']}", wb["Model"][f"F{positions['inventory_cf']}"].value)
        self.assertEqual(wb["Assumptions"]["D6"].value, 0.1)
        self.assertEqual(wb["Assumptions"]["E6"].value, 0)
        self.assertIsNone(wb["Assumptions"]["D8"].value)
        self.assertEqual(wb["Assumptions"]["D6"].font.color.rgb[-6:], "0000FF")
        self.assertEqual(result["assumption_cells"]["pricing"]["2026"], {"sheet":"Assumptions","cell":"D6"})
        self.assertEqual(wb["Drivers"]["L6"].value, "Assumptions!D6")
        self.assertTrue(wb["Drivers"]["I6"].value)
        self.assertGreater(wb["Dependencies"].max_row, 5)

    def test_historical_pulls_use_stable_lookup_and_missing_is_not_zero(self):
        result, wb = self.render()
        positions = result["model_rows"]
        observed = wb["Model"][f"E{positions['north_sales']}"]
        self.assertIn("INDEX(Historics!", observed.value)
        self.assertIn("MATCH($A", observed.value)
        self.assertEqual(observed.font.color.rgb[-6:], "D63384")
        self.assertEqual(result["expected_model_cells"][f"D{positions['north_sales']}"], "N/C")
        self.assertEqual(result["expected_model_cells"][f"E{positions['pricing']}"], 0)
        missing = wb["Model"][f"F{positions['missing_result']}"]
        self.assertIn("COUNT(", missing.value)
        self.assertIn("NA()", missing.value)
        self.assertIsNone(result["expected_model_cells"][missing.coordinate])

    def test_historical_calculated_rows_and_live_checks_are_formulas(self):
        result, wb = self.render([dict(id="source_review", label="Source facts", status="PENDING", detail="Review is frozen")])
        positions = result["model_rows"]
        self.assertNotIn("INDEX", wb["Model"][f"E{positions['north_profit']}"].value)
        self.assertIn(f"E{positions['north_sales']}", wb["Model"][f"E{positions['north_profit']}"].value)
        self.assertIn("INDEX", wb["Model"][f"E{positions['group_sales']}"].value)
        self.assertEqual(result["expected_check_cells"]["D7"], 0)
        self.assertEqual(result["expected_check_cells"]["F7"], "PASS")
        self.assertIn("'Model'!", wb["Checks"]["D7"].value)
        self.assertIn('"FAIL"', wb["Checks"]["F7"].value)
        self.assertTrue(any("FROZEN VALIDATION" in str(c.value) for row in wb["Checks"] for c in row))

    def test_bunzl_layout_and_supporting_tables(self):
        result, wb = self.render()
        self.assertFalse(wb["Model"].sheet_view.showGridLines)
        self.assertEqual(wb["Model"].sheet_view.zoomScale, 70)
        self.assertEqual(wb["Model"].freeze_panes, "D6")
        self.assertTrue(wb["Model"].column_dimensions["A"].hidden)
        self.assertEqual(wb["Model"]["D5"].value, "2024A")
        self.assertEqual(wb["Model"]["F5"].value, "2026E")
        header = wb["Model"].cell(result["section_rows"]["north"], 2)
        self.assertEqual(header.fill.fgColor.rgb[-6:], "183B56")
        self.assertEqual(wb["Interims"]["C6"].value, 75)
        self.assertIn("Historical Segments", wb.sheetnames)
        self.assertEqual(wb["Reporting Changes"]["B6"].value, "Synthetic segment basis change")
        self.assertGreater(wb["Sources"].max_row, 5)
        self.assertIn("Assumption history: initialization values, not reported assumptions", wb["Model"]["B4"].value)
        self.assertTrue(wb["Model"]["B4"].alignment.wrap_text)

    def test_spec_can_generate_another_segment_arrangement(self):
        self.spec["sections"][0]["label"] = "INDUSTRIAL DIVISION"
        self.spec["company"] = "SYNTHETIC Different Company"
        result, wb = self.render()
        self.assertEqual(wb["Model"].cell(result["section_rows"]["north"], 2).value, "INDUSTRIAL DIVISION")
        self.assertEqual(wb["Model"]["B2"].value, "SYNTHETIC Different Company")

    def test_historical_formulas_respect_explicit_reporting_periods(self):
        self.spec["rows"][3]["historical_expression_periods"] = ["2025"]
        result, wb = self.render()
        r = result["model_rows"]["north_profit"]
        self.assertNotIn("INDEX", wb["Model"][f"E{r}"].value)
        self.assertIn("INDEX", wb["Model"][f"D{r}"].value)
        self.assertEqual(result["expected_model_cells"][f"D{r}"], "N/C")

    def test_check_display_roundoff_does_not_change_values_or_formula(self):
        self.spec["rows"][3]["check_tolerance"] = 0.000001
        result, wb = self.render()
        r = result["model_rows"]["north_profit"]
        cell = wb["Model"][f"F{r}"]
        self.assertEqual(cell.number_format, '#,##0.0000000;(#,##0.0000000);"–"')
        rules = wb["Model"].conditional_formatting[f"D{r}:G{r}"]
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0].formula, [f"AND(ISNUMBER(D{r}),ABS(D{r})<=1e-06)"])
        self.assertEqual(rules[0].dxf.numFmt.formatCode, '"–";"–";"–"')
        self.assertTrue(rules[0].stopIfTrue)
        self.assertEqual(result["expected_model_cells"][cell.coordinate], 50)
        self.assertIn(f"F{result['model_rows']['north_sales']}", cell.value)
        self.assertEqual(self.spec["rows"][3]["check_tolerance"], 0.000001)

    def test_broken_mapping_and_nonfinite_inputs_fail(self):
        self.spec["rows"][0]["expression"] = "unknown_input+1"
        with self.assertRaisesRegex(ValueError, "Unknown expression dependency"):
            export_integrated(self.path, self.spec, self.values, [])
        self.spec, self.values = example()
        self.spec["rows"][1]["inputs"]["2026"] = float("nan")
        with self.assertRaisesRegex(ValueError, "Non-numeric assumption"):
            export_integrated(self.path, self.spec, self.values, [])


def periodic_example():
    periods = build_periods([2025, 2026], {"fiscal_year_end": "06-30"})
    for period in periods:
        period["status"] = "actual" if period["fiscal_year"] == 2025 else "forecast"
    def bound(expression, pairs, status="derived_actual", **extra):
        return dict(expression=expression, bindings={symbol:dict(row_id=row,period_id=pid) for symbol,(row,pid) in pairs.items()}, status=status, **extra)
    sales = {"2025-Q1":dict(value=10,status="actual",source_refs=[dict(file="original.xlsx",cell="C2")]),
             "2025-H1":dict(value=30,status="actual"),"2025-9M":dict(value=60,status="actual"),"2025-FY":dict(value=100,status="actual")}
    for target,left,right in [("Q2","H1","Q1"),("Q3","9M","H1"),("Q4","FY","9M"),("H2","FY","H1")]:
        sales["2025-"+target]=bound("a-b",{"a":("sales","2025-"+left),"b":("sales","2025-"+right)})
    growth={pid:dict(status="not_applicable",reason="Forecast-only control") for pid in sales}
    for quarter in ["Q1","Q2","Q3","Q4"]:
        pid="2026-"+quarter
        growth[pid]=dict(value=0.1 if quarter=="Q1" else 0,status="forecast",editable=True)
        sales[pid]=bound("prior*(1+growth)",{"prior":("sales","2025-"+quarter)},status="forecast")
    for kind,parts in {"H1":["Q1","Q2"],"H2":["Q3","Q4"],"9M":["Q1","Q2","Q3"],"FY":["Q1","Q2","Q3","Q4"]}.items():
        sales["2026-"+kind]=bound("+".join(parts),{part:("sales","2026-"+part) for part in parts},status="forecast")
        growth["2026-"+kind]=dict(status="not_applicable",reason="Quarterly controls are authoritative")
    seasonal={"2026-Q1":bound("quarter/annual",{"quarter":("sales","2025-Q1"),"annual":("sales","2025-FY")},status="forecast",editable=False),
              "2026-Q2":dict(value=None,status="not_disclosed",editable=True,required=False,reason="Optional future input")}
    missing={"2025-Q1":dict(status="not_disclosed",reason="Sales-only release"),"2025-H1":dict(status="not_comparable",reason="Scope changed"),
             "2025-FY":dict(status="not_applicable",reason="Forecast-only control"),"2026-Q1":dict(status="error",reason="Unresolved mapping")}
    spec=dict(company="SYNTHETIC June Year-end USD",currency="USD",periods=periods,sections=[dict(id="segment_north",label="North division")],
              rows=[dict(id="sales",label="Sales",section="segment_north",unit="USDm",kind="calculated",cells=sales,style="subtotal"),
                    dict(id="growth",label="Quarterly year-on-year growth",section="segment_north",unit="ratio",kind="assumption",cells=growth),
                    dict(id="seasonality",label="Historical first-quarter share",section="segment_north",unit="ratio",kind="assumption",cells=seasonal),
                    dict(id="missing",label="Disclosed gap classifications",section="segment_north",unit="USDm",kind="reported",cells=missing)],
              identities=[dict(id="annual_sales",label="Quarterly sales aggregate",periods=["2026-FY"],tolerance=1e-6,expression="total-q1-q2-q3-q4",bindings={symbol:dict(row_id="sales",period_id=pid) for symbol,pid in {"total":"2026-FY","q1":"2026-Q1","q2":"2026-Q2","q3":"2026-Q3","q4":"2026-Q4"}.items()})],
              observations=[dict(id="legacy_1",entity_id="north_old",model_entity_id="north",basis="Original northern scope",metric="sales",label="Sales",period_id="2025-FY",start="2024-07-01",end="2025-06-30",unit="USDm",value=100,status="reported",source_refs=[dict(file="original.xlsx",cell="C2")])],
              reporting_bridges=[dict(predecessor_ids=["north_old"],successor_ids=["north"],effective_date="2025-07-01",scope_change="Same business; label change",relationship="exact",source_refs=[dict(file="original.xlsx")])],
              fx_observations=[dict(report_currency="USD",source_currency="EUR",period_id="2025-Q1",value=1.1,status="actual")],
              missing_value_audit={"before":{"generic_nc":4},"after":{"not_disclosed":1,"not_applicable":1,"not_comparable":1,"error":1}},sources=[],gaps=[])
    return spec,calculate_periodic(spec)


class PeriodicExcelTests(unittest.TestCase):
    def test_readonly_support_rows_collapse_without_hiding_editable_missing_controls(self):
        spec,values=periodic_example()
        for row in spec['rows']:
            if row['id'] in ('missing','seasonality'):row['support_matrix']=True
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'periods.xlsx';inspection=export_integrated(path,spec,values,[])
            workbook=load_workbook(path);self.addCleanup(workbook.close)
            self.assertIn('missing',inspection['hidden_assumption_rows'])
            self.assertNotIn('seasonality',inspection['hidden_assumption_rows'])
            row=inspection['hidden_assumption_rows']['missing']
            self.assertTrue(workbook['Assumptions'].row_dimensions[row].hidden)
            self.assertEqual(workbook['Assumptions'].row_dimensions[row].outlineLevel,1)
            self.assertIn('missing',inspection['assumption_cells'])

    def test_large_comment_evidence_is_preserved_outside_excel_and_panes_are_unique(self):
        import json
        spec,values=periodic_example()
        long_evidence='Complete source limitation. '*2000
        spec['gaps']=[long_evidence]
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'model.xlsx'
            export_integrated(path,spec,values,[])
            audit=json.loads(path.with_suffix('.audit.json').read_text(encoding='utf-8'))
            self.assertIn(long_evidence,audit['extended_cell_comments']['Model!B7'])
            with path.open('rb') as stream:
                wb=load_workbook(stream)
                self.assertIn('complete text is preserved',wb['Model']['B7'].comment.text)
                for sheet in wb:
                    panes=[selection.pane for selection in sheet.sheet_view.selection]
                    self.assertEqual(len(panes),len(set(panes)))
                    self.assertTrue(all(not cell.comment or len(cell.comment.text)<=32767 for row in sheet for cell in row))
                wb.close()

    def test_supporting_inputs_remain_in_graph_but_are_expandable_hidden_rows(self):
        spec, values = periodic_example()
        spec["rows"][0]["supporting_input"] = True
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "periods.xlsx"
            result = export_integrated(path, spec, values, [])
            wb = load_workbook(path)
            self.addCleanup(wb.close)
            r = result["model_rows"]["sales"]
            self.assertTrue(wb["Model"].row_dimensions[r].hidden)
            self.assertEqual(wb["Model"].row_dimensions[r].outlineLevel, 1)
            address = result["period_columns"]["2025-Q1"] + str(r)
            self.assertEqual(result["expected_model_cells"][address], 10)
            self.assertIn("INDEX(Historics!", wb["Model"][address].value)
            seasonal = result["assumption_cells"]["seasonality"]["2026-Q1"]["cell"]
            self.assertIn("'Model'!" + address, wb["Assumptions"][seasonal].value)
            for sheet in wb:
                self.assertGreaterEqual(sheet.page_margins.bottom, .5)
                self.assertLess(sheet.page_margins.footer, sheet.page_margins.bottom)

    def test_missing_audit_has_individual_counts_and_complete_scope_text(self):
        spec, values = periodic_example()
        spec["model_configuration"] = {"diagnostic_export": True}
        scope = "Same row IDs and fiscal years only; quarterly additions are excluded. " * 20
        spec["missing_value_audit"]["previous_workbook_comparison"] = {
            "before_nc": 1447, "remaining_nc": 749,
            "after_statuses_for_same_cells": {"actual": 85, "derived_actual": 210},
            "scope": scope}
        spec["missing_value_audit"]["reasons"] = {"Supported original scope differs from current basis": 749}
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "periods.xlsx"
            export_integrated(path, spec, values, [])
            wb = load_workbook(path)
            self.addCleanup(wb.close)
            audit = wb["Missing Data Audit"]
            rows = list(audit.iter_rows(min_row=6, values_only=True))
            self.assertIn(("Before", "Generic nc", 4, None), rows)
            self.assertIn(("After", "Not disclosed", 1, None), rows)
            self.assertIn(("Previous workbook comparison", "Before nc", 1447, None), rows)
            self.assertIn(("Previous workbook comparison / After statuses for same cells", "Actual", 85, None), rows)
            chunks = [row[3] for row in rows if row[0] == "Previous workbook comparison" and row[1].startswith("Scope")]
            self.assertEqual("".join(chunks), scope)
            for row_number, row in enumerate(rows, 6):
                if row[3] and len(row[3]) > 200:
                    self.assertGreater(audit.row_dimensions[row_number].height, 30)
                    self.assertLessEqual(audit.row_dimensions[row_number].height, 409)

    def test_noncalendar_period_graph_uses_authoritative_assumptions_and_live_derivations(self):
        spec,values=periodic_example()
        with tempfile.TemporaryDirectory() as temporary:
            result=export_integrated(Path(temporary)/"periods.xlsx",spec,values,[])
            with (Path(temporary)/"periods.xlsx").open("rb") as stream:
                wb=load_workbook(stream)
                self.addCleanup(wb.close)
                model=wb["Model"]
                positions=result["model_rows"];columns=result["period_columns"]
                q1=columns["2026-Q1"]+str(positions["growth"])
                control=result["assumption_cells"]["growth"]["2026-Q1"]["cell"]
                self.assertIn("'Assumptions'!"+control,model[q1].value)
                self.assertEqual(wb["Assumptions"][control].value,0.1)
                self.assertIsNone(model[q1].hyperlink)
                self.assertTrue(model.cell(positions['growth'],2).hyperlink.target.startswith("#'Assumptions'!B"))
                source=result["assumption_cells"]["seasonality"]["2026-Q1"]["cell"]
                self.assertIn("'Model'!",wb["Assumptions"][source].value)
                self.assertNotIn("IFERROR",wb["Assumptions"][source].value)
                self.assertIn("'Assumptions'!",model[columns["2026-Q1"]+str(positions["sales"])].value)
                derived=model[columns["2025-Q2"]+str(positions["sales"])].value
                self.assertIn(columns["2025-H1"]+str(positions["sales"]),derived)
                self.assertIn("-",derived)
                self.assertIn("INDEX(Historics!",model[columns["2025-Q1"]+str(positions["sales"])].value)
                self.assertEqual(model.cell(6,4).value,"Q1A")
                self.assertIn("2024-07-01",model.cell(6,4).comment.text)
                self.assertNotIn('2025-9M',columns)
                self.assertFalse(any('9M' in str(c.value) for sheet in (model,wb['Assumptions']) for row in (5,6) for c in sheet[row]))
                internal=result['internal_period_columns']['2025-9M']
                self.assertTrue(wb['Historics'].column_dimensions[internal].hidden)
                q3=model[columns['2025-Q3']+str(positions['sales'])]
                self.assertIn("'Historics'!"+internal,q3.value)
                self.assertEqual(result["expected_check_cells"]["F6"],"PASS")
                self.assertEqual(result["expected_model_cells"][columns["2026-FY"]+str(positions["sales"])],101)
                self.assertIn("Historical Segments",wb.sheetnames)
                self.assertNotIn("Reporting Bridges",wb.sheetnames)
                self.assertIn("FX",wb.sheetnames)
                self.assertNotIn("Missing Data Audit",wb.sheetnames)
                self.assertEqual([sh.title for sh in wb if sh.sheet_state == "visible"], ["Model", "Assumptions", "Historics", "FX", "Historical Segments"])
                self.assertIn("#'Historical Segments'!",model.cell(result["section_rows"]["segment_north"],2).hyperlink.target)
                self.assertEqual(validate_workbook_layout(Path(temporary)/'periods.xlsx',result)['status'],'passed')

    def test_missing_causes_remain_distinct_from_zero(self):
        spec,values=periodic_example()
        with tempfile.TemporaryDirectory() as temporary:
            result=export_integrated(Path(temporary)/"periods.xlsx",spec,values,[])
            wb=load_workbook(Path(temporary)/"periods.xlsx")
            self.addCleanup(wb.close)
            row=result["model_rows"]["missing"]
            for pid,expected in {"2025-Q1":None,"2025-H1":None,"2025-FY":None,"2026-Q1":"ERR"}.items():
                cell=wb["Model"][result["period_columns"][pid]+str(row)]
                self.assertEqual(cell.value,expected)
                self.assertIn("reason",cell.comment.text)
            zero=result["assumption_cells"]["growth"]["2026-Q2"]["cell"]
            self.assertEqual(wb["Assumptions"][zero].value,0)
            optional=result["assumption_cells"]["seasonality"]["2026-Q2"]["cell"]
            self.assertIsNone(wb["Assumptions"][optional].value)
            display=wb["Model"][result["period_columns"]["2026-Q2"]+str(result["model_rows"]["seasonality"])]
            self.assertEqual(display.value,f'=IF(ISBLANK(\'Assumptions\'!{optional}),"",\'Assumptions\'!{optional})')
            self.assertIsNone(result["expected_model_cells"][display.coordinate])
            self.assertIsNone(result["expected_assumption_cells"][optional])


class ValuationExcelTests(unittest.TestCase):
    def render(self, *, transition=0, market=False, event=False):
        from test_valuation import fixture
        from research_platform.valuation import value_company
        spec, values, inputs = fixture()
        spec.update(historical_years=[2024],forecast_years=[2025],sections=[dict(id='financials',label='Synthetic financial statements')],identities=[],sources=[],gaps=[])
        for row in spec['rows']:row.update(label=row['id'].replace('_',' ').capitalize(),section='financials',kind='reported')
        inputs['assumptions']['transition_years']['value']=transition
        if market:inputs['market'].update(price=dict(value=1000.,as_of='2024-12-30',status='observed',method='Synthetic quote'),price_unit_scale=.01)
        if event:inputs['events']=[dict(id='acquisition',date='2025-12-31',effective_date='2025-12-31',net_cash=dict(value=-80.),cash_tax=dict(value=0.),balance_movements={'associates':-7.,'nonlease_debt':20.,'nci_equity':3.})]
        valuation=value_company(spec,values,inputs);valuation['input_snapshot']=inputs
        temporary=tempfile.TemporaryDirectory();self.addCleanup(temporary.cleanup)
        path=Path(temporary.name)/'model.xlsx'
        inspection=export_integrated(path,spec,values,[],valuation=valuation)
        workbook=load_workbook(path,data_only=False);self.addCleanup(workbook.close)
        return workbook,inspection,valuation

    def evaluate_sheet(self, workbook, inspection):
        """Small independent Excel arithmetic interpreter, without Python equations."""
        sheet=workbook['Valuation'];cache={}
        def cell(address):
            if address in cache:return cache[address]
            value=sheet[address].value
            if isinstance(value,str) and value.startswith('='):
                expression=value[1:].replace('^','**')
                def external(match):
                    title, coordinate=match.groups()
                    namespace={'Model':'expected_model_cells','Assumptions':'expected_assumption_cells','FX':'expected_fx_cells','Historics':'expected_history_cells'}[title]
                    return repr(inspection[namespace][coordinate])
                expression=re.sub(r"'([^']+)'!([A-Z]+[0-9]+)",external,expression)
                expression=re.sub(r'\b[A-Z]+[0-9]+\b',lambda match:repr(cell(match[0])),expression)
                def calc(node):
                    if isinstance(node,ast.Expression):return calc(node.body)
                    if isinstance(node,ast.Constant):return node.value
                    if isinstance(node,ast.UnaryOp):return -calc(node.operand) if isinstance(node.op,ast.USub) else calc(node.operand)
                    if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id=='SUM':return sum(calc(a) for a in node.args)
                    if isinstance(node,ast.BinOp):
                        a,b=calc(node.left),calc(node.right)
                        if isinstance(node.op,ast.Add):return a+b
                        if isinstance(node.op,ast.Sub):return a-b
                        if isinstance(node.op,ast.Mult):return a*b
                        if isinstance(node.op,ast.Div):return a/b
                        if isinstance(node.op,ast.Pow):return a**b
                    raise AssertionError('Unsupported Excel test expression '+ast.dump(node))
                value=calc(ast.parse(expression,mode='eval'))
            cache[address]=value;return value
        return cell

    def test_live_model_links_shared_equations_and_independent_known_value(self):
        workbook,inspection,valuation=self.render()
        self.assertEqual(workbook.sheetnames[:2],['Model','Valuation'])
        locations=inspection['valuation_cells']['2025-FY']
        self.assertIn("'Model'!",workbook['Valuation'][locations['operating_profit']].value)
        self.assertIn(locations['cash_lease_cost'],workbook['Valuation'][locations['operating_profit_after_cash_leases']].value)
        evaluate=self.evaluate_sheet(workbook,inspection)
        expected=(63.5/1.1**(183/365)+765/1.1**.5)/10
        self.assertAlmostEqual(evaluate(inspection['valuation_cells']['fair_value_per_share']),expected)
        for coordinate,target in inspection['expected_valuation_cells'].items():self.assertAlmostEqual(evaluate(coordinate),target,places=7,msg=coordinate)
        self.assertTrue(inspection['valuation_print_area'].startswith('B2:'))
        self.assertGreater(inspection['valuation_formula_count'],20)
        sheet=workbook['Valuation']
        claim_row=sheet[inspection['valuation_assumption_cells']['assumptions.pension_claim']].row
        dilution_row=sheet[inspection['valuation_assumption_cells']['assumptions.dilution_shares']].row
        self.assertEqual(sheet.cell(claim_row,3).value,'GBPm')
        self.assertEqual(sheet.cell(dilution_row,3).value,'m shares')
        self.assertTrue(sheet.cell(claim_row,6).alignment.wrap_text)
        self.assertEqual(sheet.cell(claim_row,6).comment.text,valuation['input_snapshot']['assumptions']['pension_claim']['status'])
        share_row=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,2).value=='Valuation-date basic outstanding shares')
        self.assertEqual(sheet.cell(share_row,3).value,'m shares')
        self.assertTrue(inspection['valuation_page_breaks'])
        self.assertEqual(inspection['valuation_page_breaks'],[item.id for item in sheet.row_breaks.brk])
        self.assertIn('0.00',sheet[inspection['valuation_cells']['fair_value_per_share']].number_format)

    def test_transition_event_market_and_all_export_targets_match_backend(self):
        workbook,inspection,valuation=self.render(transition=2,market=True,event=True)
        evaluate=self.evaluate_sheet(workbook,inspection)
        for coordinate,target in inspection['expected_valuation_cells'].items():self.assertAlmostEqual(evaluate(coordinate),target,places=7,msg=coordinate)
        self.assertAlmostEqual(valuation['portfolio_events'][0]['present_value'],-110/1.1)
        self.assertEqual(workbook['Valuation'][inspection['valuation_assumption_cells']['market.price']].value,1000.)
        self.assertTrue(any('backend snapshots' in str(c.value).lower() for row in workbook['Valuation'] for c in row))
        self.assertTrue(all((d.height or 15)<=72 for d in workbook['Valuation'].row_dimensions.values()))
        sheet=workbook['Valuation']
        event_start=next(r for r in range(1,sheet.max_row+1) if sheet.cell(r,2).value=='DATED PORTFOLIO CASH AND RETAINED OWNERSHIP')
        factor_row=next(r for r in range(event_start,sheet.max_row+1) if sheet.cell(r,2).value=='Discount factor')
        self.assertEqual(sheet.cell(factor_row,3).value,'%')
        self.assertIn('%',sheet.cell(factor_row,4).number_format)
        sensitivity_row=next(r for r in range(1,sheet.max_row+1) if str(sheet.cell(r,2).value).startswith('VALUATION SENSITIVITIES'))
        self.assertTrue(sheet.cell(sensitivity_row+1,5).alignment.wrap_text)
        for heading in inspection['valuation_section_rows']:
            self.assertNotIn(heading,inspection['valuation_page_breaks'])
            self.assertNotIn(heading+1,inspection['valuation_page_breaks'])

    def test_incomplete_valuation_is_visible_for_annual_and_periodic_exports(self):
        for periodic in (False,True):
            spec,values=periodic_example() if periodic else example()
            with tempfile.TemporaryDirectory() as directory:
                path=Path(directory)/'model.xlsx'
                inspection=export_integrated(path,spec,values,[],valuation=dict(status='incomplete',diagnostics=[dict(message='No supported lease principal cash flow')]))
                workbook=load_workbook(path);self.addCleanup(workbook.close)
                self.assertIn('Valuation',workbook)
                self.assertEqual(inspection['valuation_status'],'incomplete')
                self.assertEqual(inspection['expected_valuation_cells'],{})
                self.assertIn('No supported lease',workbook['Valuation']['B9'].value)


if __name__ == "__main__":
    unittest.main()
