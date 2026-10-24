"""Independent synthetic source-format tests; no private archive or API calls."""
from datetime import date
import hashlib
from pathlib import Path
import tempfile
import unittest

import openpyxl

from research_platform.fx import load_fx_source, fx_for_period


def source_workbook(path, shift=0):
    """Different currencies, layout, vintage and unequal quarterly day counts."""
    book=openpyxl.Workbook();book.remove(book.active)
    daily=book.create_sheet('Synthetic daily quotes')
    daily.cell(2+shift,2+shift,'Source convention: currency units per 1 EUR.')
    daily.cell(3+shift,2+shift,'Publication-date weighting. Prior local quotes retain their source dates.')
    header=7+shift;column=2+shift
    labels=['Date','JPY','CHF','EUR','GBP','USD','JPY source date','JPY provider']
    for i,label in enumerate(labels):daily.cell(header,column+i,label)
    days=[date(2024,1,2),date(2024,1,3),date(2024,3,28),date(2024,6,28),date(2024,9,27)]
    for i,day in enumerate(days):
        values=[day,100,1+i,1,4,2,days[0] if i==1 else day,'Synthetic prior quote' if i==1 else 'Synthetic central bank']
        for j,value in enumerate(values):daily.cell(header+1+i,column+j,value)
    for reporting in ['GBP','CHF']:
        sheet=book.create_sheet(reporting+' reporting')
        sheet.cell(2+shift,column,f'1 source-currency unit in {reporting}.')
        sheet.cell(3+shift,column,'Historical through Q2:24. Q3:24E projects the final publication day. Q4:24E stays at spot.')
        for i,label in enumerate(['Code','Currency','Q1:24','Q2:24','Q3:24E','Q4:24E']):
            sheet.cell(header,column+i,label)
        # Q1 has three observations, Q2 one. These are independent expected values.
        rates={'CHF':{'USD':[1,2,2.5,2.5],'JPY':[.02,.04,.05,.05],'CHF':[1,1,1,1],'GBP':[.5,1,1.25,1.25]},
               'GBP':{'USD':[2,2,2,2],'JPY':[.04,.04,.04,.04],'CHF':[22/9,1,.8,.8],'GBP':[1,1,1,1]}}
        for row,(code,quarters) in enumerate(rates[reporting].items(),header+1):
            for offset,value in enumerate([code,code+' synthetic',*quarters]):sheet.cell(row,column+offset,value)
    assumptions=book.create_sheet('Projection controls')
    for row,label,value,text in [(3,'Spot reference date',days[-1],None),
                               (4,'Q3 quarter end',date(2024,9,30),None),
                               (5,'Q3 actual publication days',1,None),
                               (6,'Q3 projected publication days',1,'30 September 2024')]:
        assumptions.cell(row+shift,column,label);assumptions.cell(row+shift,column+1,value)
        if text:assumptions.cell(row+shift,column+3,text)
    book.save(path);book.close()


class FXSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'different-source.xlsx'
        source_workbook(self.path)
        self.source=load_fx_source(self.path)

    def test_codes_dates_and_quote_direction_survive_shifted_layout(self):
        shifted=Path(self.temp.name)/'shifted.xlsx';source_workbook(shifted,shift=4)
        other=load_fx_source(shifted)
        for source in (self.source,other):
            result=fx_for_period(source,'GBP','JPY','2024-01-01','2024-03-31')
            self.assertAlmostEqual(result['value'],.04)
            self.assertEqual(result['status'],'actual')
            self.assertEqual(result['quote_direction'],'reporting-currency units per one source-currency unit')
            self.assertEqual(source['metadata']['spot_reference_date'],'2024-09-27')
        self.assertNotEqual(self.source['observations'][0]['provenance'][0]['cell'],other['observations'][0]['provenance'][0]['cell'])

    def test_unequal_quarter_weights_and_daily_crosses_determine_half_year(self):
        half=fx_for_period(self.source,'CHF','USD','2024-01-01','2024-06-30')
        self.assertAlmostEqual(half['value'],1.25)  # (.5+1+1.5+2)/4; not (1+2)/2
        self.assertEqual(half['coverage'],dict(available_days=4,expected_days=4,complete=True))
        inverse=fx_for_period(self.source,'GBP','CHF','2024-01-01','2024-06-30')
        self.assertAlmostEqual(inverse['value'],(4+2+4/3+1)/4)
        self.assertNotAlmostEqual(inverse['value'],4/2.5)  # cross first, never divide averages

    def test_noncalendar_period_and_source_carry_dates_are_preserved(self):
        result=fx_for_period(self.source,'CHF','USD','2024-02-01','2024-06-30')
        self.assertAlmostEqual(result['value'],1.75)
        carried=fx_for_period(self.source,'CHF','JPY','2024-01-01','2024-06-30')
        self.assertTrue(any('prior source dates' in note for note in carried['qualifications']))
        point=self.source['daily_series']['JPY']['points'][1]
        self.assertEqual(point['date'],'2024-01-03');self.assertEqual(point['source_date'],'2024-01-02')

    def test_estimates_retain_source_vintage_after_calendar_end(self):
        for start,end in [('2024-07-01','2024-09-30'),('2024-10-01','2024-12-31')]:
            result=fx_for_period(self.source,'CHF','USD',start,end)
            self.assertEqual(result['status'],'estimate');self.assertEqual(result['value'],2.5)
        self.assertEqual(self.source['metadata']['actual_through'],'2024-06-30')
        self.assertEqual(self.source['metadata']['projected_publication_dates'],{'2024-Q3':['2024-09-30']})
        combined=fx_for_period(self.source,'CHF','USD','2024-01-01','2024-12-31')
        self.assertEqual(combined['status'],'missing');self.assertIsNone(combined['value'])
        self.assertIn('2024-Q4',combined['qualifications'][-1])

    def test_explicit_projected_calendar_supports_aggregate_estimate(self):
        source=load_fx_source(self.path,adapter={'projected_publication_dates':{'2024-Q4':['2024-10-01','2024-12-31']},
            'projection_calendar_assumptions':{'2024-Q4':{'description':'SYNTHETIC explicit two-date projection; not an observed calendar.'}}})
        result=fx_for_period(source,'CHF','USD','2024-01-01','2024-12-31')
        self.assertEqual(result['status'],'estimate')
        self.assertAlmostEqual(result['value'],(0.5+1+1.5+2+2.5+2.5+2.5+2.5)/8)
        self.assertEqual(result['coverage']['expected_days'],8)
        self.assertEqual(source['observations'],self.source['observations'])
        self.assertEqual(source['metadata']['source_projected_publication_dates'],{'2024-Q3':['2024-09-30']})
        self.assertEqual(source['metadata']['projection_calendar_provenance']['2024-Q4']['origin'],'adapter_configuration')
        self.assertEqual(source['metadata']['projection_calendar_provenance']['2024-Q3']['origin'],'source_workbook')
        self.assertTrue(any('configured projected publication calendar' in note for note in result['qualifications']))
        self.assertTrue(any('SYNTHETIC explicit two-date projection' in note for note in result['qualifications']))
        with self.assertRaisesRegex(ValueError,'must not be empty'):
            load_fx_source(self.path,adapter={'projected_publication_dates':{'2024-Q4':[]}})

    def test_missing_past_future_currency_and_closing_are_not_filled(self):
        for currency,start,end in [('USD','2023-01-01','2023-12-31'),('USD','2025-01-01','2025-12-31'),('CAD','2024-01-01','2024-06-30')]:
            result=fx_for_period(self.source,'CHF',currency,start,end)
            self.assertEqual(result['status'],'missing');self.assertIsNone(result['value'])
        closing=fx_for_period(self.source,'CHF','USD','2024-01-01','2024-03-28',kind='closing')
        self.assertEqual(closing['value'],1.5)
        for end in ['2024-03-31','2024-09-30','2024-12-31']:
            self.assertEqual(fx_for_period(self.source,'CHF','USD','2024-01-01',end,kind='closing')['status'],'missing')

    def test_hash_original_bytes_and_observations_are_preserved(self):
        before=self.path.read_bytes()
        source=load_fx_source(self.path)
        self.assertEqual(source['source']['sha256'],hashlib.sha256(before).hexdigest())
        self.assertEqual(before,self.path.read_bytes())
        result=fx_for_period(source,'CHF','USD','2024-01-01','2024-03-31')
        result['provenance'][-1]['cell']='changed by caller'
        self.assertNotEqual(source['observations'][0]['provenance'][0]['cell'],'changed by caller')

    def test_bad_quote_convention_and_inconsistent_cached_quarter_are_visible(self):
        book=openpyxl.load_workbook(self.path);book['CHF reporting'].cell(8,4,999)
        book.save(self.path);book.close()
        source=load_fx_source(self.path)
        result=fx_for_period(source,'CHF','USD','2024-01-01','2024-03-31')
        self.assertEqual(result['status'],'missing')
        self.assertIn('disagrees',result['qualifications'][-1])
        book=openpyxl.load_workbook(self.path);book['CHF reporting'].cell(2,2,'Undocumented inverse quote')
        book.save(self.path);book.close()
        with self.assertRaisesRegex(ValueError,'Quote direction'):
            load_fx_source(self.path)

    def test_partial_source_coverage_is_labelled_without_filling_missing_quotes(self):
        book=openpyxl.load_workbook(self.path)
        book['Synthetic daily quotes'].cell(8,3).value=None
        # CHF per JPY Q1 uses only the two supported daily crosses .02 and .03.
        book['CHF reporting'].cell(9,4,.025)
        book.save(self.path);book.close()
        source=load_fx_source(self.path)
        result=fx_for_period(source,'CHF','JPY','2024-01-01','2024-03-31')
        self.assertAlmostEqual(result['value'],.025)
        self.assertEqual(result['coverage'],dict(available_days=2,expected_days=3,complete=False))
        self.assertTrue(any('Partial source coverage' in note for note in result['qualifications']))

    def test_configured_sheet_name_and_missing_formula_cache_remain_explicit(self):
        book=openpyxl.load_workbook(self.path);book['CHF reporting'].title='Rates for company reporting'
        book['Rates for company reporting'].cell(8,4,'=1')
        book.save(self.path);book.close()
        source=load_fx_source(self.path,adapter={'reporting_sheets':{'CHF':'Rates for company reporting','GBP':'GBP reporting'}})
        result=fx_for_period(source,'CHF','USD','2024-01-01','2024-03-31')
        self.assertEqual(result['status'],'missing');self.assertIsNone(result['value'])
        self.assertIn('saved formula cache',result['qualifications'][-1])
        self.assertEqual(result['provenance'][-1]['formula'],'=1')

    def test_duplicate_periods_and_out_of_period_projection_are_rejected(self):
        with self.assertRaisesRegex(ValueError,'outside its declared quarter'):
            load_fx_source(self.path,adapter={'projected_publication_dates':{'2024-Q4':['2025-01-01']}})
        book=openpyxl.load_workbook(self.path);book['CHF reporting'].cell(7,5,'Q1:24')
        book.save(self.path);book.close()
        with self.assertRaisesRegex(ValueError,'Duplicate quarter'):
            load_fx_source(self.path)


if __name__=='__main__':
    unittest.main()
