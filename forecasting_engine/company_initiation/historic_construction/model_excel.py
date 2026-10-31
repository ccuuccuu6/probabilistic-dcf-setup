"""Historical-stage workbook, without importing forecast or valuation code.

The saved historical projection owns the facts. Excel links those facts and
reproduces only the compatible historical derivations already in that projection.
"""
import json
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from company_initiation.shared_excel import (
    STYLE_VERSION, BLUE, PINK, INK, NAVY, PALE, _bar, _cell,
    _display_unit, _format, _number, _style_sheet, _table,
)
from company_initiation.shared_artifacts import register_stage_output
from company_initiation.shared_runtime import check, expression_tree, formula_excel, formula_value, write_json
from .model_contract import historical_model


_PERIOD_ORDER = {kind: index for index, kind in enumerate(('Q1', 'Q2', 'Q3', 'Q4', 'H1', 'H2', 'FY'))}


def _export_value_equal(actual, expected):
    """Allow Excel's numeric serialization precision; keep structure exact."""
    if isinstance(actual,bool) or isinstance(expected,bool):return type(actual)==type(expected) and actual==expected
    if isinstance(actual,(int,float)) and isinstance(expected,(int,float)):
        return math.isfinite(actual) and math.isfinite(expected) and math.isclose(actual,expected,rel_tol=5e-15,abs_tol=0)
    return actual==expected


def export_historical_model(path, spec, *, sources=None, review=None, incomplete_reason=None):
    """Write a history-only audit workbook and verify the saved presentation.

    This verifies saved values, formulas, visibility and stage boundaries. It
    does not claim desktop Excel recalculation or independent source acceptance.
    """
    path = Path(path)
    workbook = Workbook()
    model = workbook.active
    model.title = 'Model'
    workbook.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    _style_sheet(model)
    model.column_dimensions['A'].width = 3
    model.column_dimensions['B'].width = 52
    model.column_dimensions['C'].width = 15
    _cell(model, 2, 2, spec['company'], bold=True)
    _cell(model, 3, 2, 'Historical model through ' + spec['latest_reporting_date'], bold=True)
    _cell(model, 4, 2, ('INCOMPLETE: '+incomplete_reason) if incomplete_reason else
          'Reported and compatible derived history; missing values remain blank.')

    observations = spec.get('observations', [])
    if any(o.get('historical_display_eligible') and not o.get('eligible',True) for o in observations):
        note='PROVISIONAL vendor history; provenance/basis limitations remain visible. No calibration or independent-source acceptance.'
        _cell(model,4,2,('INCOMPLETE. '+note) if incomplete_reason else note)
    by_id = {observation['id']: observation for observation in observations}
    if len(by_id) != len(observations):
        raise ValueError('Historical workbook requires unique observation identities')
    if any(observation.get('status') not in ('actual', 'reported', 'derived', 'derived_actual')
           for observation in observations):
        raise ValueError('Historical workbook cannot contain forecast or estimated observations')
    if any(period['end'] > spec['latest_reporting_date'] or period.get('status') == 'forecast'
           for period in spec['periods']):
        raise ValueError('Historical workbook periods must end by the latest reporting date')
    positions = {observation['id']: index for index, observation in enumerate(observations, 6)}
    columns = ('Observation ID', 'Financial row', 'Period', 'Definition', 'Value', 'Unit', 'Status',
               'Historical derivation', 'Source references', 'Basis', 'Entity / scope', 'Published at',
               'Original/restated status','Definition status','Period basis status','Provisional display eligible',
               'Calibration eligible','Independent primary verified','Publication vintage')
    records = [[item['id'], item['row_id'], item['period_id'], item.get('definition', item.get('label')),
                item.get('value'), item.get('unit'), item.get('status'), item.get('expression'),
                item.get('source_refs', []), item.get('basis'),
                item.get('entity_id', item.get('scope')), item.get('published_at'),item.get('restatement'),
                item.get('definition_status'),item.get('basis_status'),item.get('historical_display_eligible',False),
                item.get('calibration_eligible',item.get('eligible',True)),item.get('direct_company_disclosure_verified',False),
                item.get('publication_vintage_status','unknown')]
               for item in observations]
    historics = _table(workbook, 'Historics', columns, records, {1: 38, 2: 30, 4: 48, 8: 30, 9: 60})
    expected = {}
    formulas = 0
    for item in observations:
        row = positions[item['id']]
        value = item.get('value')
        if item.get('expression'):
            tree = expression_tree(item['expression'])

            def binding(name, lag):
                if lag:
                    raise ValueError('Historical derivations require explicit observation bindings')
                oid = item.get('bindings', {}).get(name, {}).get('observation_id')
                if oid not in by_id:
                    raise ValueError('Missing historical derivation observation: ' + str(oid))
                return oid

            calculated = formula_value(tree, lambda name, lag: by_id[binding(name, lag)].get('value'))
            if not (_number(calculated) and _number(value)
                    and math.isclose(calculated, value, rel_tol=1e-10, abs_tol=1e-8)):
                raise ValueError('Historical formula does not reproduce the saved observation: ' + item['id'])
            value = '=' + formula_excel(tree, lambda name, lag: 'E' + str(positions[binding(name, lag)]))
            formulas += 1
        cell = _cell(historics, row, 5, value, color=INK if item.get('expression') else BLUE,
                     number_format=_format(item.get('unit', '')))
        expected[('Historics', cell.coordinate)] = value

    rows = [row for row in spec['rows'] if row.get('cells') or row.get('show_when_missing')]
    used_periods = {pid for row in rows for pid in row['cells']}
    periods = sorted((period for period in spec['periods']
                      if period['id'] in used_periods and period['kind'] in _PERIOD_ORDER),
                     key=lambda period: (int(period['fiscal_year']), _PERIOD_ORDER[period['kind']]))
    for column, period in enumerate(periods, 4):
        _cell(model, 5, column, str(period['fiscal_year']) + ' ' + period['kind'], bold=True,
              color='FFFFFF' if period['kind'] == 'FY' else INK,
              fill=NAVY if period['kind'] == 'FY' else PALE)
        model.column_dimensions[get_column_letter(column)].width = 16
    _cell(model, 5, 2, 'Financial line', bold=True)
    _cell(model, 5, 3, 'Unit', bold=True)
    row_number = 6
    sections = list(dict.fromkeys(row.get('section', 'historical') for row in rows))
    model_positions={};next_row=row_number
    for section in sections:
        next_row+=1
        for item in (r for r in rows if r.get('section','historical')==section):
            model_positions[item['id']]=next_row;next_row+=1
        next_row+=1
    for section in sections:
        _bar(model, row_number, section.replace('segment_', 'Segment: ').replace('_', ' ').title(),
             max(4, len(periods) + 3))
        row_number += 1
        for row in (item for item in rows if item.get('section', 'historical') == section):
            _cell(model, row_number, 2, row.get('label', row['id']))
            _cell(model, row_number, 3, _display_unit(row.get('unit', '')))
            for column, period in enumerate(periods, 4):
                fact = row['cells'].get(period['id'])
                if fact is None:
                    continue
                oid = fact.get('observation_id')
                value = None
                if fact.get('model_expression'):
                    tree=expression_tree(fact['model_expression'])
                    by_row={r['id']:r for r in rows}
                    def model_binding(name,lag):
                        if lag:raise ValueError('Reporting totals require the same fiscal period')
                        return fact['model_bindings'][name]
                    calculated=formula_value(tree,lambda name,lag:by_row[model_binding(name,lag)]['cells'][period['id']]['value'])
                    if not (_number(calculated) and math.isclose(calculated,fact['value'],rel_tol=5e-15,abs_tol=0)):
                        raise ValueError('Reporting total differs from its source-bound segment inputs')
                    value='='+formula_excel(tree,lambda name,lag:get_column_letter(column)+str(model_positions[model_binding(name,lag)]))
                elif fact.get('value') is not None:
                    if oid not in positions:
                        raise ValueError('Model cell has no retained historical observation: ' + str(oid))
                    if by_id[oid].get('value') != fact['value']:
                        raise ValueError('Model cell differs from its historical observation: ' + str(oid))
                    value = "='Historics'!E" + str(positions[oid])
                cell = _cell(model, row_number, column, value, color=PINK,
                             number_format=_format(row.get('unit', '')))
                cell.comment = Comment(json.dumps(dict(status=fact.get('status'), observation_id=oid,
                    model_expression=fact.get('model_expression'),model_bindings=fact.get('model_bindings'),
                    source_refs=fact.get('source_refs', []), conflicts=fact.get('observation_ids', [])),
                    ensure_ascii=False, indent=2), 'Historical construction')
                expected[('Model', cell.coordinate)] = value
            row_number += 1
        row_number += 1
    model.print_area = f'B2:{get_column_letter(max(4, len(periods) + 3))}{max(6, row_number - 1)}'

    source_items = list(sources.values()) if isinstance(sources, dict) else list(sources or [])
    _table(workbook, 'Sources', ('ID', 'Title', 'File / URL', 'SHA256', 'Published at'),
           [[source.get('id'), source.get('title', source.get('name')),
             source.get('local_path', source.get('path', source.get('url'))),
             source.get('sha256'), source.get('published_at')] for source in source_items],
           {2: 50, 3: 70, 4: 70})
    checks = list(spec.get('derivation_checks', []))
    native=spec.get('native_source_rows',[])
    if native:
        from .evidence import parse_number
        raw=workbook.create_sheet('VA Source');_style_sheet(raw)
        _cell(raw,2,2,'Imported VA actuals in native source units; mapping/definition acceptance is separate.',bold=True)
        pmap={p['id']:p for p in spec['periods']}
        pmap.update({c['period_id']:c['period_metadata'] for r in native for c in r['cells']})
        pids=sorted({c['period_id'] for r in native for c in r['cells']},key=lambda pid:(pmap[pid]['end'],pid))
        for col,title in enumerate(['Source row','Label','Native unit']+pids,1):_cell(raw,5,col,title,bold=True)
        raw.column_dimensions['B'].width=55
        for row_no,item in enumerate(native,6):
            _cell(raw,row_no,1,item['source_id']+' / '+item['unit_id']);_cell(raw,row_no,2,item['label']);_cell(raw,row_no,3,item['reported_unit'])
            cells={c['period_id']:c for c in item['cells']}
            for col,pid in enumerate(pids,4):
                if pid not in cells:continue
                source=cells[pid];value=parse_number(source['source_raw'])
                cell=_cell(raw,row_no,col,value,number_format=source.get('display_format') or 'General')
                expected[('VA Source',cell.coordinate)]=value
        raw.freeze_panes='D6'
    if incomplete_reason:checks.append(check('historical_construction','completion','MISSING',incomplete_reason))
    checks.extend(check('historical_source_review', item.get('topic', str(index)), item.get('status', 'unknown'),
                        item.get('detail', '')) for index, item in enumerate((review or {}).get('checks', [])))
    checks.append(check('workbook', 'acceptance', 'INFO',
        'Historical audit export only. Desktop Excel recalculation and independent source acceptance are separate gates.'))
    _table(workbook, 'Validation', ('Area', 'Check', 'Status', 'Detail', 'Difference', 'Tolerance'),
           [[item.get(key) for key in ('area', 'name', 'status', 'detail', 'difference', 'tolerance')]
            for item in checks], {2: 40, 4: 95})
    if spec.get('entity_versions'):
        _table(workbook, 'Reporting versions', ('Identity', 'Reporting version'),
               [[item.get('id'), item] for item in spec['entity_versions']], {2: 100})
    if spec.get('reporting_bridges'):
        _table(workbook, 'Reporting bridges', ('Identity', 'Disclosed bridge'),
               [[item.get('id'), item] for item in spec['reporting_bridges']], {2: 100})
    workbook.active = 0
    model.sheet_view.tabSelected = True
    temporary=path.with_name('.history_export.tmp.xlsx')
    workbook.save(temporary)
    workbook.close()

    saved = load_workbook(temporary, data_only=False)
    try:
        mismatches = [f'{sheet}!{address}' for (sheet, address), value in expected.items()
                      if not _export_value_equal(saved[sheet][address].value,value)]
        layout_errors=[]
        if saved.active.title != 'Model' or any(sheet.sheet_state != 'visible' for sheet in saved):
            layout_errors.append('Historical workbook opening view or sheet visibility changed')
        if any(str(saved['Model'].cell(5, column).value).endswith(' 9M')
               for column in range(4, saved['Model'].max_column + 1)):
            layout_errors.append('Historical Model must not display cumulative 9M columns')
        sheets = list(saved.sheetnames)
    finally:
        saved.close()
    temporary.replace(path)
    return dict(stage='historical', status='exported', style_version=STYLE_VERSION,
                workbook=path.name, latest_reporting_date=spec['latest_reporting_date'],
                sheets=sheets, visible_periods=[period['id'] for period in periods],
                observations=len(observations), historical_formulas=formulas,
                saved_value_and_formula_readback='failed' if mismatches else 'passed',
                saved_layout='failed' if layout_errors else 'passed',mismatches=mismatches,layout_errors=layout_errors,
                numeric_rel_tolerance=5e-15,numeric_abs_tolerance=0,
                source_review_status=(review or {}).get('status', 'not_recorded'),
                desktop_excel_recalculation='not_performed', checks=checks)


def write_historical_output(directory, data, *, incomplete=False):
    """Write a historical checkpoint locally, without a provider or run journal."""
    directory = Path(directory)
    spec = historical_model(data)
    audit=data.get('model_configuration',{}).get('va_unresolved_series',[])
    if audit:write_json(directory/'va_unresolved_series.json',audit)
    stem='historical_model_incomplete' if incomplete else 'historical_model'
    validation_name=stem+'_workbook_validation.json' if incomplete else 'historical_workbook_validation.json'
    write_json(directory / (stem+'.json'), spec)
    path = directory / (stem+'.xlsx')
    try:
        validation = export_historical_model(path, spec, sources=data.get('sources'),
            review=data.get('combined_history_review'),
            incomplete_reason=data['model_configuration'].get('historical_incomplete_reason','Historical work is unfinished') if incomplete else None)
    except Exception as exc:
        validation=dict(status='failed',error=str(exc),saved_value_and_formula_readback='failed',saved_layout='unknown')
    write_json(directory / validation_name, validation)
    failed=validation.get('saved_value_and_formula_readback')!='passed' or validation.get('saved_layout')!='passed'
    return dict(
        status='failed' if failed else 'incomplete' if incomplete else 'exported', workbook=path.name, model=stem+'.json',
        validation=validation_name,
        files=[name for name in (stem+'.json',path.name,validation_name,*(['va_unresolved_series.json'] if audit else [])) if (directory/name).is_file()],
        latest_reporting_date=spec['latest_reporting_date'])


def save_historical_output(journal, data, *, incomplete=False):
    """Save and hash a stage-one checkpoint in the current run, never its parent."""
    output = write_historical_output(journal.directory, data,incomplete=incomplete)
    register_stage_output(journal, 'historical', output)
    if output['status']=='failed':raise ValueError('Historical workbook validation failed; inspect '+str(Path(journal.directory)/output['validation']))
    return Path(journal.directory) / output['workbook']
