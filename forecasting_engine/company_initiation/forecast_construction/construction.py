"""Local, cheap-first model construction. No generated Python is executed.

The saved specification is the authority for both Python and Excel. Historical
acceptance remains distinct from a provisional, executable forecasting framework.
"""
from __future__ import annotations

import ast
import copy
import hashlib
import json
import math
import os
from pathlib import Path
import re
from datetime import date, datetime, timezone
from contextlib import nullcontext
from time import monotonic, sleep
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from company_initiation.shared_ai import OpenAIProvider, AIError, DEFAULT_MODEL, aggregate_usage, response_usage, format_usage
from company_initiation.historic_construction.evidence import SourceStore, inventory_local_sources, parse_number
from company_initiation.historic_construction.initiation import _object, _array, _enum, TEXT, INTEGER, UNIT_REF, NUMBER_REF, PERIOD, validate_shape, _initiation_lock
from company_initiation.shared_artifacts import code_snapshot
from company_initiation.shared_runtime import (BLIND, HISTORY, HISTORY_EQUATIONS, HISTORY_PROMPT, HISTORY_SCOPE, IDENTITY, MAPPING_BATCH, MAPPING_ROW, PERIOD_SCHEMA, REFS, REVIEW, RULES, RunJournal, SELECT, SELECT_PROMPT, VERSION, _LOADED_CODE_IDENTITY, _current_code_identity, check, dependencies, digest, dimension, expression_tree, formula_excel, formula_value, response_output_limit, saved_output_limit_matches, unit_dimension, verify_source_bytes, write_json)










DRIVER = _object(id=TEXT, label=TEXT, definition=TEXT, kind=_enum(
    ("business", "accounting_ratio", "accounting_identity")), measure=_enum(
    ("level", "growth_rate", "ratio", "index")), unit=TEXT, scale={"type": "number"},
    frequency=_enum(("annual",)), entity=TEXT, materiality=TEXT, priority=_enum(("high", "medium", "low")),
    relationship=_enum(("disclosed", "derived", "hypothesis")), references=REFS,
    historical_row=TEXT, limitations=_array(TEXT),
    effects=_array(_object(line=TEXT, direction={"type": "integer", "enum": [-1, 0, 1]}, lag=INTEGER,
                           explanation=TEXT)),
    components=_array(TEXT), baseline=_object(rule=_enum(("zero", "last", "ratio", "missing")),
        numerator=TEXT, denominator=TEXT, explanation=TEXT))
DRIVERS = _object(drivers=_array(DRIVER), gaps=_array(TEXT))
FORECAST_ROW = _object(id=TEXT, label=TEXT, entity=TEXT, section=TEXT, unit=TEXT,
    scale={"type": "number"}, time_basis=_enum(("flow", "stock", "ratio", "average")),
    expression=TEXT, rationale=TEXT, references=REFS, simplifications=_array(TEXT))
FORMULAS = _object(rows=_array(FORECAST_ROW), identities=_array(IDENTITY), gaps=_array(TEXT))

DRIVER_PROMPT = RULES + """Choose material company/segment drivers using the whole
disclosures and supported historical model. Distinguish business drivers from
convenient accounting ratios. Each effect names an existing financial row (or a
clearly needed intermediate), expected direction (-1,0,1) and lag in ANNUAL periods.
Identify purchasing/hedging/inventory/pass-through limits for commodity exposures;
do not invent coefficients or lags. Unknown transmission is a hypothesis, not a
quantified verified relationship. Use components to name economic effects captured
by the driver (e.g. price, volume, mix, fx, acquisitions, commodity_cost). Organic
growth must include price and volume components; total margins include costs they
already capture. Components overlap is checked to prevent double counting.
Use annual temporary baselines only: zero for explicitly neutral GROWTH changes,
last for an existing historical row, ratio for numerator/denominator historical
rows, missing when evidence cannot initialise it. A missing physical volume/unit
cost must never become zero. baseline.explanation must justify this provisional
rule. historical_row is an existing history ID or empty. Unquantified drivers may
stay in the register with missing baselines. No assumed pass-through elasticities.
"""
FORMULA_PROMPT = RULES + """Build an annual executable specification for supported
financial lines, preserving history IDs. Use the driver interfaces supplied; never
invent additional assumptions in formulas. Every reference must be a row/driver ID.
Syntax: + - * / parentheses; lag("id", positive_integer) for prior ANNUAL periods;
literals 0 and 1 only. No Python, Excel functions, comparisons, powers or indexing.
lag uses annual periods, never the preceding interim column. Declare every forecast
line's units and display scale, rationale, evidence and simplifications. Revenue
growth conventions must follow company disclosures (additive vs multiplicative).
Prefer physical quantities/prices only if disclosed. Avoid organic plus price/volume,
mix twice, FX twice, M&A twice, or margin plus the same input-cost shock. Accounting
rollforwards may use opening debt/assets with explicitly documented conventions.
Create linked profit, cash flow, working capital, assets, financing and equity
schedules where supported. Missing schedules must be named, not plugged. No debt
or cash balancing plug. Baseline assumptions remain in drivers, never constants.
Use all supported historical identities in forecast when their components can be
modelled. Identities with no periods apply to every forecast year. Unsupported lines
remain absent with precise gaps. Drivers without quantified exposure stay unresolved.
"""




















def validate_spec(history, drivers, specification):
    """Structural failures stop execution; evidence gaps remain visible warnings."""
    issues = []
    rows, inputs = specification["rows"], drivers["drivers"]
    if not rows:
        return ["No executable forecast rows were supplied"]
    identifiers = [r["id"] for r in rows] + [d["id"] for d in inputs]
    if len(set(identifiers)) != len(identifiers):
        return ["Duplicate forecast row/driver IDs"]
    all_items = {r["id"]: r for r in history["rows"]}
    all_items.update({r["id"]: r for r in rows + inputs})
    units = {key: item["unit"] for key, item in all_items.items()}
    historic = {r["id"]: r for r in history["rows"]}
    for item in rows + inputs:
        if not re.fullmatch(r"[a-z][a-z0-9_]*", item["id"]):
            issues.append(f"Invalid ID: {item['id']}")
        if not math.isfinite(item["scale"]) or item["scale"] <= 0:
            issues.append(f"Invalid display scale: {item['id']}")
        if item["unit"] == "ratio" and item["scale"] != 1:
            issues.append(f"{item['id']}: ratios must use display scale 1; percentage formatting supplies the display conversion")
        if item["id"] in historic and any(item[k] != historic[item["id"]][k] for k in ("unit", "entity", "time_basis")):
            issues.append(f"Forecast changed history definition/unit/entity: {item['id']}")
    graph = {}
    for row in rows:
        try:
            tree = expression_tree(row["expression"])
            deps = dependencies(tree)
            unknown = {key for key, lag in deps if key not in all_items}
            if unknown:
                raise ValueError("Unknown dependencies: " + ", ".join(sorted(unknown)))
            current_missing = {key for key, lag in deps if lag == 0 and key not in identifiers}
            if current_missing:
                raise ValueError("Current-period reference has no forecast/input: " + ", ".join(sorted(current_missing)))
            expected = unit_dimension(row["unit"])
            if dimension(tree, units) != expected:
                raise ValueError("Formula result units disagree with row unit")
            graph[row["id"]] = [key for key, lag in deps if lag == 0]
            # Components are scoped by entity: price in two distinct segments is valid.
            used = [d for d in inputs if d["id"] in {key for key, lag in deps}]
            occupied = {}
            for driver in used:
                for component in driver["components"]:
                    key = (driver["entity"], component)
                    if key in occupied:
                        raise ValueError(f"Double-counted {component}: {occupied[key]} and {driver['id']}")
                    occupied[key] = driver["id"]
        except (ValueError, KeyError, RecursionError) as exc:
            issues.append(f"{row['id']}: {exc}")
    visiting, visited = set(), set()
    def visit(key):
        if key in visiting:
            raise ValueError("Circular same-period dependency: " + key)
        if key in visited:
            return
        visiting.add(key)
        for dep in graph.get(key, []):
            visit(dep)
        visiting.remove(key)
        visited.add(key)
    try:
        for key in graph:
            visit(key)
    except ValueError as exc:
        issues.append(str(exc))
    for driver in inputs:
        if driver["relationship"] != "hypothesis" and not driver["references"]:
            issues.append(f"{driver['id']}: sourced/derived relationship needs references")
        if not driver["baseline"]["explanation"] or not driver["materiality"]:
            issues.append(f"{driver['id']}: missing baseline rationale/materiality")
        if driver["baseline"]["rule"] == "zero" and driver["measure"] != "growth_rate":
            issues.append(f"{driver['id']}: neutral zero allowed only for growth changes")
        if any(effect["lag"] < 0 or effect["lag"] > 10 for effect in driver["effects"]):
            issues.append(f"{driver['id']}: invalid annual effect lag")
    for identity in specification["identities"]:
        try:
            tree = expression_tree(identity["expression"])
            if identity["target"] not in {r["id"] for r in rows}:
                raise ValueError("Reconciliation target has no forecast row")
            if not math.isfinite(identity["tolerance"]) or identity["tolerance"] < 0:
                raise ValueError("Reconciliation tolerance must be finite and nonnegative")
            if dimension(tree, units) != unit_dimension(units[identity["target"]]):
                raise ValueError("Reconciliation units disagree")
        except (ValueError, KeyError) as exc:
            issues.append(f"Identity {identity['id']}: {exc}")
    return issues


def initialise_assumptions(history, drivers, years):
    annual = [p for p in history["periods"] if p["kind"] == "annual"]
    latest = max(annual, key=lambda p: p["end"])["id"]
    rows = {row["id"]: row for row in history["rows"]}
    result = []
    def last(key):
        return rows.get(key, {}).get("values", {}).get(latest)
    for driver in drivers["drivers"]:
        rule = driver["baseline"]
        value = None
        if rule["rule"] == "zero":
            value = 0.0
        elif rule["rule"] == "last":
            value = last(driver["historical_row"] or rule["numerator"])
        elif rule["rule"] == "ratio":
            numerator, denominator = last(rule["numerator"]), last(rule["denominator"])
            if numerator is not None and denominator not in (None, 0):
                value = numerator / denominator
        for year in years:
            result.append(dict(driver=driver["id"], entity=driver["entity"], period=str(year),
                unit=driver["unit"], scale=driver["scale"], value=value,
                scenario="temporary_baseline", source="temporary_manual_assumption",
                provenance={"rule": rule, "history_period": latest},
                status="missing" if value is None else "temporary"))
    return result


def calculate(history, drivers, specification, assumptions, years):
    errors = validate_spec(history, drivers, specification)
    if errors:
        raise ValueError("; ".join(errors))
    rows = {row["id"]: row for row in specification["rows"]}
    inputs = {}
    for entry in assumptions:
        key = (entry["driver"], int(entry["period"]))
        if key in inputs:
            raise ValueError("Duplicate driver/period input")
        value = entry["value"]
        if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
            raise ValueError("Driver value must be finite or null")
        inputs[key] = value
    annual = {p["id"]: p["year"] for p in history["periods"] if p["kind"] == "annual"}
    values = {(r["id"], annual[p]): v for r in history["rows"] for p, v in r["values"].items() if p in annual}
    # A lagged driver can use its explicitly mapped historical observation only.
    history_by_id = {r["id"]: r for r in history["rows"]}
    for driver in drivers["drivers"]:
        for p, v in history_by_id.get(driver["historical_row"], {}).get("values", {}).items():
            if p in annual:
                values[driver["id"], annual[p]] = v
    trees = {key: expression_tree(row["expression"]) for key, row in rows.items()}
    evaluating = set()
    def resolve(key, year):
        pair = key, year
        if pair in inputs:
            return inputs[pair]
        if pair in values:
            return values[pair]
        if year not in years or key not in trees:
            return None
        if pair in evaluating:
            raise ValueError("Circular formula dependency")
        evaluating.add(pair)
        value = formula_value(trees[key], lambda dep, lag: resolve(dep, year - lag))
        evaluating.remove(pair)
        if value is not None and not math.isfinite(value):
            raise ValueError("Non-finite calculated result")
        values[pair] = value
        return value
    return {key: {str(year): resolve(key, year) for year in years} for key in rows}


def validate_results(history, drivers, specification, assumptions, years, values):
    checks = []
    historic_years = {p["id"]: str(p["year"]) for p in history["periods"] if p["kind"] == "annual"}
    lookup = {(r["id"], historic_years[p]): v for r in history["rows"] for p,v in r["values"].items() if p in historic_years}
    lookup.update({(key, year): v for key, row in values.items() for year,v in row.items()})
    lookup.update({(a["driver"], a["period"]): a["value"] for a in assumptions})
    for row in specification["rows"]:
        for year in years:
            missing = values[row["id"]][str(year)] is None
            checks.append(check("formula", f"{row['id']}:{year}", "MISSING" if missing else "PASS",
                "Unresolved input/lag or zero denominator" if missing else "Calculated from explicit dependencies"))
    for identity in specification["identities"]:
        try:
            tree = expression_tree(identity["expression"])
            if identity["periods"] and not set(identity["periods"]) & {str(year) for year in years}:
                checks.append(check("accounting", identity["id"], "MISSING", "Reconciliation does not apply to any forecast period"))
            for year in years:
                if identity["periods"] and str(year) not in identity["periods"]:
                    continue
                expected = formula_value(tree, lambda key, lag: lookup.get((key, str(year - lag))))
                actual = values.get(identity["target"], {}).get(str(year))
                difference = None if actual is None or expected is None else actual - expected
                status = "MISSING" if difference is None else "PASS" if abs(difference) <= identity["tolerance"] else "FAIL"
                checks.append(check("accounting", identity["id"] + ":" + str(year), status,
                    identity["rationale"], difference, identity["tolerance"]))
        except (ValueError, KeyError) as exc:
            checks.append(check("accounting", identity["id"], "FAIL", str(exc)))
    first = years[0]
    for driver in drivers["drivers"]:
        if not driver["effects"]:
            checks.append(check("behaviour", driver["id"], "MISSING", "No quantified directional/lag expectation; driver transmission remains unresolved"))
        baseline = next(a for a in assumptions if a["driver"] == driver["id"] and a["period"] == str(first))
        if baseline["value"] is None:
            checks.append(check("behaviour", driver["id"], "MISSING", "Sensitivity requires a quantified baseline"))
            continue
        modified = copy.deepcopy(assumptions)
        shock = max(abs(baseline["value"]) * .01, .01 * driver["scale"])
        next(a for a in modified if a["driver"] == driver["id"] and a["period"] == str(first))["value"] += shock
        changed = calculate(history, drivers, specification, modified, years)
        for effect in driver["effects"]:
            year = first + effect["lag"]
            if year not in years:
                checks.append(check("behaviour", driver["id"] + ":" + effect["line"], "MISSING", "Effect lag is outside configured horizon"))
                continue
            before = values.get(effect["line"], {}).get(str(year))
            after = changed.get(effect["line"], {}).get(str(year))
            delta = None if before is None or after is None else after - before
            ok = delta is not None and (abs(delta) <= 1e-9 if effect["direction"] == 0 else delta * effect["direction"] > 0)
            checks.append(check("behaviour", f"{driver['id']}:{effect['line']}:{year}",
                "MISSING" if delta is None else "PASS" if ok else "FAIL",
                f"First-year input shock {shock:g}; expected direction {effect['direction']}, lag {effect['lag']}. " + effect["explanation"], delta))
            for earlier in years:
                if earlier >= year:
                    break
                b = values.get(effect["line"], {}).get(str(earlier))
                a = changed.get(effect["line"], {}).get(str(earlier))
                if a is not None and b is not None and abs(a - b) > 1e-9:
                    checks.append(check("behaviour", f"{driver['id']}:{effect['line']}:timing", "FAIL", "Effect appears before its declared lag", a - b))
    return checks


def source_unit(label, token):
    """Only documented explicit unit strings; unsupported units remain errors."""
    label = label.strip()
    if label in ("%", "Percent", "percent"):
        return "ratio", 1 if "%" in token.get("display_format", "") or str(token["raw"]).strip().endswith("%") else .01
    if label.lower() in ("bps", "bp", "basis points"):
        return "ratio", .0001
    if label.lower() in ("x", "ratio", "multiple"):
        return "ratio", 1
    absolute = re.fullmatch(r"([A-Z]{3})Abs", label)
    if absolute:
        return absolute[1], 1
    if label == "M#":
        return "units", 1e6
    match = re.fullmatch(r"([A-Z]{3})([MBK])?", label)
    if match:
        return match[1], {None: 1, "K": 1e3, "M": 1e6, "B": 1e9}[match[2]]
    if label.lower() in ("shares", "units", "tonnes", "tonne", "index"):
        return label.lower(), 1
    if label.lower() in ("sharesm", "m shares", "million shares"):
        return "shares", 1e6
    raise ValueError("Unsupported source unit: " + label)


def metric_catalog(store, years):
    """Labels/context only for mapping; never repeatedly prompt a numerical grid."""
    result = []
    for source in store.sources:
        if source.get("category") != "visible_alpha":
            continue
        units = store._vendor_history(source, years)
        for unit in units:
            row = unit.get("vendor_row", {})
            result.append({"source_id": source["id"], "unit_id": unit["unit_id"],
                "text": unit["text"], "role": row.get("role"), "unit": row.get("unit"),
                "missing_markers": row.get("missing_markers", {}),
                "numeric_cells_available": len(unit["numbers"])})
    return result


def validate_references(items, store):
    """Ensure references exist; this does not certify semantic support."""
    if isinstance(items, dict):
        if "source_id" in items and "unit_id" in items:
            try:
                store.read([items])
                if "number_index" in items:
                    store.resolve(items)
            except ValueError as exc:
                raise ValueError(f"Invalid evidence reference {items}: {exc}") from exc
        for value in items.values():
            validate_references(value, store)
    elif isinstance(items, list):
        for value in items:
            validate_references(value, store)


def normalise_history(mapping, store, years):
    validate_shape(mapping, HISTORY)
    validate_references(mapping, store)
    history = copy.deepcopy(mapping)
    checks, observations, excluded = [], [], []
    identifiers = [r["id"] for r in mapping["rows"]]
    if len(identifiers) != len(set(identifiers)) or any(not re.fullmatch(r"[a-z][a-z0-9_]*", key) for key in identifiers):
        raise ValueError("Historical row IDs must be unique snake_case identifiers")
    periods = {p["id"]: p for p in history["periods"]}
    if len(periods) != len(mapping["periods"]):
        raise ValueError("Duplicate period IDs")
    annual = {p["year"]: p for p in periods.values() if p["kind"] == "annual"}
    if set(annual) != set(years):
        raise ValueError("Explicit annual periods must cover every requested fiscal year")
    for p in periods.values():
        if p["year"] not in years or date.fromisoformat(p["start"]) > date.fromisoformat(p["end"]):
            raise ValueError("Invalid historical fiscal period")
        anchor = annual[p["year"]]
        p.update(fiscal_year_start=anchor["start"], fiscal_year_end=anchor["end"])
        if not re.fullmatch(str(p["year"]) + (r"" if p["kind"] == "annual" else r"-(?:Q[1-4]|H[12]|9M)"), p["id"]):
            raise ValueError("Fiscal IDs must be YYYY, YYYY-Q1..Q4, YYYY-H1/H2 or YYYY-9M; do not rename vendor periods")
    for row in history["rows"]:
        if not math.isfinite(row["scale"]) or row["scale"] <= 0:
            raise ValueError("Invalid historical display scale")
        choices = {}
        row["values"] = {}
        row["origins"] = {}
        for ref in row["source_rows"]:
            source = next(s for s in store.sources if s["id"] == ref["source_id"])
            if source.get("category") != "visible_alpha":
                raise ValueError("Whole-row import requires a recognized vendor export; use token references for company documents")
            unit = store.read([ref])[0]
            vendor = unit.get("vendor_row", {})
            if vendor.get("role") != "metric" or vendor.get("section") != "company_reported":
                raise ValueError("Mapping selected excluded standardized/helper vendor row")
            headers = source["vendor_profile"]["header_sets"][vendor["header_id"]]["columns"]
            for token in unit["numbers"]:
                header = headers.get(re.sub(r"\d+$", "", token["locator"]))
                if not header or header["year"] not in years:
                    continue
                period_id = str(header["year"]) + ("" if header["slot"] == "FY" else "-" + header["slot"])
                number_ref = dict(ref, number_index=token["index"], number_format="dot_decimal")
                if not header["eligible"] or period_id not in periods or str(token["raw"]).strip() in ("-", "â€”", "â€“"):
                    excluded.append(dict(row=row["id"], reference=number_ref, locator=token["locator"],
                        period_label=header["period_label"], series=header["series"], raw=token["raw"],
                        reason="Estimate/consensus, missing marker, or no disclosed compatible fiscal period"))
                    continue
                provenance = store.validate_historical_reference(number_ref, periods[period_id])
                canonical, multiplier = source_unit(vendor["unit"], token)
                if canonical == "units" and row["unit"] == "shares" and "share" in vendor["label"].lower():
                    canonical = "shares"
                if (row["unit"] == canonical + "/shares" and
                        re.search(r"per share|\beps\b", vendor["label"], re.I)):
                    canonical = row["unit"]
                if canonical != row["unit"]:
                    raise ValueError(f"{row['id']}: mapped unit {row['unit']} disagrees with source {canonical}")
                value = store.resolve(number_ref) * multiplier * row["sign"]
                observation = dict(row=row["id"], period=period_id, value=value, raw=token["raw"],
                    reference=number_ref, source_unit=vendor["unit"], multiplier=multiplier,
                    sign=row["sign"], basis=row["basis"], original_label=vendor["label"], **provenance)
                observations.append(observation)
                choices.setdefault(period_id, []).append(observation)
        for item in row["observations"]:
            if item["period"] not in periods:
                raise ValueError("Observation has unknown period")
            store.validate_historical_reference(item["reference"], periods[item["period"]])
            source = next(s for s in store.sources if s["id"] == item["reference"]["source_id"])
            if source.get("category") == "visible_alpha":
                raise ValueError("Vendor observations must use deterministic whole-row import")
            if not math.isfinite(item["multiplier"]) or item["multiplier"] <= 0:
                raise ValueError("Invalid source scale multiplier")
            observation = dict(row=row["id"], period=item["period"], reference=item["reference"],
                raw=store.resolve(item["reference"]), value=store.resolve(item["reference"]) * item["multiplier"] * row["sign"],
                multiplier=item["multiplier"], sign=row["sign"], series="Company original", basis=row["basis"])
            choices.setdefault(item["period"], []).append(observation)
            observations.append(observation)
        for period, candidates in choices.items():
            rank = {"Company original": 0, "Filings": 1, "VA Actuals": 2}
            candidates.sort(key=lambda x: rank[x["series"]])
            best = [c for c in candidates if rank[c["series"]] == rank[candidates[0]["series"]]]
            conflict = len({c["value"] for c in best}) > 1
            if conflict:
                checks.append(check("data", row["id"] + ":" + period, "FAIL", "Conflicting equal-priority observations; value remains missing",
                                    max(c["value"] for c in best) - min(c["value"] for c in best)))
                continue
            row["values"][period] = best[0]["value"]
            row["origins"][period] = best[0]
            if len({c["value"] for c in candidates}) > 1:
                checks.append(check("data", row["id"] + ":" + period, "WARNING",
                    "Vendor/original disagreement retained; selected explicit source precedence. Reporting vintage needs verification.",
                    max(c["value"] for c in candidates) - min(c["value"] for c in candidates)))
        for period in periods:
            if period not in row["values"]:
                checks.append(check("coverage", row["id"] + ":" + period, "MISSING", "No compatible selected observation; no zero substitution"))
    by_id = {r["id"]: r for r in history["rows"]}
    units = {r["id"]: r["unit"] for r in history["rows"]}
    for identity in history["identities"]:
        tree = expression_tree(identity["expression"])
        target = by_id.get(identity["target"])
        if target is None or any(k not in by_id or lag for k, lag in dependencies(tree)):
            raise ValueError("Historical identity has an unknown or cross-period dependency")
        if identity["tolerance"] < 0 or not math.isfinite(identity["tolerance"]):
            raise ValueError("Invalid historical reconciliation tolerance")
        if dimension(tree, units) != unit_dimension(target["unit"]):
            raise ValueError("Historical identity has incompatible units")
        for period in identity["periods"] or list(periods):
            actual = target["values"].get(period)
            expected = formula_value(tree, lambda key, lag: by_id[key]["values"].get(period))
            diff = None if expected is None or actual is None else actual - expected
            checks.append(check("history_accounting", identity["id"] + ":" + period,
                "MISSING" if diff is None else "PASS" if abs(diff) <= identity["tolerance"] else "FAIL",
                identity["rationale"], diff, identity["tolerance"]))
    history.update(checks=checks, imported_observations=observations, excluded_observations=excluded)
    return history




def export_model(path, history, drivers, spec, assumptions, years, values, checks, sources, gaps):
    book = Workbook()
    book.remove(book.active)
    for name in ("Model", "Historics", "Drivers", "Assumptions", "Checks", "Sources", "Dependencies", "Gaps"):
        book.create_sheet(name)
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
    navy, blue, green, yellow = "18324F", "0000CC", "008050", "FFF2CC"
    formats = {"ratio": "0.0%;(0.0%);0.0%"}
    number_format = '#,##0.0;(#,##0.0);0.0'
    generated, locations, input_cells = set(), {}, {}
    historical = sorted(history["periods"], key=lambda p: (p["end"], p["start"]))
    annual = {p["year"]: p["id"] for p in historical if p["kind"] == "annual"}
    hist_rows = {r["id"]: r for r in history["rows"]}
    fore_rows = {r["id"]: r for r in spec["rows"]}
    all_rows = list(history["rows"]) + [r for r in spec["rows"] if r["id"] not in hist_rows]
    scale = {r["id"]: r["scale"] for r in all_rows}
    scale.update({d["id"]: d["scale"] for d in drivers["drivers"]})
    model = book["Model"]
    model.append([history["company"] + (" | TEMPORARY BASELINE" if spec["rows"] else " | HISTORICAL DRAFT - FORECAST STAGE INCOMPLETE")])
    model.append(["Provisional forecasts; inspect Checks and Gaps. Workbook edits do not update saved Python results or renew evidence review."])
    model.append(["Annual forecasts only. Historical cumulative and standalone periods overlap; never sum across columns. Blue inputs; green linked formulas; yellow editable assumptions."])
    model.append(["Line / schedule", "Entity", "Unit / display scale", *[p["id"] + " actual (" + p["kind"] + ")" for p in historical],
                  *[str(y) + " forecast TEMP" for y in years], "Stable ID", "Definition / rationale"])
    hist = book["Historics"]
    hist.append(["Line", "Entity", "Unit / display scale", *[p["id"] for p in historical], "ID", "Accounting basis", "Definition"])
    hloc = {}
    for row in history["rows"]:
        hist.append([row["label"], row["entity"], f"{row['unit']} / {row['scale']:g}",
                     *[row["values"].get(p["id"], None) / row["scale"] if row["values"].get(p["id"]) is not None else "MISSING" for p in historical],
                     row["id"], row["basis"], row["definition"]])
        for col, period in enumerate(historical, 4):
            cell = hist.cell(hist.max_row, col)
            cell.font = Font(name="Aptos", size=10, color=blue)
            cell.number_format = formats.get(row["unit"], number_format)
            cell.comment = Comment(json.dumps(row["origins"].get(period["id"], {"status": "missing"}), ensure_ascii=False), "Source provenance")
            hloc[row["id"], period["id"]] = cell.coordinate
    assumption_sheet = book["Assumptions"]
    assumption_sheet.append(["Driver", "Entity", "Unit / display scale", "Source", "Initialisation rule",
                            *[str(y) + " TEMP" for y in years], "ID", "Limitations"])
    for driver in drivers["drivers"]:
        entries = {a["period"]: a for a in assumptions if a["driver"] == driver["id"]}
        assumption_sheet.append([driver["label"], driver["entity"], f"{driver['unit']} / {driver['scale']:g}",
            "Temporary manual assumption", json.dumps(driver["baseline"], ensure_ascii=False),
            *[entries[str(y)]["value"] / driver["scale"] if entries[str(y)]["value"] is not None else None for y in years],
            driver["id"], "; ".join(driver["limitations"])])
        for col, year in enumerate(years, 6):
            cell = assumption_sheet.cell(assumption_sheet.max_row, col)
            cell.fill, cell.font = PatternFill("solid", fgColor=yellow), Font(name="Aptos", color=blue, size=10)
            cell.number_format = formats.get(driver["unit"], number_format)
            cell.comment = Comment("EDITABLE temporary input. Blank means missing, zero means an explicit zero. " + json.dumps(entries[str(year)], ensure_ascii=False), "Assumption")
            input_cells[driver["id"], year] = cell.coordinate
    for row in all_rows:
        model.append([row["label"], row["entity"], f"{row['unit']} / {row['scale']:g}",
                      *[None for _ in historical], *[None for _ in years], row["id"],
                      row.get("definition", row.get("rationale", ""))])
        if model.max_row == 5 or row["section"] != all_rows[model.max_row - 6]["section"]:
            model.cell(model.max_row, 1).value = row["section"].replace("_", " ").upper() + "\n" + row["label"]
            model.cell(model.max_row, 1).fill = PatternFill("solid", fgColor="DAE4F0")
            model.cell(model.max_row, 1).font = Font(name="Aptos", size=10, color=navy, bold=True)
            model.row_dimensions[model.max_row].height = 32
        for col, p in enumerate(historical, 4):
            locations[row["id"], p["id"]] = model.cell(model.max_row, col).coordinate
        for col, year in enumerate(years, 4 + len(historical)):
            locations[row["id"], str(year)] = model.cell(model.max_row, col).coordinate
    def put(sheet, coordinate, formula):
        cell = sheet[coordinate]
        cell.value = formula
        generated.add((sheet.title, coordinate))
        cell.font = Font(name="Aptos", size=10, color=green)
    def ref(key, year, lag=0):
        target = year - lag
        if (key, target) in input_cells:
            return f"('Assumptions'!{input_cells[key, target]}*{scale[key]:g})"
        period = str(target) if target in years else annual.get(target)
        if (key, period) in locations:
            return f"('Model'!{locations[key, period]}*{scale[key]:g})"
        driver = next((d for d in drivers["drivers"] if d["id"] == key), None)
        if driver and (driver["historical_row"], period) in locations:
            name = driver["historical_row"]
            return f"('Model'!{locations[name, period]}*{scale[name]:g})"
        return "NA()"
    def guarded(tree, resolver, display_scale=1):
        deps = list(dict.fromkeys(dependencies(tree)))
        references = [resolver(key, lag) for key, lag in deps]
        # COUNT on expressions would turn a blank into zero; check raw cells too.
        raw_refs = [re.sub(r"^\((.*)\*[-+0-9.e]+\)$", r"\1", r) for r in references]
        guard = "+".join("COUNT(" + ",".join(raw_refs[i:i+200]) + ")" for i in range(0, len(raw_refs), 200)) or "0"
        return f"=IF({guard}={len(raw_refs)},({formula_excel(tree, resolver)})/{display_scale:g},NA())"
    expected_cells = {}
    for row in all_rows:
        for p in historical:
            coordinate = locations[row["id"], p["id"]]
            if (row["id"], p["id"]) in hloc:
                raw = f"'Historics'!{hloc[row['id'], p['id']]}"
                put(model, coordinate, f"=IF(COUNT({raw})=1,{raw},NA())")
                historical_value = hist_rows[row["id"]]["values"].get(p["id"])
                if historical_value is not None:
                    expected_cells[coordinate] = historical_value / row["scale"]
            else:
                model[coordinate] = "MISSING"
        for year in years:
            coordinate = locations[row["id"], str(year)]
            if row["id"] in fore_rows:
                tree = expression_tree(fore_rows[row["id"]]["expression"])
                put(model, coordinate, guarded(tree, lambda k, lag: ref(k, year, lag), row["scale"]))
                value = values[row["id"]][str(year)]
                expected_cells[coordinate] = None if value is None else value / row["scale"]
                model[coordinate].comment = Comment(fore_rows[row["id"]]["expression"] + "\n" + fore_rows[row["id"]]["rationale"], "Shared specification")
            else:
                model[coordinate] = "UNMODELLED"
            model[coordinate].fill = PatternFill("solid", fgColor="E8EFF7")
        for p in [p["id"] for p in historical] + [str(y) for y in years]:
            model[locations[row["id"], p]].number_format = formats.get(row["unit"], number_format)
    book["Drivers"].append(["ID", "Name", "Kind", "Measure", "Unit", "Entity", "Priority", "Relationship", "Materiality", "Effects / annual lags", "Evidence", "Historical data", "Limitations"])
    for d in drivers["drivers"]:
        book["Drivers"].append([d["id"], d["label"], d["kind"], d["measure"], d["unit"], d["entity"], d["priority"], d["relationship"], d["materiality"],
            json.dumps(d["effects"]), json.dumps(d["references"]), d["historical_row"] or "Not available", "; ".join(d["limitations"])])
    book["Dependencies"].append(["ID", "Formula", "Units", "Entity", "Dependencies / annual lags", "Rationale", "Evidence", "Simplifications"])
    for row in spec["rows"]:
        book["Dependencies"].append([row["id"], row["expression"], row["unit"], row["entity"], json.dumps(dependencies(expression_tree(row["expression"]))),
            row["rationale"], json.dumps(row["references"]), "; ".join(row["simplifications"])])
    cs = book["Checks"]
    cs.append(["Area", "Check", "Frozen result", "Detail", "Python difference (base units)", "Tolerance", "Live difference (base units)", "Live result"])
    for c in checks:
        cs.append([c["area"], c["name"], c["status"], c["detail"], c["difference"], c["tolerance"]])
    expected_checks = {}
    for identity in history["identities"]:
        for period in historical:
            key = period["id"]
            if identity["periods"] and key not in identity["periods"]:
                continue
            cs.append(["Live historical accounting", identity["id"] + ":" + key,
                       "See frozen checks above", identity["rationale"], None, identity["tolerance"]])
            n = cs.max_row
            def historical_ref(name, lag):
                if lag or (name, key) not in locations:
                    return "NA()"
                return f"('Model'!{locations[name, key]}*{scale[name]:g})"
            expression = expression_tree(identity["target"] + " - (" + identity["expression"] + ")")
            put(cs, f"G{n}", guarded(expression, historical_ref))
            put(cs, f"H{n}", f'=IF(ISNUMBER(G{n}),IF(ABS(G{n})<=F{n},"PASS","FAIL"),"MISSING")')
            difference = formula_value(expression, lambda name, lag: hist_rows[name]["values"].get(key))
            expected_checks[f"G{n}"] = difference
            expected_checks[f"H{n}"] = "MISSING" if difference is None else "PASS" if abs(difference) <= identity["tolerance"] else "FAIL"
    for identity in spec["identities"]:
        tree = expression_tree(identity["expression"])
        for year in years:
            if identity["periods"] and str(year) not in identity["periods"]:
                continue
            cs.append(["Live accounting", identity["id"] + ":" + str(year), "See frozen checks above", identity["rationale"], None, identity["tolerance"]])
            n = cs.max_row
            expression = expression_tree(identity["target"] + " - (" + identity["expression"] + ")")
            put(cs, f"G{n}", guarded(expression, lambda k, lag: ref(k, year, lag)))
            put(cs, f"H{n}", f'=IF(ISNUMBER(G{n}),IF(ABS(G{n})<=F{n},"PASS","FAIL"),"MISSING")')
    book["Sources"].append(["ID", "Original file", "Title", "Publication (supplied, unverified)", "SHA256", "Frozen original", "URL", "Parser warnings"])
    for source in sources:
        book["Sources"].append([source["id"], source.get("local_path"), source["title"], source.get("published_at"), source["sha256"], source["path"], source["url"], json.dumps(source.get("extraction_warnings", []))])
    for period in historical:
        book["Sources"].append(["Fiscal period " + period["id"], period["kind"],
            period["start"] + " to " + period["end"], "Dates are historical mapping claims", None, None,
            json.dumps(period.get("references", [])), "Stocks refer to balance dates; flows to the stated duration"])
    book["Gaps"].append(["Status", "Issue / missing information"])
    for gap in gaps:
        book["Gaps"].append(["PROVISIONAL / MISSING", gap])
    for sheet in book:
        header = 4 if sheet.title == "Model" else 1
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "D5" if sheet.title == "Model" else "D2" if sheet.title in ("Historics", "Assumptions") else "B2"
        sheet.auto_filter.ref = f"A{header}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        sheet.print_title_rows = f"1:{header}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        for row in sheet:
            for cell in row:
                if (sheet.title, cell.coordinate) not in generated and isinstance(cell.value, str):
                    cell.data_type = "s"  # Untrusted labels cannot become workbook formulas.
                if cell.font.name != "Aptos":
                    font = copy.copy(cell.font)
                    font.name, font.size = "Aptos", 10
                    cell.font = font
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column <= 3 or sheet.title not in ("Model", "Historics", "Assumptions"))
        for cell in sheet[header]:
            cell.fill, cell.font = PatternFill("solid", fgColor=navy), Font(name="Aptos", color="FFFFFF", bold=True, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[header].height = 32
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 42 if col == 1 else 25 if col <= 3 else 19
        if sheet.title not in ("Model", "Historics", "Assumptions"):
            for col in range(4, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 45
        for text_value, colour in (("FAIL", "FCE4D6"), ("MISSING", "FFF2CC"), ("PASS", "E2F0D9")):
            sheet.conditional_formatting.add(sheet.dimensions, CellIsRule(operator="equal", formula=['"' + text_value + '"'], fill=PatternFill("solid", fgColor=colour)))
    for n in (1, 2, 3):
        model.merge_cells(start_row=n, start_column=1, end_row=n, end_column=model.max_column)
        model.row_dimensions[n].height = 28
    model["A1"].font = Font(name="Aptos", size=16, color=navy, bold=True)
    for i, row in enumerate(all_rows, 5):
        if row["id"] in {item["target"] for item in spec["identities"] + history["identities"]}:
            for cell in model[i]:
                cell.border = Border(top=Side(style="thin", color="A0A0A0"))
    book.save(path)
    book.close()
    checked = load_workbook(path, data_only=False)
    from openpyxl.formula.tokenizer import Tokenizer
    for sheet, coordinate in generated:
        Tokenizer(checked[sheet][coordinate].value)
    if set(checked.sheetnames) != {"Model", "Historics", "Drivers", "Assumptions", "Checks", "Sources", "Dependencies", "Gaps"}:
        raise ValueError("Workbook structure mismatch")
    checked.close()
    return dict(formula_count=len(generated), expected_model_cells=expected_cells, expected_check_cells=expected_checks,
                status="structure_and_formula_syntax_checked", recalculation="not_run")


def recalculate_workbook(path, inspection):
    """Optional real Excel calculation, bounded and hidden, on our generated book."""
    import subprocess
    import shutil
    if os.name != "nt":
        return dict(status="unavailable", detail="Automatic Excel recalculation is implemented for Windows COM only")
    import tempfile
    import time
    script = r'''$ErrorActionPreference='Stop'
$excel=$null
$book=$null
try {
 Add-Type -TypeDefinition 'using System; using System.Runtime.InteropServices; public class RecalcWindow { [DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(IntPtr h, out uint pid); }'
 $excel=New-Object -ComObject Excel.Application
 $ownedPid=[uint32]0
 [void][RecalcWindow]::GetWindowThreadProcessId([IntPtr]$excel.Hwnd,[ref]$ownedPid)
 [IO.File]::WriteAllText($env:RESEARCH_RECALC_PID,[string]$ownedPid)
 $excel.Visible=$false
 $excel.DisplayAlerts=$false
 $excel.AutomationSecurity=3
 $blank=$excel.Workbooks.Add()
 $excel.Calculation=-4135
 [IO.File]::WriteAllText($env:RESEARCH_RECALC_PHASE,'opening workbook')
 $book=$excel.Workbooks.Open($env:RESEARCH_RECALC_PATH,0,$false)
 [IO.File]::WriteAllText($env:RESEARCH_RECALC_PHASE,'calculating formulas')
 $excel.CalculateFullRebuild()
 [IO.File]::WriteAllText($env:RESEARCH_RECALC_PHASE,'saving caches')
 $book.Save()
 $book.Close($false)
 $book=$null
 $blank.Close($false)
 Write-Output 'RECALCULATED'
} finally {
 if ($null -ne $book) { $book.Close($false) }
 if ($null -ne $excel) { $excel.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) }
}'''
    started=time.monotonic()
    with tempfile.TemporaryDirectory(prefix='refractive_recalculation_') as temporary:
        temporary=Path(temporary);copied=temporary/'model.xlsx';pid_file=temporary/'owned_excel_pid';phase_file=temporary/'phase'
        shutil.copy2(path,copied)
        env=dict(os.environ,RESEARCH_RECALC_PATH=str(copied),RESEARCH_RECALC_PID=str(pid_file),RESEARCH_RECALC_PHASE=str(phase_file))
        try:
            completed=subprocess.run(['powershell','-NoProfile','-NonInteractive','-Command',script],
                env=env,capture_output=True,text=True,timeout=600 if inspection.get('periods') else 120,
                creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        except (OSError,subprocess.TimeoutExpired) as exc:
            # Terminate only the automation instance identified from this call's
            # newly created Excel window, never an existing interactive workbook.
            if pid_file.exists():
                owned_pid=int(pid_file.read_text())
                cleanup=f"$p=Get-CimInstance Win32_Process -Filter 'ProcessId={owned_pid}'; if ($p -and $p.Name -eq 'EXCEL.EXE' -and $p.CommandLine -like '*/automation*') {{ Stop-Process -Id {owned_pid} }}"
                subprocess.run(['powershell','-NoProfile','-NonInteractive','-Command',cleanup],capture_output=True,timeout=20,
                    creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            return dict(status='unavailable',detail=type(exc).__name__,phase=phase_file.read_text() if phase_file.exists() else 'starting Excel',duration_seconds=time.monotonic()-started)
        if completed.returncode or 'RECALCULATED' not in completed.stdout:
            return dict(status='unavailable',detail='Excel COM could not recalculate: '+completed.stderr[-1000:],
                phase=phase_file.read_text() if phase_file.exists() else 'starting Excel',duration_seconds=time.monotonic()-started)
        shutil.copy2(copied,path)
    workbook = load_workbook(path, data_only=True)
    differences = []
    compared = 0
    for sheet, cells in (("Model", inspection["expected_model_cells"]),
                         ("Assumptions", inspection.get("expected_assumption_cells", {})),
                         ("FX", inspection.get("expected_fx_cells", {})),
                         ("Historics", inspection.get("expected_history_cells", {})),
                         ("Valuation", inspection.get("expected_valuation_cells", {})),
                         ("Checks", inspection.get("expected_check_cells", {}))):
        for coordinate, expected in cells.items():
            compared += 1
            actual = workbook[sheet][coordinate].value
            if expected is None:
                matches = actual is None if inspection.get("periods") else actual in ("#N/A", "#DIV/0!", "#VALUE!")
            elif isinstance(expected, str):
                matches = actual == expected
            else:
                matches = isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-7)
            if not matches:
                differences.append(dict(sheet=sheet, cell=coordinate, python=expected, excel=actual))
    workbook.close()
    return dict(status="passed" if not differences else "failed", engine="Microsoft Excel COM",
                compared_cells=compared, differences=differences,duration_seconds=time.monotonic()-started)


def construct_company(root, company, source_directory, *, years=None, historical_years=10,
                      forecast_years=5, budget_usd=20., max_calls=40, max_repairs=2,
                      max_output_tokens=12000, max_generated_tokens=200000,
                      max_retries=0,
                      fresh=False, provider=None, progress=None, recalculate=False,
                      reuse_spec=None, assumptions_file=None, escalation_file=None,
                      continue_from=None):
    root, source_directory = Path(root).resolve(), Path(source_directory)
    if not source_directory.is_absolute():
        source_directory = root / source_directory
    if not company.strip() or not 0 < budget_usd <= 20 or not math.isfinite(budget_usd):
        raise ValueError("Company is required; cumulative budget must be >0 and <=20 USD")
    if (not 1 <= historical_years <= 30 or not 1 <= forecast_years <= 10 or max_calls < 1
            or not 0 <= max_repairs <= 10 or not 0 <= max_retries <= 3
            or not 1 <= max_output_tokens <= 32000 or max_generated_tokens < max_output_tokens):
        raise ValueError("Invalid construction horizon/call/repair/output-token limits")
    if years is not None and (not years or years != sorted(set(years)) or any(type(y) is not int or not 1900 <= y <= 2200 for y in years)):
        raise ValueError("Years must be distinct, increasing fiscal years")
    provider = provider or OpenAIProvider.from_env(root / ".env")
    provider.max_output_tokens = max_output_tokens
    provider.validate_configuration()
    inventory = inventory_local_sources(source_directory)
    configuration = dict(company=company, sources=str(source_directory.resolve()), manifest=inventory["manifest_sha256"],
        years=years, historical_years=historical_years, forecast_years=forecast_years,
        provider=provider.metadata(), schema=VERSION, code=_current_code_identity(),
        limits=dict(usd=budget_usd, calls=max_calls, repairs=max_repairs, retries=max_retries, generated_tokens=max_generated_tokens),
        reuse_spec=json.loads(Path(reuse_spec).read_text(encoding="utf-8")) if reuse_spec else None,
        assumptions=json.loads(Path(assumptions_file).read_text(encoding="utf-8")) if assumptions_file else None,
        escalations=json.loads(Path(escalation_file).read_text(encoding="utf-8")) if escalation_file else {},
        recalculate=recalculate)
    outputs = root / "outputs"
    with _initiation_lock(outputs):
        previous = None
        matching = []
        company_runs = []
        from ..shared_budget import build_record_paths, record_in_build_scope
        for file in build_record_paths(outputs,"construct_*/record.json"):
            saved = json.loads(file.read_text(encoding="utf-8"))
            if not record_in_build_scope(saved):continue
            config = saved.get("configuration", {})
            if config.get("company", "").strip().casefold() == company.strip().casefold():
                company_runs.append((file, saved))
            if (config.get("company") == company and config.get("years") == years
                    and config.get("historical_years") == historical_years):
                matching.append((file, saved))
        if company_runs:
            _, latest_company = max(company_runs, key=lambda item: item[0].stat().st_mtime_ns)
            if latest_company.get("pending_call") or not aggregate_usage(latest_company["calls"])["cost_is_complete"]:
                raise AIError("Previous construction has unknown call charges; --fresh or changed years cannot reset this guard. Use unpaid review-saved for inspection")
        if continue_from:
            if fresh:
                raise ValueError("Choose --fresh or explicit --continue-from, not both")
            selected = Path(continue_from).resolve()
            if selected.parent != outputs or not (selected / "record.json").is_file():
                raise ValueError("Continuation must name a preserved construction run in outputs")
            previous = json.loads((selected / "record.json").read_text(encoding="utf-8"))
            comparable_old = {k: v for k, v in previous["configuration"].items() if k not in ("code", "schema")}
            comparable_new = {k: v for k, v in configuration.items() if k not in ("code", "schema")}
            comparable_old = copy.deepcopy(comparable_old)
            comparable_old["limits"].setdefault("retries", 0)
            if digest(comparable_old) != digest(comparable_new):
                raise ValueError("Explicit code migration must retain source inventory, model settings and original limits")
            if matching and max(matching, key=lambda item: item[0].stat().st_mtime_ns)[0].parent != selected:
                raise ValueError("Continue the newest matching run; branching a paid budget is prohibited")
            if progress:
                progress("Explicit code migration: preserving parent and all charges; only identical saved requests are reused.")
        elif matching and not fresh:
            file, previous = max(matching, key=lambda item: item[0].stat().st_mtime_ns)
            if previous.get("pending_call") or not aggregate_usage(previous["calls"])["cost_is_complete"]:
                raise AIError("Previous construction has unknown call charges; no further spending. Saved work can be exported without AI using review-saved")
            if digest(previous["configuration"]) != digest(configuration):
                raise ValueError("Construction settings/source/code changed; use --fresh explicitly. Previous budget and evidence are preserved")
        if previous:
            file = Path(previous["directory"]) / "record.json"
            journal_path = file.parent / "calls.jsonl"
            # JSONL records are LF-delimited. Unicode paragraph/line separators
            # inside a valid JSON string are evidence text, not record breaks.
            saved_calls = [json.loads(line) for line in journal_path.read_text(encoding="utf-8").split('\n') if line.strip()] if journal_path.exists() else []
            if digest(saved_calls) != digest(previous["calls"]):
                # Explicit recovery of the early construction mutable-request bug.
                # The fsynced per-call journal is authoritative only if its original
                # signature verifies AND every non-request field/charge agrees.
                recoverable = (continue_from and len(saved_calls) == len(previous["calls"])
                    and all(digest(a.get("request")) == a.get("signature")
                            and digest({k:v for k,v in a.items() if k != "request"}) ==
                                digest({k:v for k,v in b.items() if k != "request"})
                            for a,b in zip(saved_calls, previous["calls"])))
                if not recoverable:
                    raise ValueError("Saved call journal disagrees with record; damaged or uncommitted charges block spending")
                previous = copy.deepcopy(previous)
                previous["calls"] = saved_calls
                previous["journal_recovery"] = dict(parent=str(file.parent),
                    reason="Recovered exact immutable requests from signature-verified durable journal; every result, response and usage field agrees. Parent unchanged.")
                if progress:
                    progress(previous["journal_recovery"]["reason"])
            for name, expected in previous.get("artifacts", {}).items():
                artifact = file.parent / name
                if not artifact.is_file() or hashlib.sha256(artifact.read_bytes()).hexdigest() != expected:
                    raise ValueError("Saved construction artifact is missing or changed: " + name)
            verify_source_bytes(previous.get("sources", []), outputs)
            if previous.get("pending_call") or not aggregate_usage(previous["calls"])["cost_is_complete"]:
                raise AIError("Previous construction has unknown call charges; no further spending")
            if previous["status"] in ("complete_provisional", "incomplete_model") and not continue_from:
                return file.parent
        slug = re.sub(r"[^a-z0-9]+", "_", company.lower()).strip("_")[:35]
        directory = outputs / f"construct_{slug}_{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}_{uuid4().hex[:8]}"
        directory.mkdir(parents=True)
        journal = RunJournal(directory, provider, configuration, progress, previous)
        journal.artifact("code_snapshot.json", code_snapshot())
        journal.artifact("inventory.json", inventory)
        store = SourceStore(directory / "sources")
        try:
            if previous:
                for source in previous.get("sources", []):
                    store.restore(source, archive_root=outputs)
            def add(source):
                if not any(s["id"] == source["id"] for s in store.sources):
                    store.add_local(source, root=Path(inventory["root"]))
                journal.state["sources"] = store.sources
                journal.save()
            vendor = [s for s in inventory["sources"] if s["category"] == "visible_alpha"]
            for source in vendor:
                add(source)
            if years is None:
                eligible = {c["year"] for source in store.sources for header in source.get("vendor_profile", {}).get("header_sets", {}).values()
                            for c in header["columns"].values() if c["eligible"] and c["slot"] == "FY" and c["year"] < date.today().year}
                if not eligible:
                    raise ValueError("Cannot infer completed years from an eligible export; supply --years explicitly")
                last = max(eligible)
                years = list(range(last - historical_years + 1, last + 1))
            company_catalog = [{k: s.get(k) for k in ("id", "local_path", "title", "category", "published_at", "filename_publication_hint")}
                               for s in inventory["sources"] if s["category"] != "visible_alpha"]
            selection = journal.stage("construction_select", SELECT_PROMPT,
                dict(company=company, years=years, catalog=company_catalog), SELECT,
                lambda answer: _valid_selection(answer, company_catalog))
            for source_id in selection["source_ids"]:
                add(next(s for s in inventory["sources"] if s["id"] == source_id))
            originals = [s for s in store.sources if s.get("category") != "visible_alpha"]
            primary = [s for s in originals if s.get("category") == "financial_reports"]
            def documents(sources):
                refs = [dict(source_id=s["id"], unit_id=next(iter(store._units[s["id"]]))) for s in sources]
                return store.read_documents(refs)
            evidence, primary_evidence = documents(originals), documents(primary)
            journal.artifact("parsed_documents.json", dict(evidence=evidence, vendor_catalog=metric_catalog(store, years)))
            coverage = journal.stage("construction_coverage", RULES + "Independently inventory the disclosed financial statements, segment versions, annual and interim coverage and reporting changes for the requested years. No candidate rows exist yet. Identify missing evidence and basis bridges; do not infer non-disclosure from an absent file.",
                dict(company=company, years=years, evidence=primary_evidence), REVIEW,
                lambda value: _valid_refs(value, store))
            journal.artifact("coverage_inventory.json", coverage)
            base = dict(company=company, years=years, evidence=evidence, vendor_catalog=metric_catalog(store, years))
            reused = configuration["reuse_spec"]
            if reused:
                if reused["company_query"] != company:
                    raise ValueError("Reusable specification belongs to a different company query")
                mapping = reused["mapping"]
                # Source identity includes bytes. Changed originals require an explicit remap.
                mapping = _rebind_mapping(mapping, reused["source_bindings"], store)
                history = normalise_history(mapping, store, years)
            else:
                scope = journal.stage("construction_history", HISTORY_PROMPT + "\nThis task establishes ONLY company identity, periods, reporting changes and gaps; rows come in later tasks. Fiscal IDs MUST be YYYY, YYYY-Q1..Q4, YYYY-H1/H2, YYYY-9M. A nine-month release usually reports cumulative 9M, not a standalone Q3. Do not invent periods. affected_rows can be empty until mappings are established.",
                    base, HISTORY_SCOPE, lambda value: _valid_scope(value, store, years))
                mapping = dict(scope, rows=[], identities=[])
                mapping_schema = copy.deepcopy(MAPPING_BATCH)
                vendor_refs = [r for r in metric_catalog(store, years) if r["role"] == "metric"]
                if vendor_refs:
                    fields = mapping_schema["properties"]["rows"]["items"]["properties"]["source_rows"]["items"]["properties"]
                    fields["source_id"] = _enum(sorted({r["source_id"] for r in vendor_refs}))
                    fields["unit_id"] = _enum(sorted({r["unit_id"] for r in vendor_refs}))
                for section in ("income_statement", "segments", "balance_sheet", "cash_flow", "schedule"):
                    def valid_batch(value):
                        probe = dict(mapping, rows=mapping["rows"] + value["rows"])
                        if any(row["section"] != section for row in value["rows"]):
                            raise ValueError("Map only the requested statement section")
                        normalise_history(probe, store, years)
                        return value
                    batch = journal.stage("construction_map", HISTORY_PROMPT + "\nThis is ONE statement mapping task, not a whole-company output. Start with up to 18 material core lines for this section, using source_rows for vendor rows and exact existing catalogue unit IDs. Do not emit other sections. List omitted supporting detail as gaps. Keep explanations concise. Never invent a source row/cell. Fiscal periods are already supplied and fixed.",
                        dict(base, reporting_scope=scope, requested_section=section,
                             existing_rows=[{"id": r["id"], "label": r["label"], "section": r["section"]} for r in mapping["rows"]]),
                        mapping_schema, valid_batch)
                    mapping["rows"].extend(batch["rows"])
                    mapping["gaps"].extend(batch["gaps"])
                    journal.artifact("mapping_progress.json", mapping)
                history = normalise_history(mapping, store, years)
                if not history["rows"]:
                    raise ValueError("No supported historical rows were mapped; forecast construction cannot proceed")
                relationships = journal.stage("construction_history", HISTORY_PROMPT + "\nOnly define supported accounting identities for the mapped rows. Use existing IDs, explicit period applicability and source rounding tolerance. Preserve discrepancies, do not introduce plugs or require a complete model. Unknown statement links remain gaps.",
                    dict(company=company, years=years, historical_rows=history["rows"], evidence=evidence), HISTORY_EQUATIONS,
                    lambda value: _valid_relationships(value, mapping, store, years))
                mapping["identities"] = relationships["identities"]
                mapping["gaps"].extend(relationships["gaps"])
                history = normalise_history(mapping, store, years)
            journal.artifact("historical_model.json", history)
            journal.artifact("consensus_reference.json", consensus_observations(store))
            journal.artifact("mapping.json", mapping)
            # Blind facts are recorded before any comparison. No values, source selections,
            # candidate definitions, coverage verdicts or earlier reviews are in this input.
            blind = dict(items=[], coverage=[], gaps=[])
            for start in range(0, len(history["rows"]), 18):
                facts = journal.stage("construction_blind", RULES + "Independently reconstruct ONLY the supplied concepts and periods from the complete primary documents. concept must be the supplied ID. Return original numeric references and scale multiplier into base units; preserve the source sign and definition. Record gaps. Do not fabricate support for every concept. Keep definitions concise; this is a bounded concept task, while documents remain whole.",
                    dict(company=company, concepts=[{"id": r["id"], "label": r["label"]} for r in history["rows"][start:start+18]],
                         periods=[{k: p[k] for k in PERIOD["properties"]} for p in history["periods"]], evidence=primary_evidence), BLIND,
                    lambda value: _valid_refs(value, store))
                for key in blind:
                    blind[key].extend(facts[key])
                journal.artifact("blind_source_facts_progress.json", blind)
            journal.artifact("blind_source_facts.json", blind)
            source_checks = compare_blind(history, blind, store)
            small_history = {k: v for k, v in history.items() if k not in ("imported_observations", "excluded_observations", "checks")}
            if reused:
                drivers, spec = reused["drivers"], reused["forecast"]
                validate_references(drivers, store)
                validate_references(spec, store)
            else:
                drivers = journal.stage("construction_drivers", DRIVER_PROMPT,
                    dict(company=company, historical_model=small_history, evidence=evidence), DRIVERS,
                    lambda value: _valid_refs(value, store))
                journal.artifact("drivers.json", drivers)
                spec = journal.stage("construction_formulas", FORMULA_PROMPT,
                    dict(company=company, historical_model=small_history, drivers=drivers, evidence=evidence), FORMULAS,
                    lambda value: _valid_spec(value, history, drivers, store))
            annuals = [p for p in history["periods"] if p["kind"] == "annual"]
            future = list(range(max(p["year"] for p in annuals) + 1, max(p["year"] for p in annuals) + forecast_years + 1))
            assumptions = configuration["assumptions"] or initialise_assumptions(history, drivers, future)
            validate_assumptions(assumptions, drivers, future)
            values = calculate(history, drivers, spec, assumptions, future)
            checks = history["checks"] + source_checks + validate_results(history, drivers, spec, assumptions, future, values)
            # Structural reviewer sees original context and proposed structure, not other verdicts.
            review = journal.stage("construction_review", RULES + "Challenge economic relationships, units, basis changes, segment lineage, fiscal periods, completeness, double counting, missing schedules and the company's growth conventions. This is also the combined multi-year history challenge. Do not approve based on balanced identities. Provide precise errors with source references; absence of criticism is not proof.",
                dict(company=company, historical_model=small_history, drivers=drivers, specification=spec, evidence=evidence), REVIEW,
                lambda value: _valid_refs(value, store))
            # A semantic/economic failure gets precise bounded Luna repairs, with
            # original context retained. Numerical source disagreements are not
            # repaired by overwriting reported data.
            for repair in range(max_repairs):
                failures = [c["detail"] for c in checks if c["area"] == "behaviour" and c["status"] == "FAIL"]
                failures += [i["detail"] for i in review["issues"] if i["severity"] == "error"]
                if not failures or reused:
                    break
                revised = journal.stage("construction_formulas", FORMULA_PROMPT,
                    dict(company=company, historical_model=small_history, drivers=drivers, evidence=evidence,
                         previous_specification=spec, economic_errors=failures, semantic_repair=repair + 1), FORMULAS,
                    lambda value: _valid_spec(value, history, drivers, store))
                if digest(revised) == digest(spec):
                    break
                spec = revised
                values = calculate(history, drivers, spec, assumptions, future)
                checks = history["checks"] + source_checks + validate_results(history, drivers, spec, assumptions, future, values)
                review = journal.stage("construction_review", RULES + "Independently challenge this revised model's economics, missing schedules, reporting changes and double counting against the whole originals. Prior verdicts are intentionally withheld.",
                    dict(company=company, historical_model=small_history, drivers=drivers, specification=spec, evidence=evidence), REVIEW,
                    lambda value: _valid_refs(value, store))
            for issue in review["issues"]:
                checks.append(check("economic_review", issue["area"], "FAIL" if issue["severity"] == "error" else "WARNING", issue["detail"]))
            gaps = list(dict.fromkeys(selection["gaps"] + history["gaps"] + drivers["gaps"] + spec["gaps"]
                    + blind["gaps"] + coverage["missing_data"] + review["missing_data"]))
            gaps += ["Temporary assumptions are a demonstration baseline, not empirically calibrated forecasts.",
                     "Source support comparison checks numerical agreement; semantic comparability and model accuracy remain provisional."]
            for row in history["rows"]:
                if row["id"] not in {r["id"] for r in spec["rows"]}:
                    gaps.append("No supported forecast calculation for historical line: " + row["id"])
                    checks.append(check("completeness", row["id"], "MISSING", "Historical line has no supported forecast calculation"))
            for section in ("income_statement", "balance_sheet", "cash_flow"):
                if not any(r["section"] == section for r in spec["rows"]):
                    gaps.append("Missing forecast statement: " + section)
                    checks.append(check("completeness", section, "MISSING", "No supported forecast statement"))
            if not primary:
                gaps.append("No primary company disclosures: identity, periods, historical support and driver economics are unverified")
            bindings = {s["id"]: {"local_path": s.get("local_path"), "sha256": s["sha256"]} for s in store.sources}
            bundle = dict(version=VERSION, company_query=company, mapping=mapping, drivers=drivers, forecast=spec, source_bindings=bindings)
            journal.artifact("model_specification.json", bundle)
            journal.artifact("drivers.json", drivers)
            journal.artifact("assumptions.json", assumptions)
            journal.artifact("forecast.json", values)
            journal.artifact("validation.json", dict(checks=checks, gaps=gaps, historical_acceptance="not_accepted_by_strict_initiation",
                framework_status="incomplete" if any(c["status"] in ("FAIL", "MISSING") for c in checks) else "executable_provisional"))
            inspection = export_model(directory / "model.xlsx", history, drivers, spec, assumptions, future, values, checks, store.sources, gaps)
            if recalculate:
                if progress:
                    progress("Recalculating generated workbook in hidden Microsoft Excel and comparing with Python...")
                inspection["recalculation"] = recalculate_workbook(directory / "model.xlsx", inspection)
            journal.artifact("workbook_validation.json", inspection)
            journal.state["artifacts"]["model.xlsx"] = hashlib.sha256((directory / "model.xlsx").read_bytes()).hexdigest()
            journal.state["status"] = "incomplete_model" if any(c["status"] in ("FAIL", "MISSING") for c in checks) else "complete_provisional"
            journal.state["summary"] = dict(history_rows=len(history["rows"]), drivers=len(drivers["drivers"]), forecast_rows=len(spec["rows"]),
                checks={status: sum(c["status"] == status for c in checks) for status in ("PASS", "FAIL", "MISSING", "WARNING")}, gaps=len(gaps))
            journal.save()
        except Exception as exc:
            journal.state["status"] = "incomplete"
            journal.state["issues"].append(str(exc))
            # Preserve an executable, explicitly incomplete review artifact when a
            # later review/repair cannot run. No manufactured forecast if formula
            # construction itself has not completed.
            if "values" in locals() and "spec" in locals() and spec.get("rows"):
                partial_checks = list(checks) if "checks" in locals() else []
                partial_checks.append(check("workflow", "unfinished_stage", "FAIL", str(exc)))
                partial_gaps = history["gaps"] + drivers["gaps"] + spec["gaps"] + [str(exc),
                    "INCOMPLETE: later review/repair did not finish. Source and model acceptance are not established."]
                journal.artifact("model_specification.json", dict(version=VERSION, company_query=company,
                    mapping=mapping, drivers=drivers, forecast=spec,
                    source_bindings={s["id"]: {"local_path": s.get("local_path"), "sha256": s["sha256"]} for s in store.sources}))
                journal.artifact("assumptions.json", assumptions)
                journal.artifact("forecast.json", values)
                journal.artifact("validation.json", dict(checks=partial_checks, gaps=partial_gaps,
                    historical_acceptance="not_accepted", framework_status="incomplete"))
                inspection = export_model(directory / "model.xlsx", history, drivers, spec, assumptions,
                    future, values, partial_checks, store.sources, partial_gaps)
                if recalculate:
                    inspection["recalculation"] = recalculate_workbook(directory / "model.xlsx", inspection)
                journal.artifact("workbook_validation.json", inspection)
                journal.state["artifacts"]["model.xlsx"] = hashlib.sha256((directory / "model.xlsx").read_bytes()).hexdigest()
            journal.save()
            if progress:
                progress("Stopped incomplete: " + str(exc))
        return directory


def _valid_selection(value, catalog):
    allowed = {s["id"] for s in catalog}
    if len(value["source_ids"]) != len(set(value["source_ids"])) or not set(value["source_ids"]) <= allowed:
        raise ValueError("Document selection has unknown or duplicate source IDs")
    return value


def consensus_observations(store):
    """Keep the vendor's future/reference values out of model actuals and inputs."""
    values = []
    for source in store.sources:
        if source.get("category") != "visible_alpha":
            continue
        for unit in store._units[source["id"]].values():
            row = unit.get("vendor_row", {})
            if row.get("role") != "metric" or row.get("section") != "company_reported":
                continue
            columns = source["vendor_profile"]["header_sets"][row["header_id"]]["columns"]
            for token in unit["numbers"]:
                column = columns.get(re.sub(r"\d+$", "", token["locator"]))
                if column and (column["estimate"] or column["series"].lower() == "consensus"):
                    values.append(dict(source_id=source["id"], unit_id=unit["unit_id"],
                        number_index=token["index"], locator=token["locator"], label=row["label"],
                        unit=row["unit"], raw=token["raw"], header=column))
    return dict(status="vendor_reference_only_not_our_forecast", observations=values,
                availability="Export vintage does not establish point-in-time availability")


def _valid_refs(value, store):
    validate_references(value, store)
    return value


def _valid_relationships(value, mapping, store, years):
    normalise_history(dict(mapping, identities=value["identities"]), store, years)
    return value


def _valid_scope(value, store, years):
    # Comparative/opening periods often appear alongside requested history. Keep
    # the complete AI response in the attempt artifact and explicitly project the
    # requested scope in Python instead of repaying for a prose correction.
    value = copy.deepcopy(value)
    outside = [p["id"] for p in value["periods"] if p["year"] not in years]
    value["periods"] = [p for p in value["periods"] if p["year"] in years]
    if outside:
        value["gaps"].append("Out-of-scope comparative/opening periods retained in raw scope response, excluded from requested history: " + ", ".join(outside))
    normalise_history(dict(value, rows=[], identities=[]), store, years)
    return value


def _valid_spec(value, history, drivers, store):
    validate_references(value, store)
    issues = validate_spec(history, drivers, value)
    if issues:
        raise ValueError("; ".join(issues))
    return value


def validate_assumptions(assumptions, drivers, years):
    expected = {(d["id"], str(y)) for d in drivers["drivers"] for y in years}
    actual = {(a["driver"], a["period"]) for a in assumptions}
    if actual != expected or len(assumptions) != len(expected):
        raise ValueError("Assumption interfaces must cover each driver and forecast year exactly once")
    by_id = {d["id"]: d for d in drivers["drivers"]}
    for item in assumptions:
        driver = by_id[item["driver"]]
        if any(item[key] != driver[key] for key in ("entity", "unit", "scale")):
            raise ValueError("Assumption unit/scale/entity differs from driver interface")
        if item["source"] not in ("reported_observation", "derived_historical_value", "temporary_manual_assumption", "consensus_reference", "external_forecasting_model") or not item["provenance"]:
            raise ValueError("Assumption needs an explicit source category and provenance")
    if len({a["scenario"] for a in assumptions}) != 1:
        raise ValueError("A calculation consumes one consistent scenario path")


def compare_blind(history, blind, store):
    checks = []
    facts = {}
    for item in blind["items"]:
        if not math.isfinite(item["multiplier"]) or item["multiplier"] <= 0:
            raise ValueError("Blind reconstruction must preserve source sign and use a positive finite scale")
        source = next(s for s in store.sources if s["id"] == item["reference"]["source_id"])
        if source.get("category") != "financial_reports":
            raise ValueError("Blind reconstruction must use company financial_reports")
        facts.setdefault((item["concept"], item["period"]), []).append(item)
    for row in history["rows"]:
        for period, value in row["values"].items():
            matching = facts.get((row["id"], period), [])
            eligible = [f for f in matching if f["unit"] == row["unit"]]
            if not eligible:
                checks.append(check("source", row["id"] + ":" + period, "MISSING", "No independent primary numeric reconstruction on matching units"))
                continue
            numbers = [store.resolve(f["reference"]) * f["multiplier"] * row["sign"] for f in eligible]
            # Signs remain visible. No absolute-value matching that hides a reversed expense.
            difference = value - numbers[0]
            tolerance = row["scale"] * .5
            status = "PASS" if all(abs(value - n) <= tolerance for n in numbers) else "FAIL"
            checks.append(check("source", row["id"] + ":" + period, status,
                "Blind primary numerical comparison; definitions still require semantic review. " + eligible[0]["definition"], difference, tolerance))
    return checks


def _rebind_mapping(mapping, bindings, store):
    """Reuse unchanged mappings exactly. New bytes require semantic remapping."""
    current = {s["id"]: s for s in store.sources}
    for old_id, binding in bindings.items():
        if old_id not in current or current[old_id]["sha256"] != binding["sha256"]:
            raise ValueError("Reusable source mapping changed; refresh with a new mapping instead of applying stale row locations")
    return copy.deepcopy(mapping)




def review_saved(root, run, *, assumptions_file=None, recalculate=False):
    """Unpaid export/recalculation of saved work, including a blocked partial run.

    No provider object or credentials are read. Missing forecast specifications stay
    missing; this does not manufacture company drivers to bypass an interrupted AI.
    """
    root, run = Path(root).resolve(), Path(run).resolve()
    outputs = root / "outputs"
    if run.parent != outputs:
        raise ValueError("Select a saved construction run inside this project's outputs")
    with _initiation_lock(outputs):
        record = json.loads((run / "record.json").read_text(encoding="utf-8"))
        if not (run / "historical_model.json").exists():
            raise ValueError("No completed historical mapping is available to export")
        verify_source_bytes(record.get("sources", []), outputs)
        for name in ("historical_model.json", "mapping.json", "model_specification.json", "assumptions.json", "validation.json"):
            if (run / name).exists():
                expected = record["artifacts"].get(name)
                if expected != hashlib.sha256((run / name).read_bytes()).hexdigest():
                    raise ValueError("Saved input artifact changed: " + name)
        history = json.loads((run / "historical_model.json").read_text(encoding="utf-8"))
        last = max(p["year"] for p in history["periods"] if p["kind"] == "annual")
        horizon = record["configuration"]["forecast_years"]
        years = list(range(last + 1, last + horizon + 1))
        bundle_path = run / "model_specification.json"
        if bundle_path.exists():
            bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
            drivers, spec = bundle["drivers"], bundle["forecast"]
            assumptions = json.loads(Path(assumptions_file or run / "assumptions.json").read_text(encoding="utf-8"))
            validate_assumptions(assumptions, drivers, years)
            values = calculate(history, drivers, spec, assumptions, years)
            old_checks = json.loads((run / "validation.json").read_text(encoding="utf-8"))
            checks = [c for c in old_checks["checks"] if c["area"] not in ("formula", "accounting", "behaviour")]
            checks += validate_results(history, drivers, spec, assumptions, years, values)
            gaps = old_checks["gaps"] + ["Offline recalculation: source/AI reviews remain frozen and are not renewed by edits."]
        else:
            if assumptions_file:
                raise ValueError("Forecast construction has not completed; assumptions cannot create a missing financial specification")
            drivers, spec, assumptions, values = dict(drivers=[], gaps=[]), dict(rows=[], identities=[], gaps=[]), [], {}
            checks = list(history["checks"]) + [check("workflow", "forecast_construction", "MISSING",
                "Live workflow stopped before driver and forecast generation; forecast columns intentionally unmodelled"),
                check("source", "independent_reconstruction", "MISSING", "Source review did not finish; history is an unaccepted draft")]
            gaps = history["gaps"] + record["issues"] + ["Driver register and executable company forecast are unavailable; see the separate synthetic known-answer workbook for calculator verification."]
        mapped_segments = {r["id"] for r in history["rows"] if r["section"] == "segments"}
        if mapped_segments and not any(identity["target"] in mapped_segments for identity in history["identities"]):
            checks.append(check("coverage", "segment_reconciliation", "MISSING", "No supported segment-to-group equation was produced; overlapping segment versions must not be summed"))
        if not any(next((r["section"] for r in history["rows"] if r["id"] == i["target"]), None) == "cash_flow" for i in history["identities"]):
            checks.append(check("coverage", "cash_flow_reconciliation", "MISSING", "No supported cash-flow reconciliation equation was produced"))
        directory = outputs / f"review_saved_{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}_{uuid4().hex[:8]}"
        directory.mkdir()
        for name, value in (("historical_model.json", history), ("drivers.json", drivers), ("forecast_specification.json", spec),
                            ("assumptions.json", assumptions), ("forecast.json", values),
                            ("validation.json", dict(checks=checks, gaps=gaps, status="incomplete"))):
            write_json(directory / name, value)
        inspection = export_model(directory / "model.xlsx", history, drivers, spec, assumptions, years, values, checks, record["sources"], gaps)
        if recalculate:
            inspection["recalculation"] = recalculate_workbook(directory / "model.xlsx", inspection)
        write_json(directory / "workbook_validation.json", inspection)
        write_json(directory / "record.json", dict(status="offline_review_only", source_run=str(run),
            new_paid_calls=0, new_usage=aggregate_usage([]), source_run_usage=record["usage"],
            forecast_rows=len(spec["rows"]), historical_rows=len(history["rows"]),
            workbook_sha256=hashlib.sha256((directory / "model.xlsx").read_bytes()).hexdigest(),
            limitation="No historical/source/AI acceptance is conferred; original run and unknown charge remain unchanged"))
        return directory


def _integrated_plan(spec):
    """Validate the company-owned graph without importing company assumptions."""
    if not isinstance(spec, dict) or not isinstance(spec.get("rows"), list) or not spec["rows"]:
        raise ValueError("Integrated specification requires populated financial rows")
    histories, forecasts = spec.get("historical_years", []), spec.get("forecast_years", [])
    for label, years in (("Historical", histories), ("Forecast", forecasts)):
        if (not isinstance(years, list) or any(type(year) is not int for year in years)
                or len(set(years)) != len(years) or years != sorted(years)):
            raise ValueError(label + " years must be unique ascending annual integers")
    if not forecasts or set(histories) & set(forecasts) or histories and forecasts[0] <= histories[-1]:
        raise ValueError("Forecast years must follow historical years without overlap")
    if any(b != a + 1 for a, b in zip(forecasts, forecasts[1:])):
        raise ValueError("Forecast years must be consecutive for annual roll-forwards")
    sections = spec.get("sections", [])
    section_ids = [section["id"] for section in sections]
    if len(section_ids) != len(set(section_ids)):
        raise ValueError("Duplicate integrated section IDs")
    rows = {}
    years = histories + forecasts
    for row in spec["rows"]:
        key = row.get("id")
        if not isinstance(key, str) or not re.fullmatch(r"[a-z][a-z0-9_]*", key):
            raise ValueError("Invalid integrated row ID: " + str(key))
        if key in rows:
            raise ValueError("Duplicate integrated row ID: " + key)
        if row.get("section") not in section_ids:
            raise ValueError(key + ": unknown financial section")
        if row.get("kind") not in ("reported", "calculated", "assumption"):
            raise ValueError(key + ": kind must be reported, calculated or assumption")
        unit_dimension(row.get("unit", ""))
        for field in ("historical", "inputs"):
            series = row.get(field, {})
            if not isinstance(series, dict):
                raise ValueError(key + ": " + field + " must map period strings to finite numbers or null")
            applicable = histories if field == "historical" else forecasts
            for year, value in series.items():
                if not isinstance(year, str) or year not in {str(y) for y in applicable}:
                    raise ValueError(key + ": unexpected " + field + " period " + str(year))
                if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
                    raise ValueError(key + ": values must be finite numbers or null")
        if row.get("inputs") and row["kind"] != "assumption":
            raise ValueError(key + ": only assumption rows may contain editable forecast inputs")
        if row["kind"] == "assumption" and row.get("expression"):
            raise ValueError(key + ": an assumption cannot have both an input and a forecast formula")
        rows[key] = row
    units = {key: row["unit"] for key, row in rows.items()}
    trees = {}
    for key, row in rows.items():
        for field in ("expression", "historical_expression"):
            if not row.get(field):
                continue
            tree = expression_tree(row[field])
            missing = {dep for dep, _ in dependencies(tree)} - rows.keys()
            if missing:
                raise ValueError(key + ": unknown references " + ", ".join(sorted(missing)))
            if dimension(tree, units) != unit_dimension(row["unit"]):
                raise ValueError(key + ": formula units disagree with declared display unit")
            trees[key, field] = tree
    identity_ids = set()
    for identity in spec.get("identities", []):
        key = identity["id"]
        if key in identity_ids:
            raise ValueError("Duplicate integrated identity ID: " + key)
        identity_ids.add(key)
        tree = expression_tree(identity["expression"])
        deps = dependencies(tree)
        if not deps or {dep for dep, _ in deps} - rows.keys():
            raise ValueError(key + ": identity needs known financial references")
        dimension(tree, units)
        tolerance = identity.get("tolerance")
        if type(tolerance) not in (float, int) or not math.isfinite(tolerance) or tolerance < 0:
            raise ValueError(key + ": identity tolerance must be finite and nonnegative")
        if type(identity.get("independent")) is not bool:
            raise ValueError(key + ": explicitly distinguish independent controls from constructed identities")
        periods = identity.get("periods", [])
        if not isinstance(periods, list) or len(periods) != len(set(periods)) or any(
                not isinstance(period, str) or period not in {str(y) for y in years} for period in periods):
            raise ValueError(key + ": identity periods must belong to the modeled annual scope")
        trees[key, "identity"] = tree
    return rows, histories, forecasts, trees


def calculate_integrated(spec, overrides=None):
    """Execute a reusable annual financial graph in its declared display units.

    Historical observations take precedence over historical formulas. Overrides
    apply only to explicit forecast assumptions, never to source observations.
    Missing data and division by zero propagate as None rather than zero.
    """
    rows, histories, forecasts, trees = _integrated_plan(spec)
    overrides = overrides or {}
    if not isinstance(overrides, dict):
        raise ValueError("Overrides must map driver IDs to forecast period values")
    amended = {}
    for key, series in overrides.items():
        if key not in rows or rows[key]["kind"] != "assumption" or not isinstance(series, dict):
            raise ValueError("Override must name an editable assumption: " + str(key))
        amended[key] = {}
        for year, value in series.items():
            period = str(year)
            if period not in {str(y) for y in forecasts}:
                raise ValueError(key + ": overrides cannot change historical or unmodeled periods")
            if period in amended[key]:
                raise ValueError(key + ": duplicate normalized override period")
            if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
                raise ValueError(key + ": override must be finite or null")
            amended[key][period] = value
    values = {key: {} for key in rows}
    evaluating = set()
    modeled = set(histories + forecasts)

    def resolve(key, year):
        if year not in modeled:
            return None
        period, pair = str(year), (key, year)
        if period in values[key]:
            return values[key][period]
        if pair in evaluating:
            raise ValueError("Circular same-period integrated dependency: " + key + ":" + period)
        row = rows[key]
        historical = year in histories
        observed = row.get("historical", {}).get(period) if historical else None
        if historical and observed is not None:
            values[key][period] = observed
            return observed
        if not historical and row["kind"] == "assumption":
            values[key][period] = amended.get(key, {}).get(period, row.get("inputs", {}).get(period))
            return values[key][period]
        field = "historical_expression" if historical else "expression"
        if historical and "historical_expression_periods" in row and period not in row["historical_expression_periods"]:
            values[key][period] = None
            return None
        tree = trees.get((key, field))
        if tree is None:
            values[key][period] = None
            return None
        evaluating.add(pair)
        try:
            value = formula_value(tree, lambda dependency, lag: resolve(dependency, year - lag))
            if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
                raise ValueError(key + ": nonfinite calculated result in " + period)
            values[key][period] = value
            return value
        finally:
            evaluating.remove(pair)

    for year in histories + forecasts:
        for key in rows:
            resolve(key, year)
    return values


def validate_integrated(spec, values):
    """Check actual company graph results without treating shared arithmetic as source proof."""
    try:
        rows, histories, forecasts, trees = _integrated_plan(spec)
        # Structure/cycle verification does not replace supplied-result checks.
        calculate_integrated(spec)
    except (ValueError, KeyError, TypeError, RecursionError) as exc:
        return [check("framework", "integrated_graph", "FAIL", str(exc))]
    checks = []
    if not isinstance(values, dict):
        return [check("framework", "integrated_values", "FAIL", "Calculated values must be a row/period mapping")]
    unknown = set(values) - rows.keys()
    if unknown:
        checks.append(check("framework", "unknown_result_rows", "FAIL", ", ".join(sorted(unknown))))

    def resolve(key, year):
        series = values.get(key, {})
        value = series.get(str(year)) if isinstance(series, dict) else None
        return value if type(value) in (int, float) and math.isfinite(value) else None

    for key, row in rows.items():
        series = values.get(key, {})
        if not isinstance(series, dict):
            checks.append(check("formula", key, "FAIL", "Row values must map period strings to numbers or null"))
            continue
        for year in histories + forecasts:
            period = str(year)
            value = series.get(period)
            finite = value is None or type(value) in (float, int) and math.isfinite(value)
            if not finite:
                checks.append(check("formula", key + ":" + period, "FAIL", "Nonfinite or nonnumeric result"))
                continue
            if year in forecasts:
                status = "MISSING" if value is None and not row.get("optional", False) else "PASS"
                checks.append(check("framework", key + ":" + period, status,
                    "Required forecast input or financial line remains unresolved" if status == "MISSING" else
                    "Optional missing line" if value is None else "Populated executable forecast"))
            if row.get("nonnegative") and value is not None:
                checks.append(check("accounting", "nonnegative:" + key + ":" + period,
                    "PASS" if value >= 0 else "FAIL", "This stock must be nonnegative", min(value, 0), 0))
            if year in histories and row.get("historical", {}).get(period) is not None:
                observed = row["historical"][period]
                difference = None if value is None else value - observed
                checks.append(check("source", "preserved_observation:" + key + ":" + period,
                    "PASS" if difference == 0 else "FAIL",
                    "Compared with the saved input observation; independent original-cell validation remains separate",
                    difference, 0))
                continue
            field = "historical_expression" if year in histories else "expression"
            if year in histories and "historical_expression_periods" in row and period not in row["historical_expression_periods"]:
                continue
            tree = trees.get((key, field))
            if tree is not None:
                expected = formula_value(tree, lambda dependency, lag: resolve(dependency, year - lag))
                difference = None if value is None or expected is None else value - expected
                tolerance = 1e-9
                checks.append(check("formula", "expression:" + key + ":" + period,
                    "MISSING" if difference is None else "PASS" if abs(difference) <= tolerance else "FAIL",
                    "Result agrees with the declared financial expression; this is an arithmetic check",
                    difference, tolerance))
    for identity in spec.get("identities", []):
        tree = trees[identity["id"], "identity"]
        for period in identity.get("periods") or [str(y) for y in forecasts]:
            difference = formula_value(tree, lambda dependency, lag: resolve(dependency, int(period) - lag))
            tolerance = identity["tolerance"]
            status = "MISSING" if difference is None else "PASS" if abs(difference) <= tolerance else "FAIL"
            evidence = ("Independent control specification; its evidence must be checked separately" if identity["independent"] else
                        "Constructed accounting identity, not independent source support")
            checks.append(check("accounting", identity["id"] + ":" + period, status,
                identity.get("label", identity["id"]) + ". " + evidence, difference, tolerance))
    if not any(row["kind"] == "assumption" for row in rows.values()):
        checks.append(check("framework", "assumption_controls", "MISSING", "No editable forecast driver controls"))
    if not any(row.get("expression") for row in rows.values()):
        checks.append(check("framework", "forecast_formulas", "MISSING", "No executable forecast expressions"))
    if not spec.get("identities"):
        checks.append(check("framework", "accounting_controls", "MISSING", "No financial reconciliation identities"))
    return checks
