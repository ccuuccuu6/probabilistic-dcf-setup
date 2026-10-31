"""Shared Excel presentation: fonts, colours, number formats and table layout.

Historical, forecast and valuation workbooks use these same native Python rules.
This module contains no company inputs or financial calculation logic.
"""
import json
import math

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


STYLE_VERSION = "financial-model-5"
NAVY = "183B56"
INK = "111827"
PINK = "D63384"
BLUE = "0000FF"
GREEN = "008000"
YELLOW = "FFFFCC"
PALE = "E4EFF3"
GREY = "64748B"
NUMBER = '#,##0.0;(#,##0.0);"–"'
RATE = '0.0%;(0.0%);"–"'
BOX = Border(*(Side(style="thin", color="B2B2B2") for _ in range(4)))
TOTAL = Border(top=Side(style="thin", color="AABAC6"), bottom=Side(style="thin", color="AABAC6"))


def _text(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else value


def _number(value):
    return type(value) in (int, float) and math.isfinite(value)


def _format(unit):
    if unit == "exchange_rate":
        return '0.000000;(0.000000);"–"'
    if unit in ("ratio", "%", "rate"):
        return RATE
    if unit == "bps":
        return '0;[Red](0);"–"'
    return NUMBER


def _display_unit(unit):
    if unit.endswith("m/shares_m"):
        return unit.replace("m/shares_m", "/share")
    return {"ratio": "%", "shares_m": "m shares", "exchange_rate": "quote / unit"}.get(unit, unit)


def _style_sheet(sheet, freeze="D6"):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 70
    sheet.freeze_panes = freeze
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = False
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.print_title_rows = "1:5"
    sheet.sheet_format.defaultRowHeight = 15
    sheet.sheet_format.defaultColWidth = 14
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.page_margins.bottom = 0.5
    sheet.page_margins.footer = 1 / 6


def _cell(sheet, row, column, value=None, *, color=INK, bold=False, fill=None, number_format=None):
    cell = sheet.cell(row, column, value)
    cell.font = Font(name="Aptos", size=10, color=color, bold=bold)
    cell.alignment = Alignment(vertical="center")
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if number_format:
        cell.number_format = number_format
    return cell


def _bar(sheet, row, label, last_column):
    for column in range(2, last_column + 1):
        cell = _cell(sheet, row, column, label if column == 2 else None,
                     color="FFFFFF", bold=True, fill=NAVY)
        cell.font = Font(name="Aptos", size=11, color="FFFFFF", bold=True)
    sheet.row_dimensions[row].height = 27


def _table(workbook, name, columns, records, widths=None):
    sheet = workbook.create_sheet(name)
    _style_sheet(sheet, "B6")
    _cell(sheet, 2, 2, name, bold=True)
    for column, label in enumerate(columns, 1):
        _cell(sheet, 5, column, label, bold=True, color="FFFFFF", fill=NAVY)
        sheet.column_dimensions[get_column_letter(column)].width = (widths or {}).get(column, 24)
    for row, record in enumerate(records, 6):
        for column, value in enumerate(record, 1):
            c = _cell(sheet, row, column, _text(value))
            c.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = 30
    sheet.auto_filter.ref = f"A5:{get_column_letter(len(columns))}{max(5, sheet.max_row)}"
    return sheet
